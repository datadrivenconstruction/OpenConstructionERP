# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The progress sheet: an XLSX site staff fill in offline.

One row per activity. Identity and plan columns are locked; the white
columns (percent complete, actual start and finish, remaining duration,
notes) are the only ones a site can edit, each with data validation. Two
hidden columns carry the activity id and a stamp of the activity's progress
state at export, and a hidden sheet carries the schedule id, the export time
and the column order, so the upload is read by position, whatever language
the headers were printed in, and an activity changed in the app since export
is recognised.

Pure: builds bytes and reads bytes. Persistence lives in
:mod:`app.modules.schedule.progress_sheet_service`.
"""

from __future__ import annotations

import hashlib
import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

SHEET_FORMAT = "oe.schedule.progress_sheet"
SHEET_VERSION = 1
META_SHEET = "_oe_meta"
DATA_SHEET_INDEX = 0

#: (key, column width, editable). The order is the sheet's column order.
COLUMNS: tuple[tuple[str, int, bool], ...] = (
    ("activity_ref", 16, False),
    ("name", 42, False),
    ("wbs_code", 12, False),
    ("responsible", 22, False),
    ("planned_start", 14, False),
    ("planned_finish", 14, False),
    ("percent_complete", 13, True),
    ("actual_start", 14, True),
    ("actual_finish", 14, True),
    ("remaining_duration", 15, True),
    ("notes", 42, True),
    ("row_key", 38, False),
    ("stamp", 20, False),
)
HIDDEN_COLUMNS = frozenset({"row_key", "stamp"})
DATE_COLUMNS = frozenset({"planned_start", "planned_finish", "actual_start", "actual_finish"})
EDITABLE_COLUMNS = tuple(key for key, _w, editable in COLUMNS if editable)
MAX_ROWS = 5000

_LOCKED_FILL = "EEF0F3"
_HEADER_FILL = "D9DEE7"


class SheetFormatError(ValueError):
    """The upload is not a progress sheet this platform can read."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def progress_stamp(
    *,
    start_date: str | None,
    end_date: str | None,
    progress_pct: float,
    remaining_duration: int | None,
    status: str | None,
    actual_start: str | None,
    actual_finish: str | None,
) -> str:
    """A short digest of an activity's progress state.

    Exported with the row and recomputed on upload: a different value means
    the activity's dates or progress were changed in the app after export.
    """
    payload = json.dumps(
        [
            start_date or "",
            end_date or "",
            round(float(progress_pct or 0.0), 2),
            remaining_duration,
            status or "",
            actual_start or "",
            actual_finish or "",
        ],
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


@dataclass
class ExportRow:
    """One activity as it goes into the sheet."""

    row_key: str
    activity_ref: str
    name: str
    wbs_code: str
    responsible: str
    planned_start: str | None
    planned_finish: str | None
    percent_complete: float
    actual_start: str | None
    actual_finish: str | None
    remaining_duration: int | None
    stamp: str
    notes: str = ""


def _as_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def build_workbook(
    rows: list[ExportRow],
    *,
    headers: dict[str, str],
    sheet_title: str,
    title: str,
    details: list[str],
    instructions: str,
    meta: dict[str, Any],
    letterhead: bool = True,
) -> bytes:
    """Write the progress sheet and return the XLSX bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    from app.core.xlsx_text import store_strings_as_text

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Progress"
    keys = [key for key, _w, _e in COLUMNS]
    ws.append([headers.get(key, key) for key in keys])
    for row in rows:
        values: list[Any] = []
        for key in keys:
            value = getattr(row, key)
            if key in DATE_COLUMNS:
                value = _as_date(value)
            values.append(value)
        ws.append(values)
    store_strings_as_text(ws)

    header_row = 1
    if letterhead:
        from app.core.xlsx_branding import apply_company_header

        header_row = apply_company_header(ws, title=title, details=[*details, instructions])

    locked_fill = PatternFill("solid", fgColor=_LOCKED_FILL)
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    last_row = header_row + len(rows)
    for idx, (key, width, editable) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = width
        if key in HIDDEN_COLUMNS:
            ws.column_dimensions[letter].hidden = True
        head = ws.cell(row=header_row, column=idx)
        head.font = Font(bold=True)
        head.fill = header_fill
        head.alignment = Alignment(wrap_text=True, vertical="center")
        for r in range(header_row + 1, last_row + 1):
            cell = ws.cell(row=r, column=idx)
            if editable:
                cell.protection = Protection(locked=False)
            else:
                cell.fill = locked_fill
            if key in DATE_COLUMNS:
                cell.number_format = "yyyy-mm-dd"
            elif key == "percent_complete":
                cell.number_format = "0.##"
    if rows:
        col = {key: get_column_letter(i) for i, (key, _w, _e) in enumerate(COLUMNS, start=1)}
        first = header_row + 1

        def _range(key: str) -> str:
            return f"{col[key]}{first}:{col[key]}{last_row}"

        pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True)
        pct.error = headers.get("percent_complete", "0-100")
        pct.showErrorMessage = True
        pct.add(_range("percent_complete"))
        dates = DataValidation(
            type="date", operator="between", formula1="DATE(1990,1,1)", formula2="DATE(2100,12,31)", allow_blank=True
        )
        dates.showErrorMessage = True
        dates.add(_range("actual_start"))
        dates.add(_range("actual_finish"))
        remaining = DataValidation(type="whole", operator="between", formula1="0", formula2="100000", allow_blank=True)
        remaining.showErrorMessage = True
        remaining.add(_range("remaining_duration"))
        notes = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="2000", allow_blank=True)
        notes.add(_range("notes"))
        for dv in (pct, dates, remaining, notes):
            ws.add_data_validation(dv)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS) - len(HIDDEN_COLUMNS))}{last_row}"

    meta_ws = wb.create_sheet(META_SHEET)
    full_meta = {
        **meta,
        "format": SHEET_FORMAT,
        "version": SHEET_VERSION,
        "header_row": header_row,
        "columns": keys,
        "row_count": len(rows),
    }
    for key, value in full_meta.items():
        meta_ws.append([key, json.dumps(value, ensure_ascii=False, default=str)])
    meta_ws.sheet_state = "veryHidden"
    meta_ws.protection.sheet = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class ParsedRow:
    """One data row read back from an uploaded sheet."""

    row_number: int
    row_key: str | None
    activity_ref: str | None
    name: str | None
    percent_complete: float | None
    actual_start: str | None
    actual_finish: str | None
    remaining_duration: int | None
    notes: str | None
    stamp: str | None
    errors: list[str] = field(default_factory=list)


@dataclass
class ParsedSheet:
    """The uploaded sheet: its export metadata and its rows."""

    meta: dict[str, Any]
    rows: list[ParsedRow]
    sha256: str


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell_date(value: Any) -> tuple[str | None, bool]:
    """``(iso, ok)``; a blank cell is ``(None, True)``."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, True
    if isinstance(value, datetime):
        return value.date().isoformat(), True
    if isinstance(value, date):
        return value.isoformat(), True
    try:
        return date.fromisoformat(str(value).strip()[:10]).isoformat(), True
    except ValueError:
        return None, False


def _cell_percent(value: Any, number_format: str | None) -> tuple[float | None, bool]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, True
    if isinstance(value, bool):
        return None, False
    if isinstance(value, (int, float)):
        number = float(value)
        if number_format and "%" in number_format:
            number *= 100.0
        return number, True
    text = str(value).strip().replace(",", ".").rstrip("%").strip()
    try:
        return float(text), True
    except ValueError:
        return None, False


def _cell_whole(value: Any) -> tuple[int | None, bool]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, True
    if isinstance(value, bool):
        return None, False
    try:
        number = float(str(value).strip().replace(",", ".")) if isinstance(value, str) else float(value)
    except ValueError:
        return None, False
    if not number.is_integer():
        return None, False
    return int(number), True


def read_meta(wb: Any) -> dict[str, Any]:
    """The export metadata of a workbook, or :class:`SheetFormatError`."""
    if META_SHEET not in wb.sheetnames:
        raise SheetFormatError("not_a_progress_sheet", "The file is not a progress sheet exported from a schedule")
    meta: dict[str, Any] = {}
    for key, raw in wb[META_SHEET].iter_rows(min_row=1, max_col=2, values_only=True):
        if not key:
            continue
        try:
            meta[str(key)] = json.loads(raw) if isinstance(raw, str) else raw
        except ValueError:
            meta[str(key)] = raw
    if meta.get("format") != SHEET_FORMAT:
        raise SheetFormatError("not_a_progress_sheet", "The file is not a progress sheet exported from a schedule")
    if int(meta.get("version") or 0) > SHEET_VERSION:
        raise SheetFormatError("unsupported_sheet_version", "The sheet was written by a newer version")
    return meta


def parse_workbook(data: bytes) -> ParsedSheet:
    """Read an uploaded progress sheet; raises :class:`SheetFormatError`."""
    from openpyxl import load_workbook

    sha256 = hashlib.sha256(data).hexdigest()
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for a non-XLSX upload
        raise SheetFormatError("unreadable_file", "The file could not be read as an XLSX workbook") from exc
    try:
        meta = read_meta(wb)
        columns = [str(c) for c in meta.get("columns") or []]
        header_row = int(meta.get("header_row") or 1)
        if not columns or "row_key" not in columns:
            raise SheetFormatError("not_a_progress_sheet", "The sheet carries no column map")
        index = {key: i for i, key in enumerate(columns)}
        ws = wb.worksheets[DATA_SHEET_INDEX]
        rows: list[ParsedRow] = []
        for offset, cells in enumerate(ws.iter_rows(min_row=header_row + 1, max_col=len(columns))):
            row_number = header_row + 1 + offset
            values = [c.value for c in cells] + [None] * (len(columns) - len(cells))
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue
            if len(rows) >= MAX_ROWS:
                raise SheetFormatError("too_many_rows", f"The sheet has more than {MAX_ROWS} rows")

            def get(key: str, _values: list[Any] = values) -> Any:
                i = index.get(key)
                return _values[i] if i is not None and i < len(_values) else None

            errors: list[str] = []
            pct_cell = cells[index["percent_complete"]] if index.get("percent_complete", 99) < len(cells) else None
            percent, ok = _cell_percent(get("percent_complete"), getattr(pct_cell, "number_format", None))
            if not ok:
                errors.append("percent_not_a_number")
            actual_start, ok = _cell_date(get("actual_start"))
            if not ok:
                errors.append("actual_start_not_a_date")
            actual_finish, ok = _cell_date(get("actual_finish"))
            if not ok:
                errors.append("actual_finish_not_a_date")
            remaining, ok = _cell_whole(get("remaining_duration"))
            if not ok:
                errors.append("remaining_not_a_whole_number")
            rows.append(
                ParsedRow(
                    row_number=row_number,
                    row_key=_cell_text(get("row_key")),
                    activity_ref=_cell_text(get("activity_ref")),
                    name=_cell_text(get("name")),
                    percent_complete=percent,
                    actual_start=actual_start,
                    actual_finish=actual_finish,
                    remaining_duration=remaining,
                    notes=_cell_text(get("notes")),
                    stamp=_cell_text(get("stamp")),
                    errors=errors,
                )
            )
        return ParsedSheet(meta=meta, rows=rows, sha256=sha256)
    finally:
        wb.close()
