# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The hakediş (Turkish progress payment certificate) as an Excel workbook.

Two sheets carrying what the PDF carries, from the same print model, so the
two cannot disagree: the summary ("Hakediş Raporu") and the works list
("Yapılan İşler Listesi"). A cost control department reworks a certificate in
a spreadsheet, so every figure is a real number, stored as the exact decimal
that was computed and never through a float, under a number format with a
thousands separator and the currency's own decimals. Excel draws the
separators in the reader's regional settings, which on a Turkish machine is
``1.234,56``. Dates are written out in the country's own order.

No formulas: the totals are the computed totals, the same ones the PDF prints.
A held line has no number to give; its cell holds the same marker the PDF
prints, as text, so a sum over the column cannot quietly swallow it as a zero.

Each sheet freezes its header rows, repeats them on every printed page, sets
its print area and fits its width to one page, and gets the company letterhead
from :func:`app.core.xlsx_branding.apply_company_header`.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.currency_registry import minor_units
from app.core.regional_format import format_date
from app.core.xlsx_branding import apply_company_header
from app.core.xlsx_text import store_strings_as_text
from app.modules.contracts.hakedis import (
    Certificate,
    header_rows,
    note_text,
    printed_summary,
    printed_works,
)
from app.modules.contracts.hakedis_layout import HakedisSettings, label, locale_languages, upper_for

_ALERT = "B91C1C"
_FILL = PatternFill("solid", fgColor="F3F4F6")
_THIN = Side(style="thin", color="9CA3AF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)

#: Column widths of the works list by kind, in Excel character units.
_WIDTHS: dict[str, float] = {
    "seq": 6,
    "text": 16,
    "unit": 8,
    "quantity": 14,
    "unit_price": 14,
    "money": 18,
    "percent": 12,
}
_DESCRIPTION_WIDTH = 48


def _money_format(currency: str) -> str:
    places = minor_units(currency)
    return "#,##0" + ("." + "0" * places if places else "")


def _number_format(kind: str, value: Decimal | int, currency: str) -> str:
    """A format that shows the decimals the value carries, within the kind's range."""
    if kind == "seq" or isinstance(value, int):
        return "0"
    if kind == "money":
        return _money_format(currency)
    exponent = value.normalize().as_tuple().exponent
    carried = -exponent if isinstance(exponent, int) and exponent < 0 else 0
    least, most = {"percent": (2, 6), "unit_price": (minor_units(currency), 4)}.get(kind, (2, 4))
    places = min(max(carried, least), most)
    return "#,##0" + ("." + "0" * places if places else "")


def _join(parts: tuple[str, ...]) -> str:
    """Both languages in one cell, the second on its own line."""
    return "\n".join(part for part in parts if part)


def _page_setup(ws: Any, *, landscape: bool, last_column: int, last_row: int, title_rows: str) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{get_column_letter(last_column)}{last_row}"
    ws.print_title_rows = title_rows
    ws.print_options.horizontalCentered = True


def _summary_sheet(
    ws: Any,
    cert: Certificate,
    locale: str,
    settings: HakedisSettings | None,
    issue_date: date | None,
) -> None:
    languages = locale_languages(locale)
    overrides = settings.labels if settings is not None else None
    inp = cert.inp
    money = _money_format(inp.currency)

    def words(key: str, **params: str) -> tuple[str, ...]:
        return tuple(label(key, language, overrides, **params) for language in languages)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 22

    row = 1
    if cert.is_draft:
        cell = ws.cell(row=row, column=1, value=" / ".join(words("draft.banner")))
        cell.font = Font(bold=True, color=_ALERT)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1
    for names, value in header_rows(cert, locale, settings, issue_date):
        ws.cell(row=row, column=1, value=_join(names)).font = Font(bold=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row=row, column=3, value=value).alignment = _WRAP_TOP
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        ws.cell(row=row, column=1).alignment = _WRAP_TOP
        row += 1
    row += 1

    day_text = format_date(inp.period_end, inp.country_code)
    heading = tuple(
        upper_for(language, label("header.work_to_date", language, overrides, date=day_text)) for language in languages
    )
    ws.cell(row=row, column=1, value=_join(heading)).font = Font(bold=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.cell(row=row, column=1).alignment = _WRAP_TOP
    row += 1

    header_row = row
    ws.cell(row=row, column=3, value=f"{_join(words('header.amount'))} ({inp.currency})")
    ws.cell(row=row, column=4, value=_join(words("basis.note")))
    for column in range(1, 5):
        cell = ws.cell(row=row, column=column)
        cell.font = Font(bold=True)
        cell.fill = _FILL
        cell.border = _BORDER
        cell.alignment = _CENTRE
    row += 1

    previous_section = ""
    for line in printed_summary(cert, locale, settings):
        if line.section == "deductions" and previous_section != "deductions":
            title = tuple(
                upper_for(language, label("section.deductions", language, overrides)) for language in languages
            )
            cell = ws.cell(row=row, column=1, value=_join(title))
            cell.font = Font(bold=True)
            cell.fill = _FILL
            cell.alignment = _WRAP_TOP
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            row += 1
        previous_section = line.section
        letter = f"{line.letter})" if line.section == "deductions" and line.letter else line.letter
        text = _join(line.labels)
        if line.formulas[0]:
            first, _, rest = text.partition("\n")
            text = f"{first} {line.formulas[0]}" + (f"\n{rest}" if rest else "")
        ws.cell(row=row, column=1, value=letter).alignment = _CENTRE
        ws.cell(row=row, column=2, value=text).alignment = _WRAP_TOP
        amount = ws.cell(row=row, column=3)
        refs = "".join(f" ({number})" for number in line.notes)
        if line.status == "value" and line.amount is not None:
            amount.value = line.amount
            amount.number_format = money
        else:
            amount.value = f"{line.text}{refs}" if line.status == "held" else line.text
            amount.alignment = Alignment(horizontal="right")
            if line.status == "held":
                amount.font = Font(bold=True, color=_ALERT)
        note_cell = ws.cell(row=row, column=4, value=_join(line.details) + (refs if line.status == "value" else ""))
        note_cell.alignment = _WRAP_TOP
        for column in range(1, 5):
            cell = ws.cell(row=row, column=column)
            cell.border = _BORDER
            if line.emphasis:
                cell.font = Font(bold=True, color=cell.font.color)
                cell.fill = _FILL
        row += 1

    notes = cert.notes
    if notes:
        row += 1
        title = " / ".join(
            upper_for(language, label("draft.notes_title", language, overrides)) for language in languages
        )
        ws.cell(row=row, column=1, value=title).font = Font(bold=True)
        row += 1
        for note in notes:
            lettered = f"{note.letter}) " if note.letter else ""
            parts = tuple(
                f"{lettered}{label(f'line.{note.line_key}', language, overrides)}: "
                f"{note_text(note, cert, language, settings)}"
                for language in languages
            )
            ws.cell(row=row, column=1, value=f"({note.number})").alignment = _WRAP_TOP
            ws.cell(row=row, column=2, value=_join(parts)).alignment = _WRAP_TOP
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            row += 1

    row += 1
    for role in cert.inp.signature_roles:
        title = tuple(upper_for(language, label(f"role.{role}", language, overrides)) for language in languages)
        ws.cell(row=row, column=2, value=_join(title)).font = Font(bold=True)
        ws.cell(row=row, column=2).alignment = _WRAP_TOP
        captions = " / ".join(
            " - ".join(label(key, language, overrides) for language in languages)
            for key in ("sign.name", "sign.signature", "sign.date")
        )
        ws.cell(row=row, column=3, value=captions).alignment = _WRAP_TOP
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 42
        row += 1

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    _page_setup(ws, landscape=False, last_column=4, last_row=row - 1, title_rows=f"{header_row}:{header_row}")


def _works_sheet(ws: Any, cert: Certificate, locale: str, settings: HakedisSettings | None) -> None:
    languages = locale_languages(locale)
    overrides = settings.labels if settings is not None else None
    inp = cert.inp
    columns, rows = printed_works(cert, locale, settings)
    title_column = next((i for i, column in enumerate(columns) if column.key == "description"), 0) + 1

    for position, column in enumerate(columns, start=1):
        width = _DESCRIPTION_WIDTH if column.key == "description" else _WIDTHS[column.kind]
        ws.column_dimensions[get_column_letter(position)].width = width
        letter = ws.cell(row=1, column=position, value=column.letter)
        names = tuple(label(f"col.{inp.flavour}.{column.key}", language, overrides) for language in languages)
        name = ws.cell(row=2, column=position, value=_join(names))
        for cell in (letter, name):
            cell.font = Font(bold=True)
            cell.fill = _FILL
            cell.border = _BORDER
            cell.alignment = _CENTRE
    ws.row_dimensions[2].height = 44 if len(languages) == 1 else 66

    row = 3
    for printed in rows:
        if printed.kind == "section":
            cell = ws.cell(row=row, column=1, value=printed.title)
            cell.font = Font(bold=True)
            cell.fill = _FILL
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(columns))
            row += 1
            continue
        bold = printed.kind != "line"
        for position, (column, text, value) in enumerate(
            zip(columns, printed.cells, printed.values, strict=True), start=1
        ):
            cell = ws.cell(row=row, column=position)
            if value is not None:
                cell.value = value
                cell.number_format = _number_format(column.kind, value, inp.currency)
            elif bold and position == title_column:
                cell.value = printed.title
            elif text:
                cell.value = text
            cell.border = _BORDER
            if column.kind in ("text",):
                cell.alignment = _WRAP_TOP
            elif column.kind == "unit":
                cell.alignment = Alignment(horizontal="center", vertical="top")
            if bold:
                cell.font = Font(bold=True)
                if printed.kind == "total":
                    cell.fill = _FILL
        row += 1

    ws.freeze_panes = "A3"
    _page_setup(ws, landscape=True, last_column=len(columns), last_row=row - 1, title_rows="1:2")


def render_hakedis_xlsx(
    cert: Certificate,
    *,
    locale: str = "tr",
    settings: HakedisSettings | None = None,
    issue_date: date | None = None,
) -> bytes:
    """Render a computed certificate to ``.xlsx`` bytes.

    Args:
        cert: The certificate from :func:`app.modules.contracts.hakedis.compute_certificate`.
        locale: ``tr``, ``en`` or ``tr-en``. The sheets are named in the first
            language of the locale.
        settings: The contract's resolved configuration, as for the PDF.
        issue_date: The date of issue, printed only when given.

    Raises:
        ValueError: ``locale`` is not one of the three.
    """
    languages = locale_languages(locale)
    overrides = settings.labels if settings is not None else None
    inp = cert.inp
    workbook = Workbook()
    summary = workbook.active
    summary.title = label("title.report", languages[0], overrides)
    works = workbook.create_sheet(label("title.works_list", languages[0], overrides))

    _summary_sheet(summary, cert, locale, settings, issue_date)
    _works_sheet(works, cert, locale, settings)

    number = " / ".join(label("header.certificate_number", language, overrides) for language in languages)
    details = [inp.project_name, f"{number}: {inp.certificate_number}", inp.contract_number]
    for sheet, key in ((summary, "title.report"), (works, "title.works_list")):
        # Text that only looks like a formula (a description opening with "=")
        # is typed as text before the letterhead helper walks the cells.
        store_strings_as_text(sheet)
        title = " / ".join(upper_for(language, label(key, language, overrides)) for language in languages)
        subtitle = (
            " / ".join(label(f"subtitle.{inp.flavour}", language, overrides) for language in languages)
            if sheet is works
            else None
        )
        apply_company_header(sheet, title=title, subtitle=subtitle, details=details)

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


__all__ = ["render_hakedis_xlsx"]
