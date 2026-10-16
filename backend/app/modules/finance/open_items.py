# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Open receivables and payables carried over from another system - pure parsing.

A company moving onto the platform brings the invoices that are still open on
the day it switches over: what customers still owe it and what it still owes
its suppliers. Each one becomes an invoice here, so ageing, statements of
account and payment runs see it from the first day. The history before the
cut-over stays in the old system: what was already paid on an invoice is one
opening payment on the cut-over date, not the payments one by one.

This module reads the file and checks each row on its own. It knows nothing
about the database; matching the counterparty and refusing numbers already in
use is the service's job.

A row carries the gross ``total``, the ``retention_held`` back from it (zero
when the column is missing) and the ``open`` amount still due in cash. Left
empty, ``open`` is the whole total less the retention. What was paid is
``total - retention_held - open``, which cannot be negative.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

_ZERO = Decimal("0")
_CENTS = Decimal("0.01")

#: Accepted header spellings per column, compared case-insensitively after
#: trimming. English first, then German, the two languages accounting exports
#: we have seen most often come in.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "direction": ("direction", "type", "ar/ap", "ap/ar", "ledger", "richtung", "art"),
    "invoice_number": (
        "invoice_number",
        "invoice number",
        "invoice no",
        "invoice no.",
        "invoice #",
        "number",
        "document number",
        "rechnungsnummer",
        "belegnummer",
    ),
    "counterparty": (
        "counterparty",
        "customer/supplier",
        "customer",
        "supplier",
        "vendor",
        "name",
        "partner",
        "geschäftspartner",
        "kunde/lieferant",
    ),
    "tax_id": (
        "tax_id",
        "tax id",
        "tax number",
        "business number",
        "bn",
        "vat number",
        "vat id",
        "steuernummer",
        "ust-idnr",
        "ust-id",
    ),
    "invoice_date": ("invoice_date", "invoice date", "date", "document date", "rechnungsdatum", "belegdatum"),
    "due_date": ("due_date", "due date", "due", "fälligkeit", "fällig am", "faelligkeit"),
    "currency": ("currency", "currency_code", "ccy", "währung", "waehrung"),
    "total": ("total", "amount", "gross", "amount_total", "invoice total", "betrag", "brutto", "rechnungsbetrag"),
    "tax": ("tax", "tax_amount", "vat", "hst", "gst", "steuer", "mwst", "ust"),
    "retention_held": (
        "retention_held",
        "retention",
        "holdback",
        "retainage",
        "einbehalt",
        "sicherheitseinbehalt",
    ),
    "open": ("open", "open amount", "outstanding", "balance", "balance due", "offen", "offener betrag", "saldo"),
    "notes": ("notes", "note", "memo", "description", "bemerkung", "text"),
}

_ALIAS_TO_COLUMN = {alias: column for column, aliases in COLUMN_ALIASES.items() for alias in aliases}

#: Columns a file has to name for any row to be read.
REQUIRED_COLUMNS = ("direction", "invoice_number", "counterparty", "invoice_date", "total")

_DIRECTIONS = {
    "receivable": "receivable",
    "ar": "receivable",
    "customer": "receivable",
    "sales": "receivable",
    "debitor": "receivable",
    "forderung": "receivable",
    "payable": "payable",
    "ap": "payable",
    "supplier": "payable",
    "vendor": "payable",
    "purchase": "payable",
    "kreditor": "payable",
    "verbindlichkeit": "payable",
}

#: Header line of the template a user can download and fill in.
TEMPLATE_HEADER = (
    "direction",
    "invoice_number",
    "counterparty",
    "tax_id",
    "invoice_date",
    "due_date",
    "currency",
    "total",
    "tax",
    "retention_held",
    "open",
    "notes",
)


def match_column(header: str) -> str | None:
    """The canonical column a header names, or ``None`` when it names none."""
    return _ALIAS_TO_COLUMN.get(str(header).strip().lower())


@dataclass(frozen=True)
class OpenItem:
    """One open invoice as read from the file, checked on its own."""

    row: int
    direction: str
    invoice_number: str
    counterparty: str
    tax_id: str
    invoice_date: str
    due_date: str | None
    currency: str
    total: Decimal
    tax: Decimal
    retention_held: Decimal
    open: Decimal
    notes: str

    @property
    def paid(self) -> Decimal:
        """What was already paid before the cut-over."""
        return self.total - self.retention_held - self.open


@dataclass(frozen=True)
class RowProblem:
    """Why a row cannot be imported, in the words a user can act on.

    ``code`` and ``params`` let the interface say it in the reader's language;
    ``message`` is the English wording, for the API and the logs.
    """

    row: int
    message: str
    invoice_number: str = ""
    code: str = ""
    params: dict[str, str] = field(default_factory=dict)


def parse_amount(value: Any) -> Decimal | None:
    """A money amount from a cell, or ``None`` when the cell is not a number.

    Takes ``1234.56``, ``1,234.56``, ``1.234,56``, ``1234,56`` and a leading
    minus or brackets. A lone separator followed by exactly three digits, such
    as ``1,234``, reads as either a thousand or a decimal depending on who
    wrote it, so it is refused rather than guessed.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, Decimal)):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    text = str(value).strip().replace(" ", "").replace(" ", "").replace("'", "")
    if not text:
        return None
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative, text = True, text[1:-1]
    if text.startswith("-"):
        negative, text = not negative, text[1:]
    text = re.sub(r"^[A-Za-z$€£¥]+|[A-Za-z$€£¥]+$", "", text)
    if not re.fullmatch(r"[0-9.,]+", text or "x"):
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif text.count(",") == 1 or text.count(".") == 1:
        sep = "," if "," in text else "."
        whole, frac = text.split(sep)
        if len(frac) == 3 and whole and whole != "0":
            return None
        text = f"{whole}.{frac}"
    elif text.count(",") > 1 or text.count(".") > 1:
        sep = "," if "," in text else "."
        groups = text.split(sep)
        if any(len(g) != 3 for g in groups[1:]):
            return None
        text = "".join(groups)
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    return -amount if negative else amount


_DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%Y/%m/%d", "%d/%m/%Y")


def parse_date(value: Any) -> str | None:
    """An ISO date from a cell, or ``None`` when the cell is not a date.

    Accepts ISO, ``DD.MM.YYYY``, ``YYYY/MM/DD`` and ``DD/MM/YYYY``, plus the
    date and datetime values a spreadsheet hands over. ``MM/DD/YYYY`` is not
    read: ``03/04/2026`` would be a different day for half the users.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    if " " in text or "T" in text:
        text = re.split(r"[ T]", text, maxsplit=1)[0]
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cents(amount: Decimal) -> Decimal:
    return amount.quantize(_CENTS)


def check_row(row_no: int, raw: dict[str, Any], *, cutover_date: str) -> OpenItem | RowProblem:
    """Read and check one row of the file.

    Args:
        row_no: the row's number in the file, header included, for messages.
        raw: the row's cells keyed by canonical column.
        cutover_date: the ISO day the company switches over; nothing in the
            file can be dated after it.

    Returns:
        The open item, or the first problem found in the row.
    """
    number = _text(raw.get("invoice_number"))

    def problem(code: str, message: str, **params: object) -> RowProblem:
        return RowProblem(
            row=row_no,
            message=message,
            invoice_number=number,
            code=code,
            params={k: str(v) for k, v in params.items()},
        )

    direction_text = _text(raw.get("direction")).lower()
    direction = _DIRECTIONS.get(direction_text)
    if direction is None:
        return problem(
            "direction_unknown",
            f"Direction '{direction_text}' is neither receivable (AR) nor payable (AP).",
            value=direction_text,
        )
    if not number:
        return problem("number_missing", "The invoice number is missing.")
    if len(number) > 50:
        return problem("number_too_long", "The invoice number is longer than 50 characters.")
    counterparty = _text(raw.get("counterparty"))
    tax_id = _text(raw.get("tax_id"))
    if not counterparty and not tax_id:
        return problem("counterparty_missing", "The counterparty is missing: give its name or its tax number.")

    invoice_date = parse_date(raw.get("invoice_date"))
    if invoice_date is None:
        value = _text(raw.get("invoice_date"))
        return problem("invoice_date_invalid", f"Invoice date '{value}' is not a date (use YYYY-MM-DD).", value=value)
    if invoice_date > cutover_date:
        return problem(
            "after_cutover",
            f"The invoice is dated {invoice_date}, after the cut-over date {cutover_date}.",
            date=invoice_date,
            cutover=cutover_date,
        )
    due_raw = _text(raw.get("due_date"))
    due_date = parse_date(raw.get("due_date")) if due_raw else None
    if due_raw and due_date is None:
        return problem("due_date_invalid", f"Due date '{due_raw}' is not a date (use YYYY-MM-DD).", value=due_raw)
    if due_date is not None and due_date < invoice_date:
        return problem("due_before_invoice", "The due date is before the invoice date.")

    currency = _text(raw.get("currency")).upper()
    if currency and not re.fullmatch(r"[A-Z]{3}", currency):
        return problem("currency_invalid", f"Currency '{currency}' is not a three-letter ISO code.", value=currency)

    amounts: dict[str, Decimal] = {}
    for column in ("total", "tax", "retention_held", "open"):
        cell = raw.get(column)
        if _text(cell) == "":
            continue
        amount = parse_amount(cell)
        if amount is None:
            return problem(
                "amount_invalid",
                f"The {column.replace('_', ' ')} '{_text(cell)}' is not an amount that reads one way only "
                "(write it as 1234.56).",
                column=column,
                value=_text(cell),
            )
        if amount < _ZERO:
            return problem(
                "amount_negative",
                f"The {column.replace('_', ' ')} is negative; a credit note is not an open item here.",
                column=column,
            )
        amounts[column] = _cents(amount)
    total = amounts.get("total")
    if total is None or total == _ZERO:
        return problem("total_missing", "The total is missing or zero.")
    tax = amounts.get("tax", _ZERO)
    retention = amounts.get("retention_held", _ZERO)
    if tax > total:
        return problem("tax_too_large", "The tax is larger than the total.")
    if retention > total:
        return problem("retention_too_large", "The retention held is larger than the total.")
    open_amount = amounts.get("open", total - retention)
    if open_amount + retention > total:
        return problem(
            "open_too_large",
            f"Open {open_amount} plus retention {retention} is more than the total {total}; "
            "the open amount is what is still due in cash, without the retention.",
            open=open_amount,
            retention=retention,
            total=total,
        )
    if open_amount == _ZERO and retention == _ZERO:
        return problem("nothing_open", "Nothing is open on this invoice, so it is not carried over.")

    return OpenItem(
        row=row_no,
        direction=direction,
        invoice_number=number,
        counterparty=counterparty,
        tax_id=tax_id,
        invoice_date=invoice_date,
        due_date=due_date,
        currency=currency,
        total=total,
        tax=tax,
        retention_held=retention,
        open=open_amount,
        notes=_text(raw.get("notes"))[:2000],
    )


def check_rows(rows: Iterable[tuple[int, dict[str, Any]]], *, cutover_date: str) -> list[OpenItem | RowProblem]:
    """Check every row, and refuse an invoice number the file repeats.

    The same number may appear once as a receivable and once as a payable:
    numbering is per direction, the way invoices are numbered here.
    """
    results: list[OpenItem | RowProblem] = []
    seen: dict[tuple[str, str], int] = {}
    for row_no, raw in rows:
        result = check_row(row_no, raw, cutover_date=cutover_date)
        if isinstance(result, OpenItem):
            key = (result.direction, result.invoice_number.casefold())
            first = seen.get(key)
            if first is not None:
                result = RowProblem(
                    row=row_no,
                    message=f"Invoice number {result.invoice_number} already appears in row {first}.",
                    invoice_number=result.invoice_number,
                    code="duplicate_in_file",
                    params={"number": result.invoice_number, "first_row": str(first)},
                )
            else:
                seen[key] = row_no
        results.append(result)
    return results


def normalise_tax_id(value: str | None) -> str:
    """A tax number reduced to letters and digits, upper case, for matching."""
    return re.sub(r"[^0-9A-Za-z]", "", value or "").upper()


def read_table(content: bytes, filename: str) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    """Read a CSV or Excel file into rows keyed by canonical column.

    Returns:
        The required columns the header does not name, and the non-empty rows
        as ``(row number in the file, cells)``.

    Raises:
        ValueError: when the file cannot be read or has no header row.
    """
    if filename.lower().endswith((".xlsx", ".xls")):
        header, body = _excel_rows(content)
    else:
        header, body = _csv_rows(content)
    if header is None:
        raise ValueError("The file has no header row naming its columns.")
    column_map: dict[int, str] = {}
    for idx, cell in enumerate(header):
        if cell is None:
            continue
        column = match_column(str(cell))
        if column and column not in column_map.values():
            column_map[idx] = column
    missing = [c for c in REQUIRED_COLUMNS if c not in column_map.values()]
    rows: list[tuple[int, dict[str, Any]]] = []
    for row_no, values in body:
        cells = {column_map[i]: v for i, v in enumerate(values) if i in column_map}
        if any(_text(v) for v in cells.values()):
            rows.append((row_no, cells))
    return missing, rows


def _csv_rows(content: bytes) -> tuple[list[str] | None, list[tuple[int, list[Any]]]]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes any byte string
        raise ValueError("The CSV file is in an encoding that cannot be read.")
    try:
        dialect: Any = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    header = next(reader, None)
    body = [(idx, row) for idx, row in enumerate(reader, start=2)]
    return header, body


def _excel_rows(content: bytes) -> tuple[list[Any] | None, list[tuple[int, list[Any]]]]:
    from openpyxl import load_workbook

    from app.core.sheet_header import find_header_row

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            raise ValueError("The Excel file has no worksheet.")
        found = find_header_row(ws.iter_rows(values_only=True), match_column)
        if found.values is None:
            return None, []
        body = [(idx, list(row)) for idx, row in enumerate(found.rows, start=found.number + 1)]
        return list(found.values), body
    finally:
        wb.close()


def template_csv() -> str:
    """The import template: the header and two example rows."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(TEMPLATE_HEADER)
    writer.writerow(
        ("receivable", "INV-1042", "Lakeview Homes Inc.", "", "2026-08-15", "2026-09-14", "CAD")
        + ("11300.00", "1300.00", "1000.00", "6300.00", "Progress claim 7")
    )
    writer.writerow(
        ("payable", "LS-5531", "Lakeshore Steel", "987654321RT0001", "2026-09-02", "2026-10-02", "CAD")
        + ("8475.00", "975.00", "0.00", "", "")
    )
    return buf.getvalue()


__all__ = [
    "COLUMN_ALIASES",
    "REQUIRED_COLUMNS",
    "TEMPLATE_HEADER",
    "OpenItem",
    "RowProblem",
    "check_row",
    "check_rows",
    "match_column",
    "normalise_tax_id",
    "parse_amount",
    "parse_date",
    "read_table",
    "template_csv",
]
