# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Quote​‌‍⁠​‌‍⁠​‌‍⁠​‌‍⁠ document readers: PDF, scanned PDF or image (OCR), Excel and email text.

Every reader returns :class:`DocLine` rows: the text of one printed line (or one
spreadsheet row) and the anchor that points back into the source, so the review
screen can highlight where a value came from:

* PDF: ``{"page": 1, "bbox": [x0, y0, x1, y1]}`` in page points, top-left origin,
  plus the page size in :attr:`QuoteDocument.pages`.
* Excel: ``{"sheet": "Quote", "row": 7, "cells": {"A": "A7", ...}}``.
* Text and email: ``{"line": 12}``.

The readers reuse the platform's own PDF reader (``app.core.pdf_reader``) and,
for a scan, Tesseract through the same optional dependency the file search uses.
Nothing here raises on a bad file: an unreadable document yields no lines.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from typing import Any

logger = logging.getLogger(__name__)

MAX_LINES = 4000
MAX_PAGES = 60
OCR_DPI = 150
_EMBEDDED_TEXT_MIN = 32


@dataclass
class DocLine:
    text: str
    anchor: dict[str, Any]
    cells: list[str] = field(default_factory=list)


@dataclass
class QuoteDocument:
    kind: str  # pdf | scan | xlsx | email | text
    engine: str  # pdfium | tesseract | openpyxl | text | none
    lines: list[DocLine]
    pages: list[dict[str, float]] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(line.text for line in self.lines)


def _group_words(words: list[tuple[float, float, float, float, str]], page_no: int) -> list[DocLine]:
    """Group word boxes into printed lines by vertical centre, left to right."""
    rows: list[list[tuple[float, float, float, float, str]]] = []
    for w in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
        cy = (w[1] + w[3]) / 2
        height = max(w[3] - w[1], 1.0)
        if rows:
            last = rows[-1]
            lcy = sum((x[1] + x[3]) / 2 for x in last) / len(last)
            if abs(cy - lcy) <= height * 0.5:
                last.append(w)
                continue
        rows.append([w])
    out: list[DocLine] = []
    for row in rows:
        row.sort(key=lambda w: w[0])
        parts: list[str] = []
        cells: list[str] = []
        prev_x1: float | None = None
        for x0, y0, x1, y1, word in row:
            gap = (x0 - prev_x1) if prev_x1 is not None else 0.0
            if prev_x1 is not None and gap > max((y1 - y0) * 1.2, 6.0):
                cells.append(" ".join(parts))
                parts = []
            parts.append(word)
            prev_x1 = x1
        cells.append(" ".join(parts))
        bbox = [
            round(min(w[0] for w in row), 1),
            round(min(w[1] for w in row), 1),
            round(max(w[2] for w in row), 1),
            round(max(w[3] for w in row), 1),
        ]
        out.append(DocLine(text="  ".join(cells), anchor={"page": page_no, "bbox": bbox}, cells=cells))
    return out


def _ocr_words(page: Any) -> list[tuple[float, float, float, float, str]]:
    """OCR one rendered page; boxes are returned in page points like the PDF words."""
    try:
        import pytesseract  # type: ignore[import-not-found]
    except Exception:  # noqa: BLE001 - optional dependency
        return []
    zoom = OCR_DPI / 72.0
    rendered = page.render(zoom=zoom)
    img = rendered.pil()
    data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
    words: list[tuple[float, float, float, float, str]] = []
    for i, text in enumerate(data.get("text", [])):
        if not text or not text.strip():
            continue
        x, y, w, h = (data[k][i] / zoom for k in ("left", "top", "width", "height"))
        words.append((x, y, x + w, y + h, text.strip()))
    return words


def read_pdf(payload: bytes) -> QuoteDocument:
    try:
        from app.core import pdf_reader
    except Exception:  # noqa: BLE001 - base dependency, defensive
        return QuoteDocument("pdf", "none", [])
    lines: list[DocLine] = []
    pages: list[dict[str, float]] = []
    engine = "pdfium"
    try:
        with pdf_reader.open_pdf(payload) as doc:
            count = min(doc.page_count, MAX_PAGES)
            per_page = []
            for index in range(count):
                page = doc[index]
                pages.append({"page": index + 1, "width": round(page.width, 1), "height": round(page.height, 1)})
                per_page.append(page.words())
            scanned = sum(len(w[4]) for words in per_page for w in words) < _EMBEDDED_TEXT_MIN
            for index in range(count):
                words = per_page[index]
                if scanned:
                    try:
                        words = _ocr_words(doc[index])
                    except Exception:  # noqa: BLE001 - one bad page must not stop the rest
                        logger.warning("quote OCR failed on page %d", index + 1, exc_info=True)
                        words = []
                lines.extend(_group_words(words, index + 1))
            if scanned:
                engine = "tesseract" if lines else "none"
    except Exception:  # noqa: BLE001 - unreadable PDF yields no lines
        logger.warning("quote PDF could not be read", exc_info=True)
        return QuoteDocument("pdf", "none", [])
    return QuoteDocument("scan" if engine == "tesseract" else "pdf", engine, lines[:MAX_LINES], pages)


def _col_letter(index: int) -> str:
    out = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def read_xlsx(payload: bytes) -> QuoteDocument:
    try:
        import openpyxl
    except Exception:  # noqa: BLE001 - base dependency, defensive
        return QuoteDocument("xlsx", "none", [])
    lines: list[DocLine] = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    values = ["" if v is None else str(v).strip() for v in row]
                    if not any(values):
                        continue
                    cells = {_col_letter(i): f"{_col_letter(i)}{r_idx}" for i, v in enumerate(values) if v}
                    text = "  ".join(v for v in values if v)
                    anchor = {"sheet": ws.title, "row": r_idx, "cells": cells}
                    lines.append(DocLine(text=text, anchor=anchor, cells=values))
                    if len(lines) >= MAX_LINES:
                        break
        finally:
            wb.close()
    except Exception:  # noqa: BLE001 - unreadable workbook yields no lines
        logger.warning("quote workbook could not be read", exc_info=True)
        return QuoteDocument("xlsx", "none", [])
    return QuoteDocument("xlsx", "openpyxl", lines)


def _email_body(raw: bytes) -> str:
    from email import policy
    from email.parser import BytesParser

    msg = BytesParser(policy=policy.default).parsebytes(raw)
    head = [f"{k}: {msg[k]}" for k in ("From", "Date", "Subject") if msg[k]]
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body is not None else ""
    if body is not None and body.get_content_type() == "text/html":
        text = re.sub(r"<[^>]+>", " ", text)
    return "\n".join([*head, "", text])


def read_text(text: str, kind: str = "text") -> QuoteDocument:
    lines: list[DocLine] = []
    for number, raw in enumerate(text.splitlines(), start=1):
        if raw.strip():
            cells = [c.strip() for c in re.split(r"\t|\s{2,}|;|\|", raw) if c.strip()]
            lines.append(DocLine(text=raw.strip(), anchor={"line": number}, cells=cells))
        if len(lines) >= MAX_LINES:
            break
    return QuoteDocument(kind, "text", lines)


def read_document(payload: bytes, filename: str = "") -> QuoteDocument:
    """Pick the reader by content signature first and file name second."""
    name = filename.lower()
    if payload[:5] == b"%PDF-":
        return read_pdf(payload)
    if payload[:4] == b"PK\x03\x04" or name.endswith((".xlsx", ".xlsm")):
        return read_xlsx(payload)
    if payload[:8] in (b"\x89PNG\r\n\x1a\n",) or payload[:3] == b"\xff\xd8\xff":
        return _read_image(payload)
    if name.endswith(".eml") or re.match(rb"^(Return-Path|Received|From|MIME-Version|Date):", payload[:200]):
        try:
            return read_text(_email_body(payload), kind="email")
        except Exception:  # noqa: BLE001 - fall back to plain text
            logger.debug("email parse failed, reading as text", exc_info=True)
    return read_text(payload.decode("utf-8", errors="replace"))


def _read_image(payload: bytes) -> QuoteDocument:
    try:
        import pytesseract  # type: ignore[import-not-found]
        from PIL import Image  # type: ignore[import-not-found]
    except Exception:  # noqa: BLE001 - optional dependency
        return QuoteDocument("scan", "none", [])
    try:
        with Image.open(io.BytesIO(payload)) as img:
            data = pytesseract.image_to_data(img.convert("RGB"), output_type=pytesseract.Output.DICT)
            size = {"page": 1, "width": float(img.width), "height": float(img.height)}
    except Exception:  # noqa: BLE001 - unreadable image yields no lines
        logger.warning("quote image OCR failed", exc_info=True)
        return QuoteDocument("scan", "none", [])
    words = [
        (data["left"][i], data["top"][i], data["left"][i] + data["width"][i], data["top"][i] + data["height"][i], t)
        for i, t in enumerate(data.get("text", []))
        if t and t.strip()
    ]
    lines = _group_words([(float(a), float(b), float(c), float(d), e) for a, b, c, d, e in words], 1)
    return QuoteDocument("scan", "tesseract" if lines else "none", lines, [size])
