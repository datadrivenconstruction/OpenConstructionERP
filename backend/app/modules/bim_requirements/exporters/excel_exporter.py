# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Excel exporter -- produces a formatted Excel workbook.

Generates a professional Excel file with:
- Sheet 1: Requirements table with formatted headers
- Sheet 2: Summary by element
- Sheet 3: Legend explaining columns and datatypes
"""

import io
import logging
from typing import Any

from app.core.document_catalogue import is_rtl, merge_missing
from app.core.document_locale import normalize_document_locale, translate
from app.modules.bim_requirements.parsers.base import UniversalRequirement

logger = logging.getLogger(__name__)

# Column widths, in the order of the ``h01`` ... ``h15`` header strings.
_COLUMN_WIDTHS = (25, 20, 22, 25, 15, 13, 30, 10, 10, 10, 20, 15, 18, 20, 18)
_COLUMN_COUNT = len(_COLUMN_WIDTHS)

# Workbook text per language. English and German live here; the platform's other
# languages come from ``core/document_catalogues`` (section ``bim_requirements_excel``).
# ``h01`` ... ``h15`` are the column headers, ``d01`` ... ``d14`` the legend lines
# (the Min and Max columns share one line).
_STRINGS: dict[str, dict[str, str]] = {
    "en": {
        "sheet_requirements": "Requirements",
        "sheet_summary": "Summary",
        "sheet_legend": "Legend",
        "summary_element": "Element",
        "summary_count": "Requirements Count",
        "summary_total": "Total",
        "legend_column": "Column",
        "legend_description": "Description",
        "unspecified": "Unspecified",
        "h01": "Element / IFC Class",
        "h02": "Classification",
        "h03": "Property Group",
        "h04": "Property Name",
        "h05": "Data Type",
        "h06": "Required",
        "h07": "Allowed Values",
        "h08": "Min",
        "h09": "Max",
        "h10": "Unit",
        "h11": "Pattern",
        "h12": "Phase",
        "h13": "Actor",
        "h14": "Use Case",
        "h15": "Source",
        "d01": "IFC entity class (e.g. IFCWALL, IFCBEAM)",
        "d02": "Classification system and code (e.g. Uniclass 2015: EF_25_10)",
        "d03": "Property set name (e.g. Pset_WallCommon)",
        "d04": "Name of the required property or attribute",
        "d05": "IFC data type (IFCLABEL, IFCREAL, IFCBOOLEAN, etc.)",
        "d06": "Cardinality: required, optional, or prohibited",
        "d07": "Semicolon-separated list of permitted values",
        "d08": "Numeric range boundaries",
        "d09": "Unit of measurement (m, m2, kg, etc.)",
        "d10": "Regular expression pattern for value validation",
        "d11": "Project phase or milestone (LP3, RIBA Stage 3, etc.)",
        "d12": "Responsible party (Architect, Structural Engineer, etc.)",
        "d13": "Purpose or use case for the requirement",
        "d14": "Source standard or document",
    },
    "de": {
        "sheet_requirements": "Anforderungen",
        "sheet_summary": "Zusammenfassung",
        "sheet_legend": "Legende",
        "summary_element": "Bauteil",
        "summary_count": "Anzahl Anforderungen",
        "summary_total": "Gesamt",
        "legend_column": "Spalte",
        "legend_description": "Beschreibung",
        "unspecified": "Nicht angegeben",
        "h01": "Bauteil / IFC-Klasse",
        "h02": "Klassifikation",
        "h03": "Merkmalsgruppe",
        "h04": "Merkmal",
        "h05": "Datentyp",
        "h06": "Erforderlich",
        "h07": "Zulaessige Werte",
        "h08": "Min-Wert",
        "h09": "Max-Wert",
        "h10": "Einheit",
        "h11": "Muster (Regex)",
        "h12": "Leistungsphase",
        "h13": "Akteur",
        "h14": "Anwendungsfall",
        "h15": "Quelle",
        "d01": "IFC-Entitätsklasse (z. B. IFCWALL, IFCBEAM)",
        "d02": "Klassifikationssystem und Code (z. B. Uniclass 2015: EF_25_10)",
        "d03": "Name des Merkmalsatzes (z. B. Pset_WallCommon)",
        "d04": "Name des geforderten Merkmals oder Attributs",
        "d05": "IFC-Datentyp (IFCLABEL, IFCREAL, IFCBOOLEAN usw.)",
        "d06": "Kardinalität: erforderlich, optional oder verboten",
        "d07": "Semikolon-getrennte Liste zulässiger Werte",
        "d08": "Grenzen des Zahlenbereichs",
        "d09": "Maßeinheit (m, m2, kg usw.)",
        "d10": "Regulärer Ausdruck zur Wertprüfung",
        "d11": "Projektphase oder Meilenstein (LP3, RIBA-Stufe 3 usw.)",
        "d12": "Verantwortliche Stelle (Architekt, Tragwerksplaner usw.)",
        "d13": "Zweck oder Anwendungsfall der Anforderung",
        "d14": "Quellnorm oder Quelldokument",
    },
}
merge_missing(_STRINGS, "bim_requirements_excel")

#: The languages the workbook can be written in.
EXPORT_LANGUAGES: tuple[str, ...] = tuple(_STRINGS)


def workbook_language(language: str | None) -> str:
    """The workbook language for a request: a catalogue language, English for anything else."""
    return normalize_document_locale(language, EXPORT_LANGUAGES, "en")


def _text(language: str, key: str) -> str:
    return translate(_STRINGS, language, key, "en")


# Color constants (RGB hex without #)
_HEADER_BG = "1F3864"
_HEADER_FG = "FFFFFF"
_ALT_ROW_BG = "EBF5FB"
_REQUIRED_BG = "C6EFCE"
_OPTIONAL_BG = "FFEB9C"
_PROHIBITED_BG = "FFC7CE"


def export_excel(
    requirements: list[UniversalRequirement],
    title: str = "BIM Requirements",
    language: str = "en",
) -> bytes:
    """Export requirements to a formatted Excel workbook.

    Args:
        requirements: List of universal requirements to export.
        title: Title shown on the Summary sheet.
        language: Any platform language code; the workbook text is written in it, English when the
            language has no catalogue.

    Returns:
        Excel file content as bytes.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    language = workbook_language(language)
    rtl = is_rtl(language)
    wb = openpyxl.Workbook()

    # ── Sheet 1: Requirements ──────────────────────────────────────────
    ws = wb.active
    ws.title = "Requirements"  # type: ignore[union-attr]
    ws.sheet_view.rightToLeft = rtl  # type: ignore[union-attr]

    header_fill = PatternFill(start_color=_HEADER_BG, end_color=_HEADER_BG, fill_type="solid")
    header_font = Font(bold=True, color=_HEADER_FG, size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alt_fill = PatternFill(start_color=_ALT_ROW_BG, end_color=_ALT_ROW_BG, fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    req_fill = PatternFill(start_color=_REQUIRED_BG, end_color=_REQUIRED_BG, fill_type="solid")
    opt_fill = PatternFill(start_color=_OPTIONAL_BG, end_color=_OPTIONAL_BG, fill_type="solid")
    prohib_fill = PatternFill(start_color=_PROHIBITED_BG, end_color=_PROHIBITED_BG, fill_type="solid")

    # Write headers
    for col_idx, width in enumerate(_COLUMN_WIDTHS, start=1):
        cell = ws.cell(row=1, column=col_idx)  # type: ignore[union-attr]
        cell.value = _text(language, f"h{col_idx:02d}")
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width  # type: ignore[union-attr]

    # Set row height for header
    ws.row_dimensions[1].height = 22  # type: ignore[union-attr]

    # Freeze top row and first column
    ws.freeze_panes = "B2"  # type: ignore[union-attr]

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(_COLUMN_COUNT)}1"  # type: ignore[union-attr]

    # Write data rows
    for row_idx, req in enumerate(requirements, start=2):
        ef = req.element_filter or {}
        cd = req.constraint_def or {}
        ctx = req.context or {}

        row_data = [
            ef.get("ifc_class", ""),
            _format_classification(ef.get("classification")),
            req.property_group or "",
            req.property_name,
            cd.get("datatype", ""),
            cd.get("cardinality", ""),
            "; ".join(cd["enum"]) if "enum" in cd else "",
            cd.get("min", ""),
            cd.get("max", ""),
            cd.get("unit", ""),
            cd.get("pattern", ""),
            ctx.get("phase", ""),
            ctx.get("actor", ""),
            ctx.get("use_case", ""),
            ctx.get("source", ""),
        ]

        is_alt = row_idx % 2 == 0
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)  # type: ignore[union-attr]
            cell.value = str(value) if value != "" else ""
            cell.border = thin_border
            if is_alt:
                cell.fill = alt_fill

        # Conditional formatting for cardinality column (F)
        card_cell = ws.cell(row=row_idx, column=6)  # type: ignore[union-attr]
        card_val = str(card_cell.value).lower() if card_cell.value else ""
        if card_val == "required":
            card_cell.fill = req_fill
        elif card_val == "optional":
            card_cell.fill = opt_fill
        elif card_val == "prohibited":
            card_cell.fill = prohib_fill

    # ── Sheet 2: Summary ───────────────────────────────────────────────
    ws_summary = wb.create_sheet("Summary")
    ws_summary.sheet_view.rightToLeft = rtl
    ws_summary.cell(row=1, column=1).value = _text(language, "sheet_summary")
    ws_summary.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws_summary.cell(row=2, column=1).value = title

    # Element summary
    element_counts: dict[str, int] = {}
    for req in requirements:
        ifc_class = (req.element_filter or {}).get("ifc_class", _text(language, "unspecified"))
        element_counts[ifc_class] = element_counts.get(ifc_class, 0) + 1

    ws_summary.cell(row=4, column=1).value = _text(language, "summary_element")
    ws_summary.cell(row=4, column=1).font = Font(bold=True)
    ws_summary.cell(row=4, column=2).value = _text(language, "summary_count")
    ws_summary.cell(row=4, column=2).font = Font(bold=True)

    for i, (element, count) in enumerate(sorted(element_counts.items()), start=5):
        ws_summary.cell(row=i, column=1).value = element
        ws_summary.cell(row=i, column=2).value = count

    ws_summary.cell(row=len(element_counts) + 6, column=1).value = _text(language, "summary_total")
    ws_summary.cell(row=len(element_counts) + 6, column=1).font = Font(bold=True)
    ws_summary.cell(row=len(element_counts) + 6, column=2).value = len(requirements)
    ws_summary.column_dimensions["A"].width = 30
    ws_summary.column_dimensions["B"].width = 20

    # ── Sheet 3: Legend ────────────────────────────────────────────────
    ws_legend = wb.create_sheet("Legend")
    ws_legend.sheet_view.rightToLeft = rtl
    ws_legend.cell(row=1, column=1).value = _text(language, "sheet_legend")
    ws_legend.cell(row=1, column=1).font = Font(bold=True, size=14)

    # The Min and Max columns share one legend line.
    legend_columns = [*(f"h{n:02d}" for n in range(1, 8)), None, *(f"h{n:02d}" for n in range(10, 16))]
    legend_rows = [(_text(language, "legend_column"), _text(language, "legend_description"))]
    for number, header_key in enumerate(legend_columns, start=1):
        name = (
            f"{_text(language, 'h08')} / {_text(language, 'h09')}"
            if header_key is None
            else _text(language, header_key)
        )
        legend_rows.append((name, _text(language, f"d{number:02d}")))

    for i, (col_name, desc) in enumerate(legend_rows, start=3):
        ws_legend.cell(row=i, column=1).value = col_name
        ws_legend.cell(row=i, column=2).value = desc
        if i == 3:
            ws_legend.cell(row=i, column=1).font = Font(bold=True)
            ws_legend.cell(row=i, column=2).font = Font(bold=True)

    ws_legend.column_dimensions["A"].width = 25
    ws_legend.column_dimensions["B"].width = 60

    # Stamp document properties so downloaded XLSX carries our authorship.
    try:
        wb.properties.creator = "OpenConstructionERP · DataDrivenConstruction"
        wb.properties.lastModifiedBy = "OpenConstructionERP"
        wb.properties.description = (
            "Generated by OpenConstructionERP (https://openconstructionerp.com) · DDC-CWICR-OE-2026"
        )
        wb.properties.keywords = "DDC-CWICR-OE,OpenConstructionERP,BIM Requirements"
    except Exception:  # noqa: BLE001 - best-effort metadata stamp
        pass

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _format_classification(classif: Any) -> str:
    """Format a classification dict as a display string."""
    if not classif:
        return ""
    if isinstance(classif, dict):
        system = classif.get("system", "")
        value = classif.get("value", "")
        if system and value:
            return f"{system}: {value}"
        return value or system
    return str(classif)


def generate_template() -> bytes:
    """Generate a downloadable Excel template with headers and example rows.

    Returns:
        Excel file content as bytes.
    """
    example_reqs = [
        UniversalRequirement(
            element_filter={"ifc_class": "IFCWALL"},
            property_group="Pset_WallCommon",
            property_name="FireRating",
            constraint_def={
                "datatype": "IFCLABEL",
                "cardinality": "required",
                "enum": ["REI60", "REI90", "REI120"],
            },
            context={"phase": "LP3", "actor": "Architect", "use_case": "Fire Safety"},
        ),
        UniversalRequirement(
            element_filter={"ifc_class": "IFCWALL"},
            property_group="Pset_WallCommon",
            property_name="ThermalTransmittance",
            constraint_def={
                "datatype": "IFCREAL",
                "cardinality": "required",
                "max": 0.28,
                "unit": "W/(m2K)",
            },
            context={"phase": "LP5", "actor": "Energy Consultant"},
        ),
        UniversalRequirement(
            element_filter={"ifc_class": "IFCDOOR"},
            property_group="Pset_DoorCommon",
            property_name="IsExternal",
            constraint_def={"datatype": "IFCBOOLEAN", "cardinality": "required"},
            context={"actor": "Architect"},
        ),
        UniversalRequirement(
            element_filter={"ifc_class": "IFCBEAM"},
            property_group="Pset_BeamCommon",
            property_name="LoadBearing",
            constraint_def={"datatype": "IFCBOOLEAN", "cardinality": "required"},
            context={"phase": "LP3", "actor": "Structural Engineer"},
        ),
        UniversalRequirement(
            element_filter={"ifc_class": "IFCSPACE"},
            property_group="Pset_SpaceCommon",
            property_name="Reference",
            constraint_def={
                "datatype": "IFCLABEL",
                "cardinality": "required",
                "pattern": "^[A-Z]{2}\\d{3}$",
            },
            context={"use_case": "Space Management"},
        ),
    ]

    return export_excel(example_reqs, title="BIM Requirements Template (Example Data)")
