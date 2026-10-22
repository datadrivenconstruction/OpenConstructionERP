# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Bulk tender recipients from a CSV or XLSX file.

Parsing and the per-row checks are pure (standard library plus openpyxl for
XLSX). The checks are also wrapped as a validation rule in the
``tender_recipient_import`` rule set, so the import is held to the same engine
as every other workflow.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass, field
from typing import Any

from app.core.validation.engine import (
    RuleCategory,
    RuleResult,
    Severity,
    ValidationContext,
    ValidationRule,
    rule_registry,
)

logger = logging.getLogger(__name__)

RECIPIENT_IMPORT_RULE_SET = "tender_recipient_import"
MAX_IMPORT_ROWS = 1000
MAX_IMPORT_BYTES = 2 * 1024 * 1024

_EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")

# Header aliases (lower-cased, spaces/underscores/dashes removed) -> canonical field.
_ALIASES: dict[str, str] = {
    "company": "company",
    "companyname": "company",
    "firm": "company",
    "contact": "contact_name",
    "contactname": "contact_name",
    "contactperson": "contact_name",
    "name": "contact_name",
    "email": "email",
    "emailaddress": "email",
    "mail": "email",
    "trade": "trade",
    "vat": "vat_id",
    "vatid": "vat_id",
    "vatnumber": "vat_id",
    "taxid": "vat_id",
}


class RecipientFileError(ValueError):
    """The file cannot be read as a recipient list; the message is a stable code."""


@dataclass
class ImportRow:
    """One parsed row and its findings (``issues`` holds stable codes)."""

    line: int
    company: str = ""
    contact_name: str = ""
    email: str = ""
    trade: str = ""
    vat_id: str = ""
    issues: list[str] = field(default_factory=list)

    @property
    def valid(self) -> bool:
        return not self.issues


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_recipient_file(filename: str, content: bytes) -> list[ImportRow]:
    """Read a CSV or XLSX upload into rows. Raises :class:`RecipientFileError`."""
    if len(content) > MAX_IMPORT_BYTES:
        raise RecipientFileError("file_too_large")
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        grid = _read_xlsx(content)
    elif name.endswith((".csv", ".txt")):
        grid = _read_csv(content)
    else:
        raise RecipientFileError("unsupported_type")
    grid = [r for r in grid if any(_clean(c) for c in r)]
    if not grid:
        raise RecipientFileError("empty_file")
    columns: dict[int, str] = {}
    for idx, head in enumerate(grid[0]):
        key = re.sub(r"[\s_\-]+", "", _clean(head).lower())
        if key in _ALIASES and _ALIASES[key] not in columns.values():
            columns[idx] = _ALIASES[key]
    if "email" not in columns.values() or "company" not in columns.values():
        raise RecipientFileError("missing_columns")
    body = grid[1:]
    if len(body) > MAX_IMPORT_ROWS:
        raise RecipientFileError("too_many_rows")
    rows: list[ImportRow] = []
    for offset, raw in enumerate(body):
        row = ImportRow(line=offset + 2)
        for idx, key in columns.items():
            setattr(row, key, _clean(raw[idx]) if idx < len(raw) else "")
        rows.append(row)
    return rows


def _read_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook  # noqa: PLC0415

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise RecipientFileError("unreadable_file") from exc
    try:
        return [[_clean(c) for c in r] for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


def check_rows(rows: list[ImportRow], existing_emails: set[str]) -> list[ImportRow]:
    """Attach issue codes to each row; returns the same rows.

    Codes: ``company_missing``, ``email_invalid``, ``duplicate_in_file`` (a
    later repeat of an address seen earlier in the file) and ``already_invited``.
    """
    existing = {e.strip().lower() for e in existing_emails}
    seen: set[str] = set()
    for row in rows:
        row.issues = []
        if not row.company:
            row.issues.append("company_missing")
        if not _EMAIL_RE.match(row.email):
            row.issues.append("email_invalid")
            continue
        email = row.email.lower()
        if email in seen:
            row.issues.append("duplicate_in_file")
        elif email in existing:
            row.issues.append("already_invited")
        seen.add(email)
    return rows


class TenderRecipientRowsValid(ValidationRule):
    """Every row of a recipient import needs a company, a valid and new email."""

    rule_id = "tendering.recipient_import_rows_valid"
    name = "Tender Recipient Import Rows"
    standard = "tendering"
    severity = Severity.ERROR
    category = RuleCategory.COMPLETENESS
    description = "Flags import rows with a missing company, a bad or repeated email, or a firm already invited."

    async def validate(self, context: ValidationContext) -> list[RuleResult]:
        data = context.data if isinstance(context.data, dict) else {}
        rows = [
            ImportRow(line=int(r.get("line", 0)), company=_clean(r.get("company")), email=_clean(r.get("email")))
            for r in data.get("rows", [])
            if isinstance(r, dict)
        ]
        check_rows(rows, set(data.get("existing_emails", [])))
        bad = [r for r in rows if r.issues]

        def result(passed: bool, message: str, row: ImportRow | None = None) -> RuleResult:
            return RuleResult(
                rule_id=self.rule_id,
                rule_name=self.name,
                severity=self.severity,
                category=self.category,
                passed=passed,
                message=message,
                element_ref=None if row is None else f"line:{row.line}",
                suggestion=None if row is None else "Fix or remove the row, then upload the file again.",
                details={} if row is None else {"line": row.line, "issues": list(row.issues)},
            )

        if not bad:
            return [result(True, "Every row can be added.")]
        return [result(False, f"Row {r.line}: {', '.join(r.issues)}.", r) for r in bad]


def register_recipient_import_rules() -> None:
    """Register the import rule set; idempotent."""
    rule_registry.register(TenderRecipientRowsValid(), [RECIPIENT_IMPORT_RULE_SET])
