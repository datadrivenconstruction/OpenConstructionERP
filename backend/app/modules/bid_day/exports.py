# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Export a submitted bid snapshot to Excel and PDF.

Both read only the snapshot's payload, so an export made a year later shows
exactly what was submitted, whatever happened to the board since.
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from app.modules.bid_day.logic import to_decimal

_TOTAL_LABELS = (
    ("direct", "Direct cost"),
    ("overhead", "Overhead"),
    ("risk", "Risk"),
    ("profit", "Profit"),
    ("bid_total", "Bid total"),
)


def _rows(payload: dict[str, Any]) -> list[list[Any]]:
    out = []
    for r in payload.get("rows") or []:
        ch = r.get("chosen") or {}
        out.append(
            [
                r.get("name") or "",
                r.get("kind") or "",
                ch.get("vendor") or "(estimate)",
                to_decimal(ch.get("as_quoted")) if ch else None,
                to_decimal(ch.get("adjusted")) if ch else None,
                to_decimal(r.get("plug_amount")),
                to_decimal(r.get("value")),
            ]
        )
    return out


HEADERS = ["Package", "Kind", "Chosen quote", "As quoted", "Adjusted", "Estimate", "In bid"]


def snapshot_xlsx(payload: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from app.core.xlsx_text import store_strings_as_text

    wb = Workbook()
    ws = wb.active
    ws.title = "Submitted bid"
    tender = payload.get("tender") or {}
    ws.append([f"Tender: {tender.get('name', '')}"])
    ws.append([f"Submitted {payload.get('submitted_at', '')} by {payload.get('submitted_by', '')}"])
    ws.append([f"Currency: {payload.get('currency', '')}"])
    ws.append([])
    ws.append(HEADERS)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in _rows(payload):
        ws.append(row)
    ws.append([])
    totals = payload.get("totals") or {}
    marks = payload.get("markups") or {}
    for key, label in _TOTAL_LABELS:
        pct = marks.get(f"{key}_pct")
        ws.append(
            [
                f"{label} ({pct}%)" if pct is not None else label,
                None,
                None,
                None,
                None,
                None,
                to_decimal(totals.get(key)),
            ]
        )
        if key == "bid_total":
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            ws.cell(row=ws.max_row, column=7).font = Font(bold=True)
    for col in "DEFG":
        for cell in ws[col]:
            if isinstance(cell.value, Decimal):
                cell.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["C"].width = 30
    store_strings_as_text(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt(v: Decimal | None) -> str:
    return "" if v is None else f"{v:,.2f}"


def snapshot_pdf(payload: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from app.core.pdf_fonts import pdf_font

    body, bold = pdf_font("Helvetica"), pdf_font("Helvetica-Bold")
    styles = getSampleStyleSheet()
    for s in styles.byName.values():
        s.fontName = body
    tender = payload.get("tender") or {}
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=f"Submitted bid {tender.get('name', '')}")
    story: list[Any] = [
        Paragraph(f"Submitted bid: {tender.get('name', '')}", styles["Title"]),
        Paragraph(
            f"Submitted {payload.get('submitted_at', '')} by {payload.get('submitted_by', '')}. "
            f"Currency {payload.get('currency', '')}.",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]
    data: list[list[str]] = [HEADERS]
    for name, kind, vendor, quoted, adj, plug, value in _rows(payload):
        data.append([name, kind, vendor, _fmt(quoted), _fmt(adj), _fmt(plug), _fmt(value)])
    totals = payload.get("totals") or {}
    marks = payload.get("markups") or {}
    for key, label in _TOTAL_LABELS:
        pct = marks.get(f"{key}_pct")
        data.append(
            [f"{label} ({pct}%)" if pct is not None else label, "", "", "", "", "", _fmt(to_decimal(totals.get(key)))]
        )
    table = Table(data, repeatRows=1, colWidths=[190, 70, 150, 80, 80, 80, 90])
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), body),
                ("FONTNAME", (0, 0), (-1, 0), bold),
                ("FONTNAME", (0, -1), (-1, -1), bold),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
                ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.grey),
                ("LINEABOVE", (0, -len(_TOTAL_LABELS)), (-1, -len(_TOTAL_LABELS)), 0.6, colors.grey),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buf.getvalue()
