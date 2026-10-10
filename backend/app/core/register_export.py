# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Printable registers and record forms shared by the project-control modules.

A register (the RFI log, the submittal log, the correspondence log, the
variation log) is a contractual record: it is printed, attached to a letter
and filed. Every module that keeps one needs the same sheet: what the
register is, which project it belongs to, when it was produced, a table that
repeats its header row, and "page x of y" so a missing sheet is noticed.
A single record (one submittal, one letter, one change order) needs the same
form the RFI already prints: title, number, status, a label and value grid,
boxed text, an item table and signature lines.

A module's own words are not here. Each module keeps its catalogue next to
its renderer, the way :mod:`app.modules.rfi.pdf_translations` does, and hands
this module finished labels. The only strings held here are the page
furniture every one of these sheets carries ("Project", "Page 2 of 5",
"Generated ..."), in :data:`FURNITURE`, so six catalogues cannot drift into
six spellings of the same footer. :class:`DocumentCatalogue` joins the two
and refuses, at import, a language that is declared without a complete
table: a locale listed as supported with no dictionary behind it is how a
reader ends up holding an English document after asking for their own.

What is shared beyond that is the layout and the three rules that are easy
to get wrong one module at a time:

* **Capitals follow the language.** ``str.upper`` turns a Turkish "i" into
  "I", so "Geçersiz" would print as "GEÇERSIZ". :func:`caps` is the one place
  that knows.
* **Amounts are written the way the project's market writes them**, by
  :mod:`app.core.regional_format`, from the ``Decimal`` and never through a
  float. In a workbook an amount stays a number with a number format, so the
  column still sorts and sums.
* **Every cell is a Paragraph**, so text wraps instead of running over the
  next column, user text is escaped instead of parsed as markup, and each
  string gets a face that can draw it (:mod:`app.core.pdf_fonts`).

The page count is found by building the document twice: the first pass
counts, the second prints the total. A register is a few hundred rows, so
the second pass costs less than a wrong "page 1" on a three page log.
"""

from __future__ import annotations

import html
import io
import unicodedata
import uuid
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    CondPageBreak,
    Frame,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from app.core.document_locale import normalize_document_locale, resolve_document_locale, translate
from app.core.pdf_branding import (
    branded_appearance,
    branded_cover_brand,
    branded_doc_metadata,
    branded_header_logo,
    branded_letterhead,
)
from app.core.pdf_fonts import (
    BODY_FONT,
    BOLD_FONT,
    pdf_fit_line,
    pdf_fitted_style,
    pdf_room_beside,
    pdf_style_for_text,
    pdf_table_paragraph_rows,
    register_pdf_fonts,
)
from app.core.regional_format import NumberStyle, format_number, number_style

__all__ = [
    "EMPTY",
    "FURNITURE",
    "DocumentCatalogue",
    "ProjectHeader",
    "RecordBlock",
    "RecordDocument",
    "RegisterColumn",
    "RegisterDocument",
    "build_record_pdf",
    "build_register_pdf",
    "build_register_xlsx",
    "caps",
    "export_filename",
    "format_amount",
    "format_count",
    "format_stored_date",
    "parse_stored_date",
    "person_name",
    "record_text_lines",
    "register_text_rows",
    "xlsx_date_format",
]

register_pdf_fonts()

#: What an absent value prints as, on every register and form.
EMPTY = "-"

_INK = colors.HexColor("#16213e")
_MUTED = colors.HexColor("#666666")
_RULE = colors.HexColor("#cccccc")
_LABEL_FILL = colors.HexColor("#f4f5f7")
_ZEBRA = colors.HexColor("#f8f9fb")

_MARGIN_SIDE = 15 * mm
_MARGIN_TOP = 18 * mm
_MARGIN_BOTTOM = 18 * mm
_FORM_MARGIN_SIDE = 20 * mm
_BLANK_BOX_HEIGHT = 30 * mm
_SIGNATURE_ROW_HEIGHT = 10 * mm
_LONG_SECTION_MIN_ROOM = 30 * mm


# ── Project header ───────────────────────────────────────────────────────


@dataclass(frozen=True)
class ProjectHeader:
    """The project facts a register or a record form is headed with.

    Attributes:
        name: The project's name; empty when the project is gone.
        code: The project number, when it has one.
        currency: ISO 4217 code, upper-cased; empty when none is set.
        country: ISO 3166-1 alpha-2 of the project's country, when recorded.
    """

    name: str = ""
    code: str | None = None
    currency: str = ""
    country: str | None = None

    @property
    def label(self) -> str:
        """``"Name (CODE)"``, the name alone without a code, or a dash."""
        name = self.name or EMPTY
        return f"{name} ({self.code})" if self.code else name

    @property
    def number_style(self) -> NumberStyle:
        """How this project's market writes a number: its country decides, then its currency."""
        return number_style(self.country, self.currency)


def person_name(value: Any, people: Mapping[str, str] | None) -> str:
    """The name to print for a stored person or company reference.

    The column holds one of three things: an id the lookup resolved, an id it
    could not (a deleted account), or a name somebody typed. A resolved id
    prints the name, an unresolved one its first eight characters rather than
    36 characters of UUID on a document meant for people, and typed text
    prints as typed.
    """
    if value is None or str(value).strip() == "":
        return EMPTY
    key = str(value).strip()
    names = people or {}
    if key in names:
        return names[key]
    try:
        canonical = str(uuid.UUID(key))
    except ValueError:
        return key
    return names.get(canonical) or canonical[:8]


# ── Catalogue ────────────────────────────────────────────────────────────

#: Page furniture shared by every register and record form. English is the
#: source; a language is added here first and to the module catalogues after.
#: ``days_one`` and ``days_other`` exist because English counts "1 day" and
#: "3 days" while Turkish keeps the noun singular after a numeral.
FURNITURE: dict[str, dict[str, str]] = {
    "en": {
        "project": "Project",
        "project_no": "Project no.",
        "generated_on": "Date generated",
        "records": "Records",
        "currency": "Currency",
        "status": "Status",
        "name": "Name",
        "signature": "Signature",
        "date": "Date",
        "signatures": "Signatures",
        "yes": "Yes",
        "no": "No",
        "days_one": "{n} day",
        "days_other": "{n} days",
        "empty_register": "No records.",
        "footer_generated": "Generated {timestamp}",
        "footer_page": "Page {page} of {total}",
        "date_format": "%Y-%m-%d",
        "datetime_format": "%Y-%m-%d %H:%M UTC",
    },
    "tr": {
        "project": "Proje",
        "project_no": "Proje no.",
        "generated_on": "Oluşturulma tarihi",
        "records": "Kayıt sayısı",
        "currency": "Para birimi",
        "status": "Durum",
        "name": "Ad soyad",
        "signature": "İmza",
        "date": "Tarih",
        "signatures": "İmzalar",
        "yes": "Evet",
        "no": "Hayır",
        "days_one": "{n} gün",
        "days_other": "{n} gün",
        "empty_register": "Kayıt yok.",
        "footer_generated": "Oluşturulma: {timestamp}",
        "footer_page": "Sayfa {page} / {total}",
        "date_format": "%d.%m.%Y",
        "datetime_format": "%d.%m.%Y %H:%M UTC",
    },
}


class DocumentCatalogue:
    """One module's printed strings, joined with the shared page furniture.

    Args:
        strings: The module's own strings, keyed by language then by key.
        labels: Named code tables (statuses, types), each keyed by language
            then by stored code.
        default: The fallback language; it must be in every table.

    Raises:
        ValueError: At construction, when a language is declared in one table
            and missing from another, when a table lacks a key the default
            language has, or when the page furniture has no such language.
            A catalogue that imports is complete.
    """

    def __init__(
        self,
        strings: dict[str, dict[str, str]],
        labels: dict[str, dict[str, dict[str, str]]] | None = None,
        *,
        default: str = "en",
    ) -> None:
        self.default = default
        self.labels = labels or {}
        self.supported: tuple[str, ...] = tuple(strings)
        problems: list[str] = []
        for locale in self.supported:
            if locale not in FURNITURE:
                problems.append(f"{locale}: no page furniture")
            missing = set(strings[default]) - set(strings[locale])
            if missing:
                problems.append(f"{locale}: strings missing {sorted(missing)}")
        for name, table in self.labels.items():
            if set(table) != set(self.supported):
                problems.append(f"label table {name!r} covers {sorted(table)}, not {sorted(self.supported)}")
                continue
            for locale in self.supported:
                missing = set(table[default]) - set(table[locale])
                if missing:
                    problems.append(f"{locale}: label table {name!r} missing {sorted(missing)}")
        if problems:
            raise ValueError("incomplete document catalogue: " + "; ".join(problems))
        self.strings = {locale: {**FURNITURE[locale], **strings[locale]} for locale in self.supported}

    def normalize(self, value: str | None) -> str:
        """Reduce a locale-ish value to a language this catalogue has."""
        return normalize_document_locale(value, self.supported, self.default)

    def resolve(self, locale_param: str | None, accept_language: str | None) -> str:
        """Pick the language for a request: ``?locale=``, then ``Accept-Language``, then the default."""
        return resolve_document_locale(locale_param, accept_language, self.supported, self.default)

    def tr(self, locale: str | None, key: str, **params: Any) -> str:
        """Resolve ``key``, falling back to the default language, then the key."""
        return translate(self.strings, self.normalize(locale), key, self.default, **params)

    def label(self, table: str, code: Any, locale: str | None) -> str:
        """The word for a stored code; an unknown code prints as stored, an absent one as a dash."""
        text = "" if code is None else str(code).strip()
        if not text:
            return EMPTY
        by_locale = self.labels[table]
        key = text.lower()
        return by_locale[self.normalize(locale)].get(key) or by_locale[self.default].get(key) or text

    def days(self, count: Any, locale: str | None) -> str:
        """A number of days with the noun in the form the language needs, or a dash."""
        if count is None or isinstance(count, bool) or str(count).strip() == "":
            return EMPTY
        try:
            number = int(count)
        except (TypeError, ValueError):
            return str(count)
        return self.tr(locale, "days_one" if abs(number) == 1 else "days_other", n=number)

    def yes_no(self, flag: Any, locale: str | None) -> str:
        """ "Yes" or "No" in the document language."""
        return self.tr(locale, "yes" if flag else "no")

    def register_details(
        self, locale: str | None, project: ProjectHeader, *, count: int, with_currency: bool = False
    ) -> list[tuple[str, str]]:
        """The identifying block above a register: project, number, date, record count."""
        today = datetime.now(tz=UTC).date().strftime(self.tr(locale, "date_format"))
        details = [
            (self.tr(locale, "project"), project.name or EMPTY),
            (self.tr(locale, "project_no"), project.code or EMPTY),
            (self.tr(locale, "generated_on"), today),
            (self.tr(locale, "records"), str(count)),
        ]
        if with_currency:
            details.append((self.tr(locale, "currency"), project.currency or EMPTY))
        return details

    def furniture(self, locale: str | None) -> dict[str, str]:
        """The footer templates and date formats a document builder passes on."""
        return {
            "page_label": self.tr(locale, "footer_page"),
            "generated_label": self.tr(locale, "footer_generated"),
            "date_format": self.tr(locale, "date_format"),
        }


# ── Value formatting ─────────────────────────────────────────────────────


def caps(label: str, locale: str | None) -> str:
    """Upper-case a label the way its language writes capitals.

    Turkish keeps the dot on a capital i ("Geçersiz" is "GEÇERSİZ") and its
    dotless ı becomes a plain I, which ``str.upper`` already does. Greek drops
    the stress accent in capitals. Everything else is ``str.upper``.

    Args:
        label: The localised word or phrase.
        locale: The document language; a regional tag reads its language.

    Returns:
        The label in capitals.
    """
    language = (locale or "").strip().lower().replace("_", "-").split("-", 1)[0]
    if language in ("tr", "az"):
        return label.replace("i", "İ").upper()
    upper = label.upper()
    if language == "el":
        stripped = "".join(ch for ch in unicodedata.normalize("NFD", upper) if ch != "́")
        return unicodedata.normalize("NFC", stripped)
    return upper


def parse_stored_date(value: Any) -> date | None:
    """Read a date out of the shapes the modules store one in.

    A ``date``, a ``datetime``, ``YYYY-MM-DD`` and a full ISO timestamp are
    all in use, sometimes in the same column. Anything else is ``None``.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        try:
            return date.fromisoformat(text[:10])
        except ValueError:
            return None


def format_stored_date(value: Any, date_format: str) -> str:
    """A stored date in ``date_format``, as stored when unparseable, or a dash.

    Args:
        value: The stored value, in any shape :func:`parse_stored_date` reads.
        date_format: A ``strftime`` pattern from the module's catalogue.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return EMPTY
    parsed = parse_stored_date(value)
    if parsed is None:
        return str(value).strip()
    return parsed.strftime(date_format)


def _as_decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value if value.is_finite() else None
    text = str(value).strip()
    if not text:
        return None
    try:
        amount = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    return amount if amount.is_finite() else None


def format_amount(value: Any, style: NumberStyle, currency: str | None = None, *, decimals: int = 2) -> str:
    """A money amount in the market's separators, with its currency code.

    The ISO code is printed rather than a symbol: a register travels between
    companies and countries, and "TRY" cannot be misread the way a symbol can.

    Args:
        value: A ``Decimal``, or anything ``Decimal`` reads. Never a float.
        style: The separators, from :func:`app.core.regional_format.number_style`.
        currency: ISO 4217 code appended after a no-break space, when given.
        decimals: Fraction digits.

    Returns:
        ``"1.234,56 TRY"``, the bare number without a currency, or a dash when
        there is no amount.
    """
    amount = _as_decimal(value)
    if amount is None:
        return EMPTY
    text = format_number(amount, decimals, style)
    code = (currency or "").strip().upper()
    return f"{text} {code}" if code else text


def format_count(value: Any) -> str:
    """A whole number (days, revision, item count), or a dash when absent."""
    if value is None or isinstance(value, bool) or str(value).strip() == "":
        return EMPTY
    try:
        return str(int(value))
    except (TypeError, ValueError):
        return str(value)


def xlsx_date_format(date_format: str) -> str:
    """The spreadsheet number format for a ``strftime`` date pattern."""
    return date_format.replace("%d", "DD").replace("%m", "MM").replace("%Y", "YYYY")


def export_filename(stem: str | None, fallback: str, extension: str) -> str:
    """A download filename every operating system accepts.

    Everything outside letters, digits, dot, dash and underscore becomes an
    underscore, so a number typed with a slash or a space still gives a name.
    """
    cleaned = "".join(ch if ch.isalnum() or ch in "-_." else "_" for ch in (stem or "").strip())
    return f"{cleaned or fallback}.{extension}"


# ── Register (a log of many records) ─────────────────────────────────────


@dataclass(frozen=True)
class RegisterColumn:
    """One column of a register.

    Attributes:
        header: The translated column heading.
        kind: How the cell is written: ``"text"`` as it is, ``"date"`` from
            any stored date shape, ``"money"`` from a ``Decimal`` with two
            decimals, ``"count"`` as a whole number.
        weight: Relative width on the printed page.
        xlsx_width: Column width in the workbook, in characters.
        wrap: Wrap the workbook cell (long free text).
    """

    header: str
    kind: str = "text"
    weight: float = 1.0
    xlsx_width: float = 16.0
    wrap: bool = False


@dataclass(frozen=True)
class RegisterDocument:
    """Everything a register prints, already translated.

    Attributes:
        title: The register's name, e.g. "Submittal Register".
        columns: The columns, left to right.
        rows: One list of raw values per record, in column order.
        details: Label and value pairs printed above the table: project,
            project number, generated date and whatever else identifies the
            sheet. Printed in the order given.
        locale: The document language, for capitals and metadata.
        date_format: ``strftime`` pattern for date cells.
        number_style: Separators for money cells.
        page_label: Footer template with ``{page}`` and ``{total}``.
        generated_label: Footer template with ``{timestamp}``.
        generated: The timestamp printed in the footer.
        empty_text: Printed instead of the table when there are no records.
        sheet_title: Workbook sheet name; cut to the 31 characters a sheet
            name may have.
        doc_type: Key in :data:`app.core.pdf_appearance.DOCUMENT_TYPES`, when
            the workspace styles this document type separately.
    """

    title: str
    columns: Sequence[RegisterColumn]
    rows: Sequence[Sequence[Any]]
    details: Sequence[tuple[str, str]] = ()
    locale: str = "en"
    date_format: str = "%Y-%m-%d"
    number_style: NumberStyle = field(default_factory=number_style)
    page_label: str = "Page {page} of {total}"
    generated_label: str = "Generated {timestamp}"
    generated: str = ""
    empty_text: str = ""
    sheet_title: str = ""
    doc_type: str | None = None


def _cell_text(value: Any, kind: str, document: RegisterDocument) -> str:
    if kind == "date":
        return format_stored_date(value, document.date_format)
    if kind == "money":
        return format_amount(value, document.number_style)
    if kind == "count":
        return format_count(value)
    text = "" if value is None else str(value).strip()
    return text or EMPTY


def register_text_rows(document: RegisterDocument) -> list[list[str]]:
    """The register's cells as the strings the PDF prints, header row first.

    This is the layer just before rendering: what a test reads to check a
    register without extracting text from a PDF, and what the renderer wraps.
    """
    rows = [[column.header for column in document.columns]]
    for row in document.rows:
        rows.append(
            [_cell_text(value, column.kind, document) for value, column in zip(row, document.columns, strict=True)]
        )
    return rows


def _footer_callback(
    *,
    page_label: str,
    generated_label: str,
    generated: str,
    total_pages: int,
    page_width: float,
    margin_side: float,
    doc_type: str | None,
    letterhead_on_first_page: bool,
) -> Callable[[Any, Any], None]:
    """``onPage`` callback: translated footer, workspace logo top right.

    Follows the workspace's document appearance the way the RFI form does: a
    saved footer line replaces the brand and timestamp, the footer colour
    colours it, and page numbers can be switched off.
    """
    look = branded_appearance(doc_type=doc_type) or {}
    style = ParagraphStyle(
        "RegisterFooter",
        fontName=BODY_FONT,
        fontSize=7,
        leading=8,
        textColor=colors.HexColor(look.get("footer_color") or "#999999"),
    )
    right_style = ParagraphStyle("RegisterFooterRight", parent=style, alignment=TA_RIGHT)
    custom_footer = str(look.get("footer_text") or "").strip()
    show_page_numbers = look.get("show_page_numbers", True) is not False
    usable = page_width - 2 * margin_side

    def _draw(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setStrokeColor(_RULE)
        canvas.setLineWidth(0.5)
        canvas.line(margin_side, 13 * mm, page_width - margin_side, 13 * mm)
        page_text = page_label.format(page=doc.page, total=max(total_pages, doc.page)) if show_page_numbers else ""
        # One line, fitted into the room beside the page number: the paragraph
        # is anchored by its top, so a wrapped line would fall off the sheet.
        left_text, _face, left_size = pdf_fit_line(
            custom_footer or branded_cover_brand(),
            pdf_room_beside(usable, page_text),
            suffix="" if custom_footer else f"  |  {generated_label.format(timestamp=generated)}",
            base=BODY_FONT,
        )
        left = Paragraph(html.escape(left_text, quote=True), pdf_fitted_style(style, left_text, left_size))
        _, left_h = left.wrapOn(canvas, usable, 20)
        left.drawOn(canvas, margin_side, 9 * mm - left_h + 2)
        if page_text:
            right = Paragraph(html.escape(page_text, quote=True), pdf_style_for_text(right_style, page_text))
            _, right_h = right.wrapOn(canvas, usable, 20)
            right.drawOn(canvas, margin_side, 9 * mm - right_h + 2)
        canvas.restoreState()
        if not (letterhead_on_first_page and doc.page == 1):
            branded_header_logo(canvas, doc)

    return _draw


def _build_twice(
    make_story: Callable[[], tuple[list[Any], bool]],
    *,
    pagesize: tuple[float, float],
    margin_side: float,
    title: str,
    subject: str,
    page_label: str,
    generated_label: str,
    generated: str,
    doc_type: str | None,
) -> bytes:
    """Build the document, then build it again knowing how many pages it has."""
    page_width, page_height = pagesize
    meta = branded_doc_metadata()
    total_pages = 0
    output = b""
    for _pass in range(2):
        story, has_letterhead = make_story()
        buffer = io.BytesIO()
        doc = BaseDocTemplate(
            buffer,
            pagesize=pagesize,
            leftMargin=margin_side,
            rightMargin=margin_side,
            topMargin=_MARGIN_TOP,
            bottomMargin=_MARGIN_BOTTOM,
            title=title,
            author=meta["author"],
            subject=subject,
            creator=meta["creator"],
            producer=meta["producer"],
            keywords=meta["keywords"],
        )
        frame = Frame(
            margin_side,
            _MARGIN_BOTTOM,
            page_width - 2 * margin_side,
            page_height - _MARGIN_TOP - _MARGIN_BOTTOM,
            leftPadding=0,
            rightPadding=0,
            topPadding=0,
            bottomPadding=0,
            id="body",
        )
        on_page = _footer_callback(
            page_label=page_label,
            generated_label=generated_label,
            generated=generated,
            total_pages=total_pages,
            page_width=page_width,
            margin_side=margin_side,
            doc_type=doc_type,
            letterhead_on_first_page=has_letterhead,
        )
        doc.addPageTemplates([PageTemplate(id="body", frames=[frame], onPage=on_page)])
        doc.build(story)
        total_pages = doc.page
        output = buffer.getvalue()
    return output


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()["Normal"]
    return {
        "title": ParagraphStyle("RegTitle", parent=base, fontName=BOLD_FONT, fontSize=16, leading=20, textColor=_INK),
        "sub": ParagraphStyle("RegSub", parent=base, fontName=BODY_FONT, fontSize=10, leading=13, textColor=_MUTED),
        "status": ParagraphStyle(
            "RegStatus", parent=base, fontName=BOLD_FONT, fontSize=10, leading=12, textColor=_INK, alignment=TA_CENTER
        ),
        "subject": ParagraphStyle(
            "RegSubject",
            parent=base,
            fontName=BOLD_FONT,
            fontSize=13,
            leading=17,
            textColor=_INK,
            spaceBefore=5 * mm,
            spaceAfter=4 * mm,
        ),
        "section": ParagraphStyle(
            "RegSection",
            parent=base,
            fontName=BOLD_FONT,
            fontSize=11,
            leading=14,
            textColor=_INK,
            spaceBefore=4 * mm,
            spaceAfter=1.8 * mm,
        ),
        "label": ParagraphStyle(
            "RegLabel", parent=base, fontName=BODY_FONT, fontSize=8.5, leading=11, textColor=_MUTED
        ),
        "value": ParagraphStyle("RegValue", parent=base, fontName=BODY_FONT, fontSize=9.5, leading=12, textColor=_INK),
        "body": ParagraphStyle(
            "RegBody", parent=base, fontName=BODY_FONT, fontSize=10, leading=14, textColor=colors.black
        ),
        "muted": ParagraphStyle("RegMuted", parent=base, fontName=BODY_FONT, fontSize=9, leading=12, textColor=_MUTED),
        "head": ParagraphStyle("RegHead", parent=base, fontName=BOLD_FONT, fontSize=8, leading=10, textColor=_INK),
        "cell": ParagraphStyle(
            "RegCell", parent=base, fontName=BODY_FONT, fontSize=8, leading=10, textColor=colors.black
        ),
        "cell_right": ParagraphStyle(
            "RegCellRight",
            parent=base,
            fontName=BODY_FONT,
            fontSize=8,
            leading=10,
            textColor=colors.black,
            alignment=TA_RIGHT,
        ),
        "head_right": ParagraphStyle(
            "RegHeadRight",
            parent=base,
            fontName=BOLD_FONT,
            fontSize=8,
            leading=10,
            textColor=_INK,
            alignment=TA_RIGHT,
        ),
    }


def _para(text: Any, style: ParagraphStyle) -> Paragraph:
    """A Paragraph for user-supplied text: escaped, newlines kept, face per script."""
    rendered = "" if text is None else str(text)
    markup = html.escape(rendered, quote=True).replace("\n", "<br/>")
    return Paragraph(markup, pdf_style_for_text(style, rendered))


def _label_value_table(rows: Sequence[Sequence[str]], styles: dict[str, ParagraphStyle], widths: list[float]) -> Table:
    """A grid whose even columns are labels and odd columns are values."""

    def _style_for(_row: int, col: int) -> ParagraphStyle:
        return styles["label"] if col % 2 == 0 else styles["value"]

    table = Table(pdf_table_paragraph_rows(rows, styles["value"], style_for=_style_for), colWidths=widths)
    commands: list[tuple[Any, ...]] = [
        ("GRID", (0, 0), (-1, -1), 0.5, _RULE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 1.8 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8 * mm),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
    ]
    for col in range(0, len(widths), 2):
        commands.append(("BACKGROUND", (col, 0), (col, -1), _LABEL_FILL))
    table.setStyle(TableStyle(commands))
    return table


def _data_table(
    rows: Sequence[Sequence[str]],
    weights: Sequence[float],
    right_aligned: Sequence[bool],
    usable_width: float,
    styles: dict[str, ParagraphStyle],
) -> Table:
    """A header row plus data rows, the header repeated on every page."""
    total = sum(weights) or 1.0
    widths = [usable_width * weight / total for weight in weights]

    def _style_for(row: int, col: int) -> ParagraphStyle:
        right = right_aligned[col] if col < len(right_aligned) else False
        if row == 0:
            return styles["head_right"] if right else styles["head"]
        return styles["cell_right"] if right else styles["cell"]

    table = Table(
        pdf_table_paragraph_rows(rows, styles["cell"], style_for=_style_for),
        colWidths=widths,
        repeatRows=1,
        splitInRow=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, _RULE),
                ("BACKGROUND", (0, 0), (-1, 0), _LABEL_FILL),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, _ZEBRA]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 1.3 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1.3 * mm),
                ("LEFTPADDING", (0, 0), (-1, -1), 1.6 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 1.6 * mm),
            ]
        )
    )
    return table


def build_register_pdf(document: RegisterDocument) -> bytes:
    """Render a register as a landscape A4 PDF.

    Args:
        document: The translated register.

    Returns:
        The PDF as bytes, starting with ``b"%PDF"``.
    """
    pagesize = landscape(A4)
    usable = pagesize[0] - 2 * _MARGIN_SIDE
    text_rows = register_text_rows(document)
    right_aligned = [column.kind in ("money", "count") for column in document.columns]
    weights = [column.weight for column in document.columns]

    def _story() -> tuple[list[Any], bool]:
        styles = _styles()
        flow: list[Any] = []
        letterhead = branded_letterhead(usable, doc_type=document.doc_type)
        if letterhead is not None:
            flow.append(letterhead)
        flow.append(_para(document.title, styles["title"]))
        flow.append(Spacer(1, 2 * mm))
        if document.details:
            # Two label and value pairs per row, so six details take three
            # lines instead of six and the table starts higher on page one.
            pairs = list(document.details)
            grid_rows = []
            for index in range(0, len(pairs), 2):
                left = pairs[index]
                right = pairs[index + 1] if index + 1 < len(pairs) else ("", "")
                grid_rows.append([left[0], left[1], right[0], right[1]])
            label_w, value_w = 38 * mm, usable / 2 - 38 * mm
            flow.append(_label_value_table(grid_rows, styles, [label_w, value_w, label_w, value_w]))
            flow.append(Spacer(1, 4 * mm))
        flow.append(_data_table(text_rows, weights, right_aligned, usable, styles))
        if len(text_rows) == 1:
            flow.append(Spacer(1, 3 * mm))
            flow.append(_para(document.empty_text, styles["muted"]))
        return flow, letterhead is not None

    return _build_twice(
        _story,
        pagesize=pagesize,
        margin_side=_MARGIN_SIDE,
        title=document.title,
        subject=document.title,
        page_label=document.page_label,
        generated_label=document.generated_label,
        generated=document.generated,
        doc_type=document.doc_type,
    )


def build_register_xlsx(document: RegisterDocument) -> io.BytesIO:
    """Render a register as a print-ready workbook.

    Dates are written as dates and amounts as numbers, each with a number
    format, so the sheet sorts, filters and sums. Text cells go through the
    formula neutraliser: a subject line is user text, and a cell that starts
    with ``=`` would run when a colleague opens the log.

    Args:
        document: The translated register.

    Returns:
        A buffer positioned at the start of the ``.xlsx`` file.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    from app.core.csv_safety import neutralise_formula
    from app.core.xlsx_branding import apply_company_header
    from app.core.xlsx_text import store_strings_as_text

    wb = Workbook()
    ws = wb.active
    # A sheet name holds 31 characters and none of these seven.
    sheet_title = "".join(" " if ch in "[]:*?/\\" else ch for ch in (document.sheet_title or document.title))
    ws.title = sheet_title[:31].strip() or "Register"

    for col, column in enumerate(document.columns, 1):
        ws.cell(row=1, column=col, value=column.header).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col)].width = column.xlsx_width

    date_number_format = xlsx_date_format(document.date_format)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    for row_index, row in enumerate(document.rows, 2):
        for col, (value, column) in enumerate(zip(row, document.columns, strict=True), 1):
            cell = ws.cell(row=row_index, column=col)
            cell.alignment = wrap_top if column.wrap else top
            if column.kind == "date":
                parsed = parse_stored_date(value)
                if parsed is not None:
                    cell.value = parsed
                    cell.number_format = date_number_format
                elif value is not None and str(value).strip():
                    cell.value = neutralise_formula(str(value).strip())
            elif column.kind == "money":
                amount = _as_decimal(value)
                if amount is not None:
                    cell.value = amount
                    cell.number_format = "#,##0.00"
            elif column.kind == "count":
                if value is not None and not isinstance(value, bool) and str(value).strip() != "":
                    try:
                        cell.value = int(value)
                    except (TypeError, ValueError):
                        cell.value = neutralise_formula(str(value))
            elif value is not None and str(value).strip() not in ("", EMPTY):
                # A dash stands for "nobody" on paper; in a workbook the cell
                # is left empty so a filter on blanks finds it.
                cell.value = neutralise_formula(str(value).strip())

    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    store_strings_as_text(ws)
    # The document block says which project and when, with or without a
    # company letterhead: a register detached from its letter must still say
    # what it is.
    apply_company_header(
        ws,
        title=document.title,
        details=[f"{label}: {value}" for label, value in document.details],
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ── Record (one item, as a form) ─────────────────────────────────────────


@dataclass(frozen=True)
class RecordBlock:
    """One section of a record form.

    Attributes:
        title: The section heading.
        kind: ``"text"`` for boxed free text, ``"grid"`` for label and value
            pairs, ``"table"`` for a header row plus rows, ``"signatures"``
            for rows to sign in.
        text: The body of a ``"text"`` block.
        empty_text: Printed in a blank box when a ``"text"`` block has no
            body, so a printed copy can be filled in by hand.
        rows: Rows of a ``"grid"`` (label, value[, label, value]), a
            ``"table"`` (header row first) or ``"signatures"`` (header row
            first, then one row per signer).
        weights: Relative column widths of a ``"table"``.
        right_aligned: Which ``"table"`` columns are numbers.
    """

    title: str
    kind: str = "text"
    text: str = ""
    empty_text: str = ""
    rows: Sequence[Sequence[str]] = ()
    weights: Sequence[float] = ()
    right_aligned: Sequence[bool] = ()


@dataclass(frozen=True)
class RecordDocument:
    """Everything a single-record form prints, already translated.

    Attributes:
        title: The document title, e.g. "Submittal".
        number: The record's own number.
        project_label: Project name, with its code when it has one.
        status_text: The status, already in capitals (see :func:`caps`).
        subject: The record's subject or title line.
        grid: Rows of ``[label, value, label, value]`` under the subject.
        blocks: The sections, top to bottom.
        page_label: Footer template with ``{page}`` and ``{total}``.
        generated_label: Footer template with ``{timestamp}``.
        generated: The timestamp printed in the footer.
        doc_type: Key in :data:`app.core.pdf_appearance.DOCUMENT_TYPES`.
    """

    title: str
    number: str
    project_label: str
    status_text: str
    subject: str
    grid: Sequence[Sequence[str]] = ()
    blocks: Sequence[RecordBlock] = ()
    page_label: str = "Page {page} of {total}"
    generated_label: str = "Generated {timestamp}"
    generated: str = ""
    doc_type: str | None = None


def record_text_lines(document: RecordDocument) -> list[str]:
    """Every string the record form prints, in reading order.

    The layer just before rendering, for tests that check a form without
    extracting text from the PDF.
    """
    lines = [document.title, f"{document.number} · {document.project_label}", document.status_text, document.subject]
    for row in document.grid:
        lines.extend(row)
    for block in document.blocks:
        lines.append(block.title)
        if block.kind == "text":
            lines.append(block.text if block.text.strip() else block.empty_text)
        for row in block.rows:
            lines.extend(row)
    return [line for line in lines if line]


def _boxed(flowables: list[Any], width: float, *, min_height: float | None = None) -> Table:
    """One bordered cell; long content splits across pages."""
    table = Table([[flowables]], colWidths=[width], rowHeights=[min_height] if min_height else None, splitInRow=1)
    table.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.6, _RULE),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 3 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3 * mm),
            ]
        )
    )
    return table


def _block_flowables(block: RecordBlock, usable: float, styles: dict[str, ParagraphStyle]) -> list[Any]:
    heading = Paragraph(html.escape(block.title), pdf_style_for_text(styles["section"], block.title))
    if block.kind == "grid":
        columns = max((len(row) for row in block.rows), default=2)
        if columns >= 4:
            widths = [32 * mm, usable / 2 - 32 * mm, 32 * mm, usable / 2 - 32 * mm]
        else:
            widths = [45 * mm, usable - 45 * mm]
        return [KeepTogether([heading, _label_value_table(block.rows, styles, widths)])]
    if block.kind == "table":
        weights = list(block.weights) or [1.0] * (len(block.rows[0]) if block.rows else 1)
        table = _data_table(block.rows, weights, list(block.right_aligned), usable, styles)
        return [CondPageBreak(_LONG_SECTION_MIN_ROOM), heading, table]
    if block.kind == "signatures":

        def _sign_style(row: int, col: int) -> ParagraphStyle:
            if row == 0:
                return styles["head"]
            return styles["label"] if col == 0 else styles["value"]

        table = Table(
            pdf_table_paragraph_rows(block.rows, styles["value"], style_for=_sign_style),
            colWidths=[usable * 0.2, usable * 0.32, usable * 0.3, usable * 0.18],
            rowHeights=[None, *([_SIGNATURE_ROW_HEIGHT] * (len(block.rows) - 1))],
        )
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, _RULE),
                    ("BACKGROUND", (0, 0), (-1, 0), _LABEL_FILL),
                    ("BACKGROUND", (0, 1), (0, -1), _LABEL_FILL),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ]
            )
        )
        return [KeepTogether([heading, table])]
    if block.text.strip():
        box = _boxed([_para(block.text, styles["body"])], usable)
    else:
        box = _boxed([_para(block.empty_text, styles["muted"])], usable, min_height=_BLANK_BOX_HEIGHT)
    return [CondPageBreak(_LONG_SECTION_MIN_ROOM), heading, box]


def build_record_pdf(document: RecordDocument) -> bytes:
    """Render one record as a portrait A4 form.

    Args:
        document: The translated record.

    Returns:
        The PDF as bytes, starting with ``b"%PDF"``.
    """
    usable = A4[0] - 2 * _FORM_MARGIN_SIDE

    def _story() -> tuple[list[Any], bool]:
        styles = _styles()
        flow: list[Any] = []
        letterhead = branded_letterhead(usable, doc_type=document.doc_type)
        if letterhead is not None:
            flow.append(letterhead)
        header = Table(
            [
                [
                    [
                        _para(document.title, styles["title"]),
                        _para(f"{document.number} · {document.project_label}", styles["sub"]),
                    ],
                    _para(document.status_text, styles["status"]),
                ]
            ],
            colWidths=[usable - 44 * mm, 44 * mm],
        )
        header.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (0, 0), 0),
                    ("BOX", (1, 0), (1, 0), 0.8, _INK),
                    ("TOPPADDING", (1, 0), (1, 0), 2 * mm),
                    ("BOTTOMPADDING", (1, 0), (1, 0), 2 * mm),
                    ("LINEBELOW", (0, 0), (-1, 0), 1.5, _INK),
                    ("BOTTOMPADDING", (0, 0), (0, 0), 3 * mm),
                ]
            )
        )
        flow.append(header)
        flow.append(_para(document.subject, styles["subject"]))
        if document.grid:
            half = usable / 2
            flow.append(_label_value_table(document.grid, styles, [32 * mm, half - 32 * mm, 32 * mm, half - 32 * mm]))
        for block in document.blocks:
            flow.extend(_block_flowables(block, usable, styles))
        return flow, letterhead is not None

    return _build_twice(
        _story,
        pagesize=A4,
        margin_side=_FORM_MARGIN_SIDE,
        title=f"{document.title} {document.number}".strip(),
        subject=document.subject,
        page_label=document.page_label,
        generated_label=document.generated_label,
        generated=document.generated,
        doc_type=document.doc_type,
    )
