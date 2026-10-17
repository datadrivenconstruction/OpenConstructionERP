# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Code lists: the dictionaries a cost structure is built from.

A code list is a flat set of ``(code, title, parent)`` entries that together
form a tree: the cost groups of DIN 276, the elements of NRM, the divisions of
MasterFormat, the major groups of UniFormat, or a list a company keeps for its
own projects. The core never interprets a standard by name. It answers four
questions about any list:

* which lists exist (shipped with the core, carried by the active pack, or
  imported by a user, the last being stored by the cost model module);
* where a code sits in a list, including a code the list does not carry
  (``331`` under a list that stops at ``330`` goes under ``330``);
* in which order codes sort, digits before letters and numbers by value, so
  ``2.9`` comes before ``2.10`` and ``A2`` before ``A10``;
* whether rows a user uploads form a valid list, with one issue per row that
  is wrong rather than a single refusal.

Shipped lists carry exactly the levels and titles the product carried before
they moved here from the picker in the BOQ editor. Deeper levels of the
licensed standards are not bundled; a user who holds them imports them.

Titles are data, not interface text: they are shown as written in the list.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

from app.modules.economics.classification import covers, normalise

logger = logging.getLogger(__name__)

__all__ = [
    "MAX_CODE_LENGTH",
    "MAX_ENTRIES",
    "MAX_TITLE_LENGTH",
    "SOURCE_BUILTIN",
    "SOURCE_CUSTOM",
    "SOURCE_PACK",
    "CodeList",
    "CodeListEntry",
    "CodeListIssue",
    "CodeListParseError",
    "available_code_lists",
    "build_entries",
    "builtin_code_lists",
    "code_sort_key",
    "get_code_list",
    "pack_code_lists",
    "parent_for",
    "parse_code_list_file",
    "reset_cache",
    "same_code",
    "specificity",
    "validate_code",
]

SOURCE_BUILTIN = "builtin"
SOURCE_PACK = "pack"
SOURCE_CUSTOM = "custom"

MAX_CODE_LENGTH = 80
MAX_TITLE_LENGTH = 255
#: Rows one import may carry. A national standard to its deepest level is a few thousand.
MAX_ENTRIES = 20000

_DATA_DIR = Path(__file__).parent / "data"

# Letters (any script) and digits, joined by single separators. ``03 30 00``,
# ``2.6.1``, ``B2010``, ``KG-300`` and ``ГЭСН.01`` are codes; ``3..1``, `` 300``
# and ``300/`` are not.
_CODE_RE = re.compile(r"^[^\W_]+(?:[ .\-/_][^\W_]+)*$")
_SEPARATORS = re.compile(r"[\s.\-_/]+")
_TOKENS = re.compile(r"\d+|[^\W\d_]+")


# ── Entries and lists ────────────────────────────────────────────────────────


@dataclass(frozen=True, slots=True)
class CodeListEntry:
    """One code of a list: its title, its parent code (None at the top) and its depth (1 at the top)."""

    code: str
    title: str
    parent: str | None
    level: int


@dataclass(frozen=True, slots=True)
class CodeList:
    """A whole code list, entries in tree order (each parent before its children)."""

    key: str
    name: str
    standard: str
    source: str
    entries: tuple[CodeListEntry, ...]
    max_depth: int | None = None
    pack_slug: str | None = None
    _by_norm: dict[str, CodeListEntry] = field(default_factory=dict, compare=False, repr=False)

    def __post_init__(self) -> None:
        index = {normalise(e.code): e for e in self.entries}
        object.__setattr__(self, "_by_norm", index)

    @property
    def depth(self) -> int:
        """The deepest level any entry sits on."""
        return max((e.level for e in self.entries), default=0)

    def find(self, code: object) -> CodeListEntry | None:
        """The entry for ``code``, matched as :func:`same_code` matches.

        An exact match (after trimming, case and separators) wins. Failing that, a
        run-together digit code matches the one entry with the same significant
        digits, so ``03 00 00`` finds the division the list spells ``03``. When two
        entries share those digits nothing is guessed.
        """
        norm = normalise(code)
        if not norm:
            return None
        hit = self._by_norm.get(norm)
        if hit is not None:
            return hit
        if norm.isdigit():
            stem = norm.rstrip("0")
            twins = [e for n, e in self._by_norm.items() if n.isdigit() and n.rstrip("0") == stem]
            if len(twins) == 1:
                return twins[0]
        return None

    def ancestors(self, code: str) -> list[CodeListEntry]:
        """The entries above ``code`` in the list, top first. ``code`` itself is not included."""
        chain: list[CodeListEntry] = []
        entry = self.find(code)
        seen: set[str] = set()
        parent = entry.parent if entry is not None else None
        while parent is not None and parent not in seen:
            seen.add(parent)
            hit = self.find(parent)
            if hit is None:
                break
            chain.append(hit)
            parent = hit.parent
        chain.reverse()
        return chain


@dataclass(frozen=True, slots=True)
class CodeListIssue:
    """One thing wrong with uploaded rows. ``problem`` is a stable key the UI translates."""

    row: int | None
    code: str
    problem: str
    params: dict[str, Any] = field(default_factory=dict)


class CodeListParseError(ValueError):
    """The file as a whole cannot be read as a code list. ``problem`` is a stable key."""

    def __init__(self, problem: str, **params: Any) -> None:
        super().__init__(problem)
        self.problem = problem
        self.params = params


# ── Code shape ───────────────────────────────────────────────────────────────


def validate_code(code: object) -> str | None:
    """Why ``code`` cannot be an account code, as a problem key, or None when it can."""
    text = "" if code is None else str(code)
    if not text.strip():
        return "code_missing"
    if len(text) > MAX_CODE_LENGTH:
        return "code_too_long"
    if text != text.strip() or not _CODE_RE.match(text):
        return "code_invalid"
    return None


def code_sort_key(code: object) -> tuple[tuple[tuple[int, int | str, int], ...], ...]:
    """Sort key for codes: segment by segment, digits before letters, numbers by value.

    ``2.9 < 2.10``, ``A2 < A10``, ``03 10 00 < 03 20 00`` and ``300 < 1000``.
    Leading zeros do not change the order of a number but break a tie, so ``01``
    and ``1`` sort next to each other in a fixed order. Letters compare without
    case.
    """
    text = "" if code is None else str(code).strip()
    key: list[tuple[tuple[int, int | str, int], ...]] = []
    for segment in _SEPARATORS.split(text):
        if not segment:
            continue
        tokens: list[tuple[int, int | str, int]] = []
        for token in _TOKENS.findall(segment):
            if token.isdigit():
                tokens.append((0, int(token), len(token)))
            else:
                tokens.append((1, token.casefold(), 0))
        key.append(tuple(tokens))
    return tuple(key)


def specificity(code: object) -> int:
    """How deep a code reaches by its shape: segments, significant digits or characters."""
    norm = normalise(code)
    if "." in norm:
        return len([s for s in norm.split(".") if s])
    if norm.isdigit():
        return len(norm.rstrip("0"))
    return len(norm)


def same_code(a: object, b: object) -> bool:
    """Do two spellings name the same code (case, spaces and dashes ignored)?"""
    na, nb = normalise(a), normalise(b)
    return bool(na) and na == nb


def parent_for(code: object, candidates: Iterable[str]) -> str | None:
    """The most specific candidate that covers ``code`` and is strictly coarser than it.

    ``parent_for("331", ["300", "330", "340"]) == "330"``. Codes are compared by
    shape only (see :mod:`app.modules.economics.classification`), so the answer
    is the same for every standard of the same shape.
    """
    own = specificity(code)
    best: str | None = None
    best_depth = -1
    for candidate in candidates:
        if same_code(candidate, code) or not covers(candidate, code):
            continue
        depth = specificity(candidate)
        if depth >= own:
            continue
        if depth > best_depth or (depth == best_depth and best is not None and len(candidate) < len(best)):
            best, best_depth = candidate, depth
    return best


# ── Building and validating a list ───────────────────────────────────────────


def build_entries(
    rows: Iterable[tuple[int | None, object, object, object]],
) -> tuple[list[CodeListEntry], list[CodeListIssue]]:
    """Turn ``(row, code, title, parent)`` rows into tree-ordered entries, or say what is wrong.

    A row without a parent is placed under the most specific other code of the
    list that covers it by shape, so a list of ``300, 330, 331`` needs no parent
    column. An explicit parent must be a code of the same list. Every problem is
    reported once, against the row that has it; when there is any problem the
    entries returned are those of the valid rows only and the caller decides
    whether a partial list is acceptable (the import refuses it).
    """
    issues: list[CodeListIssue] = []
    seen: dict[str, int | None] = {}
    accepted: list[tuple[int | None, str, str, str | None]] = []
    count = 0
    for row, raw_code, raw_title, raw_parent in rows:
        count += 1
        if count > MAX_ENTRIES:
            issues.append(CodeListIssue(row=row, code="", problem="too_many_rows", params={"max": MAX_ENTRIES}))
            break
        code = "" if raw_code is None else str(raw_code).strip()
        title = "" if raw_title is None else str(raw_title).strip()
        parent = None if raw_parent in (None, "") else str(raw_parent).strip() or None
        problem = validate_code(code)
        if problem is not None:
            issues.append(CodeListIssue(row=row, code=code, problem=problem))
            continue
        if len(title) > MAX_TITLE_LENGTH:
            issues.append(CodeListIssue(row=row, code=code, problem="title_too_long", params={"max": MAX_TITLE_LENGTH}))
            continue
        norm = normalise(code)
        if norm in seen:
            issues.append(CodeListIssue(row=row, code=code, problem="duplicate_code", params={"first_row": seen[norm]}))
            continue
        seen[norm] = row
        if parent is not None and same_code(parent, code):
            issues.append(CodeListIssue(row=row, code=code, problem="parent_self"))
            continue
        accepted.append((row, code, title, parent))

    by_norm = {normalise(c): c for _, c, _, _ in accepted}
    codes = [c for _, c, _, _ in accepted]
    parents: dict[str, str | None] = {}
    rows_of: dict[str, int | None] = {}
    titles: dict[str, str] = {}
    for row, code, title, parent in accepted:
        rows_of[code] = row
        titles[code] = title
        if parent is not None:
            resolved = by_norm.get(normalise(parent))
            if resolved is None:
                issues.append(CodeListIssue(row=row, code=code, problem="parent_unknown", params={"parent": parent}))
                parents[code] = None
                continue
            parents[code] = resolved
        else:
            parents[code] = parent_for(code, codes)

    # A cycle (a -> b -> a) is a list nobody can draw. Report each code on one once.
    levels: dict[str, int] = {}
    cyclic: set[str] = set()
    for code in codes:
        chain: list[str] = []
        current: str | None = code
        while current is not None and current not in levels:
            if current in chain:
                cyclic.update(chain[chain.index(current) :])
                break
            chain.append(current)
            current = parents.get(current)
        base = levels.get(current, 0) if current is not None and current not in cyclic else 0
        for depth, member in enumerate(reversed(chain), start=1):
            if member in cyclic:
                continue
            levels[member] = base + depth
    for code in sorted(cyclic, key=code_sort_key):
        issues.append(CodeListIssue(row=rows_of.get(code), code=code, problem="parent_cycle"))

    children: dict[str | None, list[str]] = {}
    for code in codes:
        if code in cyclic:
            continue
        parent = parents.get(code)
        if parent in cyclic:
            parent = None
        children.setdefault(parent, []).append(code)
    ordered: list[CodeListEntry] = []

    def _walk(parent: str | None, level: int) -> None:
        for code in sorted(children.get(parent, []), key=code_sort_key):
            ordered.append(CodeListEntry(code=code, title=titles[code], parent=parent, level=level))
            _walk(code, level + 1)

    _walk(None, 1)
    issues.sort(key=lambda i: (i.row if i.row is not None else 0, i.problem))
    return ordered, issues


def _from_document(doc: dict[str, Any], *, source: str, pack_slug: str | None = None) -> CodeList:
    key = str(doc.get("key") or "").strip()
    if not key:
        raise ValueError("code list without a key")
    rows = [
        (index, item.get("code"), item.get("title"), item.get("parent"))
        for index, item in enumerate(doc.get("entries") or [], start=1)
        if isinstance(item, dict)
    ]
    entries, issues = build_entries(rows)
    if issues:
        raise ValueError(f"code list {key!r} is not valid: {issues[0].problem} at row {issues[0].row}")
    max_depth = doc.get("max_depth")
    return CodeList(
        key=key,
        name=str(doc.get("name") or key),
        standard=str(doc.get("standard") or key),
        source=source,
        entries=tuple(entries),
        max_depth=int(max_depth) if isinstance(max_depth, int) and max_depth > 0 else None,
        pack_slug=pack_slug,
    )


@lru_cache(maxsize=1)
def builtin_code_lists() -> tuple[CodeList, ...]:
    """The lists shipped with the core, read once per process."""
    lists: list[CodeList] = []
    for path in sorted(_DATA_DIR.glob("*.json")):
        try:
            doc = json.loads(path.read_text(encoding="utf-8"))
            lists.append(_from_document(doc, source=SOURCE_BUILTIN))
        except (OSError, ValueError) as exc:
            # A broken shipped file is a build defect; say so and keep the others.
            logger.error("Shipped code list %s cannot be read: %s", path.name, exc)
    return tuple(lists)


def pack_code_lists() -> tuple[CodeList, ...]:
    """The lists the active pack declares under ``code_lists`` in its manifest."""
    try:
        from app.core.partner_pack.discovery import get_active_pack, read_pack_file
    except ImportError:  # pragma: no cover - the pack system is part of the core
        return ()
    pack = get_active_pack()
    if pack is None:
        return ()
    lists: list[CodeList] = []
    for rel in getattr(pack, "code_lists", None) or []:
        raw = read_pack_file(pack.slug, rel)
        if raw is None:
            logger.warning("Pack %s names code list %s but does not carry it", pack.slug, rel)
            continue
        try:
            lists.append(_from_document(json.loads(raw.decode("utf-8")), source=SOURCE_PACK, pack_slug=pack.slug))
        except (UnicodeDecodeError, ValueError) as exc:
            logger.warning("Pack %s code list %s cannot be read: %s", pack.slug, rel, exc)
    return tuple(lists)


def available_code_lists() -> list[CodeList]:
    """Shipped lists, then the active pack's. A pack cannot replace a shipped key."""
    lists = list(builtin_code_lists())
    taken = {c.key for c in lists}
    for code_list in pack_code_lists():
        if code_list.key in taken:
            logger.warning("Pack code list %r ignored: the key is a shipped list", code_list.key)
            continue
        taken.add(code_list.key)
        lists.append(code_list)
    return lists


def get_code_list(key: str) -> CodeList | None:
    """A shipped or pack list by key."""
    for code_list in available_code_lists():
        if code_list.key == key:
            return code_list
    return None


def reset_cache() -> None:
    """Forget the shipped lists read so far (tests)."""
    builtin_code_lists.cache_clear()


# ── Reading an uploaded file ─────────────────────────────────────────────────

_HEADER_NAMES: dict[str, str] = {"code": "code", "title": "title", "name": "title", "parent": "parent"}


def _rows_from_table(table: Sequence[Sequence[object]]) -> list[tuple[int, object, object, object]]:
    """Map a table to ``(row, code, title, parent)``, reading a header row when there is one.

    With a header the columns are found by name (``code``, ``title`` or ``name``,
    ``parent``) in any order. Without one the first three columns are taken as
    code, title, parent. Blank rows are skipped. Row numbers are the file's own,
    counting the header, so a person can find the row in their spreadsheet.
    """
    if not table:
        raise CodeListParseError("empty_file")
    first = [str(c or "").strip().casefold() for c in table[0]]
    columns: dict[str, int] = {}
    for index, name in enumerate(first):
        target = _HEADER_NAMES.get(name)
        if target is not None and target not in columns:
            columns[target] = index
    has_header = "code" in columns
    if not has_header:
        if any(name in _HEADER_NAMES for name in first):
            raise CodeListParseError("missing_code_column")
        columns = {"code": 0, "title": 1, "parent": 2}
    start = 1 if has_header else 0

    def _cell(cells: Sequence[object], name: str) -> object:
        index = columns.get(name)
        if index is None or index >= len(cells):
            return None
        value = cells[index]
        if isinstance(value, float) and value.is_integer():
            # A spreadsheet stores 300 as 300.0; the code is "300".
            return str(int(value))
        if isinstance(value, int | float):
            return str(value)
        return value

    rows: list[tuple[int, object, object, object]] = []
    for offset, cells in enumerate(table[start:], start=start + 1):
        if not any(str(c or "").strip() for c in cells):
            continue
        rows.append((offset, _cell(cells, "code"), _cell(cells, "title"), _cell(cells, "parent")))
    if not rows:
        raise CodeListParseError("empty_file")
    return rows


def parse_code_list_file(filename: str, content: bytes) -> list[tuple[int, object, object, object]]:
    """Read a CSV or Excel upload into ``(row, code, title, parent)`` rows.

    Raises:
        CodeListParseError: the file is not a readable CSV or workbook, is empty,
            or names its columns without a ``code`` column.
    """
    name = (filename or "").casefold()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises a zoo of types on a bad file
            raise CodeListParseError("unreadable_file") from exc
        try:
            sheet = workbook.worksheets[0] if workbook.worksheets else None
            if sheet is None:
                raise CodeListParseError("empty_file")
            table = [list(r) for r in sheet.iter_rows(values_only=True, max_row=MAX_ENTRIES + 2)]
        finally:
            workbook.close()
        return _rows_from_table(table)
    if name.endswith((".csv", ".txt", ".tsv")) or not name:
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:  # pragma: no cover - cp1252 decodes almost anything
            raise CodeListParseError("unreadable_file")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
        table = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
        return _rows_from_table(table[: MAX_ENTRIES + 2])
    raise CodeListParseError("unsupported_file_type")
