"""Every BOQ export format honours the sale-rate (embedded markups) mode.

One bill, one line of 2 x 100.00, overhead 10 % and a cumulative tax of 20 %.
In the embedded mode every format must carry the sale rate 110.00 and leave
the tax as the only markup line; the stored bill must keep its 100.00.

Run::

    cd backend
    python -m pytest tests/integration/test_export_embedded_markups.py -v
"""

from __future__ import annotations

import asyncio
import io
import uuid

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient

from app.main import create_app


@pytest_asyncio.fixture(scope="module")
async def client():
    app = create_app()
    async with app.router.lifespan_context(app):
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
            yield ac


@pytest_asyncio.fixture(scope="module")
async def auth(client: AsyncClient) -> dict[str, str]:
    unique = uuid.uuid4().hex[:8]
    email = f"embed-{unique}@test.io"
    password = f"EmbedMarkups{unique}9"
    reg = await client.post(
        "/api/v1/users/auth/register",
        json={"email": email, "password": password, "full_name": "Embed Tester", "role": "admin"},
    )
    assert reg.status_code == 201, reg.text
    from ._auth_helpers import promote_to_admin

    await promote_to_admin(email)
    token = ""
    for attempt in range(3):
        data = (await client.post("/api/v1/users/auth/login", json={"email": email, "password": password})).json()
        token = data.get("access_token", "")
        if token or "Too many login attempts" not in data.get("detail", ""):
            break
        await asyncio.sleep(5 * (attempt + 1))
    assert token
    return {"Authorization": f"Bearer {token}"}


@pytest_asyncio.fixture(scope="module")
async def boq_id(client: AsyncClient, auth: dict[str, str]) -> str:
    resp = await client.post(
        "/api/v1/projects/",
        json={"name": f"Embed {uuid.uuid4().hex[:6]}", "description": "", "currency": "EUR"},
        headers=auth,
    )
    assert resp.status_code == 201, resp.text
    resp = await client.post(
        "/api/v1/boq/boqs/",
        json={"project_id": resp.json()["id"], "name": "Sale rates", "description": ""},
        headers=auth,
    )
    assert resp.status_code == 201, resp.text
    bid = resp.json()["id"]
    listed = (await client.get(f"/api/v1/boq/boqs/{bid}/markups/", headers=auth)).json()
    if isinstance(listed, dict):
        listed = listed.get("items") or listed.get("markups") or []
    for markup in listed:
        await client.delete(f"/api/v1/boq/boqs/{bid}/markups/{markup['id']}", headers=auth)
    resp = await client.post(
        f"/api/v1/boq/boqs/{bid}/positions/",
        json={
            "boq_id": bid,
            "ordinal": "01.001",
            "description": "Concrete slab",
            "unit": "m3",
            "quantity": 2,
            "unit_rate": 100,
        },
        headers=auth,
    )
    assert resp.status_code == 201, resp.text
    for body in (
        {"name": "Overhead", "category": "overhead", "percentage": 10, "sort_order": 1},
        {"name": "VAT", "category": "tax", "percentage": 20, "apply_to": "cumulative", "sort_order": 2},
    ):
        resp = await client.post(f"/api/v1/boq/boqs/{bid}/markups/", json=body, headers=auth)
        assert resp.status_code in (200, 201), resp.text
    return bid


@pytest.mark.asyncio
async def test_preview_reports_k_net_and_gap(client: AsyncClient, auth: dict[str, str], boq_id: str) -> None:
    resp = await client.get(f"/api/v1/boq/boqs/{boq_id}/export/embedding-preview/", headers=auth)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["buckets"][0]["k"] == "1.100000"
    assert body["sale_net"] == "220.00" and body["editor_net"] == "220.00"
    assert body["rounding_gap"] == "0.00" and body["within_tolerance"] is True
    assert [m["name"] for m in body["visible_markups"]] == ["VAT"]
    assert body["grand_total"] == "264.00"


def _text(path: str, content: bytes) -> str:
    if path.startswith("export/excel"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(content)).active
        return " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None)
    if path.startswith("export/pdf"):
        from pypdf import PdfReader

        return " ".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    return content.decode("utf-8", errors="replace")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "path", ["export/csv/", "export/excel/", "export/pdf/", "export/gaeb/?format=x84", "export/bc3/"]
)
async def test_every_format_carries_the_sale_rate(
    client: AsyncClient, auth: dict[str, str], boq_id: str, path: str
) -> None:
    sep = "&" if "?" in path else "?"
    resp = await client.get(f"/api/v1/boq/boqs/{boq_id}/{path}{sep}embed_markups=true", headers=auth)
    assert resp.status_code == 200, resp.text
    text = _text(path, resp.content)
    assert "110" in text, text[:2000]
    assert "Overhead" not in text, text[:2000]

    plain = await client.get(f"/api/v1/boq/boqs/{boq_id}/{path}", headers=auth)
    assert plain.status_code == 200
    assert "110" not in _text(path, plain.content).replace("1100", "")


@pytest.mark.asyncio
async def test_bill_setting_turns_the_mode_on_and_the_stored_bill_is_untouched(
    client: AsyncClient, auth: dict[str, str], boq_id: str
) -> None:
    resp = await client.patch(
        f"/api/v1/boq/boqs/{boq_id}", json={"metadata": {"export_embed_markups": True}}, headers=auth
    )
    assert resp.status_code == 200, resp.text
    csv = (await client.get(f"/api/v1/boq/boqs/{boq_id}/export/csv/", headers=auth)).text
    assert "110" in csv
    off = (await client.get(f"/api/v1/boq/boqs/{boq_id}/export/csv/?embed_markups=false", headers=auth)).text
    assert "Overhead" in off

    structured = (await client.get(f"/api/v1/boq/boqs/{boq_id}/structured/", headers=auth)).json()
    rates = [p["unit_rate"] for p in structured["positions"]] + [
        p["unit_rate"] for s in structured["sections"] for p in s["positions"]
    ]
    assert [float(r) for r in rates] == [100.0]
    assert {m["name"] for m in structured["markups"]} == {"Overhead", "VAT"}
