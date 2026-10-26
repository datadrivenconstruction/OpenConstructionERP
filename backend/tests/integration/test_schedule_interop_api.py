"""Schedule interop over HTTP: progress sheet round-trip, scheduling profiles, imported source.

Runs the real app against the test PostgreSQL cluster with two tenants: A owns
a project and a schedule, B is a viewer of nothing. 2026-05-04 is a Monday.
"""

from __future__ import annotations

import io
import itertools
import json
import uuid

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.modules.schedule.progress_sheet import COLUMNS

_WBS = itertools.count(1)
COL = {key: i + 1 for i, (key, _w, _e) in enumerate(COLUMNS)}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

XER = (
    "ERMHDR\t19.12\n"
    "%T\tTASK\n"
    "%F\ttask_id\ttask_code\ttask_name\ttarget_start_date\ttarget_end_date\ttarget_drtn_hr_cnt\ttask_type\tphys_complete_pct\n"
    "%R\t501\tA1000\tExcavation\t2026-05-04 08:00\t2026-05-08 17:00\t40\tTT_Task\t0\n"
    "%R\t502\tA1010\tFoundations\t2026-05-11 08:00\t2026-05-15 17:00\t40\tTT_Task\t0\n"
    "%T\tTASKPRED\n"
    "%F\ttask_pred_id\ttask_id\tpred_task_id\tpred_type\tlag_hr_cnt\n"
    "%R\t1\t502\t501\tPR_FS\t0\n"
    "%E\n"
)

MSPDI = """<?xml version="1.0" encoding="UTF-8"?>
<Project xmlns="http://schemas.microsoft.com/project">
  <Tasks>
    <Task><UID>0</UID><Name>Root</Name></Task>
    <Task><UID>10</UID><Name>Walls</Name><Start>2026-05-04T08:00:00</Start><Finish>2026-05-08T17:00:00</Finish>
      <Duration>PT40H0M0S</Duration></Task>
    <Task><UID>20</UID><Name>Roof</Name><Start>2026-05-11T08:00:00</Start><Finish>2026-05-15T17:00:00</Finish>
      <Duration>PT40H0M0S</Duration>
      <PredecessorLink><PredecessorUID>10</PredecessorUID><Type>1</Type></PredecessorLink></Task>
  </Tasks>
</Project>
"""


@pytest_asyncio.fixture(scope="module")
async def app_instance():
    from app.config import get_settings

    get_settings.cache_clear()
    from app.main import create_app

    app = create_app()
    async with app.router.lifespan_context(app):
        from app.database import Base, engine
        from app.modules.schedule import models as _schedule_models  # noqa: F401

        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        yield app


@pytest_asyncio.fixture(scope="module")
async def client(app_instance):
    async with AsyncClient(transport=ASGITransport(app=app_instance), base_url="http://test") as ac:
        yield ac


async def _user(client: AsyncClient, tenant: str, *, admin: bool) -> dict[str, str]:
    from sqlalchemy import update

    from app.database import async_session_factory
    from app.modules.users.models import User

    email = f"{tenant}-{uuid.uuid4().hex[:8]}@schedule-interop.io"
    password = f"Interop{uuid.uuid4().hex[:6]}9"
    reg = await client.post(
        "/api/v1/users/auth/register", json={"email": email, "password": password, "full_name": tenant}
    )
    assert reg.status_code in (200, 201), reg.text
    async with async_session_factory() as s:
        values = {"is_active": True, **({"role": "admin"} if admin else {})}
        await s.execute(update(User).where(User.email == email.lower()).values(**values))
        await s.commit()
    login = await client.post("/api/v1/users/auth/login", json={"email": email, "password": password})
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


@pytest_asyncio.fixture(scope="module")
async def tenants(client):
    a = await _user(client, "a", admin=True)
    b = await _user(client, "b", admin=False)
    proj = await client.post(
        "/api/v1/projects/", json={"name": f"Interop {uuid.uuid4().hex[:6]}", "currency": "EUR"}, headers=a
    )
    assert proj.status_code == 201, proj.text
    return {"a": a, "b": b, "project_id": proj.json()["id"]}


async def _schedule(client: AsyncClient, tenants: dict, name: str) -> str:
    res = await client.post(
        "/api/v1/schedule/schedules/",
        json={"project_id": tenants["project_id"], "name": name, "start_date": "2026-05-04"},
        headers=tenants["a"],
    )
    assert res.status_code == 201, res.text
    return res.json()["id"]


async def _activity(client: AsyncClient, tenants: dict, schedule_id: str, name: str, start: str, end: str) -> str:
    res = await client.post(
        f"/api/v1/schedule/schedules/{schedule_id}/activities/",
        json={
            "name": name,
            "wbs_code": f"01.{next(_WBS):02d}",
            "start_date": start,
            "end_date": end,
            "activity_type": "task",
        },
        headers=tenants["a"],
    )
    assert res.status_code == 201, res.text
    return res.json()["id"]


async def _export(client: AsyncClient, tenants: dict, schedule_id: str) -> bytes:
    res = await client.get(f"/api/v1/schedule/schedules/{schedule_id}/progress-sheet?locale=de", headers=tenants["a"])
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith(XLSX)
    return res.content


def _fill(data: bytes, by_key: dict[str, dict[str, object]]) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    header_row = next(int(json.loads(v)) for k, v in wb["_oe_meta"].iter_rows(values_only=True) if k == "header_row")
    for r in range(header_row + 1, ws.max_row + 1):
        key = ws.cell(row=r, column=COL["row_key"]).value
        for field, value in by_key.get(str(key), {}).items():
            ws.cell(row=r, column=COL[field]).value = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _files(data: bytes) -> dict:
    return {"file": ("progress.xlsx", data, XLSX)}


@pytest.mark.asyncio
async def test_progress_sheet_round_trip_applies_confirmed_rows(client, tenants) -> None:
    sid = await _schedule(client, tenants, "Sheet round trip")
    a1 = await _activity(client, tenants, sid, "Excavation", "2026-05-04", "2026-05-08")
    a2 = await _activity(client, tenants, sid, "Foundations", "2026-05-11", "2026-05-15")
    data = await _export(client, tenants, sid)
    wb = load_workbook(io.BytesIO(data))
    assert wb.worksheets[0].protection.sheet is True

    filled = _fill(data, {a1: {"percent_complete": 60, "actual_start": "2026-05-04", "notes": "on track"}})
    preview = await client.post(
        f"/api/v1/schedule/schedules/{sid}/progress-sheet/preview",
        files=_files(filled),
        data={"status_date": "2026-05-06"},
        headers=tenants["a"],
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    rows = {r["activity_id"]: r for r in body["rows"]}
    assert rows[a1]["applicable"] is True
    assert rows[a1]["changes"]["percent_complete"] == {"old": 0.0, "new": 60.0}
    assert rows[a2]["applicable"] is False
    assert body["changed_count"] == 1

    applied = await client.post(
        f"/api/v1/schedule/schedules/{sid}/progress-sheet/apply",
        files=_files(filled),
        data={
            "expected_sha256": body["sha256"],
            "status_date": "2026-05-06",
            "row_numbers": json.dumps([rows[a1]["row_number"]]),
        },
        headers=tenants["a"],
    )
    assert applied.status_code == 200, applied.text
    result = applied.json()
    assert result["applied_count"] == 1
    assert result["data_date_advanced"] is True

    gantt = (await client.get(f"/api/v1/schedule/schedules/{sid}/gantt/", headers=tenants["a"])).json()
    progress = {a["id"]: float(a["progress_pct"]) for a in gantt["activities"]}
    assert progress[a1] == 60.0
    updates = (
        await client.get(
            f"/api/v1/schedule/progress-updates/?project_id={tenants['project_id']}&activity_id={a1}",
            headers=tenants["a"],
        )
    ).json()
    assert updates[0]["actual_start"] == "2026-05-04"
    assert updates[0]["status"] == "approved"
    meta = updates[0].get("metadata") or updates[0].get("metadata_") or {}
    assert meta["source"] == "progress_sheet"
    schedule = (await client.get(f"/api/v1/schedule/schedules/{sid}", headers=tenants["a"])).json()
    assert schedule["data_date"] == "2026-05-06"

    # The same sheet again: its stamp no longer matches, so the row is flagged stale.
    again = await client.post(
        f"/api/v1/schedule/schedules/{sid}/progress-sheet/preview", files=_files(filled), headers=tenants["a"]
    )
    stale = {r["activity_id"]: r for r in again.json()["rows"]}[a1]
    assert stale["stale"] is True
    assert any(i["rule_id"].endswith("unchanged_since_export") for i in stale["issues"])


@pytest.mark.asyncio
async def test_progress_sheet_rejects_bad_rows_other_schedules_and_changed_files(client, tenants) -> None:
    sid = await _schedule(client, tenants, "Sheet errors")
    a1 = await _activity(client, tenants, sid, "Masonry", "2026-05-04", "2026-05-08")
    other = await _schedule(client, tenants, "Sheet other")
    data = await _export(client, tenants, sid)
    bad = _fill(data, {a1: {"percent_complete": 150, "actual_start": "2026-05-07", "actual_finish": "2026-05-05"}})
    preview = await client.post(
        f"/api/v1/schedule/schedules/{sid}/progress-sheet/preview",
        files=_files(bad),
        data={"status_date": "2026-05-10"},
        headers=tenants["a"],
    )
    row = preview.json()["rows"][0]
    assert row["has_errors"] is True and row["applicable"] is False
    codes = {i["code"] for i in row["issues"]}
    assert {"percent_out_of_range", "finish_before_start"} <= codes

    skipped = await client.post(
        f"/api/v1/schedule/schedules/{sid}/progress-sheet/apply",
        files=_files(bad),
        data={
            "expected_sha256": preview.json()["sha256"],
            "status_date": "2026-05-10",
            "row_numbers": json.dumps([row["row_number"]]),
        },
        headers=tenants["a"],
    )
    assert skipped.json()["applied_count"] == 0
    assert skipped.json()["skipped"][0]["code"] == "has_errors"

    wrong = await client.post(
        f"/api/v1/schedule/schedules/{other}/progress-sheet/preview", files=_files(bad), headers=tenants["a"]
    )
    assert wrong.status_code == 409
    assert wrong.json()["detail"]["code"] == "sheet_for_another_schedule"

    mismatch = await client.post(
        f"/api/v1/schedule/schedules/{sid}/progress-sheet/apply",
        files=_files(bad),
        data={"expected_sha256": "0" * 64, "status_date": "2026-05-10", "row_numbers": "[]"},
        headers=tenants["a"],
    )
    assert mismatch.status_code == 409
    assert mismatch.json()["detail"]["code"] == "preview_mismatch"


@pytest.mark.asyncio
async def test_other_tenant_cannot_use_any_interop_endpoint(client, tenants) -> None:
    sid = await _schedule(client, tenants, "Private")
    await _activity(client, tenants, sid, "Secret", "2026-05-04", "2026-05-08")
    b = tenants["b"]
    for method, url in (
        ("get", f"/api/v1/schedule/schedules/{sid}/progress-sheet"),
        ("get", f"/api/v1/schedule/schedules/{sid}/scheduling-profile"),
        ("post", f"/api/v1/schedule/schedules/{sid}/recalculate/preview"),
        ("get", f"/api/v1/schedule/schedules/{sid}/import-source"),
        ("get", f"/api/v1/schedule/scheduling-profiles/?project_id={tenants['project_id']}"),
    ):
        res = await (client.get(url, headers=b) if method == "get" else client.post(url, json={}, headers=b))
        assert res.status_code in (403, 404), (url, res.status_code, res.text)


@pytest.mark.asyncio
async def test_scheduling_profile_recalculation_names_the_setting(client, tenants) -> None:
    a = tenants["a"]
    sid = await _schedule(client, tenants, "Profiles")
    pred = await _activity(client, tenants, sid, "Pour", "2026-05-04", "2026-05-08")
    succ = await _activity(client, tenants, sid, "Strip formwork", "2026-05-11", "2026-05-12")
    rel = await client.post(
        f"/api/v1/schedule/schedules/{sid}/relationships/",
        json={"predecessor_id": pred, "successor_id": succ, "relationship_type": "FS", "lag_days": 6},
        headers=a,
    )
    assert rel.status_code == 201, rel.text

    created = await client.post(
        "/api/v1/schedule/scheduling-profiles/",
        json={
            "project_id": tenants["project_id"],
            "name": f"Working-day lags {uuid.uuid4().hex[:4]}",
            "lag_calendar": "project",
            "critical_rule": "longest_path",
            "is_default": True,
        },
        headers=a,
    )
    assert created.status_code == 201, created.text
    profile = created.json()
    dup = await client.post(
        "/api/v1/schedule/scheduling-profiles/",
        json={"project_id": tenants["project_id"], "name": profile["name"]},
        headers=a,
    )
    assert dup.status_code == 409

    effective = (await client.get(f"/api/v1/schedule/schedules/{sid}/scheduling-profile", headers=a)).json()
    assert effective["source"] == "project_default"
    assert effective["settings"]["lag_calendar"] == "project"

    preview = await client.post(f"/api/v1/schedule/schedules/{sid}/recalculate/preview", json={}, headers=a)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    rows = {r["activity_id"]: r for r in body["activities"]}
    # Mon 11 May + six working days = Tue 19 May (calendar days would give Mon 18 May).
    assert rows[succ]["new_start"] == "2026-05-19"
    assert "lag_calendar" in rows[succ]["changed_by"]
    assert body["setting_effects"]["lag_calendar"] >= 1
    assert body["applied"] is False

    applied = await client.post(f"/api/v1/schedule/schedules/{sid}/recalculate/apply", json={}, headers=a)
    assert applied.status_code == 200 and applied.json()["applied"] is True
    gantt = (await client.get(f"/api/v1/schedule/schedules/{sid}/gantt/", headers=a)).json()
    assert {x["id"]: x["start_date"] for x in gantt["activities"]}[succ] == "2026-05-19"

    assigned = await client.put(
        f"/api/v1/schedule/schedules/{sid}/scheduling-profile", json={"profile_id": profile["id"]}, headers=a
    )
    assert assigned.json()["source"] == "schedule"
    deleted = await client.delete(f"/api/v1/schedule/scheduling-profiles/{profile['id']}", headers=a)
    assert deleted.status_code == 204
    fallback = (await client.get(f"/api/v1/schedule/schedules/{sid}/scheduling-profile", headers=a)).json()
    assert fallback["source"] == "builtin"


@pytest.mark.asyncio
async def test_xer_import_keeps_the_source_until_accepted(client, tenants) -> None:
    a = tenants["a"]
    sid = await _schedule(client, tenants, "Imported")
    res = await client.post(
        f"/api/v1/schedule/schedule/import/xer/?schedule_id={sid}",
        files={"file": ("programme.xer", XER.encode("utf-8"), "text/plain")},
        headers=a,
    )
    assert res.status_code == 201, res.text
    assert res.json()["activities_imported"] == 2
    assert res.json()["source_document_id"]

    comparison = (await client.get(f"/api/v1/schedule/schedules/{sid}/import-source", headers=a)).json()
    assert comparison["source"]["format"] == "xer"
    assert comparison["source"]["status"] == "pending_review"
    assert comparison["deviating_count"] == 0
    by_code = {r["activity_code"]: r for r in comparison["activities"]}
    assert by_code["A1010"]["source_id"] == "502"
    assert by_code["A1010"]["source_finish"] == "2026-05-15"

    moved = await client.patch(
        f"/api/v1/schedule/activities/{by_code['A1010']['activity_id']}",
        json={"start_date": "2026-05-12", "end_date": "2026-05-18"},
        headers=a,
    )
    assert moved.status_code == 200, moved.text
    comparison = (await client.get(f"/api/v1/schedule/schedules/{sid}/import-source", headers=a)).json()
    assert comparison["deviating_count"] == 1
    assert {r["activity_code"]: r for r in comparison["activities"]}["A1010"]["finish_deviation_days"] == 3

    restored = await client.post(f"/api/v1/schedule/schedules/{sid}/import-source/restore", headers=a)
    assert restored.json()["moved"] == 1
    accepted = await client.post(f"/api/v1/schedule/schedules/{sid}/import-source/accept", headers=a)
    assert accepted.json()["status"] == "accepted"
    again = await client.post(f"/api/v1/schedule/schedules/{sid}/import-source/restore", headers=a)
    assert again.status_code == 409

    document = await client.get(f"/api/v1/documents/{comparison['source']['document_id']}", headers=a)
    assert document.status_code == 200, document.text


@pytest.mark.asyncio
async def test_mspdi_reexport_keeps_the_source_uids(client, tenants) -> None:
    a = tenants["a"]
    sid = await _schedule(client, tenants, "Round trip ids")
    res = await client.post(
        f"/api/v1/schedule/schedule/import/msp-xml/?schedule_id={sid}",
        files={"file": ("programme.xml", MSPDI.encode("utf-8"), "application/xml")},
        headers=a,
    )
    assert res.status_code == 201, res.text
    exported = await client.get(f"/api/v1/schedule/schedule/export/msp-xml/?schedule_id={sid}", headers=a)
    assert exported.status_code == 200
    xml = exported.text
    assert "<UID>10</UID>" in xml and "<UID>20</UID>" in xml
    assert "<PredecessorUID>10</PredecessorUID>" in xml
