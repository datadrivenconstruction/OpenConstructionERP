# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The HSE frame exports on a seeded German demo project: both registers, both formats.

The spreadsheet is opened and read back, so a header that lost its German
wording, a register without rows or a sheet title that fell back to English
fails here; the PDF must be a real, non-trivial document. The register sheet
keeps its header on row 1 so it sorts and filters, and carries the register
name as the sheet title.
"""

from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import load_workbook

from app.modules.hse_frame.export import render, rendered_language
from app.modules.hse_frame.seed import seed_hse_frame_demo
from app.modules.hse_frame.service import HseFrameService
from tests.pg.governance_fixtures import add_project, add_user

pytestmark = pytest.mark.asyncio

HEADERS = {
    "register": ["Thema", "Anforderung", "Rechtsgrundlage", "Verbindlichkeit", "Verantwortlich", "Fällig", "Status"],
    "roles": ["Rolle", "Person", "Organisation", "Status", "Bestellt am", "Qualifikation", "Gültig bis"],
}
TITLES = {"register": "Register der HSE-Anforderungen", "roles": "HSE-Verantwortlichkeiten"}


async def _seeded(session) -> tuple[HseFrameService, uuid.UUID]:
    from app.modules.projects.models import Project

    owner = await add_user(session, "manager")
    pid = await add_project(session, owner, [])
    project = await session.get(Project, pid)
    project.country_code, project.subdivision_code = "DE", "DE-BY"
    await session.flush()
    assert (await seed_hse_frame_demo(session, [pid]))["projects"] == 1
    return HseFrameService(session), pid


def _sheet(data: bytes) -> tuple[str, list[list[str]]]:
    sheet = load_workbook(io.BytesIO(data)).active
    return sheet.title, [[("" if c is None else str(c)) for c in row] for row in sheet.iter_rows(values_only=True)]


@pytest.mark.parametrize("what", ["register", "roles"])
async def test_the_spreadsheet_is_german_and_has_the_rows(pg_session, what: str) -> None:
    svc, pid = await _seeded(pg_session)
    data = await render(svc, pid, what=what, fmt="xlsx", locale="de")
    title, rows = _sheet(data)
    flat = [cell for row in rows for cell in row]
    assert title == TITLES[what][:31]
    header = rows[0]
    assert [h for h in HEADERS[what] if h not in header] == []
    body = rows[rows.index(header) + 1 :]
    assert len([row for row in body if any(row)]) >= (10 if what == "register" else 2)
    # No status word is left in English.
    assert not {"Open", "Met", "Appointed", "Mandatory"} & set(flat)


@pytest.mark.parametrize("what", ["register", "roles"])
async def test_the_pdf_is_a_document(pg_session, what: str) -> None:
    svc, pid = await _seeded(pg_session)
    data = await render(svc, pid, what=what, fmt="pdf", locale="de")
    assert data.startswith(b"%PDF") and len(data) > 2000


async def test_english_is_the_fallback_and_is_declared(pg_session) -> None:
    svc, pid = await _seeded(pg_session)
    title, rows = _sheet(await render(svc, pid, what="roles", fmt="xlsx", locale="en"))
    assert title == "HSE responsibilities"
    assert rows[0][0] == "Role"
    assert rendered_language("de-AT") == "de"
    assert rendered_language("xx") == "en"
