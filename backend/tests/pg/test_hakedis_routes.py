# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The payment certificate (hakediş) through its routes: who may, and what is printed.

Two projects with two owners. The certificate of a claim on the first is read,
edited, taxed and printed by its owner, in Turkish, in English and in both.
Then somebody from the second project asks every route for it, and asks the
statutory tax routes to file the claim's taxes under a project of their own.
Each of those answers 404 and leaves the owner's sheet as it was.

Every rate in this file is SYNTHETIC and none may be read as a Turkish rate.
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass, replace
from datetime import date
from decimal import Decimal

import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.payment_taxes import RateRow
from app.modules.contracts import certificate_taxes, hakedis_document
from app.modules.contracts.models import Contract, ContractLine, ProgressClaim
from app.modules.contracts.schemas import AutoGenerateClaimRequest
from app.modules.contracts.service import ContractsService
from app.modules.projects.models import Project
from app.modules.tax_withholding import service as tax_service
from app.modules.tax_withholding import source_owners
from app.modules.tax_withholding.models import StatutoryTaxCalc
from app.modules.users.models import User
from tests._pg import transactional_session

pytestmark = pytest.mark.asyncio

D = Decimal
CLAIMS = "/api/v1/contracts/progress-claims"
STATUTORY = "/api/v1/tax-withholding/statutory"

# SYNTHETIC rates, illustrative only.
SYNTHETIC_VAT_PCT = D("20")


def _row(**changes) -> RateRow:
    base = RateRow(
        country_code="TR",
        kind="vat_withholding",
        code="W1",
        labels={"en": "Synthetic work", "tr": "Sentetik iş"},
        base="vat",
        rate_pct=None,
        numerator=1,
        denominator=2,
        threshold_amount=None,
        threshold_currency="",
        threshold_scope="",
        threshold_measure="",
        cap_amount=None,
        effective_from=date(2020, 1, 1),
        effective_to=None,
        legal_reference="Synthetic Act art. 1",
        source_url="https://example.invalid/act/1",
        read_date="2026-01-01",
        review_status="confirmed",
    )
    return replace(base, **changes)


ROWS: tuple[RateRow, ...] = (
    _row(),
    _row(kind="income_withholding", code="I1", base="net", rate_pct=D("3"), numerator=None, denominator=None),
    _row(kind="stamp_duty", code="S1", base="net", rate_pct=D("0.5"), numerator=None, denominator=None),
)

TAXES = {
    "vat_withholding": {"state": "selected", "code": "W1"},
    "income_withholding": {"state": "selected", "code": "I1"},
    "stamp_duty": {"state": "not_applicable", "reason": "Kağıt düzenlenmedi"},
}


@dataclass
class World:
    session: object
    app: FastAPI
    owner_id: uuid.UUID
    project_id: uuid.UUID
    outsider_id: uuid.UUID
    outsider_project_id: uuid.UUID
    claim_id: uuid.UUID
    acting: dict

    def act_as(self, user_id: uuid.UUID, role: str = "manager") -> None:
        self.acting.clear()
        self.acting.update({"sub": str(user_id), "role": role})

    def client(self) -> AsyncClient:
        return AsyncClient(transport=ASGITransport(app=self.app), base_url="http://test")

    @property
    def url(self) -> str:
        return f"{CLAIMS}/{self.claim_id}/hakedis"


@pytest_asyncio.fixture
async def world(monkeypatch):
    from app.dependencies import get_current_user_payload, get_session
    from app.modules.contracts import router as contracts_router
    from app.modules.contracts.permissions import register_contracts_permissions
    from app.modules.contracts.validators import register_contracts_validation_rules
    from app.modules.tax_withholding import router as tax_router
    from app.modules.tax_withholding.permissions import register_tax_withholding_permissions
    from app.modules.tax_withholding.validators import register_tax_withholding_rules

    async def vat_rate(_session, _country, _subdivision, _on):
        return SYNTHETIC_VAT_PCT

    monkeypatch.setattr(tax_service, "certificate_rows", lambda country: ROWS)
    monkeypatch.setattr(hakedis_document, "tax_rows", lambda country: ROWS)
    monkeypatch.setattr(hakedis_document, "vat_rate_source", vat_rate)

    provider, writer = certificate_taxes._provider, certificate_taxes._writer
    owners = dict(source_owners._resolvers)
    register_contracts_permissions()
    register_tax_withholding_permissions()
    register_tax_withholding_rules()
    # The production registrar: it is what claims progress claims as a tax source.
    register_contracts_validation_rules()
    tax_service.register_certificate_tax_provider()
    try:
        async with transactional_session() as session:
            tag = uuid.uuid4().hex[:8]
            owner = User(email=f"hk-owner-{tag}@reference.example", hashed_password="x", full_name="Owner")
            outsider = User(email=f"hk-outsider-{tag}@reference.example", hashed_password="x", full_name="Outsider")
            session.add_all([owner, outsider])
            await session.flush()
            project = Project(name=f"Kule {tag}", owner_id=owner.id, currency="TRY", country_code="TR")
            other = Project(name=f"Elsewhere {tag}", owner_id=outsider.id, currency="TRY", country_code="TR")
            session.add_all([project, other])
            await session.flush()
            contract = Contract(
                code=f"S-{tag}",
                title="Mekanik tesisat",
                project_id=project.id,
                contract_type="unit_price",
                counterparty_type="client",
                currency="TRY",
                total_value=D("500000"),
                original_contract_value=D("500000"),
                retention_percent=D("10"),
                status="active",
                start_date="2026-01-05",
                end_date="2027-03-31",
                terms={},
            )
            session.add(contract)
            await session.flush()
            line = ContractLine(
                contract_id=contract.id,
                code="M-25.101",
                description="Siyah çelik boru DN50, yangın tesisatı",
                unit="m",
                quantity=D("1000"),
                unit_rate=D("500"),
                total_value=D("500000"),
                metadata_={},
            )
            claim = ProgressClaim(
                contract_id=contract.id,
                claim_number="1",
                currency="TRY",
                status="draft",
                period_start="2026-02-01",
                period_end="2026-02-28",
                period_from=date(2026, 2, 1),
                period_to=date(2026, 2, 28),
            )
            session.add_all([line, claim])
            await session.flush()
            await ContractsService(session).auto_generate_claim_lines(
                claim.id, AutoGenerateClaimRequest(measurements={str(line.id): D("200")})
            )

            app = FastAPI()
            app.include_router(contracts_router.router, prefix="/api/v1/contracts")
            app.include_router(tax_router.router, prefix="/api/v1/tax-withholding")
            acting: dict = {"sub": str(owner.id), "role": "manager"}

            async def current_session():
                yield session

            app.dependency_overrides[get_session] = current_session
            app.dependency_overrides[get_current_user_payload] = lambda: dict(acting)
            app.dependency_overrides[tax_router.statutory_row_source] = lambda: lambda country: ROWS
            yield World(
                session=session,
                app=app,
                owner_id=owner.id,
                project_id=project.id,
                outsider_id=outsider.id,
                outsider_project_id=other.id,
                claim_id=claim.id,
                acting=acting,
            )
    finally:
        certificate_taxes._provider, certificate_taxes._writer = provider, writer
        source_owners._resolvers.clear()
        source_owners._resolvers.update(owners)


def _line(view: dict, key: str) -> dict:
    return next(line for line in view["summary"] if line["key"] == key)


async def _complete(client: AsyncClient, world: World) -> dict:
    """Decide every manual line and store the taxes; the sheet as it then reads."""
    view = (await client.get(world.url)).json()
    for line in view["summary"]:
        if line["op"] == "manual":
            entry = {"state": "not_applicable", "note": "Sözleşmede yok"}
            saved = await client.put(f"{world.url}/lines/{line['key']}", json=entry)
            assert saved.status_code == 200, saved.text
    taxed = await client.put(f"{world.url}/taxes", json=TAXES)
    assert taxed.status_code == 200, taxed.text
    return taxed.json()


async def test_the_owner_reads_edits_and_taxes_the_sheet(world: World) -> None:
    async with world.client() as client:
        opened = await client.get(world.url)
        assert opened.status_code == 200, opened.text
        view = opened.json()
        assert view["locale"] == "tr"
        assert view["editable"] is True and view["frozen"] is False
        assert view["can_certify"] is False
        assert D(_line(view, "work_done")["amount"]) == D("100000.00")
        assert _line(view, "advance_recovery")["enterable"] is True
        assert _line(view, "advance_recovery")["accepts_percent"] is True
        assert _line(view, "vat")["enterable"] is False
        assert {"W1"} == {category["code"] for category in view["taxes"]["categories"]["vat_withholding"]}

        # A percent needs its base: while the price adjustment above it is
        # undecided, line E has no figure and neither has a share of it.
        early = await client.put(f"{world.url}/lines/advance_recovery", json={"state": "value", "pct": "10"})
        assert early.status_code == 200, early.text
        assert _line(early.json(), "advance_recovery")["reason_key"] == "base_held"
        decided = await client.put(
            f"{world.url}/lines/price_adjustment", json={"state": "not_applicable", "note": "Sözleşmede yok"}
        )
        assert decided.status_code == 200, decided.text

        # A percent of the line's base, an amount, and taking the entry back.
        by_percent = await client.put(f"{world.url}/lines/advance_recovery", json={"state": "value", "pct": "10"})
        assert by_percent.status_code == 200, by_percent.text
        assert D(_line(by_percent.json(), "advance_recovery")["amount"]) == D("10000.00")
        by_amount = await client.put(
            f"{world.url}/lines/advance_recovery", json={"state": "value", "amount": "2500.50"}
        )
        assert D(_line(by_amount.json(), "advance_recovery")["amount"]) == D("2500.50")
        unset = await client.put(f"{world.url}/lines/advance_recovery", json={"state": "unset"})
        assert _line(unset.json(), "advance_recovery")["status"] == "held"

        # What cannot be entered says so, each with its own code.
        unknown = await client.put(f"{world.url}/lines/no_such_line", json={"state": "unset"})
        assert (unknown.status_code, unknown.json()["detail"]["error"]) == (404, "hakedis_line_unknown")
        computed = await client.put(f"{world.url}/lines/vat", json={"state": "value", "amount": "1.00"})
        assert (computed.status_code, computed.json()["detail"]["error"]) == (409, "hakedis_line_not_manual")
        unexplained = await client.put(f"{world.url}/lines/delay_penalty", json={"state": "not_applicable"})
        assert (unexplained.status_code, unexplained.json()["detail"]["error"]) == (422, "hakedis_line_entry_invalid")
        assert (await client.get(world.url, params={"locale": "de"})).status_code == 422

        taxed = await _complete(client, world)
        assert taxed["taxes"]["stored"] is True
        assert taxed["taxes"]["status"] == "draft"
        assert taxed["taxes"]["vat_rate_pct"] is not None
        assert D(_line(taxed, "vat")["amount"]) == D("20000.00")
        assert D(_line(taxed, "vat_withholding")["amount"]) == D("10000.00")
        assert _line(taxed, "stamp_duty")["status"] == "not_applicable"

        final = await client.put(f"{world.url}/options", json={"is_final": True})
        assert final.status_code == 200, final.text
        assert final.json()["header"]["is_final"] is True

        # The taxes themselves are confirmed on the tax routes, by the same person.
        confirmed = await client.post(
            f"{STATUTORY}/progress_claim/{world.claim_id}/confirm", json={"project_id": str(world.project_id)}
        )
        assert confirmed.status_code == 200, confirmed.text
        ready = (await client.get(world.url)).json()
        assert ready["taxes"]["status"] == "confirmed"
        assert ready["can_certify"] is True, ready["findings"]


@pytest.mark.parametrize(("locale", "languages"), [("tr", 1), ("en", 1), ("tr-en", 2)])
async def test_the_sheet_prints_in_turkish_english_and_both(world: World, locale: str, languages: int) -> None:
    async with world.client() as client:
        await _complete(client, world)
        view = (await client.get(world.url, params={"locale": locale})).json()
        pdf = await client.get(f"{world.url}/pdf", params={"locale": locale})
        xlsx = await client.get(f"{world.url}/xlsx", params={"locale": locale})

    assert view["locale"] == locale
    assert all(len(line["labels"]) == languages for line in view["summary"])
    assert all(len(column["labels"]) == languages for column in view["works"]["columns"])

    assert pdf.status_code == 200, pdf.text
    assert pdf.headers["content-type"].startswith("application/pdf")
    assert pdf.content.startswith(b"%PDF")
    assert len(pdf.content) > 2000
    assert f"-01-{locale}-draft.pdf" in pdf.headers["content-disposition"]

    assert xlsx.status_code == 200, xlsx.text
    workbook = load_workbook(io.BytesIO(xlsx.content))
    printed = {
        str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value
    }
    text = "\n".join(printed)
    # Every label of the summary, in every language asked for, is on the sheet,
    # Turkish with its own letters.
    for line in view["summary"]:
        for label in line["labels"]:
            assert label in text, (line["key"], label)
    assert ("Siyah çelik boru DN50, yangın tesisatı" in text) is True
    if locale != "en":
        assert "HAKEDİŞ" in text.upper() or "Hakediş" in text


async def test_a_viewer_reads_and_prints_but_does_not_edit(world: World) -> None:
    world.act_as(world.owner_id, "viewer")
    async with world.client() as client:
        assert (await client.get(world.url)).status_code == 200
        assert (await client.get(f"{world.url}/pdf")).status_code == 200
        for target, body in (
            (f"{world.url}/lines/advance_recovery", {"state": "unset"}),
            (f"{world.url}/options", {"is_final": True}),
            (f"{world.url}/taxes", TAXES),
        ):
            assert (await client.put(target, json=body)).status_code == 403, target


async def test_somebody_from_another_project_reaches_nothing(world: World) -> None:
    statutory = f"{STATUTORY}/progress_claim/{world.claim_id}"
    async with world.client() as client:
        before = await _complete(client, world)
        world.act_as(world.outsider_id, "manager")

        # The certificate's own routes.
        for method, target, kwargs in (
            ("GET", world.url, {}),
            ("GET", f"{world.url}/pdf", {}),
            ("GET", f"{world.url}/xlsx", {}),
            ("PUT", f"{world.url}/lines/advance_recovery", {"json": {"state": "value", "amount": "1.00"}}),
            ("PUT", f"{world.url}/options", {"json": {"is_final": True}}),
            ("PUT", f"{world.url}/taxes", {"json": TAXES}),
        ):
            response = await client.request(method, target, **kwargs)
            assert response.status_code == 404, (method, target, response.text)

        # The tax routes, naming first the claim's real project and then a
        # project the caller does own. Neither reads, changes, confirms,
        # reopens or voids the claim's taxes, and the claim cannot be filed
        # under the second project.
        body = {
            "direction": "withheld_by_us",
            "source_reference": "taken",
            "country_code": "TR",
            "currency_code": "TRY",
            "document_date": "2026-02-28",
            "net_amount": "1.00",
            "vat_rate_pct": "20",
            "vat_withholding": {"state": "not_applicable", "reason": "x"},
            "income_withholding": {"state": "not_applicable", "reason": "x"},
            "stamp_duty": {"state": "not_applicable", "reason": "x"},
        }
        for project in (str(world.project_id), str(world.outsider_project_id)):
            for method, target, kwargs in (
                ("GET", statutory, {"params": {"project_id": project}}),
                ("PUT", statutory, {"json": {**body, "project_id": project}}),
                (
                    "POST",
                    f"{statutory}/override",
                    {"json": {"project_id": project, "kind": "vat_withheld", "amount": "1.00", "reason": "x"}},
                ),
                ("DELETE", f"{statutory}/override/vat_withheld", {"params": {"project_id": project}}),
                ("POST", f"{statutory}/confirm", {"json": {"project_id": project}}),
                ("POST", f"{statutory}/reopen", {"json": {"project_id": project, "reason": "x"}}),
                ("POST", f"{statutory}/void", {"json": {"project_id": project, "reason": "x"}}),
            ):
                response = await client.request(method, target, **kwargs)
                assert response.status_code == 404, (method, target, project, response.text)

        # A claim id nobody has: there is no document to file taxes on, even
        # under the caller's own project.
        invented = await client.put(
            f"{STATUTORY}/progress_claim/{uuid.uuid4()}", json={**body, "project_id": str(world.outsider_project_id)}
        )
        assert invented.status_code == 404, invented.text

        world.act_as(world.owner_id, "manager")
        after = (await client.get(world.url)).json()
        read = await client.get(statutory, params={"project_id": str(world.project_id)})

    assert after["summary"] == before["summary"]
    assert after["taxes"] == before["taxes"]
    assert read.status_code == 200, read.text
    assert read.json()["project_id"] == str(world.project_id)
    assert read.json()["direction"] == "borne_by_us"
    assert read.json()["status"] == "draft"
    stored = await world.session.scalar(
        select(func.count()).select_from(StatutoryTaxCalc).where(StatutoryTaxCalc.source_id == world.claim_id)
    )
    assert stored == 1
