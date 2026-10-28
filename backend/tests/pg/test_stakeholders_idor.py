# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""​‌‍⁠​‌‍⁠​‌‍⁠​‌‍⁠Who may read, mask and decide in the stakeholders API, over HTTP against PostgreSQL."""

from __future__ import annotations

import io
import uuid
from contextlib import asynccontextmanager
from datetime import timedelta

import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.modules.stakeholders.permissions import register_stakeholder_permissions
from app.modules.stakeholders.router import router
from app.modules.stakeholders.service import EVENT_TRIGGERS, GATE_BOUND_EVENT_TRIGGERS
from tests._pg import isolated_engine
from tests.pg.stakeholders_fixtures import TODAY, World, build_world, commitment, contact, seat, stakeholder

pytestmark = [pytest.mark.asyncio, pytest.mark.tenant_isolation]

BASE = "/api/v1/stakeholders"


@pytest_asyncio.fixture
async def world():
    register_stakeholder_permissions()
    async with isolated_engine() as engine:
        factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
        yield await build_world(factory)


@asynccontextmanager
async def client_as(w: World, user_id: uuid.UUID, role: str):
    from app.dependencies import get_current_user_id, get_current_user_payload, get_session

    app = FastAPI()
    app.include_router(router, prefix=BASE)

    async def _session():
        async with w.factory() as s:
            yield s
            await s.commit()

    app.dependency_overrides[get_session] = _session
    app.dependency_overrides[get_current_user_id] = lambda: str(user_id)
    app.dependency_overrides[get_current_user_payload] = lambda: {
        "sub": str(user_id),
        "role": role,
        "permissions": [],
        "tenant_id": str(user_id),
    }
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://stk-test") as c:
        yield c


def _code(response) -> str:
    return response.json()["detail"]["code"]


# ── project and tenant boundaries ─────────────────────────────────────────


async def test_a_stranger_from_another_tenant_gets_404_everywhere(world: World) -> None:
    sid = await stakeholder(world, "Jane Private", category="neighbour")
    p = f"{BASE}/projects/{world.project_id}"
    async with client_as(world, world.outsider, "manager") as c:
        listed = await c.get(f"{p}/stakeholders")
        read = await c.get(f"{p}/stakeholders/{sid}")
        assessed = await c.post(f"{p}/stakeholders/{sid}/assessments", json={"scores": {"power": 3, "interest": 3}})
        exported = await c.get(f"{p}/export")
        swept = await c.post(f"{p}/sweep")
    assert [r.status_code for r in (listed, read, assessed, exported, swept)] == [404] * 5


async def test_another_projects_stakeholder_is_404_inside_my_project(world: World) -> None:
    other = await build_world(world.factory)
    foreign = await stakeholder(other, "Elsewhere")
    async with client_as(world, world.editor, "editor") as c:
        r = await c.get(f"{BASE}/projects/{world.project_id}/stakeholders/{foreign}")
    assert r.status_code == 404


async def test_a_real_meeting_of_another_project_cannot_be_linked_to_an_event(world: World) -> None:
    from app.modules.meetings.models import Meeting

    other = await build_world(world.factory)
    async with world.factory() as s:
        meeting = Meeting(
            project_id=other.project_id,
            meeting_number="M-001",
            meeting_type="progress",
            title="Their progress meeting",
            meeting_date="2026-10-01",
        )
        s.add(meeting)
        await s.commit()
        meeting_id = meeting.id
    async with client_as(world, world.editor, "editor") as c:
        r = await c.post(
            f"{BASE}/projects/{world.project_id}/events",
            json={"title": "Info evening", "participation_level": "inform", "meeting_id": str(meeting_id)},
        )
    assert (r.status_code, _code(r)) == (404, "meeting_not_found")


async def test_references_outside_the_callers_reach_are_refused_by_name(world: World) -> None:
    theirs = await contact(world, world.outsider, "Hidden", "Person")
    p = f"{BASE}/projects/{world.project_id}/stakeholders"
    async with client_as(world, world.editor, "editor") as c:
        foreign_contact = await c.post(p, json={"category": "neighbour", "contact_id": str(theirs)})
        foreign_risk = await c.post(
            p, json={"category": "utility", "display_name": "Utility", "risk_id": str(uuid.uuid4())}
        )
        stranger = await c.post(
            p,
            json={"category": "utility", "display_name": "Utility", "relationship_owner_user_id": str(world.outsider)},
        )
    assert (foreign_contact.status_code, _code(foreign_contact)) == (404, "contact_not_found")
    assert (foreign_risk.status_code, _code(foreign_risk)) == (404, "risk_not_found")
    assert (stranger.status_code, _code(stranger)) == (422, "user_not_project_member")
    assert "Hidden" not in foreign_contact.text


# ── rule 17: a private person's name ──────────────────────────────────────


async def test_an_editor_without_read_personal_never_receives_a_private_name(world: World) -> None:
    from openpyxl import load_workbook

    p = f"{BASE}/projects/{world.project_id}"
    async with client_as(world, world.editor, "editor") as c:
        created = await c.post(
            f"{p}/stakeholders",
            json={"category": "neighbour", "display_name": "Jane Private", "geo": {"address": "1 Garden Lane"}},
        )
        sid = created.json()["id"]
        listed = await c.get(f"{p}/stakeholders")
        read = await c.get(f"{p}/stakeholders/{sid}")
        patched = await c.patch(f"{p}/stakeholders/{sid}", json={"interests": "Quiet mornings"})
        grid = await c.get(f"{p}/grid")
        exported = await c.get(f"{p}/export", params={"format": "xlsx"})
    assert [r.status_code for r in (created, listed, read, patched, grid, exported)] == [201, 200, 200, 200, 200, 200]
    for r in (created, listed, read, patched, grid):
        assert "Jane Private" not in r.text, r.request.url
        assert "Garden Lane" not in r.text, r.request.url
    row = listed.json()["items"][0]
    assert (row["masked"], row["name"], row["display_name"], row["geo"]) == (True, None, "", None)
    wb = load_workbook(io.BytesIO(exported.content))
    cells = " ".join(str(c.value) for ws in wb for r in ws.iter_rows() for c in r if c.value is not None)
    assert "Jane Private" not in cells
    assert row["code"] in cells


async def test_a_manager_with_read_personal_sees_the_name(world: World) -> None:
    await stakeholder(world, "Jane Private", category="neighbour")
    async with client_as(world, world.manager, "manager") as c:
        r = await c.get(f"{BASE}/projects/{world.project_id}/stakeholders")
    row = r.json()["items"][0]
    assert (row["masked"], row["name"]) == (False, "Jane Private")


# ── who may decide ────────────────────────────────────────────────────────


async def test_a_viewer_cannot_put_mandates(world: World) -> None:
    rows = [{"decision_type": "scope_change", "holder_kind": "role", "holder_code": "client", "letter": "D"}]
    async with client_as(world, world.viewer, "viewer") as c:
        r = await c.put(f"{BASE}/projects/{world.project_id}/organisation/mandates", json={"rows": rows})
    assert r.status_code == 403


async def test_a_date_sent_by_the_caller_does_not_revive_an_expired_mandate(world: World) -> None:
    sid = await stakeholder(world, "Residents east", is_group=True)
    await seat(
        world, "client_pm", world.manager, valid_from=TODAY - timedelta(days=60), valid_to=TODAY - timedelta(days=1)
    )
    cid = await commitment(world, sid)
    p = f"{BASE}/projects/{world.project_id}"
    back_then = {"today": (TODAY - timedelta(days=10)).isoformat()}
    async with client_as(world, world.manager, "manager") as c:
        r = await c.post(f"{p}/commitments/{cid}/transition", params=back_then, json={"target": "agreed"})
        swept = await c.post(f"{p}/sweep", params={"today": "2099-01-01"})
        after = await c.get(f"{p}/commitments/{cid}")
    assert (r.status_code, _code(r)) == (403, "not_decider")
    assert swept.status_code == 200
    assert after.json()["status"] == "proposed"


async def test_only_a_manager_moves_the_owner_of_an_agreed_commitment(world: World) -> None:
    sid = await stakeholder(world, "Residents east", is_group=True)
    await seat(world, "client_pm", world.manager)
    cid = await commitment(world, sid)
    p = f"{BASE}/projects/{world.project_id}/commitments/{cid}"
    async with client_as(world, world.manager, "manager") as c:
        agreed = await c.post(f"{p}/transition", json={"target": "agreed"})
    async with client_as(world, world.editor, "editor") as c:
        by_editor = await c.patch(p, json={"owner_user_id": str(world.manager2)})
        frozen = await c.patch(p, json={"text": "Something else"})
    async with client_as(world, world.manager, "manager") as c:
        by_manager = await c.patch(p, json={"owner_user_id": str(world.manager2)})
        again = await c.post(f"{p}/transition", json={"target": "agreed"})
    assert agreed.status_code == 200
    assert (by_editor.status_code, _code(by_editor)) == (403, "owner_change_needs_manage")
    assert (frozen.status_code, _code(frozen)) == (409, "commitment_frozen")
    assert by_manager.status_code == 200
    assert by_manager.json()["owner_user_id"] == str(world.manager2)
    assert again.status_code == 409


# ── input the server must refuse, never crash on ──────────────────────────


async def test_malformed_input_is_422(world: World) -> None:
    sid = await stakeholder(world, "Residents east", is_group=True)
    p = f"{BASE}/projects/{world.project_id}"
    plan = {"title": "Note", "audience_stakeholder_ids": [str(sid)], "participation_level": "inform"}
    async with client_as(world, world.manager, "manager") as c:
        answers = {
            "group size": await c.post(
                f"{p}/stakeholders",
                json={
                    "category": "neighbour",
                    "display_name": "Crowd",
                    "is_group": True,
                    "group_size_estimate": 10**12,
                },
            ),
            "attendance": await c.post(
                f"{p}/events", json={"title": "Evening", "participation_level": "inform", "attendance_count": 10**12}
            ),
            "interval zero": await c.post(
                f"{p}/plan-items", json={**plan, "channel": "newsletter", "frequency": "FREQ=DAILY;INTERVAL=0"}
            ),
            "bad id": await c.post(f"{p}/plan-items", json={**plan, "audience_stakeholder_ids": ["not-an-id"]}),
            "negative amount": await c.post(
                f"{p}/commitments",
                json={"stakeholder_ids": [str(sid)], "text": "Pay", "category": "working_hours", "cost_estimate": "-5"},
            ),
            "role code": await c.put(f"{p}/organisation/roles/{'x' * 80}", json={"holders": []}),
        }
    assert {k: r.status_code for k, r in answers.items()} == dict.fromkeys(answers, 422)


async def test_an_explicit_null_on_a_required_field_is_never_a_500(world: World) -> None:
    sid = await stakeholder(world, "Residents east", is_group=True)
    cid = await commitment(world, sid)
    p = f"{BASE}/projects/{world.project_id}"
    nulls = {
        f"{p}/stakeholders/{sid}": ("display_name", "category", "status", "is_group", "is_confidential", "code"),
        f"{p}/commitments/{cid}": ("text", "category", "owner_user_id", "stakeholder_ids", "currency", "source_kind"),
    }
    seen: dict[str, int] = {}
    async with client_as(world, world.manager, "manager") as c:
        for url, fields in nulls.items():
            for field in fields:
                seen[f"{url.rsplit('/', 2)[1]}.{field}"] = (await c.patch(url, json={field: None})).status_code
        still = await c.get(f"{p}/stakeholders/{sid}")
    assert {k: v for k, v in seen.items() if v not in (200, 422)} == {}
    assert still.json()["name"] == "Residents east"


# ── what the screen may offer ─────────────────────────────────────────────


async def test_each_commitment_tells_the_caller_what_they_may_do_and_why_not(world: World) -> None:
    sid = await stakeholder(world, "Residents east", is_group=True)
    await seat(world, "client_pm", world.manager)
    cid = await commitment(world, sid)
    p = f"{BASE}/projects/{world.project_id}/commitments"
    async with client_as(world, world.manager, "manager") as c:
        decider = (await c.get(f"{p}/{cid}")).json()["can"]
    async with client_as(world, world.editor, "editor") as c:
        editor = (await c.get(p)).json()["items"][0]["can"]
    async with client_as(world, world.viewer, "viewer") as c:
        viewer = (await c.get(f"{p}/{cid}")).json()["can"]
    assert decider["agreed"] == {"applies": True, "allowed": True, "reason": None}
    assert editor["agreed"] == {"applies": True, "allowed": False, "reason": "not_decider"}
    assert decider["verified"]["applies"] is False
    assert editor["edit"]["allowed"] is True
    assert {v["reason"] for v in viewer.values()} == {"no_permission"}


async def test_before_the_first_pin_the_404_still_says_who_may_pin(world: World) -> None:
    from tests.pg.governance_fixtures import add_project

    async with world.factory() as s:
        bare = await add_project(s, world.manager, [world.viewer])
        await s.commit()
    url = f"{BASE}/projects/{bare}/settings"
    async with client_as(world, world.manager, "manager") as c:
        manager = await c.get(url)
    async with client_as(world, world.viewer, "viewer") as c:
        viewer = await c.get(url)
    assert (manager.status_code, _code(manager)) == (404, "settings_missing")
    assert manager.json()["detail"]["can"] == {"edit": True, "read_personal": True, "manage_organisation": True}
    assert viewer.json()["detail"]["can"] == {"edit": False, "read_personal": False, "manage_organisation": False}


async def test_who_decides_names_each_party(world: World) -> None:
    await seat(world, "client_pm", world.manager)
    async with client_as(world, world.viewer, "viewer") as c:
        r = await c.get(
            f"{BASE}/projects/{world.project_id}/organisation/who-decides",
            params={"decision_type": "stakeholder_commitment"},
        )
    body = r.json()
    assert r.status_code == 200
    assert body["user_ids"] == [str(world.manager)]
    assert [p["name"] for p in body["parties"]] == ["Team member"]


async def test_commitment_links_and_their_pickers_stay_inside_the_project(world: World) -> None:
    from app.modules.requirements.models import Requirement, RequirementSet

    other = await build_world(world.factory)
    ids: dict[str, str] = {}
    async with world.factory() as s:
        for name, w in (("mine", world), ("theirs", other)):
            rset = RequirementSet(project_id=w.project_id, name=f"Brief {name}")
            s.add(rset)
            await s.flush()
            req = Requirement(
                requirement_set_id=rset.id, entity="Site access", attribute="hours", constraint_value="07:00-18:00"
            )
            s.add(req)
            await s.flush()
            ids[name] = str(req.id)
        await s.commit()
    sid = await stakeholder(world, "Residents east", is_group=True)
    cid = await commitment(world, sid)
    p = f"{BASE}/projects/{world.project_id}"
    async with client_as(world, world.editor, "editor") as c:
        foreign_requirement = await c.patch(f"{p}/commitments/{cid}", json={"requirement_id": ids["theirs"]})
        foreign_task = await c.patch(f"{p}/commitments/{cid}", json={"task_id": str(uuid.uuid4())})
        own = await c.patch(f"{p}/commitments/{cid}", json={"requirement_id": ids["mine"]})
        options = await c.get(f"{p}/requirements")
        release = await c.get(f"{p}/release-items")
        gates = await c.get(f"{p}/gates")
    async with client_as(world, world.outsider, "manager") as c:
        strangers = [await c.get(f"{p}/{name}") for name in ("requirements", "release-items", "gates")]
    assert (foreign_requirement.status_code, _code(foreign_requirement)) == (404, "requirement_not_found")
    assert (foreign_task.status_code, _code(foreign_task)) == (404, "task_not_found")
    assert (own.status_code, own.json()["requirement_id"]) == (200, ids["mine"])
    assert [o["id"] for o in options.json()["items"]] == [ids["mine"]]
    assert release.json() == {"items": [], "available": False, "reason": "commitment_release_not_installed"}
    # This world is a project run without gates.
    assert (gates.json()["items"], gates.json()["gates_in_use"]) == ([], False)
    assert [r.status_code for r in strangers] == [404, 404, 404]


# ── every stored reference is checked against the project ─────────────────


async def _foreign_rows(world: World) -> dict[str, str]:
    """One row of each linkable kind in a second project."""
    from app.modules.correspondence.models import Correspondence
    from app.modules.risk.models import RiskItem
    from app.modules.tasks.models import Task
    from tests.pg.governance_fixtures import add_container

    other = await build_world(world.factory)
    async with world.factory() as s:
        rows = {
            "task": Task(project_id=other.project_id, task_type="task", title="Their task"),
            "risk": RiskItem(project_id=other.project_id, code="R-001", title="Their risk"),
            "correspondence": Correspondence(
                project_id=other.project_id,
                reference_number="COR-001",
                direction="incoming",
                subject="Their letter",
                correspondence_type="letter",
            ),
        }
        s.add_all(rows.values())
        await s.commit()
        ids = {kind: str(row.id) for kind, row in rows.items()}
        ids["container"] = str(await add_container(s, other.project_id, "their-container"))
        await s.commit()
    ids["stakeholder"] = str(await stakeholder(other, "Their neighbour group", is_group=True))
    return ids


async def test_a_row_of_another_project_is_refused_for_every_reference_kind(world: World) -> None:
    foreign = await _foreign_rows(world)
    sid = await stakeholder(world, "Residents east", is_group=True)
    cid = await commitment(world, sid)
    p = f"{BASE}/projects/{world.project_id}"
    new_commitment = {"stakeholder_ids": [str(sid)], "text": "Promise", "category": "working_hours"}
    new_stakeholder = {"category": "utility", "display_name": "Utility"}
    event = {"title": "Evening", "participation_level": "inform"}
    async with client_as(world, world.editor, "editor") as c:
        answers = {
            "task on create": await c.post(f"{p}/commitments", json={**new_commitment, "task_id": foreign["task"]}),
            "task on patch": await c.patch(f"{p}/commitments/{cid}", json={"task_id": foreign["task"]}),
            "risk on create": await c.post(f"{p}/commitments", json={**new_commitment, "risk_id": foreign["risk"]}),
            "risk on patch": await c.patch(f"{p}/commitments/{cid}", json={"risk_id": foreign["risk"]}),
            "stakeholder risk on create": await c.post(
                f"{p}/stakeholders", json={**new_stakeholder, "risk_id": foreign["risk"]}
            ),
            "stakeholder risk on patch": await c.patch(f"{p}/stakeholders/{sid}", json={"risk_id": foreign["risk"]}),
            "parent on create": await c.post(
                f"{p}/stakeholders", json={**new_stakeholder, "parent_stakeholder_id": foreign["stakeholder"]}
            ),
            "parent on patch": await c.patch(
                f"{p}/stakeholders/{sid}", json={"parent_stakeholder_id": foreign["stakeholder"]}
            ),
            "commitment stakeholder on create": await c.post(
                f"{p}/commitments", json={**new_commitment, "stakeholder_ids": [foreign["stakeholder"]]}
            ),
            "correspondence on create": await c.post(
                f"{p}/events", json={**event, "correspondence_ids": [foreign["correspondence"]]}
            ),
            "container on create": await c.post(
                f"{p}/events", json={**event, "cde_container_ids": [foreign["container"]]}
            ),
        }
        events = await c.get(f"{p}/events")
        listed = await c.get(f"{p}/commitments")
    expected = {
        "task on create": "task_not_found",
        "task on patch": "task_not_found",
        "risk on create": "risk_not_found",
        "risk on patch": "risk_not_found",
        "stakeholder risk on create": "risk_not_found",
        "stakeholder risk on patch": "risk_not_found",
        "parent on create": "stakeholder_not_found",
        "parent on patch": "stakeholder_not_found",
        "commitment stakeholder on create": "stakeholder_not_found",
        "correspondence on create": "correspondence_not_found",
        "container on create": "container_not_found",
    }
    assert {k: (r.status_code, _code(r)) for k, r in answers.items()} == {k: (404, v) for k, v in expected.items()}
    assert listed.json()["total"] == 1
    assert events.json()["total"] == 0


async def test_a_release_item_is_stored_only_when_its_module_confirms_it(world: World) -> None:
    from app.modules.stakeholders import authority

    sid = await stakeholder(world, "Residents east", is_group=True)
    cid = await commitment(world, sid)
    mine, theirs = uuid.uuid4(), uuid.uuid4()
    p = f"{BASE}/projects/{world.project_id}/commitments/{cid}"

    async def checker(session, project_id, item_id):
        return authority.ReleaseItemAnswer(
            has_matrix=True, item_valid=project_id == world.project_id and item_id == mine
        )

    async with client_as(world, world.editor, "editor") as c:
        unchecked = await c.patch(p, json={"release_item_id": str(mine)})
        authority.register_release_item_checker(checker)
        try:
            foreign = await c.patch(p, json={"release_item_id": str(theirs)})
            own = await c.patch(p, json={"release_item_id": str(mine)})
        finally:
            authority.register_release_item_checker(None)
        cleared = await c.patch(p, json={"release_item_id": None})
    assert (unchecked.status_code, _code(unchecked)) == (409, "release_items_unavailable")
    assert (foreign.status_code, _code(foreign)) == (404, "release_item_not_found")
    assert (own.status_code, own.json()["release_item_id"]) == (200, str(mine))
    assert (cleared.status_code, cleared.json()["release_item_id"]) == (200, None)


# ── a project run without gates ───────────────────────────────────────────


async def test_a_project_without_gates_is_walked_to_the_end_and_takes_no_gate_trigger(world: World) -> None:
    """Nothing in this module waits for a gate: a promise goes from proposed to verified by hand."""
    from datetime import UTC, datetime

    p = f"/api/v1/stakeholders/projects/{world.project_id}"
    sid = await stakeholder(world, "Residents east", is_group=True)
    await seat(world, "client_pm", world.manager)
    plan = {"title": "Briefing", "audience_stakeholder_ids": [str(sid)], "participation_level": "inform"}
    # Every trigger only a gate can cause, taken from the service's own list, not written out here.
    gate_bound = [("gate", ""), *(("event", name) for name in GATE_BOUND_EVENT_TRIGGERS)]
    gateless = [name for name in EVENT_TRIGGERS if name not in GATE_BOUND_EVENT_TRIGGERS]
    assert "project_governance.gate.decided" in GATE_BOUND_EVENT_TRIGGERS and gateless
    async with client_as(world, world.editor, "editor") as c:
        gates = await c.get(f"{p}/gates")
        dead = {
            (kind, ref): await c.post(
                f"{p}/plan-items", json={**plan, "channel": "meeting", "trigger_kind": kind, "trigger_ref": ref}
            )
            for kind, ref in gate_bound
        }
        # A listened event that needs no gate still works here.
        free = await c.post(
            f"{p}/plan-items",
            json={**plan, "channel": "meeting", "trigger_kind": "event", "trigger_ref": gateless[0]},
        )
        live = await c.post(
            f"{p}/plan-items",
            json={**plan, "channel": "meeting", "frequency": "FREQ=MONTHLY", "starts_on": TODAY.isoformat()},
        )
        retriggered = {
            (kind, ref): await c.patch(
                f"{p}/plan-items/{live.json().get('id')}", json={"trigger_kind": kind, "trigger_ref": ref}
            )
            for kind, ref in gate_bound
        }
        promise = await c.post(
            f"{p}/commitments",
            json={
                "stakeholder_ids": [str(sid)],
                "text": "No deliveries before seven",
                "category": "information",
                "owner_user_id": str(world.editor),
                "due_date": (TODAY + timedelta(days=10)).isoformat(),
            },
        )
        cid = promise.json().get("id")
    async with client_as(world, world.manager, "manager") as c:
        steps = [
            await c.post(f"{p}/commitments/{cid}/transition", json={"target": target, **extra})
            for target, extra in (
                ("agreed", {}),
                ("in_progress", {}),
                ("fulfilled", {}),
                ("verified", {"communicated_back_at": datetime.now(UTC).isoformat()}),
            )
        ]
        readiness = await c.get(f"{p}/readiness")
        settings = await c.get(f"{p}/settings")
        filing = await c.post(f"{p}/frame-documents/file", json={"document": "register"})
    assert gates.json() == {
        "items": [],
        "gates_in_use": False,
        "gate_bound_event_triggers": list(GATE_BOUND_EVENT_TRIGGERS),
    }
    # No frame documents without gates: the filing action does not exist, and the write agrees.
    gone = {"applies": False, "allowed": False, "reason": "gates_not_in_use"}
    assert settings.json()["can_file"] == {"register": gone, "organisation": gone}
    assert (filing.status_code, _code(filing)) == (409, "frame_document_missing")
    for trigger, response in dead.items():
        assert (response.status_code, _code(response)) == (422, "gates_not_in_use"), trigger
    assert free.status_code == 201, free.text
    assert live.status_code == 201, live.text
    for trigger, response in retriggered.items():
        assert (response.status_code, _code(response)) == (422, "gates_not_in_use"), trigger
    assert promise.status_code == 201, promise.text
    assert [(r.status_code, r.json().get("status")) for r in steps] == [
        (200, "agreed"),
        (200, "in_progress"),
        (200, "fulfilled"),
        (200, "verified"),
    ], [r.text for r in steps]
    assert readiness.status_code == 200
