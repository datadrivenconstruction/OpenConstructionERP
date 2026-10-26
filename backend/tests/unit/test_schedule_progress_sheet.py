"""The progress sheet: build, fill in like a site would, read back (pure, no database)."""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app.modules.schedule.progress_sheet import (
    COLUMNS,
    META_SHEET,
    ExportRow,
    SheetFormatError,
    build_workbook,
    parse_workbook,
    progress_stamp,
)

HEADERS = {key: key.upper() for key, _w, _e in COLUMNS}
COL = {key: i + 1 for i, (key, _w, _e) in enumerate(COLUMNS)}


def _row(key: str, ref: str, pct: float = 0.0) -> ExportRow:
    return ExportRow(
        row_key=key,
        activity_ref=ref,
        name=f"Activity {ref}",
        wbs_code="01.02",
        responsible="Site team",
        planned_start="2026-05-04",
        planned_finish="2026-05-15",
        percent_complete=pct,
        actual_start=None,
        actual_finish=None,
        remaining_duration=None,
        stamp="abc123",
    )


def _sheet(rows: list[ExportRow], **kw) -> bytes:
    return build_workbook(
        rows,
        headers=HEADERS,
        sheet_title="Progress",
        title="Progress update sheet",
        details=["Schedule: Test"],
        instructions="Fill in the white columns only.",
        meta={"schedule_id": "s-1", "exported_at": "2026-05-01T10:00:00+00:00"},
        letterhead=kw.get("letterhead", False),
    )


def _edit(data: bytes, edits: dict[tuple[int, str], object], *, percent_format: bool = False) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    for (row, key), value in edits.items():
        cell = ws.cell(row=row, column=COL[key])
        cell.value = value
        if percent_format and key == "percent_complete":
            cell.number_format = "0%"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_round_trip_reads_what_the_site_typed() -> None:
    data = _sheet([_row("id-1", "A1"), _row("id-2", "A2", 10.0)])
    filled = _edit(
        data,
        {
            (2, "percent_complete"): 45,
            (2, "actual_start"): date(2026, 5, 5),
            (2, "remaining_duration"): 6,
            (2, "notes"): "formwork done =)",
            (3, "actual_finish"): "2026-05-14",
        },
    )
    sheet = parse_workbook(filled)
    assert sheet.meta["schedule_id"] == "s-1"
    assert sheet.meta["columns"] == [key for key, _w, _e in COLUMNS]
    first, second = sheet.rows
    assert first.row_key == "id-1"
    assert first.percent_complete == 45.0
    assert first.actual_start == "2026-05-05"
    assert first.remaining_duration == 6
    assert first.notes == "formwork done =)"
    assert first.stamp == "abc123"
    assert second.percent_complete == 10.0
    assert second.actual_finish == "2026-05-14"
    assert not first.errors and not second.errors


def test_identity_columns_are_locked_and_hidden_columns_hidden() -> None:
    wb = load_workbook(io.BytesIO(_sheet([_row("id-1", "A1")])))
    ws = wb.worksheets[0]
    assert ws.protection.sheet is True
    assert ws.cell(row=2, column=COL["name"]).protection.locked is True
    assert ws.cell(row=2, column=COL["percent_complete"]).protection.locked is False
    letter = ws.cell(row=1, column=COL["row_key"]).column_letter
    assert ws.column_dimensions[letter].hidden is True
    assert wb[META_SHEET].sheet_state == "veryHidden"
    validated = {str(dv.sqref) for dv in ws.data_validations.dataValidation}
    pct_letter = ws.cell(row=1, column=COL["percent_complete"]).column_letter
    assert any(pct_letter in ref for ref in validated)


def test_percent_typed_with_percent_format_is_scaled() -> None:
    filled = _edit(_sheet([_row("id-1", "A1")]), {(2, "percent_complete"): 0.5}, percent_format=True)
    assert parse_workbook(filled).rows[0].percent_complete == 50.0


def test_text_percent_with_sign_and_comma_is_read() -> None:
    filled = _edit(_sheet([_row("id-1", "A1")]), {(2, "percent_complete"): "37,5 %"})
    assert parse_workbook(filled).rows[0].percent_complete == 37.5


def test_bad_values_become_row_errors_not_exceptions() -> None:
    filled = _edit(
        _sheet([_row("id-1", "A1")]),
        {
            (2, "percent_complete"): "half",
            (2, "actual_start"): "next monday",
            (2, "remaining_duration"): 2.5,
        },
    )
    row = parse_workbook(filled).rows[0]
    assert set(row.errors) == {"percent_not_a_number", "actual_start_not_a_date", "remaining_not_a_whole_number"}


def test_letterhead_moves_the_table_and_parse_follows_the_meta() -> None:
    data = _sheet([_row("id-1", "A1")], letterhead=True)
    sheet = parse_workbook(data)
    assert sheet.rows[0].row_key == "id-1"
    assert sheet.rows[0].row_number == int(sheet.meta["header_row"]) + 1


def test_a_workbook_without_the_meta_sheet_is_refused() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["anything"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(SheetFormatError) as exc:
        parse_workbook(buf.getvalue())
    assert exc.value.code == "not_a_progress_sheet"


def test_garbage_bytes_are_refused() -> None:
    with pytest.raises(SheetFormatError) as exc:
        parse_workbook(b"not a workbook")
    assert exc.value.code == "unreadable_file"


def test_progress_stamp_changes_with_progress_only() -> None:
    base = {
        "start_date": "2026-05-04",
        "end_date": "2026-05-15",
        "progress_pct": 10.0,
        "remaining_duration": None,
        "status": "in_progress",
        "actual_start": "2026-05-04",
        "actual_finish": None,
    }
    assert progress_stamp(**base) == progress_stamp(**base)
    assert progress_stamp(**base) != progress_stamp(**{**base, "progress_pct": 20.0})
    assert progress_stamp(**base) != progress_stamp(**{**base, "end_date": "2026-05-18"})
