# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Every platform language has a complete entry in every outgoing-document catalogue.

An email, a workbook header or a notification that a reader in Hindi or Hebrew
gets in English is a defect nobody sees until that one event fires for that one
account. So each catalogue is held to the platform's language list
(:data:`app.core.i18n.SUPPORTED_LOCALES`): every code present, the same keys as
English, the same ``{placeholders}`` per key, no empty value. A language the
platform does not carry reads English, and no request is refused for its language.

Pure: no database, no event bus.
"""

from __future__ import annotations

import inspect
import io
import re
from collections.abc import Mapping

import pytest
from openpyxl import load_workbook

from app.core.document_catalogue import is_rtl
from app.core.i18n import SUPPORTED_LOCALES
from app.modules.bim_requirements import router as bim_router
from app.modules.bim_requirements.exporters import excel_exporter
from app.modules.cost_plan import export as cost_export
from app.modules.module_builder import suggest
from app.modules.notifications import email_render, localized
from app.modules.notifications.templates import all_template_keys, english_template
from app.modules.portal import login_email
from app.modules.schedule import tabular_headers as headers
from app.modules.schedule import tabular_import_service
from app.modules.tendering import invitation_email, xlsx_translations
from tests.unit.test_cost_plan_nrm1_engine import _leaf, _plan

_FIELD = re.compile(r"\{[A-Za-z_][A-Za-z0-9_]*\}")

_TABLES: dict[str, Mapping[str, Mapping[str, str]]] = {
    "portal login email": login_email._TABLES,
    "tender invitation email": invitation_email._TABLES,
    "notification email": email_render._TABLES,
    "tender workbooks": xlsx_translations._STRINGS,
    "cost plan workbook": cost_export._LABELS,
    "BIM requirements workbook": excel_exporter._STRINGS,
}


def _fields(text: str) -> set[str]:
    return set(_FIELD.findall(text))


@pytest.mark.parametrize("name", list(_TABLES))
def test_every_platform_language_has_a_complete_entry(name: str) -> None:
    table = _TABLES[name]
    english = table["en"]
    assert set(SUPPORTED_LOCALES) <= set(table), f"{name}: no entry for {sorted(set(SUPPORTED_LOCALES) - set(table))}"
    for locale in SUPPORTED_LOCALES:
        entry = table[locale]
        assert set(entry) == set(english), f"{name}/{locale}: keys differ from English"
        for key, text in entry.items():
            assert text.strip(), f"{name}/{locale}/{key}: empty"
            assert _fields(text) == _fields(english[key]), f"{name}/{locale}/{key}: placeholders differ"
            assert "<" not in text and ">" not in text, f"{name}/{locale}/{key}: markup in a plain text value"


@pytest.mark.parametrize("locale", SUPPORTED_LOCALES)
def test_worksheet_names_fit_the_workbook_format(locale: str) -> None:
    names = [
        cost_export._LABELS[locale]["sheet_plan"],
        cost_export._LABELS[locale]["sheet_unallocated"],
        xlsx_translations._STRINGS[locale]["comparison_sheet"],
        xlsx_translations._STRINGS[locale]["bidder_sheet"],
    ]
    for name in names:
        assert 0 < len(name) <= 31, (locale, name)
        assert not set(name) & set("[]:*?/\\"), (locale, name)


def test_a_language_the_platform_does_not_carry_reads_english() -> None:
    assert login_email.email_locale("xx") == "en"
    assert login_email.email_locale("ja-JP") == "ja"
    assert invitation_email.email_locale("xx") == "en"
    assert invitation_email.email_locale("he-IL") == "he"
    assert email_render.email_locale("xx") == "en"
    assert xlsx_translations.resolve_xlsx_locale("xx", "yy") == "en"
    assert xlsx_translations.resolve_xlsx_locale(None, "xx-YY, ko;q=0.8") == "ko"
    assert cost_export.label("xx", "title") == cost_export.label("en", "title")
    assert excel_exporter.workbook_language("xx") == "en"
    assert excel_exporter.workbook_language("th-TH") == "th"
    assert localized.catalogue_locale("xx") == "en"
    assert headers.LANGUAGES[0] == "en"
    assert suggest.language_of("xx") == "en"
    assert tabular_import_service.template_language("xx") == "en"


@pytest.mark.parametrize("locale", SUPPORTED_LOCALES)
def test_every_notification_is_written_in_every_language(locale: str) -> None:
    assert localized.catalogue_locale(locale) == locale
    if locale == "en":
        return
    catalogue = localized._CATALOGUES[locale]
    english = set(all_template_keys())
    assert set(catalogue) == english, f"{locale}: {sorted(english ^ set(catalogue))[:5]}"
    for key, text in catalogue.items():
        template = english_template(key)
        assert template is not None
        assert _fields(text) == _fields(template), f"{locale}/{key}: placeholders differ"
        assert text.strip(), f"{locale}/{key}: empty"


@pytest.mark.parametrize("locale", SUPPORTED_LOCALES)
def test_a_notification_renders_in_the_readers_language(locale: str) -> None:
    key = "notifications.rfi.responded.body"
    context = {"code": "RFI-7", "title": "Shaft"}
    text = localized.render(key, context, locale)
    assert "{" not in text and "RFI-7" in text
    if locale != "en":
        assert text != localized.render(key, context, "en") or locale in {"fil", "id"}
    pattern = localized.date_pattern(locale)
    assert all(token in pattern for token in ("DD", "MM", "YYYY"))


@pytest.mark.parametrize("locale", SUPPORTED_LOCALES)
def test_the_schedule_import_knows_every_language(locale: str) -> None:
    assert locale in headers.LANGUAGES
    for field in headers.FIELDS:
        assert headers.HEADER_SYNONYMS[field][locale], (locale, field)
    template = headers.TEMPLATE_HEADERS[locale]
    assert len(template) == len(headers.FIELDS)
    # The header row the template writes reads back as the fields it was written for.
    for field, header in zip(headers.FIELDS, template, strict=True):
        assert headers.match_header(header).field == field, (locale, field, header)
    yes, no = headers.TEMPLATE_YES_NO[locale]
    assert headers.flag_value(yes) is True, (locale, yes)
    assert headers.flag_value(no) is False, (locale, no)
    assert len(headers.TEMPLATE_EXAMPLE_NAMES[locale]) == 4
    rows = tabular_import_service.template_rows(locale)
    assert rows[0] == list(template)
    for rendered in ("csv", "xlsx"):
        content, _media, filename = tabular_import_service.render_template(locale, rendered)  # type: ignore[arg-type]
        assert content and filename.endswith(rendered)


@pytest.mark.parametrize("locale", SUPPORTED_LOCALES)
def test_the_module_builder_suggests_in_every_language(locale: str) -> None:
    assert locale in suggest.LANGUAGES
    for concept, by_language in suggest.KEYWORDS.items():
        assert by_language[locale], (locale, concept)
    for target, labels in suggest.LINK_LABELS.items():
        assert labels[locale].strip(), (locale, target)
    states = suggest.STATE_LABELS[locale]
    assert len(states) == 3 and all(state.strip() for state in states)
    assert [s.label for s in suggest.default_status(locale).states] == list(states)


@pytest.mark.parametrize("locale", SUPPORTED_LOCALES)
def test_a_keyword_names_its_own_concept_only(locale: str) -> None:
    for concept, by_language in suggest.KEYWORDS.items():
        for keyword in by_language[locale]:
            sample = keyword.replace("*", "")
            assert suggest._matches(concept, sample)[0], (locale, concept, keyword)
            for other in suggest.KEYWORDS:
                if other != concept:
                    assert not suggest._matches(other, sample)[0], (locale, keyword, concept, other)


def test_right_to_left_languages_are_marked() -> None:
    assert {code for code in SUPPORTED_LOCALES if is_rtl(code)} == {"ar", "fa", "he", "ur"}
    assert is_rtl("ar-SA") and not is_rtl("en") and not is_rtl(None)


@pytest.mark.parametrize("locale", ["ar", "fa", "he", "ur"])
def test_a_right_to_left_email_sets_its_direction(locale: str) -> None:
    body = login_email.login_html(
        kind="invite", locale=locale, full_name="A", sender="B", action_url="https://x.test/a", valid_hours=12
    )
    assert f'lang="{locale}"' in body and 'dir="rtl"' in body
    assert "12" in body
    assert 'dir="rtl"' not in login_email.login_html(
        kind="invite", locale="de", full_name="A", sender="B", action_url="https://x.test/a", valid_hours=12
    )
    notification = email_render.render_notification_email(
        locale=locale, recipient_name="A", subject="S", body_text="T", action_url=None, base_url="https://x.test"
    )
    assert 'dir="rtl"' in notification


@pytest.mark.parametrize("locale", SUPPORTED_LOCALES)
def test_the_invitation_and_login_emails_render_in_every_language(locale: str) -> None:
    invitation = invitation_email.invitation_html(
        locale=locale,
        company_name="Alpha",
        package_name="Lot 2",
        project_name="Depot",
        description="",
        deadline="2026-11-30",
        custom_message=None,
        action_url="https://x.test/bid",
    )
    assert "Lot 2" in invitation and "Alpha" in invitation and "2026-11-30" in invitation
    assert "{" not in invitation
    login = login_email.login_html(
        kind="login", locale=locale, full_name="", sender="", action_url="https://x.test/a", valid_hours=24
    )
    assert "24" in login and "{" not in login
    assert "Lot 2" in invitation_email.invitation_subject(locale, "Lot 2")


@pytest.mark.parametrize("locale", ["ar", "ja", "xx"])
def test_the_cost_plan_workbook_is_written_in_any_language(locale: str) -> None:
    plan = _plan([_leaf("1000", "2.1"), _leaf("7", None)], gifa="100")
    book = load_workbook(io.BytesIO(cost_export.build_cost_plan_workbook(plan, locale=locale)))
    shown = locale if locale in cost_export.EXPORT_LOCALES else "en"
    assert book.worksheets[0].title == cost_export._LABELS[shown]["sheet_plan"][:31]
    assert bool(book.worksheets[0].sheet_view.rightToLeft) == is_rtl(shown)


@pytest.mark.parametrize("locale", ["en", "de", "ru", "ja", "he", "xx"])
def test_the_bim_requirements_workbook_follows_the_language(locale: str) -> None:
    book = load_workbook(io.BytesIO(excel_exporter.export_excel([], title="Set", language=locale)))
    shown = excel_exporter.workbook_language(locale)
    table = excel_exporter._STRINGS[shown]
    assert [c.value for c in book.worksheets[0][1]] == [table[f"h{n:02d}"] for n in range(1, 16)]
    assert book.sheetnames == ["Requirements", "Summary", "Legend"]  # tabs keep their names for other tools
    assert book["Summary"]["A1"].value == table["sheet_summary"]
    assert bool(book.worksheets[0].sheet_view.rightToLeft) == is_rtl(shown)


def test_the_bim_requirements_export_does_not_refuse_a_language() -> None:
    default = inspect.signature(bim_router.export_excel).parameters["language"].default
    assert not [m for m in default.metadata if getattr(m, "pattern", None)], "language is restricted by a pattern"
