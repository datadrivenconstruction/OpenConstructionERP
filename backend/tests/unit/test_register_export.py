"""Tests for the shared register and record form renderer.

A register is a contractual record that gets attached to a letter, so these
tests pin what makes it one: the header says which project and when, the
footer says "page x of y" with the real total, the header row repeats, and
amounts and dates are written the way the project's market writes them.

They also pin the three rules the module exists to hold in one place:

* capitals follow the language (a Turkish "i" becomes "İ", never "I");
* an amount goes from ``Decimal`` to the page without passing through a
  float, and stays a number in a workbook;
* a catalogue that declares a language without a complete table does not
  import, so "listed as supported, printed in English" cannot happen.

Text assertions go through ``pypdf`` extraction, like the RFI PDF tests.
"""

from __future__ import annotations

import io
import unicodedata
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.pdf_fonts import BODY_FONT, BOLD_FONT, font_can_draw_all, register_pdf_fonts
from app.core.regional_format import ANGLO, CONTINENTAL
from app.core.register_export import (
    FURNITURE,
    DocumentCatalogue,
    ProjectHeader,
    RecordBlock,
    RecordDocument,
    RegisterColumn,
    RegisterDocument,
    build_record_pdf,
    build_register_pdf,
    build_register_xlsx,
    caps,
    export_filename,
    format_amount,
    format_count,
    format_stored_date,
    person_name,
    record_text_lines,
    register_text_rows,
    xlsx_date_format,
)

TURKISH_LETTERS = "ç ğ ı İ ö ş ü Ç Ğ I i Ö Ş Ü"


@pytest.fixture(autouse=True)
def _no_company_profile(monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep a company profile on the machine running the tests out of them."""
    monkeypatch.setattr("app.core.company_profile.read_company_profile", lambda: {})


def _text(pdf: bytes) -> str:
    raw = " ".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(pdf)).pages)
    return " ".join(unicodedata.normalize("NFC", raw).replace(" ", " ").split())


def _register(rows: list[list[Any]], **overrides: Any) -> RegisterDocument:
    params: dict[str, Any] = {
        "title": "Ölçüm Kayıt Listesi",
        "columns": [
            RegisterColumn("Kod"),
            RegisterColumn("Açıklama", weight=3.0, wrap=True),
            RegisterColumn("Tarih", kind="date"),
            RegisterColumn("Tutar", kind="money"),
            RegisterColumn("Gün", kind="count"),
        ],
        "rows": rows,
        "details": [("Proje", "Işıklı Veri Merkezi"), ("Proje no.", "TR-001")],
        "locale": "tr",
        "date_format": "%d.%m.%Y",
        "number_style": CONTINENTAL,
        "page_label": "Sayfa {page} / {total}",
        "generated_label": "Oluşturulma: {timestamp}",
        "generated": "10.10.2026 09:00 UTC",
        "empty_text": "Kayıt yok.",
    }
    params.update(overrides)
    return RegisterDocument(**params)


# ── Fonts ─────────────────────────────────────────────────────────────────


def test_the_bundled_faces_are_registered_and_draw_every_turkish_letter() -> None:
    """The fallback face is WinAnsi Helvetica, which has no ğ, ı, İ or ş.

    A silent fallback is the realistic way a Turkish document breaks, so the
    registration itself is asserted, not only the glyph coverage.
    """
    assert register_pdf_fonts() is True
    assert font_can_draw_all(BODY_FONT, TURKISH_LETTERS)
    assert font_can_draw_all(BOLD_FONT, TURKISH_LETTERS)


# ── Capitals ──────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("label", "expected"),
    [
        ("Geçersiz", "GEÇERSİZ"),
        ("İncelemede", "İNCELEMEDE"),
        ("Arşivlendi", "ARŞİVLENDİ"),
        ("Açık", "AÇIK"),
        ("Kapatıldı", "KAPATILDI"),
        ("Revize edip yeniden sunun", "REVİZE EDİP YENİDEN SUNUN"),
    ],
)
def test_turkish_capitals_keep_the_dot_and_the_dotless_i(label: str, expected: str) -> None:
    assert caps(label, "tr") == expected
    assert caps(label, "tr-TR") == expected


def test_plain_upper_would_have_dropped_the_dot() -> None:
    """The defect ``caps`` exists for, stated as a fact about ``str.upper``."""
    assert "Geçersiz".upper() == "GEÇERSIZ"
    assert caps("Geçersiz", "tr") != "Geçersiz".upper()


def test_other_languages_use_plain_capitals() -> None:
    assert caps("Revise and resubmit", "en") == "REVISE AND RESUBMIT"
    assert caps("Ανοιχτό", "el") == "ΑΝΟΙΧΤΟ"


# ── Value formatting ──────────────────────────────────────────────────────


def test_an_amount_is_written_in_the_markets_separators_with_its_currency() -> None:
    assert format_amount(Decimal("1234.56"), CONTINENTAL, "try") == "1.234,56 TRY"
    assert format_amount(Decimal("1234.56"), ANGLO, "USD") == "1,234.56 USD"
    assert format_amount("1234567.5", CONTINENTAL) == "1.234.567,50"
    assert format_amount(Decimal("-0.004"), CONTINENTAL) == "0,00"


def test_an_amount_never_passes_through_a_float() -> None:
    """2.675 is 2.67499... as a float; a person rounding it writes 2,68."""
    assert format_amount(Decimal("2.675"), CONTINENTAL) == "2,68"
    assert format_amount(Decimal("123456789012.345"), CONTINENTAL) == "123.456.789.012,35"


@pytest.mark.parametrize("missing", [None, "", "  ", "abc", True, Decimal("NaN")])
def test_a_missing_amount_is_a_dash_not_a_zero(missing: Any) -> None:
    assert format_amount(missing, CONTINENTAL, "TRY") == "-"


def test_dates_read_every_stored_shape() -> None:
    assert format_stored_date("2026-10-10", "%d.%m.%Y") == "10.10.2026"
    assert format_stored_date("2026-10-10T08:30:00+00:00", "%d.%m.%Y") == "10.10.2026"
    assert format_stored_date(date(2026, 1, 5), "%d.%m.%Y") == "05.01.2026"
    assert format_stored_date(datetime(2026, 1, 5, 23, 0, tzinfo=UTC), "%d.%m.%Y") == "05.01.2026"
    assert format_stored_date(None, "%d.%m.%Y") == "-"
    assert format_stored_date("sonra", "%d.%m.%Y") == "sonra"
    assert xlsx_date_format("%d.%m.%Y") == "DD.MM.YYYY"
    assert xlsx_date_format("%Y-%m-%d") == "YYYY-MM-DD"


def test_counts_people_and_filenames() -> None:
    assert format_count(3) == "3"
    assert format_count(None) == "-"
    resolved = "8f6203d9-f81a-41f9-acb3-d67bc5d8187c"
    assert person_name(resolved, {resolved: "Çağrı Öztürk"}) == "Çağrı Öztürk"
    assert person_name("0b1c2d3e-4f50-4a6b-8c7d-9e0f1a2b3c4d", {}) == "0b1c2d3e"
    assert person_name("İsmail Şahin", {}) == "İsmail Şahin"
    assert person_name(None, {}) == "-"
    assert export_filename("SUB/012 a", "x", "pdf") == "SUB_012_a.pdf"
    assert export_filename("", "submittal", "pdf") == "submittal.pdf"


def test_a_project_header_labels_itself_and_picks_its_separators() -> None:
    turkish = ProjectHeader(name="Işıklı Veri Merkezi", code="TR-001", currency="TRY", country="TR")
    assert turkish.label == "Işıklı Veri Merkezi (TR-001)"
    assert turkish.number_style == CONTINENTAL
    assert ProjectHeader(name="Depot", currency="USD").number_style == ANGLO
    assert ProjectHeader().label == "-"


# ── Catalogue ─────────────────────────────────────────────────────────────


def test_the_page_furniture_is_complete_in_every_language() -> None:
    english = FURNITURE["en"]
    for locale, table in FURNITURE.items():
        assert set(table) == set(english), locale
        assert all(value.strip() for value in table.values()), locale
        assert "%" not in date(2026, 9, 10).strftime(table["date_format"]), locale


def test_a_language_declared_without_its_strings_does_not_import() -> None:
    with pytest.raises(ValueError, match="tr: strings missing"):
        DocumentCatalogue({"en": {"title": "Log", "col": "Code"}, "tr": {"title": "Liste"}})


def test_a_language_declared_without_its_labels_does_not_import() -> None:
    strings = {"en": {"title": "Log"}, "tr": {"title": "Liste"}}
    with pytest.raises(ValueError, match="label table 'status'"):
        DocumentCatalogue(strings, {"status": {"en": {"open": "Open"}}})
    with pytest.raises(ValueError, match="label table 'status' missing"):
        DocumentCatalogue(strings, {"status": {"en": {"open": "Open", "closed": "Closed"}, "tr": {"open": "Açık"}}})


def test_a_language_without_page_furniture_does_not_import() -> None:
    with pytest.raises(ValueError, match="zz: no page furniture"):
        DocumentCatalogue({"en": {"title": "Log"}, "zz": {"title": "Log"}})


def test_catalogue_lookups_fall_back_and_count_days() -> None:
    catalogue = DocumentCatalogue(
        {"en": {"title": "Log"}, "tr": {"title": "Liste"}},
        {"status": {"en": {"open": "Open"}, "tr": {"open": "Açık"}}},
    )
    assert catalogue.supported == ("en", "tr")
    assert catalogue.normalize("tr-TR") == "tr"
    assert catalogue.normalize("zz") == "en"
    assert catalogue.resolve(None, "zz-ZZ,tr;q=0.8") == "tr"
    assert catalogue.resolve("en", "tr") == "en"
    assert catalogue.tr("tr", "title") == "Liste"
    assert catalogue.tr("tr", "project") == "Proje"
    assert catalogue.label("status", "open", "tr") == "Açık"
    assert catalogue.label("status", "parked", "tr") == "parked"
    assert catalogue.label("status", None, "tr") == "-"
    assert catalogue.days(1, "en") == "1 day"
    assert catalogue.days(3, "en") == "3 days"
    assert catalogue.days(3, "tr") == "3 gün"
    assert catalogue.days(None, "tr") == "-"
    assert catalogue.yes_no(True, "tr") == "Evet"


# ── Register ──────────────────────────────────────────────────────────────


def test_register_cells_are_formatted_by_kind() -> None:
    document = _register([["İ-01", "Şantiye ölçümü", "2026-10-10", Decimal("1234.56"), 7]])
    assert register_text_rows(document) == [
        ["Kod", "Açıklama", "Tarih", "Tutar", "Gün"],
        ["İ-01", "Şantiye ölçümü", "10.10.2026", "1.234,56", "7"],
    ]


def test_a_row_with_the_wrong_number_of_cells_fails_loudly() -> None:
    with pytest.raises(ValueError, match="zip"):
        register_text_rows(_register([["İ-01", "eksik"]]))


def test_register_pdf_carries_header_rows_and_page_x_of_y() -> None:
    rows = [
        [f"K-{index:03d}", "Çığ düşmesi öncesi ölçüm", "2026-10-10", Decimal("1234.56"), index]
        for index in range(1, 120)
    ]
    pdf = build_register_pdf(_register(rows))
    assert pdf.startswith(b"%PDF")
    pages = [unicodedata.normalize("NFC", page.extract_text() or "") for page in PdfReader(io.BytesIO(pdf)).pages]
    total = len(pages)
    assert total >= 3
    for number, page in enumerate(pages, 1):
        flat = " ".join(page.replace(" ", " ").split())
        assert f"Sayfa {number} / {total}" in flat
        # The header row repeats on every sheet.
        assert "Açıklama" in flat
    text = _text(pdf)
    for expected in ("Ölçüm Kayıt Listesi", "Işıklı Veri Merkezi", "TR-001", "10.10.2026", "1.234,56", "K-119"):
        assert expected in text, expected
    assert "Oluşturulma: 10.10.2026 09:00 UTC" in text


def test_an_empty_register_says_so() -> None:
    assert "Kayıt yok." in _text(build_register_pdf(_register([])))


def test_user_text_is_printed_not_parsed_as_markup() -> None:
    text = _text(build_register_pdf(_register([["K-1", "<b>kalın</b> & <font color=red>x</font>", None, None, None]])))
    assert "<b>kalın</b> & <font color=red>x</font>" in text


def test_register_workbook_keeps_dates_and_amounts_as_values() -> None:
    document = _register(
        [
            ["İ-01", "Şantiye ölçümü", "2026-10-10", Decimal("1234.56"), 7],
            ["=cmd|'/c calc'!A0", "-", None, None, None],
        ],
        sheet_title="Çok uzun bir sayfa adı: otuz bir karakteri aşar",
    )
    ws = load_workbook(build_register_xlsx(document)).active
    assert len(ws.title) <= 31
    assert ":" not in ws.title
    cells = {cell.value: cell for row in ws.iter_rows() for cell in row if cell.value is not None}
    assert "Ölçüm Kayıt Listesi" in cells
    assert "Proje: Işıklı Veri Merkezi" in cells
    header_row = cells["Kod"].row
    assert [cell.value for cell in ws[header_row]] == ["Kod", "Açıklama", "Tarih", "Tutar", "Gün"]
    first = ws[header_row + 1]
    assert first[0].value == "İ-01"
    assert first[2].value.date() == date(2026, 10, 10)
    assert first[2].number_format == "DD.MM.YYYY"
    assert first[3].value == pytest.approx(1234.56)
    assert first[3].number_format == "#,##0.00"
    assert first[4].value == 7
    second = ws[header_row + 2]
    # A formula payload is neutralised, and a dash is an empty cell.
    assert second[0].value == "'=cmd|'/c calc'!A0"
    assert second[1].value is None
    assert ws.page_setup.orientation == "landscape"


# ── Record form ───────────────────────────────────────────────────────────


def _record(**overrides: Any) -> RecordDocument:
    params: dict[str, Any] = {
        "title": "Onay Belgesi",
        "number": "SUB-012",
        "project_label": "Işıklı Veri Merkezi (TR-001)",
        "status_text": caps("İncelemede", "tr"),
        "subject": "Soğutma grubu imalat çizimi",
        "grid": [["Proje", "Işıklı Veri Merkezi (TR-001)", "Tür", "İmalat çizimi"]],
        "blocks": [
            RecordBlock("İnceleyen yorumları", text="Kolon aksları ile çakışıyor.\nİkinci satır."),
            RecordBlock("Notlar", text="", empty_text="Not kaydedilmedi."),
            RecordBlock(
                "Kalemler",
                kind="table",
                rows=[["Açıklama", "Tutar"], ["Çelik boru", "1.234,56"]],
                weights=[3, 1],
                right_aligned=[False, True],
            ),
            RecordBlock("İmzalar", kind="signatures", rows=[["", "Ad soyad", "İmza", "Tarih"], ["Sunan", "", "", ""]]),
        ],
        "page_label": "Sayfa {page} / {total}",
        "generated_label": "Oluşturulma: {timestamp}",
        "generated": "10.10.2026 09:00 UTC",
    }
    params.update(overrides)
    return RecordDocument(**params)


def test_record_form_prints_every_section() -> None:
    document = _record()
    pdf = build_record_pdf(document)
    assert pdf.startswith(b"%PDF")
    assert len(PdfReader(io.BytesIO(pdf)).pages) == 1
    text = _text(pdf)
    for expected in (
        "Onay Belgesi",
        "SUB-012 · Işıklı Veri Merkezi (TR-001)",
        "İNCELEMEDE",
        "Soğutma grubu imalat çizimi",
        "İmalat çizimi",
        "Kolon aksları ile çakışıyor.",
        "İkinci satır.",
        "Not kaydedilmedi.",
        "Çelik boru",
        "1.234,56",
        "İmzalar",
        "Sayfa 1 / 1",
    ):
        assert expected in text, expected
    assert "İnceleyen yorumları" in record_text_lines(document)


def test_a_long_record_runs_on_and_counts_its_pages() -> None:
    pdf = build_record_pdf(_record(blocks=[RecordBlock("Açıklama", text="Uzun bir açıklama satırı. " * 900)]))
    pages = len(PdfReader(io.BytesIO(pdf)).pages)
    assert pages >= 3
    assert f"Sayfa {pages} / {pages}" in _text(pdf)
