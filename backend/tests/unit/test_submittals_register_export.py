# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The printed submittal register carries what procurement runs on.

Follows ``test_pilot_documents_render_in_turkish.py``: the documents are
rendered in Turkish and in English from rows whose text carries every letter
Turkish adds, and read back through ``pypdf`` and ``openpyxl``.

What is held here beyond that file:

* the workbook has one column per fact and the sheet a set that fits it, and
  both print the same values;
* the landscape sheet is readable: no heading word and no date is wider than
  the column it is printed in, measured in the face the table is drawn with;
* a register cut to one type is named as that type's register;
* a row written before these columns existed renders and exports, and its new
  cells are dashes.
"""

from __future__ import annotations

import io
import unicodedata
from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth

from app.core.pdf_fonts import BODY_FONT, register_pdf_fonts
from app.core.register_export import (
    ProjectHeader,
    build_record_pdf,
    build_register_pdf,
    build_register_xlsx,
    record_text_lines,
    register_layout,
    register_text_rows,
)
from app.modules.submittals import pdf_export
from app.modules.submittals.pdf_translations import (
    CATALOGUE,
    register_title,
    submittal_register_filename,
    tr,
)
from app.modules.submittals.tracking import DISCIPLINES

PROJECT = ProjectHeader(name="Işıklı Veri Merkezi", code="TR-001", currency="TRY", country="TR")
GENERATED = "10.10.2026 09:00 UTC"
AS_OF = date(2026, 10, 10)
REVIEWER = "8f6203d9-f81a-41f9-acb3-d67bc5d8187c"
SUPPLIER = "5b1f6a53-51a4-4f0e-9f0b-3d1c2a7e9a10"
DRAWING = "0c7a4d8e-52f3-4a55-8f6e-2b9f1d3c4e5f"
GONE_DRAWING = "9d2b7c1a-3e4f-4b6a-8c7d-1e2f3a4b5c6d"
PEOPLE = {REVIEWER: "Çağrı Öztürk", SUPPLIER: "Şahin İklimlendirme Ltd."}
DRAWINGS = {DRAWING: "M-201 Soğutma grubu yerleşimi"}


@pytest.fixture(autouse=True)
def _no_company_profile(monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep a company profile on the machine running the tests out of them."""
    monkeypatch.setattr("app.core.company_profile.read_company_profile", lambda: {})


def _chiller(**overrides: object) -> SimpleNamespace:
    """A long-lead material submittal, returned once and now back with the reviewer."""
    base: dict[str, object] = {
        "submittal_number": "SUB-012",
        "title": "Soğutma grubu ürün verisi, İç ünite şartnamesi",
        "submittal_type": "product_data",
        "spec_section": "23 64 00",
        "current_revision": 2,
        "status": "under_review",
        "date_submitted": "2026-10-01",
        "date_required": "2026-10-10",
        "date_returned": None,
        "reviewer_id": REVIEWER,
        "approver_id": None,
        "ball_in_court": REVIEWER,
        "submitted_by_org": "Yüklenici A.Ş.",
        "metadata_": {"review_notes": "Fan eğrileri ığüşöç ve IĞÜŞÖÇİ ile eklenmeli."},
        "discipline": "hvac",
        "manufacturer": "Soğutma Sanayi",
        "model_reference": "ÇG-450",
        "country_of_origin": "TR",
        "supplier": SUPPLIER,
        "review_outcome": None,
        "review_code": None,
        "review_period_days": 7,
        "review_history": [
            {
                "revision": 1,
                "outcome": "revise_and_resubmit",
                "code": "C",
                "date_submitted": "2026-09-10",
                "date_returned": "2026-09-24",
                "reviewer_id": REVIEWER,
                "notes": None,
                "resubmit_for_record": False,
            }
        ],
        "required_on_site_date": "2026-12-14",
        "long_lead": True,
        "lead_time_weeks": 12,
        "linked_drawing_ids": [DRAWING, GONE_DRAWING],
    }
    base.update(overrides)
    return SimpleNamespace(**base)


def _approved_drawing() -> SimpleNamespace:
    return _chiller(
        submittal_number="SUB-013",
        submittal_type="shop_drawing",
        discipline="electrical",
        status="approved_as_noted",
        review_outcome="approved_as_noted",
        review_code="2",
        date_returned="2026-10-08",
        manufacturer=None,
        model_reference=None,
        country_of_origin=None,
        supplier="Elektrik Pano İmalat",
        long_lead=False,
        lead_time_weeks=None,
        required_on_site_date=None,
        review_history=[],
        linked_drawing_ids=[],
    )


def _legacy() -> SimpleNamespace:
    """A row as an installation from before the register columns holds it."""
    return SimpleNamespace(
        submittal_number="SUB-001",
        title="Çelik konstrüksiyon imalat çizimi",
        submittal_type="shop_drawing",
        spec_section=None,
        current_revision=1,
        status="approved",
        date_submitted="2026-08-01",
        date_required="2026-08-20",
        date_returned="2026-08-15",
        reviewer_id=REVIEWER,
        approver_id=None,
        ball_in_court=None,
        submitted_by_org=None,
        metadata_={},
    )


def _context(locale: str, **extra: object) -> dict:
    return {
        "project": PROJECT,
        "locale": locale,
        "generated": GENERATED,
        "people": PEOPLE,
        "drawings": DRAWINGS,
        "as_of": AS_OF,
        **extra,
    }


def _flat(text: str) -> str:
    return " ".join(unicodedata.normalize("NFC", text).split())


def _pdf_text(pdf: bytes) -> str:
    return _flat(" ".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages))


def _column(document, key: str, locale: str) -> list[str]:
    rows = register_text_rows(document)
    index = rows[0].index(tr(locale, key))
    return [row[index] for row in rows[1:]]


def _sheet(document) -> dict[str, list]:
    sheet = load_workbook(build_register_xlsx(document)).active
    grid = list(sheet.iter_rows(values_only=True))
    header_at = next(i for i, row in enumerate(grid) if row[0] == document.columns[0].header)
    header = grid[header_at]
    return {str(name): [row[i] for row in grid[header_at + 1 :]] for i, name in enumerate(header) if name}


# ── The workbook: one column per fact ─────────────────────────────────────


def test_the_turkish_workbook_has_the_procurement_columns_with_turkish_headings() -> None:
    document = pdf_export.submittal_register([_chiller(), _approved_drawing()], **_context("tr", layout="full"))
    headings = [column.header for column in document.columns]
    for expected in (
        "Disiplin",
        "Üretici / Marka",
        "Model / Referans",
        "Menşei",
        "Tedarikçi",
        "Onay Kodu",
        "İnceleme sonucu",
        "İncelemede geçen gün",
        "Geciken gün",
        "Şantiyede Gerekli Tarih",
        "Uzun temin süreli",
        "Temin Süresi (hafta)",
        "Onay için son tarih",
        "Onayda geciken gün",
        "İlgili çizimler",
    ):
        assert expected in headings, expected
    sheet = _sheet(document)
    assert sheet["Disiplin"] == ["İklimlendirme", "Elektrik"]
    assert sheet["Üretici / Marka"][0] == "Soğutma Sanayi"
    assert sheet["Model / Referans"][0] == "ÇG-450"
    assert sheet["Menşei"][0] == "TR"
    # A contact id reads as the contact, a typed name as typed.
    assert sheet["Tedarikçi"] == ["Şahin İklimlendirme Ltd.", "Elektrik Pano İmalat"]
    # The mark is printed as the reviewer wrote it.
    assert sheet["Onay Kodu"][1] == "2"
    assert sheet["İnceleme sonucu"][1] == "Notlarla onaylandı"
    assert sheet["Uzun temin süreli"][0] == "Evet"
    assert sheet["İlgili çizimler"][0] == "M-201 Soğutma grubu yerleşimi"


def test_the_workbook_keeps_counts_as_numbers_and_dates_as_dates() -> None:
    document = pdf_export.submittal_register([_chiller()], **_context("tr", layout="full"))
    sheet = _sheet(document)
    # Submitted 01.10, read on 10.10: nine days, two past a seven day period.
    assert sheet["İncelemede geçen gün"] == [9]
    assert sheet["Geciken gün"] == [2]
    assert sheet["Temin Süresi (hafta)"] == [12]
    # Required on site 14.12 less twelve weeks, nineteen days ago.
    assert sheet["Onay için son tarih"][0].date() == date(2026, 9, 21)
    assert sheet["Onayda geciken gün"] == [19]
    assert sheet["Şantiyede Gerekli Tarih"][0].date() == date(2026, 12, 14)


def test_the_english_workbook_has_the_same_columns_in_english() -> None:
    turkish = pdf_export.submittal_register([_chiller()], **_context("tr", layout="full"))
    english = pdf_export.submittal_register([_chiller()], **_context("en", layout="full"))
    assert len(english.columns) == len(turkish.columns)
    headings = [column.header for column in english.columns]
    for expected in (
        "Discipline",
        "Manufacturer / brand",
        "Country of origin",
        "Supplier",
        "Review code",
        "Required on site",
        "Lead time (weeks)",
        "Approval needed by",
    ):
        assert expected in headings, expected
    assert _sheet(english)["Discipline"] == ["HVAC"]
    assert not set(headings) & {column.header for column in turkish.columns} - {"Rev."}


def test_the_pdf_and_workbook_builders_pick_their_own_layout() -> None:
    workbook = load_workbook(pdf_export.build_submittal_register_xlsx([_chiller()], **_context("tr"))).active
    cells = {cell for row in workbook.iter_rows(values_only=True) for cell in row if isinstance(cell, str)}
    assert "Tedarikçi" in cells
    text = _pdf_text(pdf_export.build_submittal_register_pdf([_chiller()], **_context("tr")))
    assert "Tedarikçi" not in text
    assert "Disiplin" in text


# ── The sheet: fewer columns, same facts ──────────────────────────────────


def test_the_printed_turkish_register_shows_the_procurement_facts() -> None:
    document = pdf_export.submittal_register([_chiller(), _approved_drawing()], **_context("tr"))
    assert _column(document, "col_discipline", "tr") == ["İklimlendirme", "Elektrik"]
    assert _column(document, "col_product", "tr") == ["Soğutma Sanayi, ÇG-450, TR", "-"]
    assert _column(document, "col_review_code", "tr") == ["-", "2"]
    assert _column(document, "col_submitted_returned", "tr") == ["01.10.2026 -", "01.10.2026 08.10.2026"]
    assert _column(document, "col_days_in_review_print", "tr") == ["9", "7"]
    assert _column(document, "col_on_site_needed_by", "tr") == ["14.12.2026 21.09.2026", "-"]
    assert _column(document, "col_lead_time", "tr") == ["12", "-"]
    text = _pdf_text(build_register_pdf(document))
    for expected in ("Onay Kodu", "Disiplin", "Üretici / Model / Menşei", "Temin Süresi (hafta)", "14.12.2026"):
        assert expected in text, expected
    assert "2026-12-14" not in text
    for letter in "çğıİöşü":
        assert letter in text


def test_the_printed_english_register_shows_the_same_rows() -> None:
    document = pdf_export.submittal_register([_chiller(), _approved_drawing()], **_context("en"))
    assert _column(document, "col_discipline", "en") == ["HVAC", "Electrical"]
    assert _column(document, "col_review_code", "en") == ["-", "2"]
    text = _pdf_text(build_register_pdf(document))
    for expected in ("Review code", "Discipline", "Lead time (weeks)", "14.12.2026"):
        assert expected in text, expected
    assert "Disiplin" not in text


def test_the_landscape_sheet_cuts_no_number_date_or_code() -> None:
    """Thirteen columns only help if none of them breaks a number or a date in two.

    The widths are the shared renderer's own (``register_layout``): an empty
    ``overflow`` means every keep-together cell and every heading word fits
    the column it is printed in, and the type size says the sheet did not have
    to be shrunk to get there. Every discipline is printed, not only two.
    """
    assert register_pdf_fonts() is True
    rows = [
        _chiller(discipline=item.code, submittal_number=f"SUB-{index:04d}") for index, item in enumerate(DISCIPLINES)
    ]
    rows.append(_approved_drawing())
    for locale in CATALOGUE.supported:
        document = pdf_export.submittal_register(rows, **_context(locale))
        assert len(document.columns) == 13
        layout = register_layout(document)
        assert layout.overflow == (), (locale, [document.columns[i].header for i in layout.overflow])
        assert layout.cut_words == ()
        assert layout.font_size == 8.0, locale
        by_key = {column.header: column for column in document.columns}
        for key in ("col_number", "col_review_code", "col_rev", "col_lead_time"):
            assert by_key[tr(locale, key)].keeps_together, key
        # The spec section keeps its number on one line (no breakable space
        # inside it) while a section name typed after it may wrap.
        spec = document.columns.index(by_key[tr(locale, "col_spec")])
        assert {row[spec] for row in register_text_rows(document)[1:]} == {"23\u00a064\u00a000"}
        named = pdf_export.submittal_register(
            [_chiller(spec_section="26 05 19 - Low voltage cables")], **_context(locale)
        )
        assert register_text_rows(named)[1][spec] == "26\u00a005\u00a019 - Low voltage cables"
        # A cell holding two dates may put one under the other, never cut one.
        pair = document.columns.index(by_key[tr(locale, "col_submitted_returned")])
        widest_date = max(
            stringWidth(word, BODY_FONT, layout.font_size)
            for row in register_text_rows(document)[1:]
            for word in row[pair].split()
        )
        assert widest_date <= layout.widths[pair] - 2 * 1.6 * mm


# ── Cut by type: the shop drawing register ────────────────────────────────


def test_a_register_cut_to_shop_drawings_is_the_shop_drawing_register() -> None:
    filters = {"submittal_type": "shop_drawing"}
    turkish = pdf_export.submittal_register([_approved_drawing()], **_context("tr", filters=filters))
    english = pdf_export.submittal_register([_approved_drawing()], **_context("en", filters=filters))
    assert turkish.title == "İmalat Çizimleri Kayıt Listesi"
    assert english.title == "Shop Drawing Register"
    assert ("Tür", "İmalat çizimi") in turkish.details
    assert _flat(turkish.title) in _pdf_text(build_register_pdf(turkish))
    assert submittal_register_filename("tr", "pdf", "shop_drawing") == "imalat-cizimleri-kayit-listesi.pdf"
    assert submittal_register_filename("en", "xlsx", "shop_drawing") == "shop-drawing-register.xlsx"


def test_a_material_register_and_an_unnamed_cut_keep_the_right_title() -> None:
    assert register_title("tr", "product_data") == "Malzeme Onay Kayıt Listesi"
    assert register_title("en", "product_data") == "Material Approval Register"
    # A type with no register name of its own keeps the general title.
    assert register_title("tr", "warranty") == "Onay Belgeleri Kayıt Listesi"
    assert register_title("tr") == "Onay Belgeleri Kayıt Listesi"
    assert submittal_register_filename("en", "xlsx") == "submittal-register.xlsx"


def test_every_filter_used_is_named_in_the_header() -> None:
    filters = {
        "submittal_type": None,
        "discipline": "fire_protection",
        "review_outcome": "revise_and_resubmit",
        "long_lead": True,
        "review_overdue": None,
        "approval_late": False,
    }
    document = pdf_export.submittal_register([], **_context("tr", filters=filters))
    details = dict(document.details)
    assert details["Disiplin"] == "Yangın tesisatı"
    assert details["İnceleme sonucu"] == "Revize edip yeniden sunun"
    assert details["Uzun temin süreli"] == "Evet"
    assert details["Onayda geciken gün"] == "Hayır"
    assert "Geciken gün" not in details
    assert "fire_protection" not in " ".join(f"{label} {value}" for label, value in document.details)


# ── The form of one submittal ─────────────────────────────────────────────


def test_the_turkish_form_carries_the_product_the_stamp_and_the_history() -> None:
    document = pdf_export.submittal_record(_chiller(), **_context("tr"))
    lines = " | ".join(record_text_lines(document))
    for expected in (
        "Disiplin",
        "İklimlendirme",
        "Üretici / Marka",
        "Soğutma Sanayi",
        "Menşei",
        "Tedarikçi",
        "Şahin İklimlendirme Ltd.",
        "Şantiyede Gerekli Tarih",
        "14.12.2026",
        "Onay için son tarih",
        "21.09.2026",
        "İnceleme geçmişi",
        "24.09.2026",
        "Revize edip yeniden sunun",
        "M-201 Soğutma grubu yerleşimi",
        "1 ilgili çizim artık mevcut değil",
    ):
        assert expected in lines, expected
    text = _pdf_text(build_record_pdf(document))
    assert "İnceleme geçmişi" in text
    assert "Onay Kodu" in text
    assert "2026-" not in text


def test_the_english_form_carries_the_same_facts() -> None:
    document = pdf_export.submittal_record(_chiller(), **_context("en"))
    lines = " | ".join(record_text_lines(document))
    for expected in ("Discipline", "HVAC", "Country of origin", "Review history", "Revise and resubmit", "21.09.2026"):
        assert expected in lines, expected
    assert "Disiplin" not in lines
    assert build_record_pdf(document).startswith(b"%PDF")


def test_an_approval_as_noted_that_owes_a_record_copy_says_so_on_the_form() -> None:
    owed = _approved_drawing()
    owed.review_history = [{"revision": 2, "outcome": "approved_as_noted", "code": "2", "resubmit_for_record": True}]
    lines = " | ".join(record_text_lines(pdf_export.submittal_record(owed, **_context("tr"))))
    assert "Düzeltilmiş nüsha kayıt için yeniden sunulacak" in lines


# ── An existing installation's rows ───────────────────────────────────────


@pytest.mark.parametrize("locale", ["tr", "en"])
@pytest.mark.parametrize("layout", ["print", "full"])
def test_a_row_from_before_the_register_columns_renders_and_exports(locale: str, layout: str) -> None:
    """Every new column absent, as on an installation that has just been upgraded."""
    document = pdf_export.submittal_register([_legacy()], **_context(locale, layout=layout, drawings=None))
    rows = register_text_rows(document)
    by_heading = dict(zip(rows[0], rows[1], strict=True))
    assert by_heading[tr(locale, "col_discipline")] == "-"
    assert by_heading[tr(locale, "col_lead_time")] == "-"
    # An approved row from before the stamp was stored still shows its decision.
    assert by_heading[tr(locale, "col_review_code")] == "A"
    days_key = "col_days_in_review_print" if layout == "print" else "col_days_in_review"
    assert by_heading[tr(locale, days_key)] == "14"
    assert build_register_pdf(document).startswith(b"%PDF")
    assert load_workbook(build_register_xlsx(document)).active.max_row >= 2


@pytest.mark.parametrize("locale", ["tr", "en"])
def test_the_form_of_a_row_from_before_the_register_columns_renders(locale: str) -> None:
    document = pdf_export.submittal_record(_legacy(), **_context(locale, drawings=None))
    lines = " | ".join(record_text_lines(document))
    assert tr(locale, "col_manufacturer") in lines
    assert tr(locale, "review_history") not in lines
    assert tr(locale, "col_drawings") not in lines
    assert build_record_pdf(document).startswith(b"%PDF")


def test_null_columns_as_the_database_returns_them_render_too() -> None:
    """After the boot heal the columns exist and hold NULL, not "missing"."""
    healed = _legacy()
    for name in (
        "discipline",
        "manufacturer",
        "model_reference",
        "country_of_origin",
        "supplier",
        "review_outcome",
        "review_code",
        "review_period_days",
        "required_on_site_date",
        "lead_time_weeks",
    ):
        setattr(healed, name, None)
    healed.long_lead = False
    healed.review_history = []
    healed.linked_drawing_ids = []
    for layout in ("print", "full"):
        document = pdf_export.submittal_register([healed], **_context("tr", layout=layout))
        assert build_register_pdf(document).startswith(b"%PDF")
        assert build_register_xlsx(document).getbuffer().nbytes > 0
    assert build_record_pdf(pdf_export.submittal_record(healed, **_context("tr"))).startswith(b"%PDF")
