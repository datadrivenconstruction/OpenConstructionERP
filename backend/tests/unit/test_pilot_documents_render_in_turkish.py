"""The site documents a Turkish project prints come out in Turkish, whole.

One case per document family: submittals, correspondence, change orders,
transmittals, variations with their claims and notices, the evidence pack and
the RFI log. Every case is rendered in Turkish and in English from rows whose
text carries every letter Turkish adds to the Latin alphabet, on a project in
Türkiye priced in lira, and held to four things:

* it renders, as a PDF and (for a register) as a workbook;
* the Turkish labels and the row text arrive intact, diacritics included;
* nothing from the English table shows in the Turkish document. The check is
  computed from the catalogue, so a label added in English only fails here
  without anyone having to remember to list it;
* dates and amounts are written the Turkish way: 10.10.2026 and 1.234,56.

The rows cycle through every code of every label table, so each status, type
and reason is printed at least once in each language.

Text assertions go through ``pypdf`` extraction, like the RFI PDF tests.
"""

from __future__ import annotations

import io
import re
import unicodedata
import uuid
from collections.abc import Callable
from dataclasses import dataclass, field
from decimal import Decimal
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.pdf_fonts import register_pdf_fonts
from app.core.register_export import (
    DocumentCatalogue,
    ProjectHeader,
    RecordDocument,
    RegisterDocument,
    build_record_pdf,
    build_register_pdf,
    build_register_xlsx,
    record_text_lines,
    register_text_rows,
)
from app.modules.changeorders import pdf_export as changeorders_export
from app.modules.changeorders.pdf_translations import CATALOGUE as CHANGEORDERS
from app.modules.claims_evidence import pdf_export as evidence_export
from app.modules.claims_evidence.pdf_translations import CATALOGUE as EVIDENCE
from app.modules.correspondence import pdf_export as correspondence_export
from app.modules.correspondence.pdf_translations import CATALOGUE as CORRESPONDENCE
from app.modules.rfi import register_export as rfi_export
from app.modules.rfi.intl import localize_status
from app.modules.submittals import pdf_export as submittals_export
from app.modules.submittals.pdf_translations import CATALOGUE as SUBMITTALS
from app.modules.transmittals import pdf_export as transmittals_export
from app.modules.transmittals.pdf_translations import CATALOGUE as TRANSMITTALS
from app.modules.variations import pdf_export as variations_export
from app.modules.variations.pdf_translations import CATALOGUE as VARIATIONS

PROJECT = ProjectHeader(name="Işıklı Veri Merkezi", code="TR-001", currency="TRY", country="TR")
GENERATED = "10.10.2026 09:00 UTC"

# Row text with every Turkish letter, upper and lower case.
TITLE = "Soğutma grubu ölçüm çizimi, İç tesisat şaftı"
BODY = "Çatı katındaki ölçümler ığüşöç ve IĞÜŞÖÇİ harfleriyle yazıldı."
PERSON_ID = "8f6203d9-f81a-41f9-acb3-d67bc5d8187c"
PEOPLE = {PERSON_ID: "Çağrı Öztürk", "İşveren Ltd.": "İşveren Ltd."}
AMOUNT = Decimal("1234.56")


@pytest.fixture(autouse=True)
def _no_company_profile(monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep a company profile on the machine running the tests out of them."""
    monkeypatch.setattr("app.core.company_profile.read_company_profile", lambda: {})


def _codes(catalogue: DocumentCatalogue, table: str, index: int) -> str:
    codes = list(catalogue.labels[table]["en"])
    return codes[index % len(codes)]


def _row_count(catalogue: DocumentCatalogue) -> int:
    return max([len(table["en"]) for table in catalogue.labels.values()] or [1]) + 1


# ── Rows ──────────────────────────────────────────────────────────────────


def _submittals() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            submittal_number=f"SUB-{index:03d}",
            title=TITLE,
            submittal_type=_codes(SUBMITTALS, "type", index),
            spec_section="23 64 00",
            current_revision=2,
            status=_codes(SUBMITTALS, "status", index),
            date_submitted="2026-10-01",
            date_required="2026-10-10",
            date_returned=None,
            reviewer_id=PERSON_ID,
            approver_id=None,
            ball_in_court=PERSON_ID,
            submitted_by_org="İşveren Ltd.",
            metadata_={"review_notes": BODY},
        )
        for index in range(_row_count(SUBMITTALS))
    ]


def _correspondence() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            reference_number=f"YZ-{index:03d}",
            subject=TITLE,
            direction=_codes(CORRESPONDENCE, "direction", index),
            correspondence_type=_codes(CORRESPONDENCE, "type", index),
            status=_codes(CORRESPONDENCE, "status", index),
            from_contact_id=PERSON_ID,
            to_contact_ids=["İşveren Ltd."],
            date_sent="2026-10-10",
            date_received=None,
            response_required_by="2026-10-17",
            contract_clause_ref="Madde 20.1",
            linked_document_ids=[],
            attachments=[],
            notes=BODY,
        )
        for index in range(_row_count(CORRESPONDENCE))
    ]


def _change_orders() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            code=f"CO-{index:03d}",
            title=TITLE,
            description=BODY,
            reason_category=_codes(CHANGEORDERS, "reason", index),
            variation_type=_codes(CHANGEORDERS, "instrument", index),
            status=_codes(CHANGEORDERS, "status", index),
            submitted_at="2026-10-10T08:00:00+00:00",
            approved_at=None,
            submitted_by=PERSON_ID,
            approved_by=None,
            contractor_amount=AMOUNT,
            engineer_amount=Decimal("1000"),
            approved_amount=None,
            cost_impact=AMOUNT,
            currency="TRY",
            time_impact_days=3,
            schedule_impact_days=3,
            approved_time_days=None,
            items=[
                SimpleNamespace(
                    sort_order=1,
                    description="Çelik boru DN100, ığüşöç",
                    change_type=_codes(CHANGEORDERS, "change", index),
                    unit="m",
                    original_quantity=Decimal("10"),
                    new_quantity=Decimal("12.5"),
                    original_rate=Decimal("100"),
                    new_rate=Decimal("98.765"),
                    cost_delta=AMOUNT,
                )
            ],
        )
        for index in range(_row_count(CHANGEORDERS))
    ]


def _transmittals() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            transmittal_number=f"TR-{index:03d}",
            subject=TITLE,
            purpose_code=_codes(TRANSMITTALS, "purpose", index),
            status=_codes(TRANSMITTALS, "status", index),
            issued_date="2026-10-10",
            response_due_date="2026-10-17",
            sender_org_id="İşveren Ltd.",
            cover_note=BODY,
            recipients=[
                SimpleNamespace(
                    recipient_name="Şükrü Iğdır",
                    recipient_user_id=None,
                    recipient_org_id=None,
                    recipient_email=None,
                    action_required="İnceleyip yanıtlayın",
                    acknowledged_at="2026-10-11",
                    responded_at=None,
                )
            ],
            items=[SimpleNamespace(item_number=1, description="Şaft kesiti çizimi", notes="Ölçek 1:50")],
        )
        for index in range(_row_count(TRANSMITTALS))
    ]


def _variation_requests() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            code=f"VR-{index:03d}",
            title=TITLE,
            description=BODY,
            classification=_codes(VARIATIONS, "classification", index),
            urgency=_codes(VARIATIONS, "urgency", index),
            status=_codes(VARIATIONS, "request_status", index),
            requested_by=PERSON_ID,
            requested_at="2026-10-10",
            submitted_at="2026-10-10",
            decision_at=None,
            decided_by=None,
            decision_notes="Keşif özeti bekleniyor, ığüşöç.",
            estimated_cost_impact=AMOUNT,
            agreed_cost_impact=None,
            currency="TRY",
            estimated_schedule_days=5,
            response_due_date="2026-10-17",
            contract_clause_ref="Madde 13.3",
        )
        for index in range(_row_count(VARIATIONS))
    ]


def _disruption_claims() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            raised_at="2026-10-10",
            claim_period_start="2026-09-01",
            claim_period_end="2026-09-30",
            description=BODY,
            status=status,
            cost_amount=AMOUNT,
            decided_amount=None,
            currency="TRY",
            schedule_days=4,
        )
        for status in ("draft", "submitted", "under_review", "agreed", "rejected")
    ]


def _eot_claims() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            raised_at="2026-10-10",
            claim_period_start="2026-09-01",
            claim_period_end="2026-09-30",
            description=BODY,
            status=status,
            requested_days=12,
            granted_days=None,
        )
        for status in ("draft", "submitted", "under_review", "granted", "rejected")
    ]


def _notices() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            code=f"NT-{index:03d}",
            title=TITLE,
            recipient_type=_codes(VARIATIONS, "recipient_type", index),
            recipient_name="Şükrü Iğdır",
            raised_at="2026-10-10",
            target_response_date="2026-10-17",
            response_received_at=None,
            status=_codes(VARIATIONS, "notice_status", index),
        )
        for index in range(_row_count(VARIATIONS))
    ]


def _evidence_pack() -> SimpleNamespace:
    kinds = list(EVIDENCE.labels["kind"]["en"])
    sections = []
    for position, name in enumerate(EVIDENCE.labels["section"]["en"]):
        sections.append(
            SimpleNamespace(
                name=name,
                entries=[
                    SimpleNamespace(
                        ref_id=str(uuid.UUID(int=position + 1)),
                        source_module="variations",
                        kind=kinds[position % len(kinds)],
                        title=TITLE,
                        occurred_at="2026-10-10T08:00:00+00:00",
                        actor_id=None,
                        summary=BODY,
                    )
                ],
            )
        )
    return SimpleNamespace(
        subject_ref="Şaft gecikmesi",
        basis="delay",
        entry_count=len(sections),
        date_from="2026-10-10T08:00:00+00:00",
        date_to="2026-10-10T08:00:00+00:00",
        sections=sections,
        content_digest="ab" * 32,
    )


RFI_STATUSES = ("draft", "open", "answered", "closed", "void")


def _rfis() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            id=uuid.UUID(int=index + 1),
            rfi_number=f"RFI-{index:03d}",
            subject=TITLE,
            status=status,
            raised_by=PERSON_ID,
            assigned_to=PERSON_ID,
            ball_in_court=None,
            date_required="2026-10-10",
            response_due_date="2026-10-17",
            cost_impact=index % 2 == 0,
            cost_impact_value="1234.56",
            schedule_impact=index % 2 == 0,
            schedule_impact_days=3,
            official_response=BODY,
        )
        for index, status in enumerate(RFI_STATUSES)
    ]


# ── Cases ─────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Case:
    """One document family: its catalogue and how to describe its documents."""

    name: str
    catalogue: DocumentCatalogue
    registers: Callable[[str], list[RegisterDocument]]
    records: Callable[[str], list[RecordDocument]] = lambda locale: []
    #: Turkish text that must be on the first register, besides its row text.
    turkish: tuple[str, ...] = ()
    money: bool = False
    #: English words the documents take from outside the catalogue.
    extra_english: tuple[str, ...] = field(default=())


def _context(locale: str, **extra: Any) -> dict[str, Any]:
    return {"project": PROJECT, "locale": locale, "generated": GENERATED, **extra}


CASES = [
    Case(
        "submittals",
        SUBMITTALS,
        lambda locale: [submittals_export.submittal_register(_submittals(), **_context(locale, people=PEOPLE))],
        lambda locale: [submittals_export.submittal_record(_submittals()[0], **_context(locale, people=PEOPLE))],
        turkish=("Onay Belgeleri Takip Listesi",),
    ),
    Case(
        "correspondence",
        CORRESPONDENCE,
        lambda locale: [
            correspondence_export.correspondence_register(_correspondence(), **_context(locale, people=PEOPLE))
        ],
        lambda locale: [
            correspondence_export.correspondence_record(_correspondence()[0], **_context(locale, people=PEOPLE))
        ],
        turkish=("Yazışma", "Gelen", "Giden"),
    ),
    Case(
        "changeorders",
        CHANGEORDERS,
        lambda locale: [changeorders_export.change_order_register(_change_orders(), **_context(locale))],
        lambda locale: [
            changeorders_export.change_order_record(_change_orders()[0], **_context(locale, people=PEOPLE))
        ],
        turkish=("Değişiklik Emirleri Kayıt Listesi",),
        money=True,
    ),
    Case(
        "transmittals",
        TRANSMITTALS,
        lambda locale: [transmittals_export.transmittal_register(_transmittals(), **_context(locale, people=PEOPLE))],
        lambda locale: [transmittals_export.transmittal_record(_transmittals()[0], **_context(locale, people=PEOPLE))],
        turkish=("İletim",),
    ),
    Case(
        "variations",
        VARIATIONS,
        lambda locale: [
            variations_export.variation_register(_variation_requests(), **_context(locale)),
            variations_export.claim_register(_disruption_claims(), _eot_claims(), **_context(locale)),
            variations_export.notice_register(_notices(), **_context(locale)),
        ],
        lambda locale: [
            variations_export.variation_record(_variation_requests()[0], **_context(locale, people=PEOPLE))
        ],
        turkish=("İlave İşler Kayıt Listesi", "Talep edilen tutar", "Kapsam değişikliği"),
        money=True,
    ),
    Case(
        "claims_evidence",
        EVIDENCE,
        lambda locale: [evidence_export.evidence_pack_register(_evidence_pack(), **_context(locale))],
        turkish=("Kanıt Paketi", "Gecikme / süre uzatımı", "İlave iş talebi", "Şaft gecikmesi"),
    ),
    Case(
        "rfi",
        rfi_export.CATALOGUE,
        lambda locale: [
            rfi_export.rfi_register(
                _rfis(), **_context(locale, people=PEOPLE, days_open={item.id: 4 for item in _rfis()})
            )
        ],
        turkish=("Bilgi Talepleri Kayıt Listesi (RFI)", "Sıradaki sorumlu", "Evet (1.234,56 TRY)", "Evet (3 gün)"),
        extra_english=tuple(localize_status(status, "en") for status in RFI_STATUSES),
    ),
]

CASE_IDS = [case.name for case in CASES]


# ── Helpers ───────────────────────────────────────────────────────────────


def _flat(text: str) -> str:
    return " ".join(unicodedata.normalize("NFC", text).replace(" ", " ").split())


def _pdf_text(pdf: bytes) -> str:
    return _flat(" ".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(pdf)).pages))


def _workbook_text(document: RegisterDocument) -> str:
    sheet = load_workbook(build_register_xlsx(document)).active
    cells = [str(cell) for row in sheet.iter_rows(values_only=True) for cell in row if isinstance(cell, str)]
    return _flat(" ".join([sheet.title, *cells]))


def _layer_text(document: RegisterDocument | RecordDocument) -> str:
    if isinstance(document, RegisterDocument):
        lines = [document.title, *(f"{label} {value}" for label, value in document.details)]
        lines += [cell for row in register_text_rows(document) for cell in row]
        return _flat(" ".join(lines))
    return _flat(" ".join(record_text_lines(document)))


_PLACEHOLDER = re.compile(r"\{[a-z_]+\}")
_NOT_PROSE = ("date_format", "datetime_format")


def _fragments(value: str) -> list[str]:
    """The fixed words of a string, with its ``{placeholders}`` cut out."""
    return [part.strip(" :,()/.-") for part in _PLACEHOLDER.split(value) if re.search(r"[A-Za-z]{2}", part)]


def _values(catalogue: DocumentCatalogue, locale: str) -> set[str]:
    values = {
        fragment
        for key, value in catalogue.strings[locale].items()
        if key not in _NOT_PROSE and not key.endswith("_filename")
        for fragment in _fragments(value)
    }
    for table in catalogue.labels.values():
        values.update(fragment for value in table[locale].values() for fragment in _fragments(value))
    return values


def _english_only(case: Case) -> set[str]:
    """English labels with a different Turkish form: the ones that must not show."""
    turkish = _values(case.catalogue, "tr")
    extra_turkish = {localize_status(status, "tr") for status in RFI_STATUSES}
    english = _values(case.catalogue, "en") | set(case.extra_english)
    return {value for value in english if value and value not in turkish and value not in extra_turkish}


def _leaks(case: Case, text: str) -> list[str]:
    return sorted(
        value for value in _english_only(case) if re.search(rf"(?<!\w){re.escape(value)}(?!\w)", text) is not None
    )


def _documents(case: Case, locale: str) -> list[RegisterDocument | RecordDocument]:
    return [*case.registers(locale), *case.records(locale)]


def _render(document: RegisterDocument | RecordDocument) -> bytes:
    return build_register_pdf(document) if isinstance(document, RegisterDocument) else build_record_pdf(document)


# ── Tests ─────────────────────────────────────────────────────────────────


def test_the_bundled_faces_are_in_use() -> None:
    """Without them the page falls back to a face that has no ğ, ı, İ or ş."""
    assert register_pdf_fonts() is True


@pytest.mark.parametrize("case", CASES, ids=CASE_IDS)
def test_the_english_only_set_is_not_empty(case: Case) -> None:
    """A leak check over an empty set would pass whatever the page said."""
    english_only = _english_only(case)
    assert len(english_only) >= 10, sorted(english_only)
    # The English documents do show those words, so the check can see a leak.
    english_text = " ".join(_layer_text(document) for document in _documents(case, "en"))
    seen = [value for value in english_only if re.search(rf"(?<!\w){re.escape(value)}(?!\w)", english_text)]
    assert len(seen) >= 10, sorted(english_only - set(seen))


@pytest.mark.parametrize("case", CASES, ids=CASE_IDS)
def test_turkish_documents_render_with_turkish_labels_and_intact_text(case: Case) -> None:
    documents = _documents(case, "tr")
    for document in documents:
        pdf = _render(document)
        assert pdf.startswith(b"%PDF")
        text = _pdf_text(pdf)
        pages = len(PdfReader(io.BytesIO(pdf)).pages)
        assert f"Sayfa 1 / {pages}" in text
        assert f"Oluşturulma: {GENERATED}" in text
        assert _flat(document.title) in text
        assert "Işıklı Veri Merkezi" in text
        assert "TR-001" in text
        # The letters Turkish adds survive to the page, on every document.
        for letter in "çğıİöşü":
            assert letter in text, (document.title, letter)
    first = _pdf_text(_render(documents[0])) + " " + _layer_text(documents[0])
    for expected in case.turkish:
        assert expected in first, expected
    everything = " ".join(_pdf_text(_render(document)) for document in documents)
    assert "ığüşöç" in everything
    assert "IĞÜŞÖÇİ" in everything


@pytest.mark.parametrize("case", CASES, ids=CASE_IDS)
def test_no_english_label_shows_in_a_turkish_document(case: Case) -> None:
    for document in _documents(case, "tr"):
        assert _leaks(case, _layer_text(document)) == [], document.title
        assert _leaks(case, _pdf_text(_render(document))) == [], document.title
        if isinstance(document, RegisterDocument):
            assert _leaks(case, _workbook_text(document)) == [], document.title


@pytest.mark.parametrize("case", CASES, ids=CASE_IDS)
def test_every_code_is_printed_as_a_turkish_word(case: Case) -> None:
    """No stored code such as ``revise_and_resubmit`` reaches the page."""
    text = " ".join(_layer_text(document) for document in _documents(case, "tr"))
    codes = {code for table in case.catalogue.labels.values() for code in table["en"] if "_" in code}
    assert [code for code in sorted(codes) if code in text] == []
    for table in case.catalogue.labels.values():
        for code, label in table["tr"].items():
            assert label.strip(), code


@pytest.mark.parametrize("case", CASES, ids=CASE_IDS)
def test_dates_and_amounts_are_written_the_turkish_way(case: Case) -> None:
    for document in case.registers("tr"):
        text = _pdf_text(build_register_pdf(document))
        assert "10.10.2026" in text, document.title
        assert "2026-10-10" not in text, document.title
    for document in case.records("tr"):
        assert "2026-10-" not in _pdf_text(build_record_pdf(document)), document.title
    if case.money:
        register = case.registers("tr")[0]
        assert "1.234,56" in _pdf_text(build_register_pdf(register))
        assert "1,234.56" not in _pdf_text(build_register_pdf(register))
        sheet = load_workbook(build_register_xlsx(register)).active
        amounts = [cell.value for row in sheet.iter_rows() for cell in row if cell.number_format == "#,##0.00"]
        # In the workbook an amount is a number the reader's spreadsheet formats.
        assert amounts
        assert all(isinstance(value, (int, float)) for value in amounts if value is not None)
        for record in case.records("tr"):
            assert "1.234,56 TRY" in _pdf_text(build_record_pdf(record))


@pytest.mark.parametrize("case", CASES, ids=CASE_IDS)
def test_english_documents_render_with_the_same_rows(case: Case) -> None:
    for document in _documents(case, "en"):
        text = _pdf_text(_render(document))
        pages = len(PdfReader(io.BytesIO(_render(document))).pages)
        assert f"Page 1 of {pages}" in text
        assert _flat(document.title) in text
        assert "Işıklı Veri Merkezi" in text
        assert "Oluşturulma" not in text
        assert "Sayfa" not in text
    english = " ".join(_pdf_text(_render(document)) for document in _documents(case, "en"))
    assert "ığüşöç" in english


@pytest.mark.parametrize("case", CASES, ids=CASE_IDS)
def test_an_unsupported_language_prints_english_not_half_of_each(case: Case) -> None:
    assert case.catalogue.normalize("zz") == "en"
    fallback = [_layer_text(document) for document in _documents(case, "zz")]
    assert fallback == [_layer_text(document) for document in _documents(case, "en")]


def test_the_rfi_log_keeps_its_english_headings_and_sheet_name() -> None:
    """Scripts and people read the English log by these headings; they did not move."""
    document = rfi_export.rfi_register(_rfis(), project=PROJECT, people=PEOPLE, locale="en")
    assert [column.header for column in document.columns] == [
        "RFI #",
        "Subject",
        "Status",
        "Raised By",
        "Assigned To",
        "Ball-in-Court",
        "Date Required",
        "Response Due",
        "Days Open",
        "Cost Impact",
        "Schedule Impact",
        "Response",
    ]
    assert load_workbook(build_register_xlsx(document)).active.title == "RFI Log"
    assert rfi_export.rfi_register_filename("en", "xlsx") == "rfi_log.xlsx"


def test_the_claims_register_numbers_rows_and_keeps_money_off_time_claims() -> None:
    document = variations_export.claim_register(_disruption_claims(), _eot_claims(), project=PROJECT, locale="tr")
    rows = register_text_rows(document)
    assert rows[0][:2] == ["Sıra no.", "Talep türü"]
    assert [row[0] for row in rows[1:]] == [str(number) for number in range(1, 11)]
    by_type = {row[1]: row for row in rows[1:]}
    assert by_type["Aksama"][7] == "1.234,56"
    assert by_type["Aksama"][9] == "TRY"
    # An extension of time claim asks for days, not money.
    assert by_type["Süre uzatımı"][7:10] == ["-", "-", "-"]
    assert by_type["Süre uzatımı"][10] == "12"
