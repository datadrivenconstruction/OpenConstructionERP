"""BOQ exports mark rows a model produced (EU AI Act, art. 50(2)).

One test per format proves the marker is there for an AI row, and each also
proves an ordinary bill carries none, so the marker cannot pass by being
stamped on everything.
"""

from __future__ import annotations

import io
import uuid
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pypdf
import pytest

from app.modules.boq.ai_provenance import ai_origin
from app.modules.boq.pdf_export import generate_boq_pdf
from app.modules.boq.router import _render_boq_xlsx, build_gaeb_xml

etree = pytest.importorskip("lxml.etree")

_PROFILE_DIR = Path(__file__).resolve().parents[2] / "app" / "modules" / "boq" / "gaeb_profile"


def _position(ordinal: str, source: str, metadata: dict[str, Any] | None = None) -> Any:
    return SimpleNamespace(
        id=uuid.uuid4(),
        parent_id=None,
        ordinal=ordinal,
        description=f"Concrete slab {ordinal}",
        unit="m3",
        quantity=Decimal("10"),
        unit_rate=Decimal("100"),
        total=Decimal("1000"),
        classification={},
        source=source,
        confidence=None,
        metadata=metadata or {},
        cad_element_ids=[],
        wbs_code=None,
        wbs_id=None,
    )


def _bill(*, with_ai: bool) -> Any:
    positions = [_position("01.001", "manual"), _position("01.002", "ai_copilot_accepted" if with_ai else "manual")]
    section = SimpleNamespace(
        id=uuid.uuid4(), ordinal="01", description="Structure", positions=positions, subtotal=Decimal("2000")
    )
    return SimpleNamespace(
        id=uuid.uuid4(),
        project_id=uuid.uuid4(),
        name="Harbour BOQ",
        description="",
        status="draft",
        sections=[section],
        positions=[],
        markups=[],
        direct_cost=Decimal("2000"),
        net_total=Decimal("2000"),
        grand_total=Decimal("2000"),
    )


def test_ai_origin_reads_source_and_estimator_metadata() -> None:
    assert ai_origin(_position("1", "ai_match")) == "ai_match"
    assert ai_origin(_position("1", "manual", {"ai_estimator_run_id": "r1"})) == "ai_estimator"
    assert ai_origin(_position("1", "manual")) is None
    assert ai_origin(_position("1", "cwicr")) is None


def test_pdf_marks_ai_rows_in_info_xmp_and_text() -> None:
    marked = pypdf.PdfReader(io.BytesIO(generate_boq_pdf(_bill(with_ai=True), "Harbour", "EUR")))
    assert marked.metadata["/AI-GENERATED"] == "true"
    assert marked.metadata["/AI-GENERATED-ROWS"] == "1"
    assert marked.metadata["/AI-GENERATED-SOURCES"] == "ai_copilot_accepted"
    xmp = marked.trailer["/Root"]["/Metadata"].get_object().get_data().decode("utf-8")
    assert "<oce:aiGenerated>true</oce:aiGenerated>" in xmp
    # The marker sits under the description in a narrow column and may wrap.
    text = "".join("".join(page.extract_text().split()) for page in marked.pages)
    assert "AI-generated(ai_copilot_accepted)" in text

    plain = pypdf.PdfReader(io.BytesIO(generate_boq_pdf(_bill(with_ai=False), "Harbour", "EUR")))
    assert "/AI-GENERATED" not in (plain.metadata or {})


def _xlsx(bill: Any) -> Any:
    from openpyxl import load_workbook

    flat = SimpleNamespace(name=bill.name, grand_total=bill.grand_total, positions=bill.sections[0].positions)
    data = _render_boq_xlsx(flat, bill, custom_columns=[], project_line=None, base_ccy="EUR", fx_map={})
    return load_workbook(io.BytesIO(data))


def test_xlsx_marks_ai_rows_with_comment_and_custom_property() -> None:
    wb = _xlsx(_bill(with_ai=True))
    props = {p.name: p.value for p in wb.custom_doc_props.props}
    assert props["AI-GENERATED"] is True
    assert props["AI-GENERATED-ROWS"] == 1
    comments = {c.value: c.comment.text for row in wb["BOQ"].iter_rows() for c in row if c.comment}
    assert comments == {"Concrete slab 01.002": "[AI-GENERATED source=ai_copilot_accepted]"}

    plain = _xlsx(_bill(with_ai=False))
    assert "AI-GENERATED" not in {p.name for p in plain.custom_doc_props.props}


@pytest.mark.parametrize(("gaeb_format", "dp_code", "holder"), [("x83", "83", "DetailTxt"), ("x84", "84", "BidComm")])
def test_gaeb_marks_ai_rows_and_stays_schema_valid(gaeb_format: str, dp_code: str, holder: str) -> None:
    xml = build_gaeb_xml(_bill(with_ai=True), project_name="Harbour", project_currency="EUR", gaeb_format=gaeb_format)
    doc = etree.fromstring(xml.encode("utf-8"))
    schema = etree.XMLSchema(etree.parse(str(_PROFILE_DIR / f"oce-gaeb-3.3-x{dp_code}.xsd")))
    assert schema.validate(doc), "; ".join(e.message for e in schema.error_log[:5])

    ns = {"g": f"http://www.gaeb.de/GAEB_DA_XML/DA{dp_code}/3.3"}
    marks = [el.text for el in doc.iterfind(f".//g:Item//g:{holder}//g:span", ns) if (el.text or "").startswith("[AI-")]
    assert marks == ["[AI-GENERATED source=ai_copilot_accepted]"]

    plain = build_gaeb_xml(
        _bill(with_ai=False), project_name="Harbour", project_currency="EUR", gaeb_format=gaeb_format
    )
    assert "AI-GENERATED" not in plain


def test_gaeb_import_drops_the_marker_from_the_long_text() -> None:
    from app.modules.boq.importers.gaeb_xml import _extract_long_text

    xml = build_gaeb_xml(_bill(with_ai=True), project_name="Harbour", project_currency="EUR", gaeb_format="x83")
    root = etree.fromstring(xml.encode("utf-8"))
    import xml.etree.ElementTree as ET

    items = ET.fromstring(etree.tostring(root)).iter("{http://www.gaeb.de/GAEB_DA_XML/DA83/3.3}Item")
    assert [_extract_long_text(i) for i in items] == ["Concrete slab 01.001", "Concrete slab 01.002"]
