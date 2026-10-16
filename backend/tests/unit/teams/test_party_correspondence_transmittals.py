"""Party visibility over HTTP in the correspondence and transmittal registers.

Every user here is a member of the project, so plain project access would show
each of them every record. A record missing from a list, a 404 on its id or a
gap in the export therefore comes from the party filter and nothing else. The
positive cases (internal member, owner, a project without parties) pin the
other direction: a filter that hides too much fails them.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator
from typing import Any

import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.dependencies import get_current_user_id, get_current_user_payload, get_session
from app.modules.correspondence.router import router as correspondence_router
from app.modules.teams.party_schemas import ProjectPartyCreate
from app.modules.teams.party_service import PartyService
from app.modules.transmittals.router import router as transmittals_router
from tests.unit.teams.conftest import http_client, make_membership, make_project, make_team, make_user

pytestmark = pytest.mark.asyncio

_PERMISSIONS = [
    f"{module}.{action}"
    for module in ("correspondence", "transmittals")
    for action in ("read", "create", "update", "delete")
]


def _app(db_session: AsyncSession, caller_id: uuid.UUID) -> FastAPI:
    """Both routers on one app, with the test session and a fixed non-admin caller."""
    app = FastAPI()
    app.include_router(correspondence_router, prefix="/v1/correspondence")
    app.include_router(transmittals_router, prefix="/v1/transmittals")

    async def _session_override() -> AsyncIterator[AsyncSession]:
        yield db_session

    async def _user_override() -> str:
        return str(caller_id)

    async def _payload_override() -> dict[str, Any]:
        return {"sub": str(caller_id), "role": "editor", "permissions": _PERMISSIONS}

    app.dependency_overrides[get_session] = _session_override
    app.dependency_overrides[get_current_user_id] = _user_override
    app.dependency_overrides[get_current_user_payload] = _payload_override
    return app


async def _world(session: AsyncSession) -> dict[str, Any]:
    """Owner on the host party, an unbound colleague, subcontractors A and B."""
    owner = await make_user(session)
    internal = await make_user(session)
    sub_a = await make_user(session)
    sub_b = await make_user(session)
    project = await make_project(session, owner.id)
    team = await make_team(session, project.id, is_default=True)
    for user in (internal, sub_a, sub_b):
        await make_membership(session, team.id, user.id)
    service = PartyService(session)
    host = await service.create_party(project.id, ProjectPartyCreate(name="Main", role="contractor"), actor_id=owner.id)
    party_a = await service.create_party(project.id, ProjectPartyCreate(name="Sub A"), actor_id=owner.id)
    party_b = await service.create_party(project.id, ProjectPartyCreate(name="Sub B"), actor_id=owner.id)
    await service.bind_member(project.id, owner.id, host.id, actor_id=owner.id)
    await service.bind_member(project.id, sub_a.id, party_a.id, actor_id=owner.id)
    await service.bind_member(project.id, sub_b.id, party_b.id, actor_id=owner.id)
    return {
        "owner": owner,
        "internal": internal,
        "sub_a": sub_a,
        "sub_b": sub_b,
        "project": project,
        "party_a": party_a,
    }


async def _create_correspondence(session, project_id, user, subject: str) -> str:
    async with http_client(_app(session, user.id)) as client:
        resp = await client.post(
            "/v1/correspondence/",
            json={
                "project_id": str(project_id),
                "direction": "outgoing",
                "subject": subject,
                "correspondence_type": "letter",
            },
        )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


async def _create_transmittal(session, project_id, user, subject: str) -> str:
    async with http_client(_app(session, user.id)) as client:
        resp = await client.post(
            "/v1/transmittals/",
            json={"project_id": str(project_id), "subject": subject, "purpose_code": "for_review"},
        )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _workbook_text(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content), read_only=True)
    cells = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells.extend(str(v) for v in row if v is not None)
    return "\n".join(cells)


_REGISTERS = [
    pytest.param("/v1/correspondence", _create_correspondence, id="correspondence"),
    pytest.param("/v1/transmittals", _create_transmittal, id="transmittal"),
]


@pytest.mark.parametrize(("base", "create"), _REGISTERS)
async def test_subcontractor_a_cannot_reach_b_record_by_any_path(session, base, create) -> None:
    w = await _world(session)
    pid = w["project"].id
    a_id = await create(session, pid, w["sub_a"], "Alpha letter AAA")
    b_id = await create(session, pid, w["sub_b"], "Bravo letter BBB")

    async with http_client(_app(session, w["sub_a"].id)) as client:
        listed = await client.get(f"{base}/", params={"project_id": str(pid)})
        assert listed.status_code == 200, listed.text
        ids = {item["id"] for item in listed.json()["items"]}
        assert ids == {a_id}
        assert listed.json()["total"] == 1

        assert (await client.get(f"{base}/{b_id}")).status_code == 404
        assert (await client.get(f"{base}/{b_id}/export/pdf/")).status_code == 404
        assert (await client.patch(f"{base}/{b_id}", json={"subject": "hijack"})).status_code == 404
        assert (await client.delete(f"{base}/{b_id}")).status_code == 404

        export = await client.get(f"{base}/export/", params={"project_id": str(pid), "format": "xlsx"})
        assert export.status_code == 200, export.text
        text = _workbook_text(export.content)
        assert "Alpha letter AAA" in text
        assert "Bravo letter BBB" not in text

        # Own record stays fully reachable.
        assert (await client.get(f"{base}/{a_id}")).status_code == 200


@pytest.mark.parametrize(("base", "create"), _REGISTERS)
async def test_internal_member_and_owner_see_every_record(session, base, create) -> None:
    w = await _world(session)
    pid = w["project"].id
    a_id = await create(session, pid, w["sub_a"], "Alpha")
    b_id = await create(session, pid, w["sub_b"], "Bravo")
    internal_id = await create(session, pid, w["internal"], "Internal draft")

    for user in (w["internal"], w["owner"]):
        async with http_client(_app(session, user.id)) as client:
            listed = await client.get(f"{base}/", params={"project_id": str(pid)})
            assert {item["id"] for item in listed.json()["items"]} == {a_id, b_id, internal_id}
            assert (await client.get(f"{base}/{b_id}")).status_code == 200

    # The internal member's record carries no party row, so no external party sees it.
    async with http_client(_app(session, w["sub_a"].id)) as client:
        assert (await client.get(f"{base}/{internal_id}")).status_code == 404


@pytest.mark.parametrize(("base", "create"), _REGISTERS)
async def test_a_project_without_parties_is_unchanged(session, base, create) -> None:
    owner = await make_user(session)
    member = await make_user(session)
    project = await make_project(session, owner.id)
    team = await make_team(session, project.id, is_default=True)
    await make_membership(session, team.id, member.id)
    rec_id = await create(session, project.id, owner, "Plain")

    async with http_client(_app(session, member.id)) as client:
        listed = await client.get(f"{base}/", params={"project_id": str(project.id)})
        assert [item["id"] for item in listed.json()["items"]] == [rec_id]
        assert (await client.get(f"{base}/{rec_id}")).status_code == 200


@pytest.mark.parametrize(
    ("base", "create", "entity_type"),
    [
        pytest.param("/v1/correspondence", _create_correspondence, "correspondence", id="correspondence"),
        pytest.param("/v1/transmittals", _create_transmittal, "transmittal", id="transmittal"),
    ],
)
async def test_view_share_grants_read_but_not_write(session, base, create, entity_type) -> None:
    w = await _world(session)
    pid = w["project"].id
    b_id = await create(session, pid, w["sub_b"], "Bravo shared")
    await PartyService(session).share_record(
        pid, entity_type, uuid.UUID(b_id), w["party_a"].id, "view", actor_id=w["sub_b"].id
    )

    async with http_client(_app(session, w["sub_a"].id)) as client:
        listed = await client.get(f"{base}/", params={"project_id": str(pid)})
        assert b_id in {item["id"] for item in listed.json()["items"]}
        assert (await client.get(f"{base}/{b_id}")).status_code == 200
        assert (await client.get(f"{base}/{b_id}/export/pdf/")).status_code == 200
        assert (await client.patch(f"{base}/{b_id}", json={"subject": "edited"})).status_code == 403
        assert (await client.delete(f"{base}/{b_id}")).status_code == 403


async def test_transmittal_child_routes_answer_404_for_a_hidden_record(session) -> None:
    w = await _world(session)
    pid = w["project"].id
    b_id = await _create_transmittal(session, pid, w["sub_b"], "Bravo")
    fake = uuid.uuid4()

    async with http_client(_app(session, w["sub_a"].id)) as client:
        assert (await client.post(f"/v1/transmittals/{b_id}/issue/")).status_code == 404
        resp = await client.post(
            f"/v1/transmittals/{b_id}/recipients/",
            json={"recipient_name": "X", "recipient_email": "x@example.com"},
        )
        assert resp.status_code == 404
        assert (
            await client.post(f"/v1/transmittals/{b_id}/recipients/{fake}/acknowledge/", json={})
        ).status_code == 404
        resp = await client.post(f"/v1/transmittals/{b_id}/recipients/{fake}/respond/", json={"response": "ok"})
        assert resp.status_code == 404


async def test_correspondence_attachment_download_answers_404_for_a_hidden_record(session) -> None:
    w = await _world(session)
    b_id = await _create_correspondence(session, w["project"].id, w["sub_b"], "Bravo")
    async with http_client(_app(session, w["sub_a"].id)) as client:
        assert (await client.get(f"/v1/correspondence/{b_id}/attachments/0")).status_code == 404
