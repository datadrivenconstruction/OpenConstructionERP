"""Party visibility on the RFI register, through the real HTTP routes.

Every caller here is a member of the project, so plain project access would
show them every RFI. A hidden RFI that disappears from a response therefore
disappeared because of the party filter, and the internal-member controls
prove the same routes still answer everybody else.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator
from typing import Any

import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.global_search import global_search
from app.core.vector_index import COLLECTION_CORRESPONDENCE, COLLECTION_RFI, COLLECTION_TASKS, VectorHit
from app.dependencies import get_current_user_id, get_current_user_payload, get_session
from app.modules.rfi.router import router as rfi_router
from app.modules.rfi.schemas import RFICreate
from app.modules.rfi.service import RFIService
from app.modules.search.service import _drop_party_hidden_hits
from app.modules.teams.party_schemas import ProjectPartyCreate
from app.modules.teams.party_service import PartyService
from tests.unit.teams.conftest import http_client, make_membership, make_project, make_team, make_user

pytestmark = pytest.mark.asyncio

PREFIX = "/v1/rfi"
_RFI_PERMISSIONS = [
    "rfi.create",
    "rfi.read",
    "rfi.update",
    "rfi.delete",
    "rfi.assign",
    "rfi.respond",
    "rfi.close",
]


def _app(session: AsyncSession, caller_id: uuid.UUID) -> FastAPI:
    app = FastAPI()
    app.include_router(rfi_router, prefix=PREFIX)

    async def _session_override() -> AsyncIterator[AsyncSession]:
        yield session

    async def _user_override() -> str:
        return str(caller_id)

    async def _payload_override() -> dict[str, Any]:
        return {"sub": str(caller_id), "role": "manager", "permissions": list(_RFI_PERMISSIONS)}

    app.dependency_overrides[get_session] = _session_override
    app.dependency_overrides[get_current_user_id] = _user_override
    app.dependency_overrides[get_current_user_payload] = _payload_override
    return app


async def _world(session: AsyncSession) -> dict[str, Any]:
    """Owner on the host party, an unbound colleague, subcontractors A and B."""
    owner = await make_user(session, role="manager")
    internal = await make_user(session, role="manager")
    sub_a = await make_user(session, role="manager")
    sub_b = await make_user(session, role="manager")
    project = await make_project(session, owner.id)
    team = await make_team(session, project.id, is_default=True)
    for user in (internal, sub_a, sub_b):
        await make_membership(session, team.id, user.id)
    service = PartyService(session)
    host = await service.create_party(project.id, ProjectPartyCreate(name="Main", role="contractor"), actor_id=owner.id)
    party_a = await service.create_party(project.id, ProjectPartyCreate(name="Sub A"), actor_id=owner.id)
    party_b = await service.create_party(project.id, ProjectPartyCreate(name="Sub B"), actor_id=owner.id)
    await service.bind_member(project.id, owner.id, host.id, actor_id=owner.id)
    await service.bind_member(project.id, sub_a.id, party_a.id, actor_id=owner.id)
    await service.bind_member(project.id, sub_b.id, party_b.id, actor_id=owner.id)
    rfi_a = await RFIService(session).create_rfi(
        RFICreate(project_id=project.id, subject="Alpha footing", question="Rebar?"), user_id=str(sub_a.id)
    )
    rfi_b = await RFIService(session).create_rfi(
        RFICreate(project_id=project.id, subject="Bravo footing", question="Cover?"), user_id=str(sub_b.id)
    )
    return {
        "owner": owner,
        "internal": internal,
        "sub_a": sub_a,
        "sub_b": sub_b,
        "project": project,
        "party_a": party_a,
        "rfi_a": rfi_a,
        "rfi_b": rfi_b,
    }


async def _list_ids(session, user, project_id) -> tuple[set[str], int]:
    async with http_client(_app(session, user.id)) as client:
        resp = await client.get(f"{PREFIX}/", params={"project_id": str(project_id)})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    return {item["id"] for item in body["items"]}, body["total"]


def _workbook_text(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content))
    return " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)


# ── List, stats and exports ──────────────────────────────────────────────────


async def test_subcontractor_list_and_total_exclude_the_other_company(session) -> None:
    w = await _world(session)
    ids, total = await _list_ids(session, w["sub_a"], w["project"].id)
    assert ids == {str(w["rfi_a"].id)}
    assert total == 1


async def test_internal_member_and_owner_see_both(session) -> None:
    w = await _world(session)
    both = {str(w["rfi_a"].id), str(w["rfi_b"].id)}
    for user in (w["internal"], w["owner"]):
        ids, total = await _list_ids(session, user, w["project"].id)
        assert ids == both
        assert total == 2


async def test_stats_count_only_visible_rfis(session) -> None:
    w = await _world(session)
    async with http_client(_app(session, w["sub_a"].id)) as client:
        resp = await client.get(f"{PREFIX}/stats/", params={"project_id": str(w["project"].id)})
    assert resp.status_code == 200, resp.text
    assert resp.json()["total"] == 1


async def test_export_log_excludes_the_hidden_rfi(session) -> None:
    w = await _world(session)
    async with http_client(_app(session, w["sub_a"].id)) as client:
        resp = await client.get(f"{PREFIX}/export/", params={"project_id": str(w["project"].id)})
    assert resp.status_code == 200, resp.text
    text = _workbook_text(resp.content)
    assert "Alpha footing" in text
    assert "Bravo footing" not in text


async def test_register_export_excludes_the_hidden_rfi(session) -> None:
    w = await _world(session)
    async with http_client(_app(session, w["sub_a"].id)) as client:
        resp = await client.get(
            f"{PREFIX}/export/register/", params={"project_id": str(w["project"].id), "format": "xlsx"}
        )
    assert resp.status_code == 200, resp.text
    text = _workbook_text(resp.content)
    assert "Alpha footing" in text
    assert "Bravo footing" not in text


# ── By id ────────────────────────────────────────────────────────────────────


async def test_hidden_rfi_is_404_on_every_by_id_route(session) -> None:
    w = await _world(session)
    rid = w["rfi_b"].id
    async with http_client(_app(session, w["sub_a"].id)) as client:
        assert (await client.get(f"{PREFIX}/{rid}")).status_code == 404
        assert (await client.get(f"{PREFIX}/{rid}/export/pdf/")).status_code == 404
        assert (await client.get(f"{PREFIX}/{rid}/activity/")).status_code == 404
        assert (await client.get(f"{PREFIX}/{rid}/route/")).status_code == 404
        assert (await client.get(f"{PREFIX}/{rid}/attachments/0")).status_code == 404
        assert (await client.patch(f"{PREFIX}/{rid}", json={"subject": "hijack"})).status_code == 404
        assert (await client.post(f"{PREFIX}/{rid}/respond/", json={"official_response": "x"})).status_code == 404
        assert (await client.delete(f"{PREFIX}/{rid}")).status_code == 404


async def test_project_without_parties_is_unchanged(session) -> None:
    owner = await make_user(session, role="manager")
    member = await make_user(session, role="manager")
    project = await make_project(session, owner.id)
    team = await make_team(session, project.id, is_default=True)
    await make_membership(session, team.id, member.id)
    rfi = await RFIService(session).create_rfi(
        RFICreate(project_id=project.id, subject="Plain", question="?"), user_id=str(owner.id)
    )
    ids, total = await _list_ids(session, member, project.id)
    assert ids == {str(rfi.id)}
    assert total == 1
    async with http_client(_app(session, member.id)) as client:
        assert (await client.get(f"{PREFIX}/{rfi.id}")).status_code == 200


async def test_view_share_allows_read_but_not_write(session) -> None:
    w = await _world(session)
    await PartyService(session).share_record(
        w["project"].id, "rfi", w["rfi_b"].id, w["party_a"].id, "view", actor_id=w["sub_b"].id
    )
    rid = w["rfi_b"].id
    async with http_client(_app(session, w["sub_a"].id)) as client:
        assert (await client.get(f"{PREFIX}/{rid}")).status_code == 200
        assert (await client.patch(f"{PREFIX}/{rid}", json={"subject": "edit"})).status_code == 403
    ids, _ = await _list_ids(session, w["sub_a"], w["project"].id)
    assert str(rid) in ids


# ── Batch ────────────────────────────────────────────────────────────────────


async def test_batch_status_skips_the_hidden_rfi(session) -> None:
    w = await _world(session)
    # The batch route fences by projects the caller owns, or by every project
    # for a JWT admin role. The payload claims admin so that fence lets both
    # ids through; the party filter reads the role from the database (manager),
    # so it alone has to drop the hidden id.
    app = _app(session, w["sub_a"].id)

    async def _admin_payload() -> dict[str, Any]:
        return {"sub": str(w["sub_a"].id), "role": "admin", "permissions": list(_RFI_PERMISSIONS)}

    app.dependency_overrides[get_current_user_payload] = _admin_payload
    async with http_client(app) as client:
        resp = await client.patch(
            f"{PREFIX}/batch/status/",
            json={"ids": [str(w["rfi_a"].id), str(w["rfi_b"].id)], "status": "open"},
        )
    assert resp.status_code == 200, resp.text
    assert resp.json()["updated"] == 1


# ── Global search ────────────────────────────────────────────────────────────


async def test_global_search_hides_the_other_company_rfi(session) -> None:
    w = await _world(session)
    hits = await global_search(
        session, query="footing", project_id=str(w["project"].id), limit=50, user_id=str(w["sub_a"].id)
    )
    rfi_ids = {h["id"] for h in hits if h["type"] == "rfi"}
    assert rfi_ids == {str(w["rfi_a"].id)}

    hits = await global_search(
        session, query="footing", project_id=str(w["project"].id), limit=50, user_id=str(w["internal"].id)
    )
    assert {h["id"] for h in hits if h["type"] == "rfi"} == {str(w["rfi_a"].id), str(w["rfi_b"].id)}


# ── Vector / chat search post-filter ─────────────────────────────────────────


def _hit(entity_id, project_id, collection) -> VectorHit:
    return VectorHit(
        id=str(entity_id),
        score=0.9,
        text="t",
        module="m",
        project_id=str(project_id) if project_id else "",
        tenant_id="",
        payload={},
        collection=collection,
    )


async def test_search_hits_on_shared_registers_respect_party_visibility(session) -> None:
    w = await _world(session)
    pid = w["project"].id
    rankings = [
        [_hit(w["rfi_a"].id, pid, COLLECTION_RFI), _hit(w["rfi_b"].id, pid, COLLECTION_RFI)],
        [_hit(uuid.uuid4(), None, COLLECTION_CORRESPONDENCE)],
        [_hit(uuid.uuid4(), pid, COLLECTION_TASKS)],
    ]
    collections = [COLLECTION_RFI, COLLECTION_CORRESPONDENCE, COLLECTION_TASKS]

    kept = await _drop_party_hidden_hits(session, rankings, collections, str(w["sub_a"].id))
    assert [h.id for h in kept[0]] == [str(w["rfi_a"].id)]
    # A scoped hit without a project cannot be checked: dropped for a bound caller.
    assert kept[1] == []
    # Collections outside the shared registers are untouched.
    assert len(kept[2]) == 1

    internal = await _drop_party_hidden_hits(session, rankings, collections, str(w["internal"].id))
    assert [len(r) for r in internal] == [2, 1, 1]
