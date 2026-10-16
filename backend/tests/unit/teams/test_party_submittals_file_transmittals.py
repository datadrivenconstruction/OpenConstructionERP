"""Party visibility in the submittal and file transmittal registers, over HTTP.

Every user in these cases is a member of the project, so plain project access
would show them every record. A record missing from a list, or a 404 on its id,
is therefore the party filter at work and not the project guard. The positive
cases (internal member, owner, a project without parties) keep a filter that
hides too much from passing.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator
from typing import Any

import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.dependencies import get_current_user_id, get_current_user_payload, get_session
from app.modules.file_transmittals.permissions import register_file_transmittal_permissions
from app.modules.file_transmittals.router import router as file_transmittals_router
from app.modules.submittals.permissions import register_submittals_permissions
from app.modules.submittals.router import router as submittals_router
from app.modules.teams.party_schemas import ProjectPartyCreate
from app.modules.teams.party_service import PartyService
from tests.unit.teams.conftest import make_membership, make_project, make_team, make_user

pytestmark = pytest.mark.asyncio

SUB = "/v1/submittals"
FT = "/v1/file-transmittals"


def _client(session: AsyncSession, caller_id: uuid.UUID) -> AsyncClient:
    """Both routers, the test session and a caller with a manager's role.

    The role only passes the permission gates; who sees what is decided by the
    real project and party checks against the rows in the database.
    """
    register_submittals_permissions()
    register_file_transmittal_permissions()
    app = FastAPI()
    app.include_router(submittals_router, prefix=SUB)
    app.include_router(file_transmittals_router, prefix=FT)

    async def _session_override() -> AsyncIterator[AsyncSession]:
        yield session

    async def _user_override() -> str:
        return str(caller_id)

    async def _payload_override() -> dict[str, Any]:
        return {"sub": str(caller_id), "role": "manager", "permissions": []}

    app.dependency_overrides[get_session] = _session_override
    app.dependency_overrides[get_current_user_id] = _user_override
    app.dependency_overrides[get_current_user_payload] = _payload_override
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


async def _world(session: AsyncSession) -> dict[str, Any]:
    """Owner in the host party, an unbound colleague, subcontractors A and B."""
    owner = await make_user(session)
    internal = await make_user(session)
    sub_a = await make_user(session)
    sub_b = await make_user(session)
    project = await make_project(session, owner.id)
    team = await make_team(session, project.id, is_default=True)
    for user in (internal, sub_a, sub_b):
        await make_membership(session, team.id, user.id)
    service = PartyService(session)
    host = await service.create_party(project.id, ProjectPartyCreate(name="Main", role="contractor"), actor_id=owner.id)
    party_a = await service.create_party(project.id, ProjectPartyCreate(name="Sub A"), actor_id=owner.id)
    party_b = await service.create_party(project.id, ProjectPartyCreate(name="Sub B"), actor_id=owner.id)
    await service.bind_member(project.id, owner.id, host.id, actor_id=owner.id)
    await service.bind_member(project.id, sub_a.id, party_a.id, actor_id=owner.id)
    await service.bind_member(project.id, sub_b.id, party_b.id, actor_id=owner.id)
    return {
        "owner": owner,
        "internal": internal,
        "sub_a": sub_a,
        "sub_b": sub_b,
        "project": project,
        "party_a": party_a,
        "party_b": party_b,
    }


async def _create_submittal(session: AsyncSession, user_id: uuid.UUID, project_id: uuid.UUID, title: str) -> str:
    async with _client(session, user_id) as client:
        response = await client.post(
            f"{SUB}/", json={"project_id": str(project_id), "title": title, "submittal_type": "product_data"}
        )
    assert response.status_code == 201, response.text
    return response.json()["id"]


async def _create_transmittal(session: AsyncSession, user_id: uuid.UUID, project_id: uuid.UUID, subject: str) -> str:
    async with _client(session, user_id) as client:
        response = await client.post(
            f"{FT}/", json={"project_id": str(project_id), "subject": subject, "reason_code": "for_review"}
        )
    assert response.status_code == 201, response.text
    return response.json()["id"]


async def _submittal_ids(session: AsyncSession, user_id: uuid.UUID, project_id: uuid.UUID) -> set[str]:
    async with _client(session, user_id) as client:
        response = await client.get(f"{SUB}/", params={"project_id": str(project_id), "limit": 100})
    assert response.status_code == 200, response.text
    return {row["id"] for row in response.json()}


async def _transmittal_ids(session: AsyncSession, user_id: uuid.UUID, project_id: uuid.UUID) -> set[str]:
    async with _client(session, user_id) as client:
        response = await client.get(f"{FT}/", params={"project_id": str(project_id)})
    assert response.status_code == 200, response.text
    return {row["id"] for row in response.json()}


def _workbook_text(content: bytes) -> str:
    sheet = load_workbook(io.BytesIO(content), read_only=True).worksheets[0]
    return " ".join(str(cell) for row in sheet.iter_rows(values_only=True) for cell in row if cell is not None)


# ── Submittals ───────────────────────────────────────────────────────────────


async def test_submittal_of_b_is_hidden_from_a_in_list_detail_pdf_export_and_summary(session) -> None:
    w = await _world(session)
    pid = w["project"].id
    mine = await _create_submittal(session, w["sub_a"].id, pid, "A pump data")
    theirs = await _create_submittal(session, w["sub_b"].id, pid, "B chiller data")

    assert await _submittal_ids(session, w["sub_a"].id, pid) == {mine}
    async with _client(session, w["sub_a"].id) as client:
        assert (await client.get(f"{SUB}/{theirs}")).status_code == 404
        assert (await client.get(f"{SUB}/{theirs}/export/pdf/")).status_code == 404
        assert (await client.get(f"{SUB}/{theirs}/attachments/")).status_code == 404
        assert (await client.get(f"{SUB}/{theirs}/validate/")).status_code == 404
        assert (await client.patch(f"{SUB}/{theirs}", json={"title": "x"})).status_code == 404
        assert (await client.delete(f"{SUB}/{theirs}")).status_code == 404
        assert (await client.get(f"{SUB}/{mine}")).status_code == 200

        export = await client.get(f"{SUB}/export/", params={"project_id": str(pid)})
        assert export.status_code == 200, export.text
        text = _workbook_text(export.content)
        assert "A pump data" in text
        assert "B chiller data" not in text

        summary = await client.get(f"{SUB}/summary/", params={"project_id": str(pid)})
        assert summary.status_code == 200, summary.text
        assert summary.json()["total"] == 1


async def test_internal_member_and_owner_see_every_submittal(session) -> None:
    w = await _world(session)
    pid = w["project"].id
    a = await _create_submittal(session, w["sub_a"].id, pid, "A")
    b = await _create_submittal(session, w["sub_b"].id, pid, "B")
    legacy = await _create_submittal(session, w["internal"].id, pid, "Internal")

    for user in (w["internal"], w["owner"]):
        assert await _submittal_ids(session, user.id, pid) == {a, b, legacy}
        async with _client(session, user.id) as client:
            assert (await client.get(f"{SUB}/{b}")).status_code == 200
            summary = await client.get(f"{SUB}/summary/", params={"project_id": str(pid)})
            assert summary.json()["total"] == 3
    # The internal member's record has no party rows, so neither subcontractor sees it.
    assert legacy not in await _submittal_ids(session, w["sub_b"].id, pid)


async def test_view_share_lets_a_read_a_submittal_but_not_change_it(session) -> None:
    w = await _world(session)
    pid = w["project"].id
    theirs = await _create_submittal(session, w["sub_b"].id, pid, "B shared")
    await PartyService(session).share_record(
        pid, "submittal", uuid.UUID(theirs), w["party_a"].id, "view", actor_id=w["sub_b"].id
    )

    assert theirs in await _submittal_ids(session, w["sub_a"].id, pid)
    async with _client(session, w["sub_a"].id) as client:
        assert (await client.get(f"{SUB}/{theirs}")).status_code == 200
        assert (await client.get(f"{SUB}/{theirs}/export/pdf/")).status_code == 200
        assert (await client.patch(f"{SUB}/{theirs}", json={"title": "changed"})).status_code == 403
        assert (await client.post(f"{SUB}/{theirs}/submit/")).status_code == 403


async def test_submittals_of_a_project_without_parties_are_not_filtered(session) -> None:
    owner = await make_user(session)
    member = await make_user(session)
    project = await make_project(session, owner.id)
    team = await make_team(session, project.id, is_default=True)
    await make_membership(session, team.id, member.id)
    first = await _create_submittal(session, owner.id, project.id, "Owner's")
    second = await _create_submittal(session, member.id, project.id, "Member's")

    assert await _submittal_ids(session, member.id, project.id) == {first, second}
    async with _client(session, member.id) as client:
        assert (await client.get(f"{SUB}/{first}")).status_code == 200
        summary = await client.get(f"{SUB}/summary/", params={"project_id": str(project.id)})
        assert summary.json()["total"] == 2


# ── File transmittals ────────────────────────────────────────────────────────


async def test_file_transmittal_of_b_is_hidden_from_a_in_list_detail_and_cover(session) -> None:
    w = await _world(session)
    pid = w["project"].id
    mine = await _create_transmittal(session, w["sub_a"].id, pid, "A drawings")
    theirs = await _create_transmittal(session, w["sub_b"].id, pid, "B drawings")

    assert await _transmittal_ids(session, w["sub_a"].id, pid) == {mine}
    async with _client(session, w["sub_a"].id) as client:
        assert (await client.get(f"{FT}/{theirs}")).status_code == 404
        assert (await client.get(f"{FT}/{theirs}/cover/")).status_code == 404
        assert (await client.post(f"{FT}/{theirs}/send/")).status_code == 404
        recipient = {"email": "someone@example.com"}
        assert (await client.post(f"{FT}/{theirs}/recipients/", json=recipient)).status_code == 404
        assert (await client.get(f"{FT}/{mine}")).status_code == 200


async def test_internal_member_and_owner_see_every_file_transmittal(session) -> None:
    w = await _world(session)
    pid = w["project"].id
    a = await _create_transmittal(session, w["sub_a"].id, pid, "A")
    b = await _create_transmittal(session, w["sub_b"].id, pid, "B")

    for user in (w["internal"], w["owner"]):
        assert await _transmittal_ids(session, user.id, pid) == {a, b}
        async with _client(session, user.id) as client:
            assert (await client.get(f"{FT}/{a}")).status_code == 200


async def test_view_share_lets_a_read_a_file_transmittal_but_not_change_it(session) -> None:
    w = await _world(session)
    pid = w["project"].id
    theirs = await _create_transmittal(session, w["sub_b"].id, pid, "B shared")
    await PartyService(session).share_record(
        pid, "file_transmittal", uuid.UUID(theirs), w["party_a"].id, "view", actor_id=w["owner"].id
    )

    assert theirs in await _transmittal_ids(session, w["sub_a"].id, pid)
    async with _client(session, w["sub_a"].id) as client:
        assert (await client.get(f"{FT}/{theirs}")).status_code == 200
        recipient = {"email": "someone@example.com"}
        assert (await client.post(f"{FT}/{theirs}/recipients/", json=recipient)).status_code == 403


async def test_file_transmittals_of_a_project_without_parties_are_not_filtered(session) -> None:
    owner = await make_user(session)
    member = await make_user(session)
    project = await make_project(session, owner.id)
    team = await make_team(session, project.id, is_default=True)
    await make_membership(session, team.id, member.id)
    first = await _create_transmittal(session, owner.id, project.id, "Owner's")

    assert await _transmittal_ids(session, member.id, project.id) == {first}
    async with _client(session, member.id) as client:
        assert (await client.get(f"{FT}/{first}")).status_code == 200
