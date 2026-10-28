# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""​‌‍⁠​‌‍⁠​‌‍⁠​‌‍⁠Pure-function pins for defects found in review: masking, hosting by tag, quorum, money, export, routes."""

from __future__ import annotations

import ast
import io
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.modules.stakeholders import export, hosting
from app.modules.stakeholders import mandates as md
from app.modules.stakeholders.catalogue import DEFAULT_FRAMEWORK_ID, FRAMEWORK_KIND, register_kind
from app.modules.stakeholders.errors import StakeholderError
from app.modules.stakeholders.evidence import placement
from app.modules.stakeholders.service import COMMITMENT_FROZEN, StakeholderService, _amount

_MODULE = Path(__file__).resolve().parents[2] / "app" / "modules" / "stakeholders"
TODAY = date(2026, 10, 12)
R = md.MandateRow


@pytest.fixture(scope="module")
def fw():
    from app.modules.project_governance.catalogue_kinds import get_catalogue_kind

    register_kind()
    return get_catalogue_kind(FRAMEWORK_KIND).get_shipped(DEFAULT_FRAMEWORK_ID)


# ── rule 17: what a masked reader gets ────────────────────────────────────


def test_masked_fields_blank_the_name_carrier_and_the_address() -> None:
    row = SimpleNamespace(
        id="i",
        project_id="p",
        code="S-001",
        display_name="Jane Private",
        category="neighbour",
        contact_id=None,
        is_group=False,
        group_size_estimate=None,
        parent_stakeholder_id=None,
        user_id=None,
        interests="",
        expectations="",
        concerns="",
        relationship_owner_user_id=None,
        geo={"address": "1 Garden Lane"},
        is_confidential=True,
        status="active",
        risk_id=None,
    )
    view = {**StakeholderService.stakeholder_out(row), **StakeholderService.masked_fields()}
    assert "Jane Private" not in str(view)
    assert "Garden Lane" not in str(view)


# ── hosting by catalogue tag, never by code ───────────────────────────────


def _phase_models():
    from app.modules.project_governance.catalogue_kinds import PHASE_MODEL_KIND, get_catalogue_kind

    return list(get_catalogue_kind(PHASE_MODEL_KIND).shipped())


def test_every_shipped_phase_model_hosts_both_subjects() -> None:
    models = _phase_models()
    assert len(models) >= 5
    for model in models:
        area_codes = {a.code for a in model.areas}
        for subject in hosting.SUBJECTS:
            found = hosting.host(model, subject)
            assert found is not None, (model.id, subject)
            assert found.area_code in area_codes, (model.id, subject)


def test_the_two_subjects_may_live_in_different_areas_or_in_one() -> None:
    split = [m.id for m in _phase_models() if len({hosting.host(m, s).area_code for s in hosting.SUBJECTS}) == 2]
    joint = [m.id for m in _phase_models() if len({hosting.host(m, s).doc_type for s in hosting.SUBJECTS}) == 1]
    # Both shapes ship, so neither may be assumed.
    assert split and joint


def test_a_model_without_the_tag_hosts_nothing() -> None:
    model = _phase_models()[0]
    bare = model.model_copy(
        update={"document_types": tuple(d.model_copy(update={"hosts": ()}) for d in model.document_types)}
    )
    assert [hosting.host(bare, s) for s in hosting.SUBJECTS] == [None, None]


def test_each_check_is_reported_under_a_subject_this_module_owns() -> None:
    from app.modules.stakeholders.evidence import CHECK_SUBJECT, host_code

    assert set(CHECK_SUBJECT.values()) == set(hosting.SUBJECTS)
    assert set(hosting.EXPORT_TAGS.values()) == set(hosting.SUBJECTS)
    assert host_code("organisation") == "stakeholders:host:organisation"


def test_the_evidence_link_opens_the_stakeholders_screen_of_the_project() -> None:
    from app.modules.project_governance.evidence import evidence_link_href, unregister_evidence_link
    from app.modules.stakeholders.evidence import LINK_KIND, register_gate_evidence

    unregister_evidence_link(LINK_KIND)
    assert evidence_link_href(LINK_KIND, "p1", "p1") is None
    register_gate_evidence()
    assert evidence_link_href(LINK_KIND, "p1", "p1") == "/projects/p1/stakeholders"
    declared = (_MODULE.parents[3] / "frontend" / "src" / "app" / "App.tsx").read_text(encoding="utf-8")
    # The registered path is a route the web app declares.
    assert 'path="/projects/:projectId/stakeholders"' in declared


def test_every_unknown_item_has_a_stable_code_and_says_what_to_do_in_each_locale() -> None:
    from app.modules.stakeholders.evidence import NO_DOCUMENT_REASON, NO_HOST_REASON, unknown_item

    codes = set()
    for subject in hosting.SUBJECTS:
        for reason in (NO_HOST_REASON, NO_DOCUMENT_REASON):
            item = unknown_item(subject, reason, ("stakeholders", "p"))
            codes.add(item.code)
            assert (item.state, item.detail) == ("unknown", reason)
            assert {"en", "de"} <= set(item.labels)
            for text in item.labels.values():
                # A sentence, not a raw key, and it ends with an instruction.
                assert text != item.label_key and "stakeholders." not in text
            assert "before you decide" in item.labels["en"]
            assert "bevor Sie entscheiden" in item.labels["de"]
    assert codes == {
        "stakeholders:host:stakeholders",
        "stakeholders:host:organisation",
        "stakeholders:document:stakeholders",
        "stakeholders:document:organisation",
    }


def test_no_module_file_matches_a_governance_code() -> None:
    literals = ("m01_project_management", "project_handbook", "stakeholder_register", "projekthandbuch")
    offenders = [
        (f.name, code) for f in _MODULE.glob("*.py") for code in literals if code in f.read_text(encoding="utf-8")
    ]
    assert offenders == []


def test_evidence_without_a_host_is_reported_once_and_never_dropped() -> None:
    assert placement("a", "a", None) == "items"
    assert placement("b", "a", None) == "nothing"
    # No area hosts the module: one unknown item, on the first area only.
    assert placement("a", None, "a") == "unplaced"
    assert placement("b", None, "a") == "nothing"
    # No governance model at all: there is no gate to report to.
    assert placement("a", None, None) == "nothing"


# ── rules 11, 12, 15: who decides in a body ───────────────────────────────


def test_a_guest_without_a_vote_does_not_decide() -> None:
    body = md.Body("board", "majority", members=(md.Member(user_id="u1"), md.Member(user_id="guest", voting=False)))
    res = md.who_decides([R("x", "body", "board", "D")], "x", "", [], [body], TODAY)
    assert res.user_ids == ["guest", "u1"]
    assert res.voting_user_ids == ["u1"]


def test_an_expired_seat_does_not_raise_the_quorum() -> None:
    body = md.Body(
        "board", "unanimous", members=(md.Member(user_id="u1"), md.Member(user_id="u2", valid_to=date(2026, 10, 1)))
    )
    assert md.quorum_needed(body, TODAY) == 1
    assert "body_inquorate" not in md.who_decides([R("x", "body", "board", "D")], "x", "", [], [body], TODAY).flags
    # While the second seat was still valid, unanimity needed both.
    assert md.quorum_needed(body, date(2026, 9, 30)) == 2


def test_a_body_with_no_valid_voting_seat_is_never_quorate() -> None:
    for kind in ("majority", "unanimous", "min_members"):
        body = md.Body("board", kind, quorum_min=1, members=(md.Member(user_id="u1", valid_to=date(2026, 1, 1)),))
        res = md.who_decides([R("x", "body", "board", "D")], "x", "", [], [body], TODAY)
        assert "body_inquorate" in res.flags, kind
    empty = md.Body("board", "unanimous")
    assert "body_inquorate" in md.who_decides([R("x", "body", "board", "D")], "x", "", [], [empty], TODAY).flags


def test_a_non_voting_seat_does_not_make_a_self_approval_conflict(fw) -> None:
    body = md.Body("board", "majority", members=(md.Member(user_id="pm", voting=False), md.Member(user_id="c")))
    rows = [R("design_decision", "role", "lead_designer", "P"), R("design_decision", "body", "board", "D")]
    holders = [md.Holder("lead_designer", "pm")]
    assert md.conflicts(fw, rows, holders, [body], TODAY) == []
    voting = md.Body("board", "majority", members=(md.Member(user_id="pm"), md.Member(user_id="c")))
    assert [c.user_id for c in md.conflicts(fw, rows, holders, [voting], TODAY)] == ["pm"]


# ── money ────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("bad", ["-1", "1.00001", "100000000000000", "NaN", "abc", 1.5, True])
def test_amounts_outside_the_column_are_refused(bad) -> None:
    with pytest.raises(StakeholderError) as exc:
        _amount(bad)
    assert exc.value.code == "invalid_amount"


def test_amounts_inside_the_column_pass_as_decimal() -> None:
    assert _amount("1250.5000") == Decimal("1250.5000")
    assert _amount("99999999999999.9999") == Decimal("99999999999999.9999")
    assert _amount("") is None


def test_currency_is_frozen_with_the_amount() -> None:
    assert {"text", "cost_estimate", "currency"} <= COMMITMENT_FROZEN


# ── write routes take no client date ─────────────────────────────────────


def test_no_write_route_takes_today_from_the_caller() -> None:
    tree = ast.parse((_MODULE / "router.py").read_text(encoding="utf-8"))
    offenders = []
    for fn in tree.body:
        if not isinstance(fn, ast.AsyncFunctionDef):
            continue
        methods = {
            d.func.attr for d in fn.decorator_list if isinstance(d, ast.Call) and isinstance(d.func, ast.Attribute)
        }
        if methods & {"post", "put", "patch", "delete"} and "today" in {a.arg for a in fn.args.args}:
            offenders.append(fn.name)
    assert offenders == []


def test_no_raw_role_string_is_passed_as_a_permission() -> None:
    src = (_MODULE / "router.py").read_text(encoding="utf-8")
    assert '"role": "viewer"' not in src


# ── export: one meaning per column, one sheet per table ───────────────────


def _register_docs(fw):
    stakeholders = [
        {
            "id": "s1",
            "code": "S-001",
            "name": None,
            "masked": True,
            "anonymised": False,
            "category": "neighbour",
            "classification": "manage_closely",
            "status": "active",
            "review_due": True,
            "open_commitments": 2,
        },
        {
            "id": "s2",
            "code": "S-002",
            "name": "Road authority",
            "masked": False,
            "anonymised": False,
            "category": "permitting_authority",
            "classification": "not_assessed",
            "status": "active",
        },
    ]
    plan = [{"title": "Newsletter", "audience_stakeholder_ids": ["s1", "s2"], "status": "active", "channel": None}]
    commitments = [
        {
            "code": "C-001",
            "text": "No early deliveries",
            "stakeholder_ids": ["s1"],
            "category": "noise",
            "status": "agreed",
        }
    ]
    return export.register_documents(
        fw, stakeholders=stakeholders, plan_items=plan, commitments=commitments, project_name="P", locale="en"
    )


def _organisation_docs(fw):
    roles = {
        "holders": [{"role_code": "client_pm", "holder_name": "A. Lead", "valid_from": "2026-01-01", "valid_to": None}],
        "vacant": ["client"],
    }
    bodies = [
        {
            "code": "steering_committee",
            "name": "Steering committee",
            "members": [{"role_code": "client_pm", "valid_today": True}, {"user_id": "u", "valid_today": True}],
            "quorum_needed": 2,
            "quorum_present": 2,
        }
    ]
    matrix = {
        "rows": [
            {"decision_type": "scope_change", "holder_kind": "body", "holder_code": "steering_committee", "letter": "D"}
        ]
    }
    return export.organisation_documents(fw, roles=roles, bodies=bodies, matrix=matrix, project_name="P", locale="en")


@pytest.mark.parametrize("build", [_register_docs, _organisation_docs])
def test_every_table_has_its_own_columns_and_every_row_fits_them(fw, build) -> None:
    docs = build(fw)
    assert len(docs) == 3
    assert len({tuple(c.header for c in d.columns) for d in docs}) == 3
    for d in docs:
        headers = [c.header for c in d.columns]
        assert len(headers) == len(set(headers)), headers
        assert d.rows and all(len(r) == len(headers) for r in d.rows), d.sheet_title


def test_the_workbook_has_one_sheet_per_table_and_masks_the_private_name(fw) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(export.render(_register_docs(fw), "xlsx")))
    assert wb.sheetnames == ["Stakeholders", "Communication plan", "Commitments"]
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert "(confidential)" in text
    assert "Road authority" in text
    assert "No early deliveries" in text


def test_the_pdf_is_the_three_tables_in_one_file(fw) -> None:
    from pypdf import PdfReader

    data = export.render(_organisation_docs(fw), "pdf")
    assert data.startswith(b"%PDF")
    assert len(PdfReader(io.BytesIO(data)).pages) >= 3


# ── the daily sweep is a registered process ──────────────────────────────


def test_the_sweep_is_registered_as_a_process() -> None:
    from app.core.processes import process_registry
    from app.modules.stakeholders.jobs import PROCESS_ID, register_sweep_process

    register_sweep_process()
    register_sweep_process()
    assert process_registry.ids().count(PROCESS_ID) == 1
    assert PROCESS_ID in process_registry.processes_for_module("stakeholders")


# ── what the caller may do, computed on the server ───────────────────────


def _can(**kw):
    from app.modules.stakeholders.service import commitment_actions

    base = {
        "status": "proposed",
        "pending": False,
        "has_owner_and_due": True,
        "is_owner": False,
        "can_edit": True,
        "routed": False,
        "flags": [],
        "is_decider": True,
    }
    return commitment_actions(**{**base, **kw})


def test_every_action_carries_allowed_and_a_reason_code() -> None:
    can = _can()
    assert set(can) == {"agreed", "in_progress", "fulfilled", "verified", "withdrawn", "breached", "edit"}
    assert can["agreed"] == {"applies": True, "allowed": True, "reason": None}
    # An action that does not exist in this state is hidden, not shown disabled.
    assert can["verified"] == {"applies": False, "allowed": False, "reason": "invalid_transition"}
    # A missing right or a precondition that can be met stays visible with its reason.
    assert _can(is_decider=False)["agreed"] == {"applies": True, "allowed": False, "reason": "not_decider"}
    assert _can(can_edit=False)["agreed"]["applies"] is True
    # A closed commitment offers nothing, editing included; an open one always offers the edit.
    assert not any(v["applies"] for v in _can(status="verified").values())
    assert not any(v["applies"] for v in _can(status="withdrawn").values())
    assert _can(status="breached")["edit"] == {"applies": True, "allowed": True, "reason": None}


def test_the_reason_is_the_code_the_write_would_answer() -> None:
    assert _can(is_decider=False)["agreed"]["reason"] == "not_decider"
    assert _can(is_decider=False, routed=True)["agreed"]["allowed"] is True
    assert _can(flags=["body_inquorate"])["agreed"]["reason"] == "body_inquorate"
    assert _can(flags=["mandate_missing"], routed=True)["agreed"]["reason"] == "mandate_missing"
    assert _can(has_owner_and_due=False)["agreed"]["reason"] == "owner_and_due_required"
    assert _can(status="fulfilled", is_owner=True)["verified"]["reason"] == "owner_cannot_verify"
    assert _can(status="fulfilled")["verified"]["allowed"] is True
    assert {v["reason"] for v in _can(pending=True).values()} == {"approval_pending"}
    assert {v["reason"] for v in _can(can_edit=False).values()} == {"no_permission"}
    assert not any(v["allowed"] for v in _can(status="verified").values() if v["reason"] == "invalid_transition")


# ── a stakeholder's place: named keys only ───────────────────────────────


def test_geo_takes_the_named_keys_and_drops_empty_values() -> None:
    from app.modules.stakeholders.schemas import StakeholderIn

    body = StakeholderIn(geo={"parcel_ref": " 12/7 ", "lat": "52.52", "lon": "13.405", "address": ""})
    assert body.geo == {"parcel_ref": "12/7", "lat": "52.52", "lon": "13.405"}
    assert StakeholderIn(geo={"address": " "}).geo is None


@pytest.mark.parametrize("geo", [{"postcode": "10115"}, {"lat": "91"}, {"lon": "-181"}, {"lat": "north"}])
def test_geo_refuses_unknown_keys_and_points_off_the_globe(geo) -> None:
    from pydantic import ValidationError

    from app.modules.stakeholders.schemas import StakeholderIn

    with pytest.raises(ValidationError):
        StakeholderIn(geo=geo)


def test_the_release_item_picker_says_when_nothing_provides_it() -> None:
    from app.modules.stakeholders import authority

    assert authority.release_item_lister() is None
    authority.register_release_item_lister(lambda session, project_id: [])
    try:
        assert authority.release_item_lister() is not None
    finally:
        authority.register_release_item_lister(None)
    assert authority.release_item_lister() is None


# ── demo seed: the site runs only after the gate into construction ───────


def test_the_seed_finds_the_undecided_gate_into_construction_by_capability() -> None:
    from app.modules.stakeholders.seed import gate_before_construction

    def gate(to_phase: str, status: str):
        return SimpleNamespace(to_phase=to_phase, status=status)

    gates = (gate("design", "decided"), gate("build", "in_preparation"), gate("use", "not_started"))
    assert gate_before_construction(gates, ("build",)) is gates[1]
    decided = (gate("design", "decided"), gate("build", "decided"))
    assert gate_before_construction(decided, ("build",)) is None
    # No phase model, or a model that names no construction start: the seed tells its usual story.
    assert gate_before_construction((), ("build",)) is None
    assert gate_before_construction(gates, ()) is None


# ── an overdue promise warns, a silent breach fails ──────────────────────


def test_an_overdue_promise_fails_the_check_only_while_it_cannot_warn() -> None:
    from app.modules.stakeholders.evidence import commitments_tracked

    old = commitments_tracked([], ["C-003"], warnings=False)
    assert (old["status"], old["refs"]) == ("not_met", ["C-003"])
    new = commitments_tracked([], ["C-003"], warnings=True)
    assert (new["status"], new["refs"], new["overdue_refs"]) == ("met", [], ["C-003"])
    # A breach nobody explained fails it either way.
    assert commitments_tracked(["C-001"], ["C-003"], warnings=True)["status"] == "not_met"
    assert commitments_tracked([], [], warnings=False)["status"] == "met"


def test_each_overdue_promise_is_a_warning_with_a_stable_code_and_a_sentence(monkeypatch) -> None:
    from app.modules.project_governance import evidence as governance_evidence
    from app.modules.stakeholders.evidence import overdue_items

    # The state arrives with a later governance version; the item must be valid there.
    monkeypatch.setattr(
        governance_evidence, "EVIDENCE_STATES", (*governance_evidence.EVIDENCE_STATES, "warning"), raising=True
    )
    overdue = [{"code": "C-003", "due_date": "2026-10-09"}]
    (item,) = overdue_items(overdue, ("stakeholders", "p"))
    assert (item.code, item.state) == ("stakeholders:commitment_overdue:C-003", "warning")
    assert overdue_items(overdue, ("stakeholders", "p"))[0].code == item.code
    assert "C-003" in item.labels["en"] and "2026-10-09" in item.labels["en"]
    assert "before you decide" in item.labels["en"]
    assert "C-003" in item.labels["de"] and "bevor Sie entscheiden" in item.labels["de"]
    assert overdue_items([], ("stakeholders", "p")) == []
