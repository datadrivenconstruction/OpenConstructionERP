# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The hakediş workbook holds the same figures as the PDF, as numbers.

A cost control department reworks a certificate in a spreadsheet. So the
workbook has to carry the computed totals exactly (no float, no formula that
could recompute them differently), keep a held line visible as a held line,
and print: frozen header rows, a print area, landscape for the works list.

The fixture and its SYNTHETIC rates are those of ``test_hakedis_math``.
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import load_workbook

from app.modules.contracts.hakedis import Certificate
from app.modules.contracts.hakedis_xlsx import render_hakedis_xlsx
from tests.unit.test_hakedis_math import BILL, EXPECTED, certify

D = Decimal


def sheets(cert: Certificate, **kwargs: Any) -> tuple[Any, Any]:
    workbook = load_workbook(io.BytesIO(render_hakedis_xlsx(cert, **kwargs)))
    return workbook.worksheets[0], workbook.worksheets[1]


def rows_of(ws: Any) -> list[tuple[Any, ...]]:
    return [row for row in ws.iter_rows(values_only=True) if any(cell is not None for cell in row)]


def find_row(ws: Any, text: str, column: int = 1) -> int:
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=column).value
        if isinstance(value, str) and text in value:
            return row
    raise AssertionError(f"{text!r} not found in column {column} of {ws.title!r}")


def exact(value: Any) -> Decimal:
    """A numeric cell as the two-decimal amount it holds."""
    assert isinstance(value, (int, float)) and not isinstance(value, bool), value
    return Decimal(str(value)).quantize(D("0.01"))


@pytest.fixture(scope="module")
def first() -> Certificate:
    return certify(1, D("0"))


@pytest.fixture(scope="module")
def book(first: Certificate) -> tuple[Any, Any]:
    return sheets(first)


def test_the_sheets_carry_the_documents_turkish_names(book: tuple[Any, Any]) -> None:
    summary, works = book
    assert summary.title == "Hakediş Raporu"
    assert works.title == "Yapılan İşler Listesi"


def test_every_lettered_line_is_a_number_equal_to_the_computed_amount(book: tuple[Any, Any]) -> None:
    summary, _ = book
    seen: dict[str, Decimal | None] = {}
    for row in range(1, summary.max_row + 1):
        letter = summary.cell(row=row, column=1).value
        if not isinstance(letter, str) or len(letter.rstrip(")")) != 1:
            continue
        cell = summary.cell(row=row, column=3)
        key = letter.rstrip(")")
        if cell.value == "-":
            seen[key] = None
        else:
            seen[key] = exact(cell.value)
            assert cell.number_format == "#,##0.00", key
            assert cell.data_type == "n", key
    payable_row = find_row(summary, "Yükleniciye Ödenecek Tutar ( G - H )", column=2)
    seen["payable"] = exact(summary.cell(row=payable_row, column=3).value)
    assert seen == EXPECTED[1]


def test_the_summary_prints_the_forms_labels_and_formulas(book: tuple[Any, Any]) -> None:
    summary, _ = book
    labels = [row[1] for row in rows_of(summary) if isinstance(row[1], str)]
    for text in (
        "Sözleşme Fiyatları İle Yapılan İş",
        "Toplam Tutar ( A + B )",
        "Bu Hakedişin Tutarı ( C - D )",
        "KDV ( E x %10 )",
        "Damga Vergisi ( E - g x %1 )",
        "KDV Tevkifatı ( F x 1/2 )",
        "Kesintiler ve Mahsuplar Toplamı",
    ):
        assert text in labels, text
    assert find_row(summary, "KESİNTİLER VE MAHSUPLAR") < find_row(summary, "Gelir / Kurumlar Vergisi", column=2)
    assert find_row(summary, "31.07.2026 TARİHİNE KADAR YAPILAN İŞİN")
    header = {row[0]: row[2] for row in rows_of(summary) if isinstance(row[0], str)}
    assert header["İşin Adı"] == "Örnek Veri Merkezi İnşaatı"
    assert header["Yüklenici - Vergi Dairesi"] == "Şişli"
    assert header["Hakediş No"] == "1 (Geçici Hakediş)"


def test_the_works_list_rows_and_totals_equal_the_computed_ones(first: Certificate, book: tuple[Any, Any]) -> None:
    _, works = book
    header_row = find_row(works, "Sıra No")
    headers = [works.cell(row=header_row, column=column).value for column in range(1, works.max_column + 1)]
    assert headers == [
        "Sıra No",
        "Poz No",
        "İşin Tanımı",
        "Birimi",
        "Teklif Birim Fiyat",
        "Toplam İmalat, İhzarat Miktarı",
        "Bir Önceki Hakediş İmalat, İhzarat Miktarı",
        "Bu Hakediş İmalat, İhzarat Miktarı",
        "Toplam İmalat, İhzarat Tutarı",
        "Bir Önceki Hakediş Tutarı",
        "Bu Hakediş Tutarı",
    ]
    letters = [works.cell(row=header_row - 1, column=column).value for column in range(1, works.max_column + 1)]
    assert letters[4:] == ["A", "B", "C", "D=B-C", "E=AxB", "F=AxC", "G=E-F"]

    by_code = {row.code: row for row in first.work_lines}
    printed = {}
    for row in range(header_row + 1, works.max_row + 1):
        code = works.cell(row=row, column=2).value
        if code in by_code:
            printed[code] = row
    assert list(printed) == [bill_row[1] for bill_row in BILL]
    for code, row in printed.items():
        result = by_code[code]
        assert works.cell(row=row, column=3).value == result.description
        assert Decimal(str(works.cell(row=row, column=6).value)) == result.cumulative_quantity, code
        assert exact(works.cell(row=row, column=9).value) == result.cumulative_amount, code
        assert exact(works.cell(row=row, column=10).value) == result.previous_amount, code
        assert exact(works.cell(row=row, column=11).value) == result.period_amount, code

    total_row = find_row(works, "Toplam", column=3)
    while works.cell(row=total_row, column=3).value != "Toplam":
        total_row = find_row_after(works, "Toplam", total_row, column=3)
    assert exact(works.cell(row=total_row, column=9).value) == first.totals.cumulative_amount
    assert exact(works.cell(row=total_row, column=10).value) == first.totals.previous_amount
    assert exact(works.cell(row=total_row, column=11).value) == first.totals.period_amount
    assert first.totals.cumulative_amount == EXPECTED[1]["A"]

    # The total is the sum of the rows the sheet itself prints.
    column_sum = sum((exact(works.cell(row=row, column=9).value) for row in printed.values()), D(0))
    assert column_sum == first.totals.cumulative_amount


def find_row_after(ws: Any, text: str, after: int, column: int) -> int:
    for row in range(after + 1, ws.max_row + 1):
        value = ws.cell(row=row, column=column).value
        if isinstance(value, str) and text in value:
            return row
    raise AssertionError(f"{text!r} not found after row {after}")


def test_section_subtotals_add_up_to_the_total(first: Certificate, book: tuple[Any, Any]) -> None:
    _, works = book
    subtotals = [
        exact(works.cell(row=row, column=9).value)
        for row in range(1, works.max_row + 1)
        if isinstance(works.cell(row=row, column=3).value, str) and "Ara Toplam" in works.cell(row=row, column=3).value
    ]
    assert len(subtotals) == 2
    assert sum(subtotals, D(0)) == first.totals.cumulative_amount
    assert find_row(works, "01 - MEKANİK TESİSAT") < find_row(works, "02 - ELEKTRİK TESİSATI")


def test_the_workbook_has_no_formulas(book: tuple[Any, Any]) -> None:
    for ws in book:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", cell.coordinate


def test_both_sheets_are_set_up_to_print(book: tuple[Any, Any]) -> None:
    summary, works = book
    assert summary.page_setup.orientation == "portrait"
    assert works.page_setup.orientation == "landscape"
    for ws in book:
        assert ws.sheet_properties.pageSetUpPr.fitToPage
        assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
        assert ws.print_area, ws.title
        assert ws.print_title_rows, ws.title
        assert ws.freeze_panes, ws.title
    header_row = find_row(works, "Sıra No")
    assert works.freeze_panes == f"A{header_row + 1}"
    assert works.print_title_rows == f"${header_row - 1}:${header_row}"
    last = find_row(works, "E-39.400", column=2)
    area = works.print_area if isinstance(works.print_area, str) else works.print_area[0]
    assert int(area.replace("$", "").split(":")[1].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ")) > last


def test_a_held_line_stays_a_marker_and_the_sheet_says_it_is_a_draft() -> None:
    cert = certify(1, D("0"), taxes=None)
    summary, _ = sheets(cert, locale="tr-en")
    assert find_row(summary, "TASLAK - Tutarların tamamı kesinleşmemiştir / DRAFT - not all figures are final")
    vat_row = find_row(summary, "KDV ( E x %.. )", column=2)
    assert summary.cell(row=vat_row, column=3).value == "Beklemede / Held (1)"
    assert summary.cell(row=vat_row, column=3).data_type == "s"
    payable_row = find_row(summary, "Yükleniciye Ödenecek Tutar", column=2)
    assert summary.cell(row=payable_row, column=3).value == "Beklemede / Held"
    assert "Amount Payable to the Contractor" in summary.cell(row=payable_row, column=2).value
    notes = find_row(summary, "AÇIKLAMALAR / NOTES")
    assert "Vergi hesaplama modülü kurulu değil" in summary.cell(row=notes + 1, column=2).value
    assert "The tax calculation module is not installed" in summary.cell(row=notes + 1, column=2).value
    # The work itself is known, and is still a number.
    work_row = find_row(summary, "Sözleşme Fiyatları İle Yapılan İş", column=2)
    assert exact(summary.cell(row=work_row, column=3).value) == EXPECTED[1]["A"]


def test_a_final_certificate_has_no_draft_banner(book: tuple[Any, Any]) -> None:
    summary, _ = book
    assert not any("TASLAK" in str(cell) for row in rows_of(summary) for cell in row if cell is not None)


def test_english_names_the_sheets_and_columns_in_english(first: Certificate) -> None:
    summary, works = sheets(first, locale="en")
    assert summary.title == "Progress Payment Certificate"
    assert works.title == "List of Works Done"
    assert find_row(works, "Contract Unit Price", column=5)
    assert find_row(summary, "VAT ( E x 10% )", column=2)


def test_a_description_that_looks_like_a_formula_is_stored_as_text(first: Certificate) -> None:
    from dataclasses import replace

    from app.modules.contracts.hakedis import compute_certificate

    lines = list(first.inp.lines)
    lines[0] = replace(lines[0], description="=SUM(1;2) boru")
    cert = compute_certificate(replace(first.inp, lines=lines))
    _, works = sheets(cert)
    row = find_row(works, "M-25.101", column=2)
    assert works.cell(row=row, column=3).value == "=SUM(1;2) boru"
    assert works.cell(row=row, column=3).data_type == "s"
