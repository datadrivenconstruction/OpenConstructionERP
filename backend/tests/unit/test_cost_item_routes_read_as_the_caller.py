# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Every cost route answers with the rows its caller may read, and no others.

The rule lives in ``app.core.cost_scope`` and is pinned against query shapes
in ``test_cost_items_are_read_as_the_caller``. This file proves it over HTTP,
through the real request dependency and the real router, for each route that
returns cost item content: a manager's imported price list, a second user who
must not find it by any road, and an admin who must.

The routes are asked in an order that warms each shared cache as the owner
first and again as the other user first: what a route caches for everyone is
held to the same rule as what it answers directly.
"""

from __future__ import annotations

import io
import uuid
from typing import Any

import pytest
import pytest_asyncio
from fastapi import BackgroundTasks, Depends, FastAPI, HTTPException
from fastapi.responses import StreamingResponse
from httpx import ASGITransport, AsyncClient, Response
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import cost_scope
from app.dependencies import (
    get_current_user_payload,
    get_optional_user_payload,
    get_session,
    rls_request_context,
)
from app.modules.costs import router as costs_router
from app.modules.costs.models import CostCatalog, CostItem
from app.modules.users.models import User
from tests._pg import transactional_session

_PERMISSIONS = ["costs.create", "costs.delete", "costs.list", "costs.read", "costs.update"]

# Strings that exist only in the private catalogue. None may reach another user.
_PRIVATE_CODE = "ZZPRIVCODE"
_PRIVATE_TEXT = "zzprivatewall"
_PRIVATE_CATEGORY = "Zzprivatecategory"
_SHARED_TEXT = "zzsharedwall"


class World:
    session: AsyncSession
    users: dict[str, dict[str, Any]]
    region: str
    shared_region: str
    private_ids: list[str]
    shared_id: str
    catalog_id: str
    who: str | None = None

    @property
    def private_markers(self) -> list[str]:
        return [_PRIVATE_CODE, _PRIVATE_TEXT, _PRIVATE_CATEGORY, self.region, *self.private_ids]


async def _user(s: AsyncSession, role: str) -> dict[str, Any]:
    user = User(
        id=uuid.uuid4(),
        email=f"route-{uuid.uuid4().hex[:8]}@test.io",
        hashed_password="x" * 60,
        full_name="Route Test",
        role=role,
        locale="en",
        is_active=True,
        metadata_={},
    )
    s.add(user)
    await s.flush()
    return {"sub": str(user.id), "role": role, "permissions": list(_PERMISSIONS)}


@pytest_asyncio.fixture
async def world() -> World:
    costs_router._invalidate_cost_cache()
    async with transactional_session() as s:
        w = World()
        w.session = s
        w.users = {
            "alice": await _user(s, "manager"),
            "bob": await _user(s, "editor"),
            "admin": await _user(s, "admin"),
        }
        tag = uuid.uuid4().hex[:6].upper()
        w.region = f"ZZPRIVREGION{tag}"
        w.shared_region = f"ZZSHAREDREGION{tag}"
        catalog = CostCatalog(
            id=uuid.uuid4(), name=w.region, currency="EUR", created_by=uuid.UUID(w.users["alice"]["sub"])
        )
        s.add(catalog)
        await s.flush()
        private = [
            CostItem(
                id=uuid.uuid4(),
                code=f"{_PRIVATE_CODE}.{i:03d}",
                description=f"{_PRIVATE_TEXT} number {i}",
                unit="m2",
                rate="101.5",
                currency="EUR",
                region=w.region,
                catalog_id=catalog.id,
                classification={"collection": _PRIVATE_CATEGORY},
                is_active=True,
            )
            for i in range(3)
        ]
        shared = CostItem(
            id=uuid.uuid4(),
            code="ZZSHARED.001",
            description=f"{_SHARED_TEXT} number 0",
            unit="m2",
            rate="55",
            currency="EUR",
            region=w.shared_region,
            classification={"collection": "Zzsharedcategory"},
            is_active=True,
        )
        s.add_all([*private, shared])
        await s.flush()
        w.private_ids = [str(p.id) for p in private]
        w.shared_id = str(shared.id)
        w.catalog_id = str(catalog.id)
        s.expire_all()
        yield w
    costs_router._invalidate_cost_cache()


@pytest_asyncio.fixture
async def client(world: World) -> AsyncClient:
    """The costs router behind the same request scoping the real app mounts."""
    app = FastAPI(dependencies=[Depends(rls_request_context)])
    app.add_middleware(cost_scope.RequestScopeMiddleware)
    app.include_router(costs_router.router, prefix="/api/v1/costs")

    def _identity() -> dict[str, Any] | None:
        return world.users[world.who] if world.who else None

    async def _session() -> Any:
        yield world.session

    def _signed_in() -> dict[str, Any]:
        if world.who is None:
            raise HTTPException(status_code=401, detail="Not authenticated")
        return world.users[world.who]

    app.dependency_overrides[get_current_user_payload] = _signed_in
    app.dependency_overrides[get_optional_user_payload] = _identity
    app.dependency_overrides[get_session] = _session
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://costs.test") as c:
        yield c


def _text(resp: Response) -> str:
    if "spreadsheet" in resp.headers.get("content-type", ""):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(resp.content), read_only=True)
        return " ".join(str(cell) for sheet in book for row in sheet.iter_rows(values_only=True) for cell in row)
    return resp.text


# Routes asked for the shared row alone: nobody's private rows belong in the answer.
_SHARED_ONLY = {"search shared text"}


# (name, method, path, json body, does the answer carry the shared row's text)
def _routes(w: World) -> list[tuple[str, str, str, Any, str | None]]:
    return [
        ("search by text", "GET", f"/?q={_PRIVATE_TEXT[:9]}", None, None),
        ("search shared text", "GET", f"/?q={_SHARED_TEXT}", None, _SHARED_TEXT),
        ("search by code", "GET", f"/?q={_PRIVATE_CODE}", None, None),
        ("list a region", "GET", f"/?region={w.region}", None, None),
        ("list a catalogue", "GET", f"/?catalog_id={w.catalog_id}", None, None),
        ("list by category", "GET", f"/?category={_PRIVATE_CATEGORY}", None, None),
        ("autocomplete", "GET", f"/autocomplete/?q={_PRIVATE_TEXT[:9]}", None, None),
        ("autocomplete code", "GET", f"/autocomplete/?q={_PRIVATE_CODE}", None, None),
        ("regions", "GET", "/regions/", None, w.shared_region),
        ("region stats", "GET", "/regions/stats/", None, w.shared_region),
        ("categories", "GET", "/categories/", None, "Zzsharedcategory"),
        ("categories of a region", "GET", f"/categories/?region={w.region}", None, None),
        ("category tree", "GET", "/category-tree/", None, "Zzsharedcategory"),
        ("category tree of a region", "GET", f"/category-tree/?region={w.region}", None, None),
        ("loaded databases", "GET", "/loaded-databases/", None, None),
        ("get by id", "GET", f"/{w.private_ids[0]}", None, None),
        ("apply preview", "POST", f"/{w.private_ids[0]}/apply-preview/", {"quantity": "2"}, None),
        ("export", "GET", "/actions/export-excel/", None, "ZZSHARED.001"),
        (
            "match",
            "POST",
            "/match/",
            {"query": f"{_PRIVATE_TEXT} number 1", "unit": "m2", "mode": "lexical", "region": w.region},
            None,
        ),
        (
            "suggest for element",
            "POST",
            "/suggest-for-element/",
            {"name": f"{_PRIVATE_TEXT} number", "element_type": _PRIVATE_TEXT, "region": w.region},
            None,
        ),
        ("certainty", "GET", f"/{w.private_ids[0]}/certainty/", None, None),
        ("certainty batch", "POST", "/certainty/batch/", {"ids": [*w.private_ids, w.shared_id]}, w.shared_id),
        ("semantic search", "GET", f"/vector/search/?q={_PRIVATE_TEXT}", None, _SHARED_TEXT),
        ("semantic autocomplete", "GET", f"/autocomplete/?q={_PRIVATE_TEXT}&semantic=true", None, _SHARED_TEXT),
    ]


def _fake_vector_store(monkeypatch: pytest.MonkeyPatch, w: World) -> None:
    """A vector store that already holds every row, as an existing install does."""
    from app.core import vector

    hits = [
        {
            "id": pid,
            "score": 0.9,
            "code": f"{_PRIVATE_CODE}.{i:03d}",
            "description": f"{_PRIVATE_TEXT} number {i}",
            "unit": "m2",
            "rate": 101.5,
            "region": w.region,
        }
        for i, pid in enumerate(w.private_ids)
    ]
    hits.append(
        {
            "id": w.shared_id,
            "score": 0.8,
            "code": "ZZSHARED.001",
            "description": f"{_SHARED_TEXT} number 0",
            "unit": "m2",
            "rate": 55.0,
            "region": w.shared_region,
        }
    )
    monkeypatch.setattr(vector, "encode_texts", lambda texts: [[0.0] for _ in texts])
    monkeypatch.setattr(vector, "vector_search", lambda *a, **k: [dict(h) for h in hits])
    monkeypatch.setattr(vector, "vector_status", lambda: {"connected": True, "cost_collection": {"vectors": 4}})


@pytest.mark.asyncio
async def test_no_cost_route_shows_a_private_catalogue_to_another_user(
    world: World, client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_vector_store(monkeypatch, world)
    failures: list[str] = []
    for name, method, path, body, shared_marker in _routes(world):
        # Both cache-warming orders: the owner first, then the other user first.
        for order in (("alice", "bob", "admin", None), ("bob", "alice", None, "admin")):
            costs_router._invalidate_cost_cache()
            for who in order:
                world.who = who
                # The app gives every request its own session; one shared
                # session here must not carry loaded rows from caller to caller.
                world.session.expunge_all()
                resp = await client.request(method, f"/api/v1/costs{path}", json=body)
                label = f"{name} as {who or 'anonymous'} (order {'/'.join(str(o) for o in order)})"
                if who is None and resp.status_code in (401, 403):
                    continue
                text = _text(resp)
                private_seen = [m for m in world.private_markers if m in text]
                if resp.status_code == 404 and who not in ("alice", "admin"):
                    # "Not found" may repeat the id the caller sent; it carries nothing else.
                    private_seen = [m for m in private_seen if m not in path]
                if who in ("alice", "admin"):
                    if resp.status_code != 200:
                        failures.append(f"{label}: owner or admin got {resp.status_code} {text[:120]}")
                    elif not private_seen and name not in _SHARED_ONLY:
                        failures.append(f"{label}: owner or admin lost their own rows")
                else:
                    if private_seen:
                        failures.append(f"{label}: showed {private_seen[:2]}")
                    if resp.status_code >= 500:
                        failures.append(f"{label}: {resp.status_code} {text[:120]}")
                if shared_marker and resp.status_code == 200 and shared_marker not in text:
                    failures.append(f"{label}: lost the shared row")
    assert not failures, "\n".join(failures)


@pytest.mark.asyncio
async def test_totals_do_not_count_rows_the_caller_cannot_read(world: World, client: AsyncClient) -> None:
    """A total that includes hidden rows tells the caller they exist."""
    totals: dict[str, int] = {}
    for who in ("bob", "alice", "bob", "admin"):
        world.who = who
        # Warm the cached per-region totals the unfiltered list reads.
        await client.get("/api/v1/costs/regions/stats/")
        resp = await client.get("/api/v1/costs/?limit=1")
        assert resp.status_code == 200, resp.text
        total = resp.json()["total"]
        assert totals.setdefault(who, total) == total, f"{who}'s total changed with who asked before"
    assert totals["alice"] == totals["bob"] + 3
    assert totals["admin"] == totals["alice"]


@pytest.mark.asyncio
async def test_a_code_held_by_a_hidden_row_is_refused_not_crashed(world: World, client: AsyncClient) -> None:
    """The unique key is installation-wide; a hidden holder is a 409, not a 500."""
    world.who = "bob"
    resp = await client.post(
        "/api/v1/costs/",
        json={"code": f"{_PRIVATE_CODE}.000", "description": "mine", "unit": "m2", "rate": "1", "region": world.region},
    )
    assert resp.status_code == 409, resp.text


@pytest.mark.asyncio
async def test_a_streamed_body_and_a_background_task_still_read_as_the_caller(world: World) -> None:
    """Both run after the request dependency has exited; neither may read as the system."""
    seen: dict[str, Any] = {}
    app = FastAPI(dependencies=[Depends(rls_request_context)])
    app.add_middleware(cost_scope.RequestScopeMiddleware)

    @app.get("/probe")
    async def probe(background: BackgroundTasks) -> StreamingResponse:
        seen["route"] = (cost_scope.is_system(), cost_scope.current_caller())

        async def body() -> Any:
            seen["stream"] = (cost_scope.is_system(), cost_scope.current_caller())
            yield b"ok"

        def later() -> None:
            seen["background"] = (cost_scope.is_system(), cost_scope.current_caller())

        background.add_task(later)
        return StreamingResponse(body(), background=background)

    bob = world.users["bob"]
    app.dependency_overrides[get_optional_user_payload] = lambda: bob
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://probe.test") as c:
        assert (await c.get("/probe")).status_code == 200
    expected = (False, uuid.UUID(bob["sub"]))
    assert seen == {"route": expected, "stream": expected, "background": expected}
    assert cost_scope.is_system(), "the request left its reader bound after it ended"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "path",
    ["/", "/autocomplete/?q=wall", "/regions/", "/categories/", "/category-tree/", "/vector/search/?q=wall"],
)
async def test_browsing_the_cost_base_needs_a_signed_in_caller(world: World, client: AsyncClient, path: str) -> None:
    world.who = None
    assert (await client.get(f"/api/v1/costs{path}")).status_code == 401


@pytest.mark.asyncio
async def test_a_request_nothing_identified_reads_shared_rows_only() -> None:
    """Inside a request the default reader is anonymous, never the system."""
    app = FastAPI()
    app.add_middleware(cost_scope.RequestScopeMiddleware)

    @app.get("/probe")
    async def probe() -> dict[str, Any]:
        return {"system": cost_scope.is_system(), "caller": cost_scope.current_caller()}

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://probe.test") as c:
        assert (await c.get("/probe")).json() == {"system": False, "caller": None}
    assert cost_scope.is_system()


def test_the_real_app_scopes_every_request() -> None:
    """The middleware and the global dependency are what make it 'every route'."""
    import inspect

    from app import main

    source = inspect.getsource(main)
    assert "app.add_middleware(RequestScopeMiddleware)" in source
    assert "Depends(rls_request_context)" in source
    assert "cost_scope.bind_caller" in inspect.getsource(rls_request_context)
