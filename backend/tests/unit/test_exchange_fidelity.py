# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Exchange fidelity report: every format's loss cases, from the writers' own traces.

Each test builds a small bill, runs the real writer (GAEB, BC3) or the real
reader (GAEB, BC3, spreadsheet), and checks that what the format cannot carry
shows up as an item with the right kind, severity, field and values, and
that a clean bill reports nothing it should not.
"""

from __future__ import annotations

import io
import uuid
import zipfile
from decimal import Decimal
from types import SimpleNamespace

import pytest

from app.modules.boq.exchange_fidelity import (
    MAX_STORED_ITEMS,
    FidelityItem,
    FidelityReport,
    TreeRow,
    analyze_bc3_export,
    analyze_gaeb_export,
    analyze_import,
    analyze_sheet_export,
    flattened_levels,
    render_csv,
    render_xlsx,
    roundtrip_diff,
)
from app.modules.boq.exporters.bc3 import build_bc3
from app.modules.boq.importers import ImportedBOQ, ImportedPosition
from app.modules.boq.importers.bc3 import BC3Importer
from app.modules.boq.importers.gaeb_xml import GAEBXMLImporter
from app.modules.boq.router import build_gaeb_xml


def _pos(ordinal: str, description: str = "Concrete wall", unit: str = "m3", qty: str = "10", rate: str = "100", **kw):
    q, r = Decimal(qty), Decimal(rate)
    return SimpleNamespace(
        id=kw.pop("id", uuid.uuid4()),
        ordinal=ordinal,
        description=description,
        unit=unit,
        quantity=q,
        unit_rate=r,
        total=kw.pop("total", (q * r).quantize(Decimal("0.0001"))),
        classification=kw.pop("classification", {}),
        metadata=kw.pop("metadata", {}),
        **kw,
    )


def _markup(name: str = "Overhead", pct: str = "10", amount: str = "100"):
    return SimpleNamespace(
        id=uuid.uuid4(),
        name=name,
        percentage=Decimal(pct),
        amount=Decimal(amount),
        apply_to="direct_cost",
        markup_type="percentage",
        is_active=True,
        category="overhead",
    )


def _bill(positions, *, sections=None, markups=None, name="Shell works"):
    sections = (
        sections
        if sections is not None
        else [SimpleNamespace(id=uuid.uuid4(), ordinal="01", description="Shell", positions=positions)]
    )
    leaves = [p for s in sections for p in s.positions] + ([] if sections else positions)
    direct = sum((p.total for p in leaves), Decimal(0))
    return SimpleNamespace(
        id=uuid.uuid4(),
        project_id=uuid.uuid4(),
        name=name,
        sections=sections,
        positions=[] if sections else positions,
        markups=markups or [],
        direct_cost=direct,
        net_total=direct,
        grand_total=direct,
        metadata={},
    )


def _kinds(report: FidelityReport) -> set[tuple[str, str, str]]:
    return {(i.kind, i.field, i.severity) for i in report.items}


def _gaeb(bill, variant="x83", project="Kita Nord"):
    trace: list = []
    xml = build_gaeb_xml(bill, project_name=project, project_currency="EUR", gaeb_format=variant, trace=trace)
    return xml, analyze_gaeb_export(bill, trace, variant=variant, project_name=project)


# ── GAEB ───────────────────────────────────────────────────────────────────


def test_a_clean_gaeb_bill_reports_only_what_the_phase_omits_by_design() -> None:
    bill = _bill([_pos("01.0010", qty="12.5", rate="80")])
    _xml, report = _gaeb(bill)
    assert {(i.kind, i.field) for i in report.items} == {("dropped_attribute", "unit_rate")}
    assert report.items[0].params.get("by_design") is True


def test_gaeb_names_rounded_quantity_truncated_unit_changed_ordinal_and_dropped_resources() -> None:
    pos = _pos(
        "01.A-1/2",
        unit="pallets",
        qty="1.23456",
        metadata={"resources": [{"name": "Mason", "type": "labor", "quantity": 2, "unit_rate": 40}]},
    )
    bill = _bill([pos], name="A bill name longer than twenty characters")
    _xml, report = _gaeb(bill)
    by_field = {(i.kind, i.field): i for i in report.items}

    qty = by_field[("rounded_value", "quantity")]
    assert (qty.original, qty.written, qty.position_ref) == ("1.23456", "1.235", "01.A-1/2")
    unit = by_field[("truncated_text", "unit")]
    assert (unit.original, unit.written, unit.severity) == ("pallets", "pall", "warning")
    oz = by_field[("changed_identifier", "ordinal")]
    assert oz.written == "01.A_1_2"
    assert by_field[("dropped_attribute", "resources")].original == "1 resources"
    name = by_field[("truncated_text", "boq.name")]
    assert len(name.written) == 20


def test_gaeb_x84_flags_a_total_the_rounded_price_cannot_reproduce() -> None:
    pos = _pos("01.0010", qty="3", rate="0.3333", total=Decimal("1.00"))
    bill = _bill([pos])
    _xml, report = _gaeb(bill, variant="x84")
    by_field = {(i.kind, i.field): i for i in report.items}
    assert ("rounded_value", "unit_rate") in by_field
    assert by_field[("dropped_attribute", "description")].params["by_design"] is True


def test_gaeb_x84_total_drift_beyond_one_percent_is_critical() -> None:
    pos = _pos("01.0010", qty="0.0004", rate="1000", total=Decimal("0.40"))
    _xml, report = _gaeb(_bill([pos]), variant="x84")
    totals = [i for i in report.items if i.field == "total"]
    assert totals and totals[0].severity == "error"


def test_gaeb_quantity_beyond_the_schema_range_is_an_error() -> None:
    _xml, report = _gaeb(_bill([_pos("01.0010", qty="123456789", rate="1")]))
    assert ("value_out_of_range", "quantity", "error") in _kinds(report)


def test_gaeb_x83_lists_each_markup_as_not_carried() -> None:
    _xml, report = _gaeb(_bill([_pos("01.0010")], markups=[_markup()]))
    markups = [i for i in report.items if i.kind == "markup_not_representable"]
    assert len(markups) == 1 and markups[0].original == "10% on direct_cost"


def test_gaeb_control_characters_are_reported_as_altered_text() -> None:
    _xml, report = _gaeb(_bill([_pos("01.0010", description="Wall\x0bpanel")]))
    assert ("altered_text", "description", "info") in _kinds(report)


def test_nested_sections_and_child_lines_are_lost_hierarchy() -> None:
    rows = [
        TreeRow(id="s1", parent_id=None, ordinal="01", is_section=True),
        TreeRow(id="s2", parent_id="s1", ordinal="01.01", is_section=True),
        TreeRow(id="p1", parent_id="s2", ordinal="01.01.0010", is_section=False),
        TreeRow(id="p2", parent_id="p1", ordinal="01.01.0010.1", is_section=False),
    ]
    lost = {row.id for row, _parent in flattened_levels(rows)}
    assert lost == {"s2", "p2"}
    bill = _bill([_pos("01.0010")])
    trace: list = []
    build_gaeb_xml(bill, project_name="P", project_currency="EUR", gaeb_format="x83", trace=trace)
    report = analyze_gaeb_export(bill, trace, variant="x83", tree=rows)
    assert sum(1 for i in report.items if i.kind == "lost_hierarchy") == 2


# ── BC3 ────────────────────────────────────────────────────────────────────


def _bc3(bill):
    trace: list = []
    data, _charset = build_bc3(bill, project_name="Obra", project_currency="EUR", trace=trace)
    return data, analyze_bc3_export(bill, trace)


def test_bc3_drops_markups_and_negative_quantities_and_cleans_reserved_characters() -> None:
    bill = _bill(
        [
            _pos("01.001", description="Muro | hormigón"),
            _pos("01.002", qty="-5", rate="10"),
            _pos("01.003", description="Line one\nline two"),
        ],
        markups=[_markup("Gastos generales", "13", "130")],
    )
    _data, report = _bc3(bill)
    kinds = _kinds(report)
    assert ("markup_not_representable", "markup", "warning") in kinds
    assert ("value_out_of_range", "quantity", "error") in kinds
    altered = {i.position_ref: i for i in report.items if i.kind == "altered_text"}
    assert altered["01.001"].severity == "warning" and altered["01.001"].written == "Muro / hormigón"
    assert altered["01.003"].severity == "info"


def test_bc3_names_a_changed_code_an_unmapped_unit_and_dropped_classification() -> None:
    bill = _bill(
        [
            _pos("01 002", unit="Palette", classification={"din276": "330"}),
            _pos("01.004", metadata={"resources": [{"name": "Oficial", "quantity": 1}]}),
        ]
    )
    _data, report = _bc3(bill)
    by_field = {(i.kind, i.field): i for i in report.items}
    assert by_field[("changed_identifier", "ordinal")].written == "01002"
    assert by_field[("unmapped_unit", "unit")].original == "Palette"
    assert by_field[("dropped_attribute", "classification")].original == "din276"
    assert by_field[("dropped_attribute", "resources")].position_ref == "01.004"


def test_bc3_roundtrip_of_a_clean_bill_matches() -> None:
    import asyncio

    bill = _bill([_pos("01.001", qty="2.5", rate="12.75"), _pos("01.002", unit="kg", qty="7", rate="3")])
    data, _report = _bc3(bill)
    back = asyncio.run(BC3Importer.parse(data))
    assert roundtrip_diff(bill, back) == []


def test_bc3_roundtrip_catches_the_lost_negative_quantity() -> None:
    import asyncio

    bill = _bill([_pos("01.001", qty="-5", rate="10")])
    data, _report = _bc3(bill)
    back = asyncio.run(BC3Importer.parse(data))
    diff = roundtrip_diff(bill, back)
    assert any(i.field == "quantity" and i.severity == "error" for i in diff)


def test_gaeb_x83_roundtrip_of_a_clean_bill_matches_without_prices() -> None:
    import asyncio

    bill = _bill([_pos("01.0010", qty="12.5", rate="80"), _pos("01.0020", unit="m2", qty="3", rate="9")])
    xml, _report = _gaeb(bill)
    back = asyncio.run(GAEBXMLImporter.parse(xml.encode("utf-8")))
    assert roundtrip_diff(bill, back, compare_prices=False) == []


def test_roundtrip_reports_a_line_that_did_not_come_back() -> None:
    bill = _bill([_pos("01.001"), _pos("01.002")])
    back = ImportedBOQ(
        positions=[
            ImportedPosition(description="Concrete wall", ordinal="01.001", unit="m3", quantity=10, unit_rate=100)
        ]
    )
    diff = roundtrip_diff(bill, back)
    assert diff[0].field == "row" and diff[0].severity == "error"
    assert any(i.field == "row" and i.position_ref == "01.002" for i in diff)


# ── spreadsheets ───────────────────────────────────────────────────────────


def test_excel_flags_formula_prefix_cell_limit_precision_and_markup_rule() -> None:
    bill = _bill(
        [
            _pos("01.001", description="=SUM(A1)"),
            _pos("01.002", description="x" * 40000),
            _pos("01.003", qty="1.23456789012345678", rate="1"),
        ],
        markups=[_markup()],
    )
    report = analyze_sheet_export(bill, format_id="excel")
    kinds = _kinds(report)
    assert ("altered_text", "description", "info") in kinds
    assert ("truncated_text", "description", "warning") in kinds
    assert ("rounded_value", "quantity", "info") in kinds
    assert ("markup_not_representable", "markup", "info") in kinds


def test_csv_keeps_long_text_and_precision() -> None:
    bill = _bill([_pos("01.002", description="x" * 40000), _pos("01.003", qty="1.23456789012345678")])
    assert analyze_sheet_export(bill, format_id="csv").items == []


# ── import ─────────────────────────────────────────────────────────────────

_GAEB_83 = """<?xml version="1.0" encoding="UTF-8"?>
<GAEB xmlns="http://www.gaeb.de/GAEB_DA_XML/DA83/3.3">
  <Award><DP>83</DP><Cur>EUR</Cur>
    <BoQ><BoQBody><BoQCtgy RNoPart="01"><LblTx><p><span>Rohbau</span></p></LblTx>
      <BoQBody><Itemlist>
        <Item RNoPart="0010"><Qty>10.000</Qty><QU>m3</QU><CtlgAssign/><Description><CompleteText><DetailTxt><Text><p><span>{text}</span></p></Text></DetailTxt></CompleteText></Description></Item>
        <Remark/>
      </Itemlist></BoQBody></BoQCtgy></BoQBody></BoQ>
  </Award>
  <Attachment/>
</GAEB>"""


def test_gaeb_import_reports_unread_elements_attachments_and_cut_text() -> None:
    import asyncio

    content = _GAEB_83.format(text="y" * 6000).encode("utf-8")
    imported = asyncio.run(GAEBXMLImporter.parse(content))
    report = analyze_import("gaeb_xml", imported, content)
    tags = {(i.kind, i.field) for i in report.items}
    assert ("unsupported_element", "CtlgAssign") in tags
    assert ("unsupported_element", "Remark") in tags
    assert ("attachment_skipped", "Attachment") in tags
    assert ("truncated_text", "description") in tags


_BC3 = "\r\n".join(
    [
        "~V|Obra|FIEBDC-3/2020|Prog|01012026|Obra|ANSI|2||GBP|",
        "~K|\\2\\2\\3\\2\\2\\2\\2\\EUR|",
        "~C|OBRA##||Obra|0|||3|",
        "~C|C01#||Capitulo|0|||1|",
        "~C|C0101#||Sub|0|||1|",
        "~C|P1|m3|Muro|100|||0|",
        "~C|MO1|h|Oficial|20|||1|",
        "~D|OBRA##|C01#\\1\\1|",
        "~D|C01#|C0101#\\1\\1|",
        "~D|C0101#|P1\\1\\10|",
        "~D|P1|MO1\\1\\2|",
        "~G|P1|foto.jpg|",
        "~M|C0101#\\P1|1|10||",
    ]
)


def test_bc3_import_reports_graphics_coefficients_components_nesting_and_currency() -> None:
    import asyncio

    imported = asyncio.run(BC3Importer.parse(_BC3.encode("cp1252")))
    report = analyze_import("bc3", imported, None)
    tags = {(i.kind, i.field) for i in report.items}
    assert ("attachment_skipped", "~G") in tags
    assert ("markup_not_representable", "~K") in tags
    assert ("dropped_attribute", "resources") in tags
    assert ("lost_hierarchy", "hierarchy") in tags
    assert ("dropped_attribute", "currency") in tags


def test_import_storage_limits_rounding_units_markups_columns_and_media() -> None:
    imported = ImportedBOQ(
        positions=[
            ImportedPosition(description="Wall", ordinal="9" * 60, unit="m3", quantity=1.123456, unit_rate=10),
            ImportedPosition(
                description="Door", ordinal="2", unit="pcs", quantity=1, metadata={"gaeb_unit_original": ""}
            ),
            ImportedPosition(description="Hatch", ordinal="3", unit="Luke", quantity=1),
        ],
        errors=[{"ordinal": "4", "error": "Quantity out of range"}],
        warnings=[{"ordinal": "5", "warning": "odd"}],
        skipped=2,
        metadata={
            "markup_items": [{"ordinal": "Z1", "it": "120", "percentage": "8"}],
            "original_columns": ["Pos", "Text", "Colour"],
            "column_mapping": {"0": "ordinal", "1": "description"},
        },
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/workbook.xml", "<x/>")
        zf.writestr("xl/media/image1.png", b"png")
    report = analyze_import("excel", imported, buf.getvalue())
    tags = {(i.kind, i.field, i.severity) for i in report.items}
    assert ("exceeds_limit", "ordinal", "error") in tags
    assert ("rounded_value", "quantity", "info") in tags
    assert ("unmapped_unit", "unit", "warning") in tags
    assert ("unmapped_unit", "unit", "info") in tags
    assert ("markup_not_representable", "markup", "info") in tags
    assert ("dropped_attribute", "column", "info") in tags
    assert ("attachment_skipped", "attachment", "warning") in tags
    assert ("rejected_row", "row", "error") in tags
    assert ("parser_warning", "row", "warning") in tags
    assert ("skipped_row", "row", "info") in tags


# ── report object and downloads ────────────────────────────────────────────


def test_summary_counts_every_item_and_storage_caps_the_list() -> None:
    report = FidelityReport(direction="export", format_id="csv")
    for n in range(MAX_STORED_ITEMS + 5):
        report.add("info", "rounded_value", "quantity", position_ref=str(n))
    report.add("error", "rejected_row", "row")
    out = report.to_dict()
    assert out["summary"]["total"] == MAX_STORED_ITEMS + 6
    assert out["summary"]["critical"] == 1
    assert out["items_truncated"] is True
    assert len(out["items"]) == MAX_STORED_ITEMS
    assert out["items"][0]["severity"] == "error"


def test_an_importer_note_with_unknown_values_is_kept_as_a_reader_warning() -> None:
    item = FidelityItem.from_dict({"severity": "fatal", "kind": "mystery", "count": 0})
    assert (item.severity, item.kind, item.count) == ("warning", "parser_warning", 1)


@pytest.mark.parametrize("renderer", ["csv", "xlsx"])
def test_downloads_label_and_defuse_cells(renderer: str) -> None:
    items = [FidelityItem("error", "rejected_row", "row", position_ref="=1+1", detail="bad").to_dict()]
    headers = ["Severity", "Kind", "Position", "Field", "Original", "Written", "Count", "Detail"]

    def label(group: str, value: str) -> str:
        return f"{group}:{value}"

    if renderer == "csv":
        text = render_csv(items, headers, label).decode("utf-8")
        assert text.startswith("﻿Severity")
        assert "severity:error" in text and "'=1+1" in text
    else:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(render_xlsx(items, headers, label, title="Report")))
        ws = wb.active
        assert ws["A2"].value == "severity:error"
        assert ws["C2"].value == "'=1+1"
