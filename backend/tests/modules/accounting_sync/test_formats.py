# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Mapping​‌‍⁠​‌‍⁠​‌‍⁠​‌‍⁠, journal, DATEV EXTF and SAF-T formats, with no database."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from xml.etree import ElementTree as ET  # noqa: N817

from openpyxl import load_workbook

from app.modules.accounting_sync import templates
from app.modules.accounting_sync.exports import datev, journal, saft
from app.modules.accounting_sync.postings import (
    Mapping,
    SyncContact,
    SyncDocument,
    SyncLine,
    build_journal,
    journal_for,
    personal_accounts,
)

PERIOD = (date(2026, 1, 1), date(2026, 12, 31))


def _mapping(key: str = "DE-SKR03") -> Mapping:
    tpl = templates.get_template(key)
    return Mapping(
        accounts={c: tuple(v[:2]) for c, v in tpl["accounts"].items()},
        tax_keys={(code, d): v for d, codes in tpl["tax_keys"].items() for code, v in codes.items()},
    )


def _ar() -> SyncDocument:
    return SyncDocument(
        record_type="ar_invoice",
        record_id="inv-1",
        number="AR-2026-001",
        doc_date=date(2026, 3, 31),
        due_date=date(2026, 4, 30),
        contact=SyncContact(record_id="c1", name="Müller Bau GmbH", kind="customer", tax_id="DE123456789"),
        currency="EUR",
        lines=[
            SyncLine("Abschlag 3 Müller", Decimal("10000.00"), "revenue", "VAT", Decimal("19")),
            SyncLine("Materials", Decimal("333.33"), "revenue_reduced", "VAT_REDUCED", Decimal("7")),
        ],
        source="interim_payment",
        project_code="P-100",
        project_name="Riverside",
    )


def _ap() -> SyncDocument:
    return SyncDocument(
        record_type="ap_bill",
        record_id="bill-1",
        number="SUB-17",
        doc_date=date(2026, 4, 2),
        contact=SyncContact(record_id="s1", name="Steelworks AG", kind="supplier"),
        currency="EUR",
        lines=[SyncLine("Steel erection", Decimal("5000.10"), "subcontract_cost", "VAT", Decimal("19"))],
        source="subcontractor",
        project_code="P-100",
    )


def test_templates_come_from_data_and_tax_codes_from_the_seed() -> None:
    keys = {t["key"] for t in templates.list_templates("DE")}
    assert {"generic", "DE-SKR03", "DE-SKR04"} <= keys
    assert {t["key"] for t in templates.list_templates("FR")} == {"generic"}
    codes = {r["tax_code"]: r["rate_pct"] for r in templates.tax_codes("DE", date(2026, 1, 1))}
    assert Decimal(codes["VAT"]) == Decimal("19") and Decimal(codes["VAT_REDUCED"]) == Decimal("7")
    assert codes["RC"] == "0" and codes["ZERO"] == "0"


def test_tax_is_exact_decimal_and_rounded_half_up() -> None:
    doc = _ar()
    assert doc.line_tax(doc.lines[1]) == Decimal("23.33")
    assert doc.gross_total == Decimal("10333.33") + Decimal("1900.00") + Decimal("23.33")


def test_journal_balances_and_maps_accounts() -> None:
    lines = build_journal(_ar(), _mapping())
    assert sum(ln.debit for ln in lines) == sum(ln.credit for ln in lines)
    assert lines[0].account == "1400" and lines[0].debit == Decimal("12256.66")
    assert {ln.account for ln in lines} == {"1400", "8400", "8300", "1776"}
    bill = build_journal(_ap(), _mapping())
    assert bill[-1].account == "1576" and bill[0].credit == Decimal("5950.12")


def test_journal_csv_columns_are_configurable() -> None:
    lines = journal_for([_ar(), _ap()], _mapping())
    text = journal.to_csv(lines, ["account", "debit", "credit", "nope"]).decode("utf-8")
    assert text.startswith("﻿account,debit,credit\r\n")
    assert "1400,12256.66,0.00" in text


def test_journal_xlsx_keeps_amounts_numeric(tmp_path) -> None:
    content = journal.to_xlsx(journal_for([_ar()], _mapping()), ["account", "debit"])
    path = tmp_path / "j.xlsx"
    path.write_bytes(content)
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == ["account", "debit"]
    assert Decimal(str(ws["B2"].value)) == Decimal("12256.66")


def test_personal_accounts_stay_in_their_ranges() -> None:
    acc = personal_accounts([_ar().contact, _ap().contact], (10000, 69999), (70000, 99999))
    assert acc == {"customer:c1": 10000, "supplier:s1": 70000}


def test_datev_extf_header_field_order_and_encoding() -> None:
    docs = [_ar(), _ap()]
    accounts = personal_accounts([d.contact for d in docs], (10000, 69999), (70000, 99999))
    raw = datev.buchungsstapel(
        docs,
        _mapping(),
        accounts,
        settings={"consultant_number": "1001", "client_number": "7"},
        period=PERIOD,
        skr="03",
    )
    assert raw.endswith(b"\r\n")
    text = raw.decode("cp1252")
    assert "Müller" in text and "Müller".encode("cp1252") in raw  # cp1252, not UTF-8
    rows = text.split("\r\n")
    head = rows[0].split(";")
    assert head[:5] == ['"EXTF"', "700", "21", '"Buchungsstapel"', "13"]
    assert head[10:16] == ["1001", "7", "20260101", "4", "20260101", "20261231"]
    assert head[21] == '"EUR"' and head[26] == '"03"' and len(head) == 31
    cols = rows[1].split(";")
    assert cols[:11] == [
        "Umsatz (ohne Soll/Haben-Kz)", "Soll/Haben-Kennzeichen", "WKZ Umsatz", "Kurs", "Basis-Umsatz",
        "WKZ Basis-Umsatz", "Konto", "Gegenkonto (ohne BU-Schlüssel)", "BU-Schlüssel", "Belegdatum", "Belegfeld 1",
    ]  # fmt: skip
    assert len(cols) == 125
    first = rows[2].split(";")
    assert first[0] == "11900,00" and first[1] == '"S"' and first[6] == "10000" and first[7] == "8400"
    assert first[8] == '"3"' and first[9] == "3103" and first[10] == '"AR-2026-001"'
    assert first[cols.index("KOST1 - Kostenstelle")] == '"P-100"'
    bill = rows[4].split(";")
    assert bill[1] == '"H"' and bill[6] == "70000" and bill[7] == "3100" and bill[8] == '"9"'
    assert all(len(r.split(";")) == 125 for r in rows[2:] if r)


def test_datev_stammdaten_lists_debtors_and_creditors() -> None:
    docs = [_ar(), _ap()]
    accounts = personal_accounts([d.contact for d in docs], (10000, 69999), (70000, 99999))
    rows = datev.stammdaten([d.contact for d in docs], accounts, settings={}).decode("cp1252").split("\r\n")
    assert rows[0].split(";")[2:5] == ["16", '"Debitoren/Kreditoren"', "5"]
    assert rows[2].split(";")[0] == "10000" and '"DE"' in rows[2] and '"123456789"' in rows[2]
    assert rows[3].split(";")[0] == "70000"


def test_saft_is_balanced_oecd_xml() -> None:
    lines = journal_for([_ar(), _ap()], _mapping("generic"))
    xml = saft.build(
        lines,
        [_ar().contact, _ap().contact],
        company_name="Demo",
        country_code="NO",
        currency="NOK",
        period=PERIOD,
        tax_rows=[{"tax_code": "VAT", "name": "VAT", "rate_pct": "25"}],
    )
    root = ET.fromstring(xml)
    ns = {"s": saft.NS}
    assert root.find("s:Header/s:AuditFileCountry", ns).text == "NO"
    gle = root.find("s:GeneralLedgerEntries", ns)
    assert gle.find("s:NumberOfEntries", ns).text == "2"
    assert gle.find("s:TotalDebit", ns).text == gle.find("s:TotalCredit", ns).text
    assert root.find("s:MasterFiles/s:Customers/s:Customer/s:TaxRegistrationNumber", ns).text == "DE123456789"


def test_saft_variant_comes_from_country_data() -> None:
    no = templates.saft_variant("no")
    assert no["key"] == "NO" and no["audit_file_version"] == "1.30"
    assert templates.saft_variant("FR")["key"] == "oecd"
    xml = saft.build(
        journal_for([_ar()], _mapping("generic")),
        [_ar().contact],
        company_name="Demo",
        country_code="NO",
        currency="NOK",
        period=PERIOD,
        tax_rows=[],
        namespace=no["namespace"],
        audit_file_version=no["audit_file_version"],
    )
    root = ET.fromstring(xml)
    assert root.tag == f"{{{no['namespace']}}}AuditFile"
    assert root.find("s:Header/s:AuditFileVersion", {"s": no["namespace"]}).text == "1.30"


def test_every_template_maps_payroll() -> None:
    for key in ("generic", "DE-SKR03", "DE-SKR04"):
        accounts = templates.get_template(key)["accounts"]
        assert {"payroll_cost", "payroll_liability"} <= set(accounts)
