# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The submittal register over the API, on PostgreSQL.

What is held, end to end through the router and a real database:

* the register columns are written by create and PATCH and read back, and the
  reviewer's stamp cannot be written by a plain edit;
* a review stamps the outcome and the mark, the stamp survives closing, and a
  resubmission raises the revision, clears the stamp and keeps the replaced
  revision in the history (after revise and resubmit, and after a rejection);
* the list filters (type, discipline, outcome, mark, long lead, overdue for
  review, approval late), the sort, and the summary all stay inside the
  project asked for;
* another user's project is a 404 on the list, the summary and the export;
* a row with every new column NULL, as an upgraded installation holds it,
  lists, summarises and exports.

Pattern: the transaction-isolated session and ``httpx.AsyncClient`` over
``ASGITransport`` of ``test_review_cycle.py``. ``verify_project_access`` is the
real one, so project scoping is the application's and not a stand-in.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, datetime, timedelta

import httpx
import pytest
import pytest_asyncio
from fastapi import FastAPI
from openpyxl import load_workbook

from app.core.validation.rules import register_builtin_rules
from app.dependencies import get_current_user_id, get_current_user_payload, get_session
from app.modules.projects.models import Project
from app.modules.submittals.models import Submittal
from app.modules.submittals.permissions import register_submittals_permissions
from app.modules.submittals.router import router as submittals_router
from app.modules.submittals.rules import register_submittal_register_rules
from app.modules.users.models import User
from tests._pg import transactional_session

TODAY = datetime.now(UTC).date()


def _day(offset: int) -> str:
    return (TODAY + timedelta(days=offset)).isoformat()


@pytest_asyncio.fixture
async def db_session() -> AsyncIterator:
    async with transactional_session() as s:
        yield s


async def _make_user(session) -> str:
    user = User(email=f"u{uuid.uuid4().hex[:10]}@example.com", hashed_password="x")
    session.add(user)
    await session.flush()
    return str(user.id)


async def _make_project(session, owner_id: str, name: str = "Register Test Project") -> str:
    project = Project(name=name, owner_id=uuid.UUID(owner_id))
    session.add(project)
    await session.flush()
    return str(project.id)


def _client(db_session, caller_id: str) -> httpx.AsyncClient:
    """A client acting as ``caller_id`` with a manager's role.

    The role passes the permission and review gates. Which projects the
    caller may see is still decided by the real ``verify_project_access``
    from the users and projects in the database.
    """
    register_submittals_permissions()
    app = FastAPI()
    app.include_router(submittals_router, prefix="/v1/submittals")

    async def _session_override():
        yield db_session

    async def _user_override() -> str:
        return caller_id

    async def _payload_override() -> dict:
        return {"sub": caller_id, "role": "manager", "permissions": []}

    app.dependency_overrides[get_session] = _session_override
    app.dependency_overrides[get_current_user_id] = _user_override
    app.dependency_overrides[get_current_user_payload] = _payload_override
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")


async def _create(client: httpx.AsyncClient, project_id: str, **fields: object) -> dict:
    body = {"project_id": project_id, "title": "Chiller product data", "submittal_type": "product_data", **fields}
    response = await client.post("/v1/submittals/", json=body)
    assert response.status_code == 201, response.text
    return response.json()


async def _legacy_row(session, project_id: str, *, number: str, status: str = "approved") -> Submittal:
    """A row carrying only the columns that existed before the register ones."""
    row = Submittal(
        project_id=uuid.UUID(project_id),
        submittal_number=number,
        title="Legacy steel shop drawing",
        submittal_type="shop_drawing",
        status=status,
        current_revision=1,
        date_submitted="2026-08-01",
        date_returned="2026-08-15",
    )
    session.add(row)
    await session.flush()
    return row


async def _numbers(client: httpx.AsyncClient, project_id: str, **params: object) -> list[str]:
    response = await client.get("/v1/submittals/", params={"project_id": project_id, **params})
    assert response.status_code == 200, response.text
    return [row["submittal_number"] for row in response.json()]


# ── Fields ────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_the_register_columns_are_written_and_read_back(db_session) -> None:
    owner = await _make_user(db_session)
    project = await _make_project(db_session, owner)
    drawing = str(uuid.uuid4())
    async with _client(db_session, owner) as client:
        created = await _create(
            client,
            project,
            discipline="HVAC",
            manufacturer="Cooling Works",
            model_reference="CW-450",
            country_of_origin="tr",
            supplier="Cooling Works Trading",
            review_period_days=14,
            required_on_site_date=_day(120),
            long_lead=True,
            lead_time_weeks=12,
            linked_drawing_ids=[drawing],
        )
        assert created["discipline"] == "hvac"
        assert created["country_of_origin"] == "TR"
        assert created["supplier_name"] == "Cooling Works Trading"
        assert created["long_lead"] is True
        assert created["linked_drawing_ids"] == [drawing]
        assert created["approval_needed_by"] == _day(120 - 84)
        assert created["submit_by_date"] == _day(120 - 84 - 14)
        assert created["approval_late_days"] == 0
        # Not submitted: nothing about the review can be said yet.
        assert created["review_outcome"] is None
        assert created["review_code"] is None
        assert created["days_in_review"] is None
        assert created["review_overdue_days"] is None

        patched = await client.patch(
            f"/v1/submittals/{created['id']}",
            json={"manufacturer": "Other Works", "lead_time_weeks": 20, "long_lead": False},
        )
        assert patched.status_code == 200, patched.text
        body = patched.json()
        assert body["manufacturer"] == "Other Works"
        assert body["lead_time_weeks"] == 20
        assert body["long_lead"] is False
        # Fields left out of the PATCH are untouched.
        assert body["discipline"] == "hvac"
        assert body["model_reference"] == "CW-450"

        fetched = (await client.get(f"/v1/submittals/{created['id']}")).json()
        assert fetched["manufacturer"] == "Other Works"
        assert fetched["approval_needed_by"] == _day(120 - 140)
        assert fetched["approval_late_days"] == 20


@pytest.mark.asyncio
async def test_a_plain_edit_cannot_write_the_stamp(db_session) -> None:
    owner = await _make_user(db_session)
    project = await _make_project(db_session, owner)
    async with _client(db_session, owner) as client:
        created = await _create(client, project)
        response = await client.patch(
            f"/v1/submittals/{created['id']}",
            json={"review_outcome": "approved", "review_code": "A", "title": "Renamed"},
        )
        assert response.status_code == 200, response.text
        assert response.json()["title"] == "Renamed"
        assert response.json()["review_outcome"] is None
        assert response.json()["may_proceed"] is False


@pytest.mark.asyncio
async def test_a_malformed_register_value_is_a_422(db_session) -> None:
    owner = await _make_user(db_session)
    project = await _make_project(db_session, owner)
    async with _client(db_session, owner) as client:
        body = {"project_id": project, "title": "T", "submittal_type": "sample", "country_of_origin": "TUR"}
        assert (await client.post("/v1/submittals/", json=body)).status_code == 422
        assert (
            await client.get("/v1/submittals/", params={"project_id": project, "sort": "metadata"})
        ).status_code == 422
        assert (
            await client.get("/v1/submittals/", params={"project_id": project, "outcome": "maybe"})
        ).status_code == 422


# ── Behaviour: the outcome drives what happens next ───────────────────────


@pytest.mark.asyncio
async def test_revise_and_resubmit_then_approval_keeps_both_cycles(db_session) -> None:
    owner = await _make_user(db_session)
    project = await _make_project(db_session, owner)
    async with _client(db_session, owner) as client:
        sid = (await _create(client, project, review_period_days=14))["id"]
        first = (await client.post(f"/v1/submittals/{sid}/submit/")).json()
        assert first["current_revision"] == 1
        assert first["days_in_review"] == 0
        assert first["review_due_date"] == _day(14)
        assert first["review_overdue_days"] == 0

        returned = await client.post(
            f"/v1/submittals/{sid}/review/", json={"status": "revise_and_resubmit", "notes": "Fan curves missing"}
        )
        assert returned.status_code == 200, returned.text
        body = returned.json()
        assert body["status"] == "revise_and_resubmit"
        assert body["review_outcome"] == "revise_and_resubmit"
        assert body["review_code"] == "C"
        assert body["may_proceed"] is False
        assert [(e["revision"], e["code"], e["notes"]) for e in body["review_history"]] == [
            (1, "C", "Fan curves missing")
        ]

        again = (await client.post(f"/v1/submittals/{sid}/submit/")).json()
        assert again["status"] == "submitted"
        assert again["current_revision"] == 2
        # The new revision has not been reviewed; the old stamp is history.
        assert again["review_outcome"] is None
        assert again["review_code"] is None
        assert len(again["review_history"]) == 1

        approved = await client.post(f"/v1/submittals/{sid}/approve/", json={"notes": "OK"})
        assert approved.status_code == 200, approved.text
        body = approved.json()
        assert body["status"] == "approved"
        assert body["review_outcome"] == "approved"
        assert body["review_code"] == "A"
        assert body["may_proceed"] is True
        assert [(e["revision"], e["outcome"]) for e in body["review_history"]] == [
            (1, "revise_and_resubmit"),
            (2, "approved"),
        ]

        # The resubmission is linked to the revision it replaced, and the stamp
        # agrees with the status: neither register rule has anything to say.
        register_builtin_rules()
        register_submittal_register_rules()
        report = (await client.get(f"/v1/submittals/{sid}/validate/")).json()
        failed = {r["rule_id"] for r in report["results"] if not r["passed"]}
        assert "submittal.resubmission_linked" not in failed
        assert "submittal.outcome_matches_status" not in failed
        assert "submittal.resubmission_linked" in {r["rule_id"] for r in report["results"]}


@pytest.mark.asyncio
async def test_approved_as_noted_proceeds_keeps_the_reviewers_mark_and_survives_closing(db_session) -> None:
    owner = await _make_user(db_session)
    project = await _make_project(db_session, owner)
    async with _client(db_session, owner) as client:
        sid = (await _create(client, project))["id"]
        await client.post(f"/v1/submittals/{sid}/submit/")
        noted = await client.post(
            f"/v1/submittals/{sid}/review/",
            json={"status": "approved_as_noted", "code": "2", "resubmit_for_record": True},
        )
        assert noted.status_code == 200, noted.text
        body = noted.json()
        assert body["review_outcome"] == "approved_as_noted"
        assert body["review_code"] == "2"
        assert body["may_proceed"] is True
        assert body["resubmit_for_record"] is True

    row = await db_session.get(Submittal, uuid.UUID(sid))
    row.status = "closed"
    await db_session.flush()
    async with _client(db_session, owner) as client:
        closed = (await client.get(f"/v1/submittals/{sid}")).json()
        assert closed["status"] == "closed"
        assert closed["review_outcome"] == "approved_as_noted"
        assert closed["review_code"] == "2"
        assert await _numbers(client, project, outcome="approved_as_noted") == [closed["submittal_number"]]
        assert await _numbers(client, project, review_code="2") == [closed["submittal_number"]]


@pytest.mark.asyncio
async def test_a_rejected_revision_is_closed_and_the_next_submission_is_a_new_one(db_session) -> None:
    owner = await _make_user(db_session)
    project = await _make_project(db_session, owner)
    async with _client(db_session, owner) as client:
        sid = (await _create(client, project))["id"]
        await client.post(f"/v1/submittals/{sid}/submit/")
        rejected = (await client.post(f"/v1/submittals/{sid}/review/", json={"status": "rejected"})).json()
        assert rejected["review_code"] == "D"
        assert rejected["may_proceed"] is False

        redrafted = await client.patch(f"/v1/submittals/{sid}", json={"status": "draft"})
        assert redrafted.status_code == 200, redrafted.text
        # Back in draft, the answer to revision 1 still stands.
        assert redrafted.json()["review_outcome"] == "rejected"

        resubmitted = (await client.post(f"/v1/submittals/{sid}/submit/")).json()
        assert resubmitted["current_revision"] == 2
        assert resubmitted["review_outcome"] is None
        assert [(e["revision"], e["outcome"]) for e in resubmitted["review_history"]] == [(1, "rejected")]


# ── Filters, sorting, scoping ─────────────────────────────────────────────


async def _seeded_project(db_session, owner: str) -> tuple[str, dict[str, str]]:
    """One project with a row for every filter to find, and a second project beside it."""
    project = await _make_project(db_session, owner)
    other = await _make_project(db_session, owner, name="Other Project")
    numbers: dict[str, str] = {}
    async with _client(db_session, owner) as client:
        # A: HVAC material, long lead, under review past a 7 day period, and
        # past the date its approval was needed by.
        a = await _create(
            client,
            project,
            title="A chiller",
            discipline="hvac",
            manufacturer="Zeta",
            long_lead=True,
            lead_time_weeks=12,
            required_on_site_date=_day(60),
            review_period_days=7,
            status="submitted",
            date_submitted=_day(-10),
        )
        # B: electrical shop drawing, under review inside a 30 day period.
        b = await _create(
            client,
            project,
            title="B panel drawing",
            submittal_type="shop_drawing",
            discipline="electrical",
            manufacturer="Alpha",
            review_period_days=30,
            status="submitted",
            date_submitted=_day(-10),
        )
        # C: electrical material, approved as noted with the mark "2".
        c = await _create(client, project, title="C cable", discipline="electrical", long_lead=True, lead_time_weeks=4)
        await client.post(f"/v1/submittals/{c['id']}/submit/")
        await client.post(f"/v1/submittals/{c['id']}/review/", json={"status": "approved_as_noted", "code": "2"})
        # D: a draft with no discipline and nothing recorded.
        d = await _create(client, project, title="D draft", submittal_type="method_statement")
        # The same discipline in another project: must never be listed here.
        await _create(client, other, title="Other project HVAC", discipline="hvac", long_lead=True)
    legacy = await _legacy_row(db_session, project, number="SUB-900")
    numbers.update(a=a["submittal_number"], b=b["submittal_number"], c=c["submittal_number"], d=d["submittal_number"])
    numbers["legacy"] = legacy.submittal_number
    return project, numbers


@pytest.mark.asyncio
async def test_the_list_filters_stay_inside_the_project(db_session) -> None:
    owner = await _make_user(db_session)
    project, n = await _seeded_project(db_session, owner)
    async with _client(db_session, owner) as client:
        assert set(await _numbers(client, project)) == set(n.values())
        assert await _numbers(client, project, discipline="hvac") == [n["a"]]
        assert set(await _numbers(client, project, discipline="electrical")) == {n["b"], n["c"]}
        assert set(await _numbers(client, project, type="shop_drawing")) == {n["b"], n["legacy"]}
        assert await _numbers(client, project, type="shop_drawing", discipline="electrical") == [n["b"]]
        assert set(await _numbers(client, project, long_lead="true")) == {n["a"], n["c"]}
        assert await _numbers(client, project, outcome="approved_as_noted") == [n["c"]]
        assert await _numbers(client, project, review_code="2") == [n["c"]]
        # A row approved before the stamp was stored is found under its decision.
        assert await _numbers(client, project, outcome="approved") == [n["legacy"]]
        assert await _numbers(client, project, outcome="rejected") == []


@pytest.mark.asyncio
async def test_the_date_dependent_filters_and_their_opposites(db_session) -> None:
    owner = await _make_user(db_session)
    project, n = await _seeded_project(db_session, owner)
    async with _client(db_session, owner) as client:
        # A is three days past a 7 day period; B is inside its 30 days.
        assert await _numbers(client, project, review_overdue="true") == [n["a"]]
        assert n["b"] in await _numbers(client, project, review_overdue="false")
        # A needed approval 24 days ago (on site in 60 days, 12 weeks lead).
        assert await _numbers(client, project, approval_late="true") == [n["a"]]
        assert set(await _numbers(client, project, approval_late="false")) == set(n.values()) - {n["a"]}
        assert await _numbers(client, project, approval_late="true", discipline="electrical") == []
        # The page is cut after the date filter, not before it.
        assert await _numbers(client, project, approval_late="false", limit=2, sort="title", order="asc") == [
            n["b"],
            n["c"],
        ]
        late = (await client.get("/v1/submittals/", params={"project_id": project, "approval_late": "true"})).json()
        assert late[0]["approval_late_days"] == 24
        assert late[0]["review_overdue_days"] == 3


@pytest.mark.asyncio
async def test_the_list_sorts_by_a_register_column_with_empty_values_last(db_session) -> None:
    owner = await _make_user(db_session)
    project, n = await _seeded_project(db_session, owner)
    async with _client(db_session, owner) as client:
        by_maker = await _numbers(client, project, sort="manufacturer", order="asc")
        assert by_maker[:2] == [n["b"], n["a"]]
        assert set(by_maker[2:]) == {n["c"], n["d"], n["legacy"]}
        by_maker_desc = await _numbers(client, project, sort="manufacturer", order="desc")
        assert by_maker_desc[:2] == [n["a"], n["b"]]
        by_number = await _numbers(client, project, sort="submittal_number", order="asc")
        assert by_number == sorted(n.values())


@pytest.mark.asyncio
async def test_the_summary_counts_the_project_and_only_the_project(db_session) -> None:
    owner = await _make_user(db_session)
    project, _ = await _seeded_project(db_session, owner)
    async with _client(db_session, owner) as client:
        response = await client.get("/v1/submittals/summary/", params={"project_id": project})
        assert response.status_code == 200, response.text
        summary = response.json()
    assert summary["total"] == 5
    assert summary["as_of"] == TODAY.isoformat()
    assert summary["awaiting_review"] == 2
    assert summary["review_overdue"] == 1
    assert summary["review_period_unknown"] == 0
    assert summary["long_lead"] == 2
    assert summary["long_lead_awaiting_approval"] == 1
    assert summary["approval_late"] == 1
    assert summary["long_lead_without_lead_time"] == 0
    assert {row["code"]: row["count"] for row in summary["by_discipline"]} == {"hvac": 1, "electrical": 2, "": 2}
    by_outcome = {row["code"]: (row["count"], row["review_code"]) for row in summary["by_outcome"]}
    assert by_outcome == {
        "approved": (1, "A"),
        "approved_as_noted": (1, "B"),
        "revise_and_resubmit": (0, "C"),
        "rejected": (0, "D"),
    }
    assert {row["code"]: row["count"] for row in summary["by_type"]} == {
        "product_data": 2,
        "shop_drawing": 2,
        "method_statement": 1,
    }


@pytest.mark.asyncio
async def test_another_users_project_is_not_found(db_session) -> None:
    owner = await _make_user(db_session)
    outsider = await _make_user(db_session)
    project, _ = await _seeded_project(db_session, owner)
    async with _client(db_session, outsider) as client:
        params = {"project_id": project}
        assert (await client.get("/v1/submittals/", params=params)).status_code == 404
        assert (await client.get("/v1/submittals/", params={**params, "approval_late": "true"})).status_code == 404
        assert (await client.get("/v1/submittals/summary/", params=params)).status_code == 404
        assert (await client.get("/v1/submittals/export/", params=params)).status_code == 404
        assert (
            await client.get("/v1/submittals/export/", params={**params, "type": "shop_drawing"})
        ).status_code == 404


@pytest.mark.asyncio
async def test_the_vocabulary_is_labelled_in_the_request_language(db_session) -> None:
    owner = await _make_user(db_session)
    async with _client(db_session, owner) as client:
        turkish = (await client.get("/v1/submittals/vocabulary/", params={"locale": "tr"})).json()
        english = (await client.get("/v1/submittals/vocabulary/")).json()
    assert turkish["locale"] == "tr"
    disciplines = {row["code"]: row for row in turkish["disciplines"]}
    assert disciplines["hvac"]["label"] == "İklimlendirme"
    assert disciplines["fire_protection"]["short_code"] == "FP"
    assert [(row["code"], row["short_code"]) for row in turkish["outcomes"]] == [
        ("approved", "A"),
        ("approved_as_noted", "B"),
        ("revise_and_resubmit", "C"),
        ("rejected", "D"),
    ]
    assert {row["code"]: row["label"] for row in turkish["types"]}["shop_drawing"] == "İmalat çizimi"
    assert {row["code"]: row["label"] for row in english["disciplines"]}["plumbing"] == "Plumbing and drainage"
    assert "required_on_site_date" in english["sort_fields"]


# ── Export through the route ──────────────────────────────────────────────


def _workbook_cells(content: bytes) -> list[str]:
    sheet = load_workbook(io.BytesIO(content)).active
    return [str(cell) for row in sheet.iter_rows(values_only=True) for cell in row if cell is not None]


@pytest.mark.asyncio
async def test_the_register_export_takes_the_filters_and_names_the_shop_drawing_register(db_session) -> None:
    owner = await _make_user(db_session)
    project, n = await _seeded_project(db_session, owner)
    async with _client(db_session, owner) as client:
        whole = await client.get("/v1/submittals/export/", params={"project_id": project, "locale": "tr"})
        assert whole.status_code == 200, whole.text
        assert whole.headers["content-language"] == "tr"
        cells = _workbook_cells(whole.content)
        for number in n.values():
            assert number in cells
        for heading in ("Disiplin", "Üretici / Marka", "Menşei", "Onay Kodu", "Şantiyede Gerekli Tarih"):
            assert heading in cells, heading
        assert "Onay Belgeleri Kayıt Listesi" in cells
        assert "Other project HVAC" not in cells

        drawings = await client.get(
            "/v1/submittals/export/", params={"project_id": project, "locale": "tr", "type": "shop_drawing"}
        )
        assert drawings.status_code == 200, drawings.text
        assert "imalat-cizimleri-kayit-listesi.xlsx" in drawings.headers["content-disposition"]
        cells = _workbook_cells(drawings.content)
        assert "İmalat Çizimleri Kayıt Listesi" in cells
        assert n["b"] in cells
        assert n["legacy"] in cells
        assert n["a"] not in cells

        late = await client.get(
            "/v1/submittals/export/", params={"project_id": project, "approval_late": "true", "format": "pdf"}
        )
        assert late.status_code == 200, late.text
        assert late.content.startswith(b"%PDF")


@pytest.mark.asyncio
async def test_rows_with_every_new_column_null_list_summarise_and_export(db_session) -> None:
    """What an upgraded installation holds after the boot heal added the columns."""
    owner = await _make_user(db_session)
    project = await _make_project(db_session, owner)
    approved = await _legacy_row(db_session, project, number="SUB-001")
    waiting = await _legacy_row(db_session, project, number="SUB-002", status="under_review")
    assert approved.discipline is None
    assert approved.review_outcome is None
    assert approved.long_lead is False
    assert approved.review_history == []
    async with _client(db_session, owner) as client:
        rows = (await client.get("/v1/submittals/", params={"project_id": project, "sort": "submittal_number"})).json()
        by_number = {row["submittal_number"]: row for row in rows}
        assert by_number["SUB-001"]["review_outcome"] == "approved"
        assert by_number["SUB-001"]["review_code"] == "A"
        assert by_number["SUB-001"]["days_in_review"] == 14
        assert by_number["SUB-002"]["review_outcome"] is None
        assert by_number["SUB-002"]["review_overdue_days"] is None
        assert by_number["SUB-002"]["approval_needed_by"] is None

        summary = (await client.get("/v1/submittals/summary/", params={"project_id": project})).json()
        assert summary["total"] == 2
        assert summary["review_period_unknown"] == 1
        assert summary["review_overdue"] == 0

        for export_format in ("xlsx", "pdf"):
            register = await client.get(
                "/v1/submittals/export/", params={"project_id": project, "locale": "tr", "format": export_format}
            )
            assert register.status_code == 200, register.text
        form = await client.get(f"/v1/submittals/{waiting.id}/export/pdf/", params={"locale": "tr"})
        assert form.status_code == 200, form.text
        assert form.content.startswith(b"%PDF")
