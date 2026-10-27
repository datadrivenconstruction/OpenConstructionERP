"""CSV, Excel, JSON and DATEV LODAS writers of a payroll period."""

import io
import json
from datetime import date

import pytest

from app.modules.payroll import pay_rules, period_exports


def _row(day, pay_type, hours, multiplier, amount):
    return {
        "employee_id": "e1",
        "employee_number": "1001",
        "employee": "Anna",
        "cost_code": "KST1",
        "work_date": day,
        "pay_type": pay_type,
        "hours": hours,
        "rate": "20",
        "multiplier": multiplier,
        "amount": amount,
        "currency": "EUR",
    }


ROWS = [
    _row("2026-10-05", "regular", "8.00", "1", "160.00"),
    _row("2026-10-06", "regular", "8.00", "1", "160.00"),
    _row("2026-10-06", "overtime", "1.50", "1.25", "37.50"),
]


def _lodas():
    return pay_rules.profile_for(country="DE")["datev_lodas"]


def test_lodas_layout():
    raw = period_exports.to_lodas(
        ROWS, period_start=date(2026, 10, 1), consultant_number="1234567", client_number="12345", lodas=_lodas()
    )
    text = raw.decode("cp1252")
    assert text.endswith("\r\n")
    lines = text.split("\r\n")
    assert lines[0] == "[Allgemein]"
    assert "Ziel=LODAS" in lines
    assert "BeraterNr=1234567" in lines
    assert "[Satzbeschreibung]" in lines
    data = lines[lines.index("[Bewegungsdaten]") + 1 : -1]
    assert data == ["1;01/10/2026;1001;100;1;16,00;KST1;", "1;01/10/2026;1001;200;1;1,50;KST1;"]


def test_lodas_refuses_missing_personnel_number():
    with pytest.raises(ValueError, match="personnel number"):
        period_exports.to_lodas(
            [{**ROWS[0], "employee_number": ""}],
            period_start=date(2026, 10, 1),
            consultant_number="1",
            client_number="1",
            lodas=_lodas(),
        )


def test_csv_configurable_columns_and_unknown_column():
    out = period_exports.to_csv(ROWS, ["employee", "pay_type", "hours"])
    assert out.splitlines()[0] == "employee,pay_type,hours"
    assert out.splitlines()[3] == "Anna,overtime,1.50"
    with pytest.raises(ValueError, match="Unknown export column"):
        period_exports.to_csv(ROWS, ["salary_secret"])


def test_json_groups_by_employee():
    data = json.loads(period_exports.to_json({"id": "p1"}, ROWS))
    assert data["schema"].startswith("openconstructionerp.payroll_period")
    assert len(data["employees"]) == 1
    assert len(data["employees"][0]["lines"]) == 3


def test_xlsx_writes_numbers():
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(period_exports.to_xlsx(ROWS))).active
    assert ws.cell(1, 1).value == "employee_number"
    assert float(ws.cell(4, 10).value) == 37.5
