# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The bill's file exports are written in a worker thread, not on the event loop.

One worker serves every user of an install, and a CSV, Excel, GAEB or BC3 file
written on its event loop answers nobody else until the last line of the bill
is on the page. A bill of several thousand lines, with a workbook styled cell by
cell, holds the whole install for that long. The queries stay on the loop with
their session; only the writing of the file leaves it. The PDF export already
renders in a thread.

Each test swaps a function the writer calls for every row (or the builder
itself) for a spy that records the thread it ran on, then calls the route and
asserts that none of those calls ran on the test's own thread, which is the
loop's thread. The service and the project read are stand-ins, since the
rendering is what is under test.

Run:
    cd backend
    python -m pytest tests/modules/boq/test_bill_exports_render_off_the_event_loop.py -v
"""

from __future__ import annotations

import threading
import uuid
from collections.abc import Callable
from decimal import Decimal
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock, MagicMock

import pytest

from app.modules.boq import router


def _recording(seen: list[int], real: Callable[..., Any]) -> Callable[..., Any]:
    def spy(*args: Any, **kwargs: Any) -> Any:
        seen.append(threading.get_ident())
        return real(*args, **kwargs)

    return spy


def _position(section_id: uuid.UUID) -> SimpleNamespace:
    return SimpleNamespace(
        id=uuid.uuid4(),
        parent_id=section_id,
        ordinal="01.0010",
        description="RC wall C30/37",
        unit="m3",
        quantity=Decimal("12.5"),
        unit_rate=Decimal("185.00"),
        total=Decimal("2312.50"),
        classification={"din276": "331"},
        metadata={},
        cad_element_ids=[],
        source="manual",
        confidence=None,
        wbs_id=None,
        reference_code=None,
    )


def _bill() -> tuple[SimpleNamespace, SimpleNamespace]:
    """The flat bill and its structured form, as the service hands them over."""
    section_id = uuid.uuid4()
    position = _position(section_id)
    section_row = SimpleNamespace(
        id=section_id,
        parent_id=None,
        ordinal="01",
        description="Concrete works",
        unit="section",
        quantity=Decimal("0"),
        unit_rate=Decimal("0"),
        total=Decimal("0"),
    )
    project_id = uuid.uuid4()
    flat = SimpleNamespace(name="Tower A", project_id=project_id, positions=[section_row, position])
    structured = SimpleNamespace(
        name="Tower A",
        project_id=project_id,
        sections=[
            SimpleNamespace(
                id=section_id,
                ordinal="01",
                description="Concrete works",
                positions=[position],
                subtotal=Decimal("2312.50"),
            )
        ],
        positions=[],
        markups=[],
        direct_cost=Decimal("2312.50"),
        net_total=Decimal("2312.50"),
        grand_total=Decimal("2312.50"),
    )
    return flat, structured


def _service() -> Any:
    flat, structured = _bill()
    return SimpleNamespace(
        get_boq_structured_for_export=AsyncMock(return_value=structured),
        get_boq_with_positions=AsyncMock(return_value=flat),
        get_boq=AsyncMock(return_value=SimpleNamespace(metadata_={})),
        get_export_fx=AsyncMock(return_value=("EUR", {"USD": "0.92"})),
    )


@pytest.fixture
def no_owner_check(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.modules.projects import repository as projects_repository

    project = SimpleNamespace(name="Riverside", currency="EUR", classification_standard="din276", region="DACH")

    class _Projects:
        def __init__(self, _session: Any) -> None:
            self.get_by_id = AsyncMock(return_value=project)

    monkeypatch.setattr(router, "_verify_boq_owner", AsyncMock(return_value=None))
    monkeypatch.setattr(projects_repository, "ProjectRepository", _Projects)


async def _body(response: Any) -> bytes:
    chunks = []
    async for chunk in response.body_iterator:
        chunks.append(chunk if isinstance(chunk, bytes) else chunk.encode("utf-8"))
    return b"".join(chunks)


async def _call(handler: Callable[..., Any], **extra: Any) -> bytes:
    response = await handler(
        boq_id=uuid.uuid4(), _user_id="u1", payload={}, session=MagicMock(), service=_service(), **extra
    )
    return await _body(response)


@pytest.mark.asyncio
@pytest.mark.usefixtures("no_owner_check")
async def test_the_csv_is_written_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    seen: list[int] = []
    monkeypatch.setattr(router, "neutralise_formula", _recording(seen, router.neutralise_formula))

    body = await _call(router.export_boq_csv)

    assert b"RC wall C30/37" in body
    assert b"Grand Total" in body
    assert seen, "the CSV writer never ran"
    assert threading.get_ident() not in seen


@pytest.mark.asyncio
@pytest.mark.usefixtures("no_owner_check")
async def test_the_workbook_is_written_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    import io

    from openpyxl import load_workbook

    seen: list[int] = []
    monkeypatch.setattr(router, "neutralise_formula", _recording(seen, router.neutralise_formula))

    body = await _call(router.export_boq_excel)

    assert body[:2] == b"PK"
    values = [cell for row in load_workbook(io.BytesIO(body)).active.iter_rows(values_only=True) for cell in row]
    assert "RC wall C30/37" in values
    assert any(isinstance(v, str) and v.startswith("Project: Riverside") for v in values)
    assert seen, "the workbook writer never ran"
    assert threading.get_ident() not in seen


@pytest.mark.asyncio
@pytest.mark.usefixtures("no_owner_check")
async def test_the_gaeb_file_is_built_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    seen: list[int] = []
    monkeypatch.setattr(router, "build_gaeb_xml", _recording(seen, router.build_gaeb_xml))

    body = await _call(router.export_boq_gaeb, gaeb_format="x83", bid_type="main")

    assert b"GAEB" in body
    assert b"RC wall C30/37" in body
    assert seen, "the GAEB builder never ran"
    assert threading.get_ident() not in seen


@pytest.mark.asyncio
@pytest.mark.usefixtures("no_owner_check")
async def test_the_bc3_file_is_built_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.modules.boq.exporters import bc3

    seen: list[int] = []
    monkeypatch.setattr(bc3, "build_bc3", _recording(seen, bc3.build_bc3))

    body = await _call(router.export_boq_bc3)

    assert b"~V" in body
    assert b"RC wall C30/37" in body
    assert seen, "the BC3 builder never ran"
    assert threading.get_ident() not in seen
