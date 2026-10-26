# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Exchange fidelity reports and proportional scaling, against real rows.

* every interchange export stores its report and names it in a header; the
  register lists it, a report of another bill answers 404, the download
  carries the items as CSV and XLSX;
* a check builds the export in memory, round-trips it through the matching
  reader and is stored as a check that never feeds validation;
* a native import returns its report and stores it, and the validation rule
  turns the critical items of the latest real exchange into results;
* scaling previews without writing, applies with a snapshot that restores
  the old norms, re-derives the unit rate, skips a locked line, and refuses a
  locked bill, a stale preview and an apply without a target.

Every read-back is a fresh SELECT.
"""

from __future__ import annotations

import io
import uuid
from decimal import Decimal

import pytest
import pytest_asyncio
from fastapi import HTTPException
from sqlalchemy import select, update

from app.core.validation.engine import ValidationContext
from app.core.validation.rules import ExchangeFidelityCritical
from app.modules.boq.exchange_fidelity_router import (
    ExchangeCheckRequest,
    ResourceScaleBody,
    apply_resource_scaling,
    check_export_fidelity,
    download_exchange_report,
    exchange_fidelity_router,
    get_exchange_report,
    list_exchange_reports,
    preview_resource_scaling,
)
from app.modules.boq.exchange_fidelity_service import REPORT_HEADER, with_exchange_fidelity
from app.modules.boq.models import BOQ, BOQExchangeReport, BOQMarkup, BOQSnapshot, Position
from app.modules.boq.resource_scaling_service import SCALING_SNAPSHOT_NAME
from app.modules.boq.router import (
    _run_native_import,
    export_boq_bc3,
    export_boq_csv,
    export_boq_excel,
    export_boq_gaeb,
)
from app.modules.boq.schemas import PositionCreate
from app.modules.boq.service import BOQService
from app.modules.projects.models import Project
from app.modules.users.models import User
from tests._pg import transactional_session

pytestmark = pytest.mark.asyncio

ROLE = {"role": "editor"}


@pytest_asyncio.fixture
async def session():
    async with transactional_session() as s:
        yield s


async def _user(session) -> User:
    user = User(email=f"fid-{uuid.uuid4().hex[:8]}@test.io", hashed_password="x", full_name="Estimator", role="editor")
    session.add(user)
    await session.flush()
    return user


def _labour(qty: float, name: str = "Mason", rate: float = 40.0) -> dict:
    return {"name": name, "code": name.upper(), "type": "labor", "unit": "h", "quantity": qty, "unit_rate": rate}


async def _bill(session, owner: User) -> tuple[BOQ, dict[str, Position]]:
    project = Project(name=f"Depot {uuid.uuid4().hex[:6]}", owner_id=owner.id, currency="EUR", country_code="DE")
    session.add(project)
    await session.flush()
    boq = BOQ(project_id=project.id, name="Depot shell", status="draft", metadata_={})
    session.add(boq)
    await session.flush()
    svc = BOQService(session)
    section = await svc.add_position(
        PositionCreate(boq_id=boq.id, ordinal="01", description="Shell", unit="section", quantity=0, unit_rate=0)
    )
    rows = (
        ("01.0010", "Walls", "m3", 10, 100, [_labour(2.0)]),
        ("01.0020", "Slabs", "m2", 5, 80, [_labour(4.0, "Carpenter")]),
        ("01.0030", "Pallet racks", "pallets", 3, 50, None),
    )
    for ordinal, text, unit, qty, rate, resources in rows:
        await svc.add_position(
            PositionCreate(
                boq_id=boq.id,
                parent_id=section.id,
                ordinal=ordinal,
                description=text,
                unit=unit,
                quantity=qty,
                unit_rate=rate,
                metadata={"resources": resources} if resources else {},
            )
        )
    await session.flush()
    found = (await session.execute(select(Position).where(Position.boq_id == boq.id))).scalars().all()
    return boq, {p.ordinal: p for p in found}


async def _body(response) -> bytes:
    chunks = [chunk async for chunk in response.body_iterator]
    return b"".join(c if isinstance(c, bytes) else c.encode("utf-8") for c in chunks)


async def _reports(session, boq_id) -> list[BOQExchangeReport]:
    stmt = select(BOQExchangeReport).where(BOQExchangeReport.boq_id == boq_id).order_by(BOQExchangeReport.created_at)
    return list((await session.execute(stmt)).scalars().all())


# ── exports store their report ─────────────────────────────────────────────


async def test_each_interchange_export_stores_its_report_and_names_it_in_a_header(session) -> None:
    owner = await _user(session)
    boq, _pos = await _bill(session, owner)
    svc = BOQService(session)
    uid = str(owner.id)
    gaeb = await export_boq_gaeb(
        boq.id, uid, ROLE, session, svc, gaeb_format="x83", bid_type="main", embed_markups=None
    )
    bc3 = await export_boq_bc3(boq.id, uid, ROLE, session, svc, embed_markups=None)
    xlsx = await export_boq_excel(boq.id, uid, ROLE, session, svc, embed_markups=None)
    csv = await export_boq_csv(boq.id, uid, ROLE, session, svc, embed_markups=None)
    for response in (gaeb, bc3, xlsx, csv):
        assert response.headers.get(REPORT_HEADER)
        assert await _body(response)

    stored = await _reports(session, boq.id)
    assert [(r.format_id, r.variant, r.direction, r.is_check) for r in stored] == [
        ("gaeb_xml", "x83", "export", False),
        ("bc3", "", "export", False),
        ("excel", "", "export", False),
        ("csv", "", "export", False),
    ]
    gaeb_items = {(i["kind"], i["field"]) for i in stored[0].items}
    assert ("truncated_text", "unit") in gaeb_items
    assert ("dropped_attribute", "resources") in gaeb_items
    assert stored[0].summary["total"] == len(stored[0].items)


async def test_register_detail_scope_and_downloads(session) -> None:
    owner = await _user(session)
    boq, _pos = await _bill(session, owner)
    other, _other_pos = await _bill(session, owner)
    svc = BOQService(session)
    uid = str(owner.id)
    response = await export_boq_bc3(boq.id, uid, ROLE, session, svc, embed_markups=None)
    report_id = uuid.UUID(response.headers[REPORT_HEADER])

    listed = await list_exchange_reports(boq.id, uid, ROLE, session, direction=None, limit=50)
    assert [r["id"] for r in listed["items"]] == [str(report_id)]
    detail = await get_exchange_report(boq.id, report_id, uid, ROLE, session)
    assert detail["format_id"] == "bc3" and isinstance(detail["items"], list)
    with pytest.raises(HTTPException) as err:
        await get_exchange_report(other.id, report_id, uid, ROLE, session)
    assert err.value.status_code == 404

    from app.core.i18n import load_translations

    load_translations()
    csv = await download_exchange_report(
        boq.id, report_id, uid, ROLE, session, file_format="csv", severity=None, kind=None
    )
    text = (await _body(csv)).decode("utf-8")
    assert text.startswith("﻿Severity")
    xlsx = await download_exchange_report(
        boq.id, report_id, uid, ROLE, session, file_format="xlsx", severity="warning", kind=None
    )
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(await _body(xlsx))).active
    assert ws["A1"].value == "Severity"
    assert all(row[0].value == "Warning" for row in ws.iter_rows(min_row=2))


@pytest.mark.parametrize(
    ("format_id", "variant", "expected"),
    [
        # X83 cuts the unit code to four characters; the round trip confirms it.
        ("gaeb_xml", "x83", {("unit", "01.0030")}),
        ("gaeb_xml", "x84", set()),
        ("bc3", "", set()),
        ("excel", "", set()),
    ],
)
async def test_a_check_round_trips_the_bill_and_confirms_only_real_losses(
    session, format_id: str, variant: str, expected: set
) -> None:
    owner = await _user(session)
    boq, _pos = await _bill(session, owner)
    body = ExchangeCheckRequest(format_id=format_id, variant=variant, roundtrip=True)
    report = await check_export_fidelity(boq.id, body, str(owner.id), ROLE, session)
    assert report["roundtrip"] is True and report["is_check"] is True
    mismatches = {(i["field"], i["position_ref"]) for i in report["items"] if i["kind"] == "roundtrip_mismatch"}
    assert mismatches == expected


async def test_a_bc3_check_round_trip_catches_a_negative_quantity_and_markups_are_reported(session) -> None:
    owner = await _user(session)
    boq, pos = await _bill(session, owner)
    pos["01.0030"].quantity = "-3"
    pos["01.0030"].total = "-150"
    session.add(BOQMarkup(boq_id=boq.id, name="Overhead", markup_type="percentage", percentage="10", sort_order=1))
    await session.flush()
    report = await check_export_fidelity(boq.id, ExchangeCheckRequest(format_id="bc3"), str(owner.id), ROLE, session)
    kinds = {(i["kind"], i["field"], i["severity"]) for i in report["items"]}
    assert ("value_out_of_range", "quantity", "error") in kinds
    assert ("roundtrip_mismatch", "quantity", "error") in kinds
    assert ("markup_not_representable", "markup", "warning") in kinds


# ── validation ─────────────────────────────────────────────────────────────


async def test_critical_export_items_reach_the_validation_rule_and_checks_do_not(session) -> None:
    owner = await _user(session)
    boq, pos = await _bill(session, owner)
    pos["01.0010"].quantity = "-1"
    pos["01.0010"].total = "-100"
    await session.flush()
    await check_export_fidelity(boq.id, ExchangeCheckRequest(format_id="bc3"), str(owner.id), ROLE, session)
    data = await with_exchange_fidelity(session, boq.id, {"positions": []})
    assert data["exchange_fidelity"] == []

    await export_boq_bc3(boq.id, str(owner.id), ROLE, session, BOQService(session), embed_markups=None)
    data = await with_exchange_fidelity(session, boq.id, {"positions": []})
    (latest,) = data["exchange_fidelity"]
    assert latest["format_id"] == "bc3" and latest["items"]
    results = await ExchangeFidelityCritical().validate(ValidationContext(data=data, metadata={"locale": "en"}))
    failed = [r for r in results if not r.passed]
    assert failed and failed[0].severity.value == "warning"
    assert failed[0].element_ref == str(pos["01.0010"].id)
    assert "01.0010" in failed[0].message


async def test_the_rule_says_nothing_about_a_bill_without_reports() -> None:
    assert await ExchangeFidelityCritical().validate(ValidationContext(data={"positions": []})) == []


# ── import ─────────────────────────────────────────────────────────────────

_BC3 = "\r\n".join(
    [
        "~V|Obra|FIEBDC-3/2020|Prog|01012026|Obra|ANSI|2||EUR|",
        "~C|OBRA##||Obra|0|||3|",
        "~C|C01#||Capitulo|0|||1|",
        "~C|P1|m3|Muro|100|||0|",
        "~D|OBRA##|C01#\\1\\1|",
        "~D|C01#|P1\\1\\10|",
        "~G|P1|foto.jpg|",
        "~M|C01#\\P1|1|10||",
    ]
)


async def test_a_native_import_returns_and_stores_its_report(session) -> None:
    from app.modules.boq.importers.bc3 import BC3Importer

    owner = await _user(session)
    boq, _pos = await _bill(session, owner)
    result = await _run_native_import(
        boq.id,
        BC3Importer,
        _BC3.encode("cp1252"),
        file_name="obra.bc3",
        overrides=None,
        delete_missing=False,
        actor_id=owner.id,
        service=BOQService(session),
    )
    fidelity = result["fidelity"]
    assert fidelity["id"]
    assert any(i["kind"] == "attachment_skipped" for i in fidelity["items"])
    (stored,) = [r for r in await _reports(session, boq.id) if r.direction == "import"]
    assert stored.format_id == "bc3" and stored.file_name == "obra.bc3"


# ── scaling ────────────────────────────────────────────────────────────────


def _scale(section_id, target=None, **kw) -> ResourceScaleBody:
    return ResourceScaleBody(metric="hours", target=target, section_id=section_id, **kw)


async def _resources(session, position_id) -> tuple[list, str, str]:
    row = (
        await session.execute(
            select(Position.metadata_, Position.unit_rate, Position.total).where(Position.id == position_id)
        )
    ).one()
    return list((row[0] or {}).get("resources") or []), str(row[1]), str(row[2])


async def test_preview_measures_and_proposes_without_writing(session) -> None:
    owner = await _user(session)
    boq, pos = await _bill(session, owner)
    uid = str(owner.id)
    measured = await preview_resource_scaling(boq.id, _scale(pos["01"].id), uid, ROLE, session)
    assert Decimal(measured["current"]) == Decimal(40)
    assert measured["counts"]["scaled"] == 0
    proposed = await preview_resource_scaling(boq.id, _scale(pos["01"].id, Decimal(20)), uid, ROLE, session)
    assert Decimal(proposed["factor"]) == Decimal("0.5")
    resources, _rate, _total = await _resources(session, pos["01.0010"].id)
    assert resources[0]["quantity"] == 2.0


async def test_apply_rewrites_norms_rederives_rates_logs_and_restores_from_the_snapshot(session) -> None:
    owner = await _user(session)
    boq, pos = await _bill(session, owner)
    uid = str(owner.id)
    before = await _resources(session, pos["01.0010"].id)
    session.expunge_all()
    preview = await preview_resource_scaling(boq.id, _scale(pos["01"].id, Decimal(20)), uid, ROLE, session)
    result = await apply_resource_scaling(
        boq.id, _scale(pos["01"].id, Decimal(20), expected_factor=Decimal(preview["factor"])), uid, ROLE, session
    )
    assert result["updated"] == 2
    resources, rate, total = await _resources(session, pos["01.0010"].id)
    assert resources[0]["quantity"] == 1.0
    assert Decimal(rate) == Decimal(40)
    assert Decimal(total) == Decimal(400)

    snap = (
        await session.execute(select(BOQSnapshot).where(BOQSnapshot.id == uuid.UUID(result["snapshot_id"])))
    ).scalar_one()
    assert snap.name == SCALING_SNAPSHOT_NAME
    await BOQService(session).restore_snapshot(boq.id, snap.id)
    restored = await session.execute(
        select(Position.metadata_).where(Position.boq_id == boq.id, Position.ordinal == "01.0010")
    )
    assert restored.scalar_one()["resources"][0]["quantity"] == before[0][0]["quantity"]


async def test_a_locked_line_keeps_its_norms(session) -> None:
    owner = await _user(session)
    boq, pos = await _bill(session, owner)
    pos["01.0020"].metadata_ = {**pos["01.0020"].metadata_, "locked": True}
    await session.flush()
    session.expunge_all()
    result = await apply_resource_scaling(boq.id, _scale(pos["01"].id, Decimal(30)), str(owner.id), ROLE, session)
    # 20 h locked; the remaining 20 h must become 10 h.
    assert Decimal(result["plan"]["factor"]) == Decimal("0.5")
    locked, _r, _t = await _resources(session, pos["01.0020"].id)
    assert locked[0]["quantity"] == 4.0


async def test_apply_refuses_a_locked_bill_a_stale_preview_and_a_missing_target(session) -> None:
    owner = await _user(session)
    boq, pos = await _bill(session, owner)
    uid = str(owner.id)
    with pytest.raises(HTTPException) as err:
        await apply_resource_scaling(boq.id, _scale(pos["01"].id, None), uid, ROLE, session)
    assert err.value.status_code == 422
    with pytest.raises(HTTPException) as err:
        await apply_resource_scaling(
            boq.id, _scale(pos["01"].id, Decimal(20), expected_factor=Decimal("0.9")), uid, ROLE, session
        )
    assert err.value.status_code == 409 and err.value.detail["code"] == "scaling_preview_stale"
    with pytest.raises(HTTPException) as err:
        await preview_resource_scaling(
            boq.id, _scale(pos["01"].id, Decimal(1), resource_type="material"), uid, ROLE, session
        )
    assert err.value.detail["code"] == "scaling_nothing_to_scale"
    await session.execute(update(BOQ).where(BOQ.id == boq.id).values(is_locked=True))
    session.expunge_all()
    with pytest.raises(HTTPException) as err:
        await apply_resource_scaling(boq.id, _scale(pos["01"].id, Decimal(20)), uid, ROLE, session)
    assert err.value.status_code == 409


async def test_a_stranger_gets_404_and_apply_needs_the_update_permission(session) -> None:
    owner = await _user(session)
    stranger = await _user(session)
    boq, pos = await _bill(session, owner)
    with pytest.raises(HTTPException) as err:
        await preview_resource_scaling(boq.id, _scale(pos["01"].id), str(stranger.id), ROLE, session)
    assert err.value.status_code == 404
    perms = {}
    for route in exchange_fidelity_router.routes:
        for dep in getattr(route, "dependencies", []):
            perm = getattr(dep.dependency, "permission", None)
            if perm:
                perms[route.path] = perm
    assert perms["/boqs/{boq_id}/resources/scale/apply/"] == "boq.update"
    assert perms["/boqs/{boq_id}/resources/scale/preview/"] == "boq.read"
