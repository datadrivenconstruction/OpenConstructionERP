# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Forecast methods, documented snapshots, comparison and exports, end to end on the database.

The project is the one of ``test_cost_control_service``: the shell line (330)
is awarded at 100 000 against an estimate of 110 000 and a budget of 100 000,
with 60 000 certified. Under the best available source it reads 100 000, under
the maximum method 110 000 (its estimate).
"""

from __future__ import annotations

import io
import uuid
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.costmodel.cost_control_export import EXPORT_KINDS, render_export, rendered_language
from app.modules.costmodel.cost_control_forecast import ForecastDocsService
from app.modules.costmodel.cost_control_phases import phase_comparison
from app.modules.costmodel.cost_control_schemas import (
    CostControlSettingsRequest,
    EvidenceItem,
    ForecastOverrideRequest,
    LineMethodRequest,
    SnapshotCreateRequest,
    SnapshotNoteRequest,
)
from app.modules.costmodel.cost_control_service import CostControlService
from app.modules.costmodel.models import CostForecastOverride, FrozenSnapshotError
from tests.modules.costmodel.test_cost_control_service import (  # noqa: F401  (fixtures)
    _row,
    project,
    session,
)

pytestmark = pytest.mark.asyncio

D = Decimal
V20 = "cost_forecast.move_without_evidence"
V22 = "cost_forecast.method_override_without_comment"


def _key(line: object) -> str:
    return str(line.id)  # type: ignore[attr-defined]


async def _frozen_baseline(session: AsyncSession, pid: uuid.UUID) -> uuid.UUID:
    service = CostControlService(session)
    docs = ForecastDocsService(session)
    snap = await service.create_snapshot(pid, SnapshotCreateRequest(stage="award", label="Baseline"), "u1")
    await docs.freeze(await docs.snapshot(snap.id), "u1")
    return snap.id


async def test_the_project_method_moves_every_line_and_defaults_to_best_available(
    session: AsyncSession, project: dict[str, object]
) -> None:
    pid = project["pid"]
    docs = ForecastDocsService(session)
    settings = await docs.get_settings(pid)  # type: ignore[arg-type]
    assert (settings.forecast_method, settings.is_default) == ("best_available", True)

    report = await CostControlService(session).report(pid)  # type: ignore[arg-type]
    assert report.forecast_method == "best_available"
    assert _row(report, _key(project["shell"])).forecast == D("100000")

    saved = await docs.put_settings(pid, CostControlSettingsRequest(forecast_method="maximum"), "u1")  # type: ignore[arg-type]
    assert (saved.forecast_method, saved.is_default) == ("maximum", False)
    report = await CostControlService(session).report(pid)  # type: ignore[arg-type]
    shell = _row(report, _key(project["shell"]))
    assert (shell.forecast, shell.forecast_method, shell.method_origin) == (D("110000"), "maximum", "project")
    # A finished contract reads its final figure under every method.
    windows = _row(report, _key(project["windows"]))
    assert (windows.forecast, windows.forecast_source) == (D("43200"), "final")


async def test_a_line_method_and_a_manual_amount_exclude_each_other(
    session: AsyncSession, project: dict[str, object]
) -> None:
    pid, shell = project["pid"], project["shell"]
    docs = ForecastDocsService(session)
    control = CostControlService(session)

    await control.set_override(shell, ForecastOverrideRequest(amount=D("105000"), reason="Site estimate"), "u1")  # type: ignore[arg-type]
    choice = await docs.set_line_method(shell, LineMethodRequest(method="maximum", comment=""), "u1")  # type: ignore[arg-type]
    assert choice.method == "maximum"
    open_overrides = (
        (
            await session.execute(
                select(CostForecastOverride).where(
                    CostForecastOverride.cost_line_id == shell.id,  # type: ignore[attr-defined]
                    CostForecastOverride.closed_at.is_(None),
                )
            )
        )
        .scalars()
        .all()
    )
    assert open_overrides == [], "choosing a method closes the manual amount"

    report = await control.report(pid)  # type: ignore[arg-type]
    row = _row(report, _key(shell))
    assert (row.forecast, row.forecast_method, row.method_origin) == (D("110000"), "maximum", "line")

    # V22: another method than the project's without a word.
    panel = await docs.panel(pid)  # type: ignore[arg-type]
    assert [f.element_ref for f in panel.findings if f.rule_id == V22] == [_key(shell)]
    assert panel.line_choices == 1

    await docs.set_line_method(shell, LineMethodRequest(method="maximum", comment="Open claim, cautious view"), "u1")  # type: ignore[arg-type]
    panel = await docs.panel(pid)  # type: ignore[arg-type]
    assert [f for f in panel.findings if f.rule_id == V22] == []

    # A manual amount takes the line off its method again.
    await control.set_override(
        shell, ForecastOverrideRequest(amount=D("104000"), reason="Agreed with the client"), "u1"
    )  # type: ignore[arg-type]
    with pytest.raises(HTTPException) as gone:
        await docs.clear_line_method(shell)  # type: ignore[arg-type]
    assert gone.value.status_code == 404


async def test_a_draft_is_documented_then_frozen_and_never_changes_again(
    session: AsyncSession, project: dict[str, object]
) -> None:
    pid, shell = project["pid"], project["shell"]
    control = CostControlService(session)
    docs = ForecastDocsService(session)

    first = await _frozen_baseline(session, pid)  # type: ignore[arg-type]

    # Only frozen snapshots are records: the baseline cannot be noted, refreshed or deleted.
    frozen = await docs.snapshot(first)
    with pytest.raises(HTTPException) as refused:
        await docs.put_note(frozen, _key(shell), SnapshotNoteRequest(comment="late"), "u1")
    assert refused.value.status_code == 409
    with pytest.raises(HTTPException) as refused:
        await docs.refresh(frozen)
    assert refused.value.status_code == 409
    with pytest.raises(HTTPException) as refused:
        await control.delete_snapshot(first)
    assert refused.value.status_code == 409
    # And the database layer refuses it even when the service is bypassed.
    frozen.label = "rewritten"
    with pytest.raises(FrozenSnapshotError):
        await session.flush()
    await session.rollback()


async def test_freeze_waits_for_the_evidence_behind_a_large_move(
    session: AsyncSession, project: dict[str, object]
) -> None:
    pid, shell = project["pid"], project["shell"]
    control = CostControlService(session)
    docs = ForecastDocsService(session)
    first = await _frozen_baseline(session, pid)  # type: ignore[arg-type]

    # The shell moves 100 000 -> 110 000: above 5 %, so it needs a reason.
    await docs.put_settings(
        pid,  # type: ignore[arg-type]
        CostControlSettingsRequest(forecast_method="maximum", evidence_threshold_pct=D("5")),
        "u1",
    )
    draft = await control.create_snapshot(pid, SnapshotCreateRequest(stage="forecast"), "u1")  # type: ignore[arg-type]
    assert draft.status == "draft" and draft.previous_snapshot_id == first
    with pytest.raises(HTTPException) as twice:
        await control.create_snapshot(pid, SnapshotCreateRequest(stage="forecast"), "u1")  # type: ignore[arg-type]
    assert twice.value.status_code == 409

    snap = await docs.snapshot(draft.id)
    detail = await docs.snapshot_detail(snap)
    moved = {line.key: line for line in detail.lines if line.needs_evidence}
    assert set(moved) == {_key(shell)}
    assert moved[_key(shell)].delta == D("10000")

    with pytest.raises(HTTPException) as refused:
        await docs.freeze(snap, "u1")
    assert refused.value.status_code == 422
    assert [f["rule_id"] for f in refused.value.detail["findings"]] == [V20]  # type: ignore[index]

    # A comment alone is not enough; with the minutes it is.
    await docs.put_note(snap, _key(shell), SnapshotNoteRequest(comment="Estimate stands until the claim settles"), "u1")
    with pytest.raises(HTTPException):
        await docs.freeze(snap, "u1")
    line = await docs.put_note(
        snap,
        _key(shell),
        SnapshotNoteRequest(
            comment="Estimate stands until the claim settles",
            evidence=[EvidenceItem(kind="minutes", ref="Site meeting 14, item 3")],
        ),
        "u1",
    )
    assert line.documented is True

    # A refresh keeps the notes on their lines.
    await docs.refresh(snap)
    assert (await docs.snapshot_detail(snap)).lines

    done = await docs.freeze(snap, "u1")
    assert done.status == "frozen" and done.frozen_by == "u1"

    # Compare the two frozen snapshots with today.
    matrix = await docs.matrix(pid)  # type: ignore[arg-type]
    assert [c.status for c in matrix.columns] == ["frozen", "frozen", "live"]
    shell_row = next(r for r in matrix.rows if r.key == _key(shell))
    assert shell_row.values == [D("100000"), D("110000"), D("110000")]
    assert shell_row.delta == D("10000")

    ladder = await docs.ladder(pid, start=str(first), end=str(draft.id))  # type: ignore[arg-type]
    assert ladder.end - ladder.start == sum((s.delta for s in ladder.steps), D("0")) + ladder.other
    step = next(s for s in ladder.steps if s.key == _key(shell))
    assert (step.delta, step.cause) == (D("10000"), "method")
    assert step.comment == "Estimate stands until the claim settles"
    assert [e.ref for e in step.evidence] == ["Site meeting 14, item 3"]


async def test_the_phase_comparison_adds_up_on_every_layer(session: AsyncSession, project: dict[str, object]) -> None:
    pid = project["pid"]
    phases = await phase_comparison(session, pid)  # type: ignore[arg-type]
    report = await CostControlService(session).report(pid)  # type: ignore[arg-type]
    assert phases.totals.values["budget"] == D("140000")
    assert phases.totals.values["forecast"] == sum((r.forecast for r in report.rows), D("0"))
    # Top-level rows alone make the total: nothing counted twice.
    top = [r for r in phases.rows if r.parent_key is None]
    for layer in phases.layers:
        assert sum((r.values[layer] for r in top), D("0")) == phases.totals.values[layer]
    shell = next(r for r in phases.rows if r.key == _key(project["shell"]))
    assert shell.pct["forecast"] == "100.0"
    assert shell.values["certified"] == D("60000")


@pytest.mark.parametrize("kind", EXPORT_KINDS)
async def test_every_export_renders_as_pdf_and_workbook(
    session: AsyncSession, project: dict[str, object], kind: str
) -> None:
    pid = project["pid"]
    await _frozen_baseline(session, pid)  # type: ignore[arg-type]
    pdf, name = await render_export(session, pid, kind, "pdf", "de")  # type: ignore[arg-type]
    assert pdf.startswith(b"%PDF") and name.endswith(".pdf")
    xlsx, name = await render_export(session, pid, kind, "xlsx", "en")  # type: ignore[arg-type]
    assert name.endswith(".xlsx")
    sheet = load_workbook(io.BytesIO(xlsx)).active
    assert sheet is not None and sheet.max_row >= 1


async def test_an_export_declares_the_language_its_texts_came_out_in() -> None:
    from app.core.i18n import load_translations

    load_translations()
    assert rendered_language("de") == "de"
    # A regional code is written in its base language, and a language with no catalogue in English.
    assert rendered_language("pt-BR") == "pt"
    assert rendered_language("xx") == "en"
